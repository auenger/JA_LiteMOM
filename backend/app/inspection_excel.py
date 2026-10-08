"""Read the exact workbook layout produced by the full inspection-sheet export."""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
import xlrd


HEADERS = (
    "序号", "设备品种", "产品编号", "生产日期", "生产厂家", "容积",
    "检验日期", "下次检验日期", "检验结果", "检验人员",
)


def _has_full_inspection_headers(title_row: tuple, header_row: tuple) -> bool:
    return (
        bool(title_row)
        and str(title_row[0] or "").strip() == "充装单位名称"
        and tuple(str(value or "").strip() for value in header_row[:10]) == HEADERS
    )


def _date_text(value: object) -> str:
    if isinstance(value, (float, Decimal)):
        return f"{value:.2f}"
    return str(value).strip()


def _date(value: object, row_number: int, label: str, *, allow_month: bool = False) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _date_text(value).replace(".", "-").replace("/", "-")
    if allow_month and re.fullmatch(r"\d{4}-\d{1,2}", text):
        text += "-01"
    try:
        return date.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"Excel 第 {row_number} 行的{label}不是有效日期") from exc


def _inspection_date_matches(value: object, selected_date: date, row_number: int) -> bool:
    if not isinstance(value, (datetime, date)):
        month_text = _date_text(value).replace(".", "-").replace("/", "-")
        if re.fullmatch(r"\d{4}-\d{1,2}", month_text):
            month_date = _date(value, row_number, "检验日期", allow_month=True)
            return month_date is not None and (month_date.year, month_date.month) == (selected_date.year, selected_date.month)
    return _date(value, row_number, "检验日期") == selected_date


def parse_full_inspection_excel(
    content: bytes, inspection_date: date, filename: str, *,
    match_by_sequence: bool = False, require_matching_date: bool = True,
) -> tuple[str, str, list[dict]]:
    if filename.lower().endswith(".xls"):
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        close_workbook = workbook.release_resources
        if "检验气瓶列表" in workbook.sheet_names():
            sheet = workbook.sheet_by_name("检验气瓶列表")
        else:
            matching_sheets = [candidate for candidate in workbook.sheets() if candidate.nrows >= 2 and
                _has_full_inspection_headers(candidate.row_values(0), candidate.row_values(1))]
            if len(matching_sheets) != 1:
                close_workbook()
                raise ValueError("Excel 未找到唯一一张符合全量检验单表头的工作表，请使用导出全量检验单的格式")
            sheet = matching_sheets[0]

        def xls_rows():
            for row_number in range(sheet.nrows):
                values = []
                for column in range(10):
                    if column >= sheet.ncols:
                        values.append(None)
                        continue
                    cell = sheet.cell(row_number, column)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and cell.value.is_integer():
                        values.append(int(cell.value))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        values.append(None)
                    else:
                        values.append(cell.value)
                yield tuple(values)

        rows = xls_rows()
    else:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        close_workbook = workbook.close
        if "检验气瓶列表" in workbook.sheetnames:
            sheet = workbook["检验气瓶列表"]
        else:
            matching_sheets = []
            for candidate in workbook:
                first_rows = candidate.iter_rows(min_row=1, max_row=2, max_col=10, values_only=True)
                if _has_full_inspection_headers(next(first_rows, ()), next(first_rows, ())):
                    matching_sheets.append(candidate)
            if len(matching_sheets) != 1:
                close_workbook()
                raise ValueError("Excel 未找到唯一一张符合全量检验单表头的工作表，请使用导出全量检验单的格式")
            sheet = matching_sheets[0]
        rows = sheet.iter_rows(min_row=1, max_col=10, values_only=True)
    try:
        title_row = next(rows, None)
        header_row = next(rows, None)
        if not title_row or str(title_row[0] or "").strip() != "充装单位名称":
            raise ValueError("Excel 第 1 行不是全量检验单的客户表头")
        if tuple(str(value or "").strip() for value in (header_row or ())) != HEADERS:
            raise ValueError("Excel 第 2 行表头与导出全量检验单不一致")
        customer_name = str(title_row[2] or "").strip()
        if len(customer_name) > 160:
            raise ValueError("客户/充装单位名称不能超过 160 字")
        lines: list[dict] = []
        seen_sequences: set[int] = set()
        for row_number, row in enumerate(rows, start=3):
            if all(value is None or str(value).strip() == "" for value in row):
                continue
            if len(lines) >= 9999:
                raise ValueError("一张归档检验单最多导入 9999 行")
            try:
                sequence = int(str(row[0]).strip())
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Excel 第 {row_number} 行序号无效") from exc
            if match_by_sequence and (sequence < 1 or sequence in seen_sequences):
                raise ValueError(f"Excel 第 {row_number} 行序号重复或无效：{sequence}")
            if not match_by_sequence and sequence != len(lines) + 1:
                raise ValueError(f"Excel 第 {row_number} 行序号应为 {len(lines) + 1}，请按原表顺序连续编号")
            seen_sequences.add(sequence)
            if not isinstance(row[2], str) or not row[2].strip():
                raise ValueError(f"Excel 第 {row_number} 行产品编号必须是文本，避免前导零丢失")
            body_code = row[2].strip()
            if len(body_code) > 100:
                raise ValueError(f"Excel 第 {row_number} 行产品编号不能超过 100 字")
            device_category = str(row[1] or "").strip()
            manufacturer = str(row[4] or "").strip()
            inspector_name = str(row[9] or "").strip()
            result = str(row[8] or "").strip()
            if not device_category or len(device_category) > 80:
                raise ValueError(f"Excel 第 {row_number} 行设备品种不能为空且不能超过 80 字")
            if len(manufacturer) > 160 or len(inspector_name) > 80 or len(result) > 20:
                raise ValueError(f"Excel 第 {row_number} 行文本字段超过长度限制")
            if _date(row[6], row_number, "检验日期", allow_month=True) is None:
                raise ValueError(f"Excel 第 {row_number} 行缺少检验日期")
            if require_matching_date and not _inspection_date_matches(row[6], inspection_date, row_number):
                raise ValueError(f"Excel 第 {row_number} 行检验日期与所选日期不一致（年月格式只校验月份）")
            next_date = _date(row[7], row_number, "下次检验日期", allow_month=True)
            if next_date is None:
                raise ValueError(f"Excel 第 {row_number} 行缺少下次检验日期")
            try:
                volume = Decimal(str(row[5])) if row[5] is not None and str(row[5]).strip() else Decimal(0)
            except InvalidOperation as exc:
                raise ValueError(f"Excel 第 {row_number} 行容积不是有效数字") from exc
            if volume < 0 or volume.as_tuple().exponent < -2 or len(volume.as_tuple().digits) > 18:
                raise ValueError(f"Excel 第 {row_number} 行容积必须为非负的两位小数")
            lines.append({
                "sequence": sequence,
                "device_category": device_category,
                "body_code": body_code,
                "manufacture_date": _date(row[3], row_number, "生产日期", allow_month=True),
                "manufacturer": manufacturer,
                "volume": volume,
                "next_inspection_date": next_date,
                "inspection_result": "GOOD" if result == "合格" else result,
                "inspector_name": inspector_name,
            })
        if not lines:
            raise ValueError("Excel 中没有可导入的钢瓶明细")
        inspector_names = {line["inspector_name"] for line in lines if line["inspector_name"]}
        return customer_name, next(iter(inspector_names)) if len(inspector_names) == 1 else "", lines
    finally:
        close_workbook()
