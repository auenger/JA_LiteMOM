from __future__ import annotations

import base64
import csv
import io
import re
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal
from tempfile import SpooledTemporaryFile
from urllib.parse import quote
from uuid import uuid4
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, Response, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl.utils.exceptions import InvalidFileException
from xlrd import XLRDError
from sqlalchemy import and_, case, func, or_, select, text, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, lazyload, selectinload

from .database import get_db
from .bottle_report_word import TEMPLATE_PATH as BOTTLE_REPORT_TEMPLATE_PATH, render_bottle_report_word
from .inspection_word import TEMPLATE_LIBRARY, render_inspection_word
from .inspection_receipt_word import render_inspection_receipt_word
from .inspection_excel import parse_full_inspection_excel
from .inspection_xls import default_next_inspection_date, inspection_age_table_rows, render_inspection_xls, split_inspection_lines_by_age
from .unverified_xls import render_unverified_bottles_xls
from .models import (
    Bottle,
    BottleInbound,
    BottleInboundLine,
    BottleInventoryAdjustment,
    BottleMonthlyInventoryLog,
    BottleMonthlyInventoryOverride,
    BottleDailyInventoryOverride,
    BottleDailyInventoryLog,
    InspectionDocument,
    InspectionDocumentExportRecord,
    InspectionDocumentSplitHistory,
    InspectionWordDailyCounter,
    InspectionWordExportAssignment,
    InspectionWordReleasedPair,
    InspectionDocumentLine,
    UnverifiedBottle,
    BottleOutbound,
    BottleOutboundAllocation,
    BottleOutboundCorrection,
    BottleOutboundLine,
    BottleSource,
    BottleSpec,
    InspectionWordTemplateChoice,
    Material,
    MaterialDocument,
    MaterialDocumentLine,
    MaterialMovement,
    MaterialDailyEntry,
    PurchaseOrder,
    PurchaseOrderLine,
    QuickConfig,
    Supplier,
    SupplierMaterialPrice,
    User,
    BottleManufacturer,
)
from .permissions import ALL_PERMISSION_KEYS, PERMISSIONS, MATERIAL_DOCUMENT_PERMISSIONS
from .schemas import (
    BottleCodingCreate,
    BottleCodingUpdate,
    BottleInboundCreate,
    BottleInboundOut,
    BottleInventoryAdjustmentCreate,
    BottleInventoryAdjustmentOut,
    BottleMonthlyInventoryUpdate,
    BottleDailyInventoryUpdate,
    BottleOut,
    CodingReviewPageOut,
    BottleOutboundCreate,
    BottleOutboundCorrectionCreate,
    BottleOutboundOut,
    BottleStatusUpdate,
    ChangePasswordRequest,
    GoodBottleExportOut,
    InspectionCreate,
    InspectionDocumentCreate,
    InspectionDocumentGenerate,
    InspectionDocumentSplitPreview,
    InspectionDocumentSplitCreate,
    InspectionDocumentSplitHistoryOut,
    InspectionDocumentLineOut,
    InspectionDocumentOut,
    InspectionDocumentSummaryOut,
    InspectionDocumentLineUpdate,
    InspectionDocumentUpdate,
    InspectionWordCodeUpdate,
    InspectionWordTemplateChoiceUpdate,
    UnverifiedBottleImport,
    UnverifiedBottleImportResult,
    UnverifiedBottleBatchCreate,
    UnverifiedBottleOut,
    LoginRequest,
    MaterialCreate,
    MaterialDocumentCreate,
    MaterialPaymentUpdate,
    MaterialDocumentOut,
    MaterialOut,
    ManufacturerCreate,
    ManufacturerOptionOut,
    ManufacturerOut,
    PageOut,
    PurchaseOrderCreate,
    PurchaseOrderOut,
    PurchaseOrderPaymentUpdate,
    QuickEntryCreate,
    QuickConfigCreate,
    QuickConfigOut,
    SourceCreate,
    SourceOut,
    SpecCreate,
    SpecOptionOut,
    SpecOut,
    SupplierCreate,
    SupplierOut,
    UserCreate,
    UserOut,
    UserUpdate,
)
from .security import create_token, current_user, hash_password, require, require_any, verify_password


router = APIRouter(prefix="/api")


def china_today() -> date:
    return datetime.now(ZoneInfo("Asia/Shanghai")).date()


def not_found(label: str = "记录") -> HTTPException:
    return HTTPException(status_code=404, detail=f"{label}不存在")


def save(db: Session, conflict_message: str = "数据重复或仍被业务引用"):
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail=conflict_message)


def make_no(prefix: str, business_date: date) -> str:
    return f"{prefix}-{business_date.strftime('%Y%m%d')}-{uuid4().hex[:6].upper()}"


def page_result(db: Session, stmt, page: int, page_size: int):
    total = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery())) or 0
    items = db.scalars(stmt.offset((page - 1) * page_size).limit(page_size)).all()
    return {"items": items, "total": total, "page": page, "page_size": page_size}


@router.get("/health")
def health():
    return {"status": "ok"}


@router.post("/auth/login")
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.username == payload.username))
    if not user or not user.is_active or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="账号或密码错误")
    return {"access_token": create_token(user.id), "token_type": "bearer", "user": UserOut.model_validate(user)}


@router.get("/auth/me", response_model=UserOut)
def me(user: User = Depends(current_user)):
    return user


@router.post("/auth/change-password")
def change_password(payload: ChangePasswordRequest, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not verify_password(payload.old_password, user.password_hash):
        raise HTTPException(status_code=400, detail="原密码不正确")
    user.password_hash = hash_password(payload.new_password)
    user.must_change_password = False
    db.commit()
    return {"message": "密码已更新"}


@router.get("/permissions/catalog")
def permission_catalog(_: User = Depends(current_user)):
    return PERMISSIONS


@router.get("/options")
def business_options(user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Active reference data used inside authorized business forms."""
    return {
        "materials": [MaterialOut.model_validate(item) for item in db.scalars(select(Material).where(Material.is_active.is_(True)).order_by(Material.name)).all()],
        "suppliers": [SupplierOut.model_validate(item) for item in db.scalars(select(Supplier).where(Supplier.is_active.is_(True)).order_by(Supplier.name)).all()],
        "sources": [SourceOut.model_validate(item) for item in db.scalars(select(BottleSource).where(BottleSource.is_active.is_(True)).order_by(BottleSource.is_default.desc(), BottleSource.sort_order)).all()],
        "specs": [SpecOptionOut.model_validate(item) for item in db.scalars(select(BottleSpec).where(BottleSpec.is_active.is_(True)).order_by(BottleSpec.is_default.desc(), BottleSpec.sort_order)).all()],
        "manufacturers": [ManufacturerOptionOut.model_validate(item) for item in db.scalars(select(BottleManufacturer).where(BottleManufacturer.is_active.is_(True)).order_by(BottleManufacturer.sort_order, BottleManufacturer.id)).all()],
        "code_prefixes": [QuickConfigOut.model_validate(item) for item in db.scalars(select(QuickConfig).where(QuickConfig.category == "CODE_PREFIX", QuickConfig.is_active.is_(True)).order_by(QuickConfig.sort_order, QuickConfig.id)).all()],
        "inspection_notes": [QuickConfigOut.model_validate(item) for item in db.scalars(select(QuickConfig).where(QuickConfig.category == "INSPECTION_NOTE", QuickConfig.is_active.is_(True)).order_by(QuickConfig.sort_order, QuickConfig.id)).all()],
        "date_presets": [QuickConfigOut.model_validate(item) for item in db.scalars(select(QuickConfig).where(QuickConfig.category == "MANUFACTURE_DATE", QuickConfig.is_active.is_(True)).order_by(QuickConfig.sort_order, QuickConfig.id)).all()],
        "users": [{"id": item.id, "name": item.name, "username": item.username, "is_active": item.is_active, "is_superadmin": item.is_superadmin} for item in db.scalars(select(User).where(User.is_active.is_(True), User.is_superadmin.is_(False)).order_by(User.name, User.id)).all()],
    }


@router.get("/users", response_model=PageOut[UserOut])
def list_users(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("users:view")), db: Session = Depends(get_db)):
    return page_result(db, select(User).where(User.is_superadmin.is_(False)).order_by(User.id), page, page_size)


@router.post("/users", response_model=UserOut, status_code=status.HTTP_201_CREATED)
def create_user(payload: UserCreate, _: User = Depends(require("users:create")), db: Session = Depends(get_db)):
    unknown = set(payload.permissions) - ALL_PERMISSION_KEYS
    if unknown:
        raise HTTPException(status_code=400, detail=f"包含未知权限：{', '.join(sorted(unknown))}")
    role_name = payload.role_name.strip()
    if not role_name:
        raise HTTPException(status_code=400, detail="角色/身份名称不能为空")
    user = User(
        username=payload.username.strip(),
        name=payload.name.strip(),
        role_name=role_name,
        phone=payload.phone,
        remark=payload.remark,
        password_hash=hash_password(payload.password),
        is_active=payload.is_active,
        permissions=payload.permissions,
        must_change_password=True,
    )
    db.add(user)
    save(db, "登录账号已存在")
    db.refresh(user)
    return user


@router.put("/users/{user_id}", response_model=UserOut)
def update_user(user_id: int, payload: UserUpdate, actor: User = Depends(require("users:update")), db: Session = Depends(get_db)):
    user = db.get(User, user_id)
    if not user:
        raise not_found("人员")
    if user.is_superadmin and not actor.is_superadmin:
        raise HTTPException(403, "仅超级管理员可以编辑超级管理员")
    data = payload.model_dump(exclude_unset=True)
    if "username" in data:
        if not actor.is_superadmin:
            raise HTTPException(403, "仅超级管理员可以修改登录账号")
        if user.is_superadmin and data["username"] != user.username:
            raise HTTPException(400, "预设超级管理员账号不支持修改")
        if db.scalar(select(User.id).where(User.username == data["username"], User.id != user.id)):
            raise HTTPException(400, "登录账号已存在")
    permissions = data.pop("permissions", None)
    password = data.pop("password", None)
    if "role_name" in data:
        data["role_name"] = data["role_name"].strip()
        if not data["role_name"]:
            raise HTTPException(status_code=400, detail="角色/身份名称不能为空")
    if permissions is not None:
        unknown = set(permissions) - ALL_PERMISSION_KEYS
        if unknown:
            raise HTTPException(status_code=400, detail=f"包含未知权限：{', '.join(sorted(unknown))}")
        if not user.is_superadmin:
            user.permissions = permissions
    if password:
        user.password_hash = hash_password(password)
        user.must_change_password = True
    if user.is_superadmin:
        data.pop("is_active", None)
    for key, value in data.items():
        setattr(user, key, value)
    save(db, "登录账号已存在")
    db.refresh(user)
    return user


@router.delete("/users/{user_id}", status_code=204)
def delete_user(user_id: int, actor: User = Depends(require("users:delete")), db: Session = Depends(get_db)):
    user = db.get(User, user_id)
    if not user:
        raise not_found("人员")
    if user.is_superadmin:
        raise HTTPException(status_code=400, detail="默认超级管理员不能删除")
    if actor.id == user.id:
        raise HTTPException(status_code=400, detail="不能删除当前登录人员")
    db.delete(user)
    save(db, "该人员已有业务记录，请改为停用")


def apply_master_values(obj, payload, exclude: set[str] | None = None):
    for key, value in payload.model_dump(exclude=exclude or set()).items():
        setattr(obj, key, value)


@router.get("/materials", response_model=PageOut[MaterialOut])
def list_materials(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("materials:view")), db: Session = Depends(get_db)):
    return page_result(db, select(Material).order_by(Material.id.desc()), page, page_size)


@router.get("/materials/export")
def export_materials(_: User = Depends(require("materials:export")), db: Session = Depends(get_db)):
    output = io.StringIO()
    fields = ["名称", "类型", "规格", "物料单位", "启用单位换算", "库存单位", "换算比例", "参考价格", "启用预警", "安全库存阈值", "档案状态", "备注"]
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    for item in db.scalars(select(Material).order_by(Material.id)).all():
        writer.writerow({
            "名称": item.name,
            "类型": "废料" if item.material_type == "SCRAP" else "原料",
            "规格": item.specification,
            "物料单位": item.material_unit,
            "启用单位换算": "是" if item.conversion_enabled else "否",
            "库存单位": item.stock_unit if item.conversion_enabled else "",
            "换算比例": item.conversion_ratio if item.conversion_enabled else "",
            "参考价格": item.price,
            "启用预警": "是" if item.warning_enabled else "否",
            "安全库存阈值": item.safety_threshold if item.safety_threshold is not None else "",
            "档案状态": "启用" if item.is_active else "停用",
            "备注": item.remark or "",
        })
    return Response(
        content="\ufeff" + output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="materials.csv"'},
    )


def csv_bool(value: str, truthy: set[str]) -> bool:
    return value.strip().lower() in truthy


@router.post("/materials/import")
async def import_materials(
    file: UploadFile = File(...),
    _: User = Depends(require("materials:import")),
    db: Session = Depends(get_db),
):
    if not (file.filename or "").lower().endswith(".csv"):
        raise HTTPException(status_code=400, detail="仅支持 CSV 文件")
    try:
        content = (await file.read()).decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="CSV 文件必须使用 UTF-8 编码")
    reader = csv.DictReader(io.StringIO(content))
    required = {"名称", "物料单位"}
    if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
        raise HTTPException(status_code=400, detail="CSV 缺少名称或物料单位列")
    created = 0
    updated = 0
    try:
        for line_number, row in enumerate(reader, start=2):
            if not (row.get("名称") or "").strip():
                continue
            warning_enabled = csv_bool(row.get("启用预警") or "", {"是", "true", "1", "yes"})
            threshold = (row.get("安全库存阈值") or "").strip()
            stock_unit = (row.get("库存单位") or "").strip()
            ratio_text = (row.get("换算比例") or "").strip()
            conversion_flag = (row.get("启用单位换算") or "").strip()
            conversion_enabled = csv_bool(conversion_flag, {"是", "true", "1", "yes"}) if conversion_flag else bool(stock_unit and ratio_text)
            payload = MaterialCreate(
                name=(row.get("名称") or "").strip(),
                material_type="SCRAP" if (row.get("类型") or "").strip() == "废料" else "MATERIAL",
                specification=(row.get("规格") or "").strip(),
                material_unit=(row.get("物料单位") or "").strip(),
                conversion_enabled=conversion_enabled,
                stock_unit=stock_unit,
                conversion_ratio=Decimal(ratio_text or "1"),
                price=Decimal((row.get("参考价格") or "0").strip()),
                warning_enabled=warning_enabled,
                safety_threshold=Decimal(threshold) if threshold else None,
                is_active=not csv_bool(row.get("档案状态") or "", {"停用", "false", "0", "no"}),
                remark=(row.get("备注") or "").strip() or None,
            )
            obj = db.scalar(select(Material).where(Material.name == payload.name, Material.specification == payload.specification))
            if obj:
                apply_master_values(obj, payload)
                obj.unit = payload.material_unit
                updated += 1
            else:
                db.add(Material(**payload.model_dump(), unit=payload.material_unit))
                created += 1
        db.commit()
    except (ValueError, ArithmeticError) as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"CSV 第 {line_number} 行数据格式错误：{exc}")
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="CSV 中存在重复的原材料名称和规格")
    return {"created": created, "updated": updated, "total": created + updated}


@router.post("/materials", response_model=MaterialOut, status_code=201)
def create_material(payload: MaterialCreate, _: User = Depends(require("materials:create")), db: Session = Depends(get_db)):
    obj = Material(**payload.model_dump(), unit=payload.material_unit)
    db.add(obj)
    save(db, "相同名称和规格的原材料已存在")
    db.refresh(obj)
    return obj


@router.put("/materials/{object_id}", response_model=MaterialOut)
def update_material(object_id: int, payload: MaterialCreate, _: User = Depends(require("materials:update")), db: Session = Depends(get_db)):
    obj = db.get(Material, object_id)
    if not obj:
        raise not_found("原材料")
    apply_master_values(obj, payload)
    obj.unit = payload.material_unit
    save(db, "相同名称和规格的原材料已存在")
    db.refresh(obj)
    return obj


@router.delete("/materials/{object_id}", status_code=204)
def delete_material(object_id: int, _: User = Depends(require("materials:delete")), db: Session = Depends(get_db)):
    obj = db.get(Material, object_id)
    if not obj:
        raise not_found("原材料")
    if db.scalar(select(MaterialDailyEntry.id).where(MaterialDailyEntry.material_id == object_id).limit(1)):
        raise HTTPException(400, "原材料已有出入库记录，请改为停用")
    db.delete(obj)
    save(db, "原材料已有单据记录，请改为停用")


@router.get("/suppliers", response_model=PageOut[SupplierOut])
def list_suppliers(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("suppliers:view")), db: Session = Depends(get_db)):
    return page_result(db, select(Supplier).order_by(Supplier.id.desc()), page, page_size)


@router.post("/suppliers", response_model=SupplierOut, status_code=201)
def create_supplier(payload: SupplierCreate, _: User = Depends(require("suppliers:create")), db: Session = Depends(get_db)):
    obj = Supplier(**payload.model_dump(exclude={"material_prices"}))
    replace_supplier_material_prices(db, obj, payload)
    db.add(obj)
    save(db, "供应商名称已存在")
    db.refresh(obj)
    return obj


@router.put("/suppliers/{object_id}", response_model=SupplierOut)
def update_supplier(object_id: int, payload: SupplierCreate, _: User = Depends(require("suppliers:update")), db: Session = Depends(get_db)):
    obj = db.get(Supplier, object_id)
    if not obj:
        raise not_found("供应商")
    apply_master_values(obj, payload, exclude={"material_prices"})
    replace_supplier_material_prices(db, obj, payload)
    save(db, "供应商名称已存在")
    db.refresh(obj)
    return obj


def replace_supplier_material_prices(db: Session, supplier: Supplier, payload: SupplierCreate):
    supplier.material_prices.clear()
    if supplier.id is not None:
        db.flush()
    for item in payload.material_prices:
        material = db.get(Material, item.material_id)
        if not material or not material.is_active:
            raise HTTPException(status_code=400, detail=f"原材料 {item.material_id} 不存在或已停用")
        supplier.material_prices.append(SupplierMaterialPrice(
            material_id=material.id,
            purchase_price=item.purchase_price,
        ))


@router.delete("/suppliers/{object_id}", status_code=204)
def delete_supplier(object_id: int, _: User = Depends(require("suppliers:delete")), db: Session = Depends(get_db)):
    obj = db.get(Supplier, object_id)
    if not obj:
        raise not_found("供应商")
    db.delete(obj)
    save(db, "供应商已有入库记录，请改为停用")


def clear_default(db: Session, model, object_id: int | None = None):
    items = db.scalars(select(model).where(model.is_default.is_(True))).all()
    for item in items:
        if item.id != object_id:
            item.is_default = False


@router.get("/bottle-sources", response_model=PageOut[SourceOut])
def list_sources(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("sources:view")), db: Session = Depends(get_db)):
    return page_result(db, select(BottleSource).order_by(BottleSource.is_default.desc(), BottleSource.sort_order, BottleSource.id), page, page_size)


@router.post("/bottle-sources", response_model=SourceOut, status_code=201)
def create_source(payload: SourceCreate, _: User = Depends(require("sources:create")), db: Session = Depends(get_db)):
    if payload.is_default:
        clear_default(db, BottleSource)
    obj = BottleSource(**payload.model_dump())
    db.add(obj)
    save(db, "供瓶单位名称已存在")
    db.refresh(obj)
    return obj


@router.put("/bottle-sources/{object_id}", response_model=SourceOut)
def update_source(object_id: int, payload: SourceCreate, _: User = Depends(require("sources:update")), db: Session = Depends(get_db)):
    obj = db.get(BottleSource, object_id)
    if not obj:
        raise not_found("供瓶单位")
    if payload.is_default:
        clear_default(db, BottleSource, object_id)
    apply_master_values(obj, payload)
    save(db, "供瓶单位名称已存在")
    db.refresh(obj)
    return obj


@router.delete("/bottle-sources/{object_id}", status_code=204)
def delete_source(object_id: int, _: User = Depends(require("sources:delete")), db: Session = Depends(get_db)):
    obj = db.get(BottleSource, object_id)
    if not obj:
        raise not_found("供瓶单位")
    db.delete(obj)
    save(db, "供瓶单位已有钢瓶业务记录，请改为停用")


@router.get("/bottle-specs", response_model=PageOut[SpecOut])
def list_specs(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("specs:view")), db: Session = Depends(get_db)):
    return page_result(db, select(BottleSpec).order_by(BottleSpec.is_default.desc(), BottleSpec.sort_order, BottleSpec.id), page, page_size)


@router.post("/bottle-specs", response_model=SpecOut, status_code=201)
def create_spec(payload: SpecCreate, _: User = Depends(require("specs:create")), db: Session = Depends(get_db)):
    if payload.is_default:
        clear_default(db, BottleSpec)
    obj = BottleSpec(**payload.model_dump())
    db.add(obj)
    save(db, "瓶子规格名称已存在")
    db.refresh(obj)
    return obj


@router.put("/bottle-specs/{object_id}", response_model=SpecOut)
def update_spec(object_id: int, payload: SpecCreate, _: User = Depends(require("specs:update")), db: Session = Depends(get_db)):
    obj = db.get(BottleSpec, object_id)
    if not obj:
        raise not_found("瓶子规格")
    if payload.is_default:
        clear_default(db, BottleSpec, object_id)
    apply_master_values(obj, payload)
    save(db, "瓶子规格名称已存在")
    db.refresh(obj)
    return obj


@router.delete("/bottle-specs/{object_id}", status_code=204)
def delete_spec(object_id: int, _: User = Depends(require("specs:delete")), db: Session = Depends(get_db)):
    obj = db.get(BottleSpec, object_id)
    if not obj:
        raise not_found("瓶子规格")
    db.delete(obj)
    save(db, "瓶子规格已有抄码记录，请改为停用")


@router.get("/bottle-manufacturers", response_model=PageOut[ManufacturerOut])
def list_manufacturers(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("manufacturers:view")), db: Session = Depends(get_db)):
    return page_result(db, select(BottleManufacturer).order_by(BottleManufacturer.sort_order, BottleManufacturer.id), page, page_size)


@router.post("/bottle-manufacturers", response_model=ManufacturerOut, status_code=201)
def create_manufacturer(payload: ManufacturerCreate, _: User = Depends(require("manufacturers:create")), db: Session = Depends(get_db)):
    obj = BottleManufacturer(**payload.model_dump())
    db.add(obj)
    save(db, "生产厂家名称已存在")
    db.refresh(obj)
    return obj


@router.put("/bottle-manufacturers/{object_id}", response_model=ManufacturerOut)
def update_manufacturer(object_id: int, payload: ManufacturerCreate, _: User = Depends(require("manufacturers:update")), db: Session = Depends(get_db)):
    obj = db.get(BottleManufacturer, object_id)
    if not obj:
        raise not_found("生产厂家")
    apply_master_values(obj, payload)
    save(db, "生产厂家名称已存在")
    db.refresh(obj)
    return obj


@router.delete("/bottle-manufacturers/{object_id}", status_code=204)
def delete_manufacturer(object_id: int, _: User = Depends(require("manufacturers:delete")), db: Session = Depends(get_db)):
    obj = db.get(BottleManufacturer, object_id)
    if not obj:
        raise not_found("生产厂家")
    db.delete(obj)
    save(db, "生产厂家名称已存在")


@router.get("/quick-configs", response_model=PageOut[QuickConfigOut])
def list_quick_configs(category: str | None = None, page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("quick_configs:view")), db: Session = Depends(get_db)):
    stmt = select(QuickConfig).order_by(QuickConfig.category, QuickConfig.sort_order, QuickConfig.id)
    if category:
        stmt = stmt.where(QuickConfig.category == category)
    return page_result(db, stmt, page, page_size)


@router.post("/quick-configs", response_model=QuickConfigOut, status_code=201)
def create_quick_config(payload: QuickConfigCreate, _: User = Depends(require("quick_configs:create")), db: Session = Depends(get_db)):
    obj = QuickConfig(**payload.model_dump())
    db.add(obj)
    save(db, "相同类型的快捷内容已存在")
    db.refresh(obj)
    return obj


@router.put("/quick-configs/{object_id}", response_model=QuickConfigOut)
def update_quick_config(object_id: int, payload: QuickConfigCreate, _: User = Depends(require("quick_configs:update")), db: Session = Depends(get_db)):
    obj = db.get(QuickConfig, object_id)
    if not obj:
        raise not_found("快捷配置")
    apply_master_values(obj, payload)
    save(db, "相同类型的快捷内容已存在")
    db.refresh(obj)
    return obj


@router.delete("/quick-configs/{object_id}", status_code=204)
def delete_quick_config(object_id: int, _: User = Depends(require("quick_configs:delete")), db: Session = Depends(get_db)):
    obj = db.get(QuickConfig, object_id)
    if not obj:
        raise not_found("快捷配置")
    db.delete(obj)
    db.commit()


def replace_purchase_order_lines(db: Session, order: PurchaseOrder, payload: PurchaseOrderCreate):
    supplier = db.get(Supplier, payload.supplier_id)
    if not supplier or not supplier.is_active:
        raise HTTPException(status_code=400, detail="供应商不存在或已停用")
    prices = {
        item.material_id: item
        for item in db.scalars(select(SupplierMaterialPrice).where(
            SupplierMaterialPrice.supplier_id == supplier.id,
        )).all()
    }
    order.lines.clear()
    if order.id is not None:
        db.flush()
    for item in payload.lines:
        price = prices.get(item.material_id)
        if not price or not price.material.is_active:
            raise HTTPException(status_code=400, detail=f"所选供应商未配置原材料 {item.material_id} 的采购价格")
        if price.material.material_type != "MATERIAL":
            raise HTTPException(status_code=400, detail=f"{price.material.name} 属于废料，不能加入供应商采购单")
        order.lines.append(PurchaseOrderLine(
            material_id=price.material_id,
            material_name=price.material.name,
            specification=price.material.specification,
            unit=price.material.material_unit,
            quantity=item.quantity,
            unit_price=price.purchase_price,
            amount=(item.quantity * price.purchase_price).quantize(Decimal("0.01")),
        ))


@router.get("/purchase-orders", response_model=PageOut[PurchaseOrderOut])
def list_purchase_orders(
    date_from: date | None = None,
    date_to: date | None = None,
    supplier_id: int | None = None,
    order_status: str | None = None,
    is_paid: bool | None = None,
    available_for_inbound: bool = False,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require_any("purchase_orders:view", "material_inbound:view")),
    db: Session = Depends(get_db),
):
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    if order_status and order_status not in {"DRAFT", "CONFIRMED", "VOID"}:
        raise HTTPException(status_code=400, detail="采购单状态无效")
    stmt = select(PurchaseOrder)
    if date_from:
        stmt = stmt.where(PurchaseOrder.order_date >= date_from)
    if date_to:
        stmt = stmt.where(PurchaseOrder.order_date <= date_to)
    if supplier_id:
        stmt = stmt.where(PurchaseOrder.supplier_id == supplier_id)
    if order_status:
        stmt = stmt.where(PurchaseOrder.status == order_status)
    if is_paid is not None:
        stmt = stmt.where(PurchaseOrder.is_paid.is_(is_paid))
    if available_for_inbound:
        active_inbound = select(MaterialDocument.id).where(
            MaterialDocument.purchase_order_id == PurchaseOrder.id,
            MaterialDocument.status != "VOID",
        ).exists()
        stmt = stmt.where(PurchaseOrder.status == "CONFIRMED", ~active_inbound)
    return page_result(db, stmt.order_by(PurchaseOrder.order_date.desc(), PurchaseOrder.id.desc()), page, page_size)


@router.post("/purchase-orders", response_model=PurchaseOrderOut, status_code=201)
def create_purchase_order(payload: PurchaseOrderCreate, user: User = Depends(require("purchase_orders:create")), db: Session = Depends(get_db)):
    order = PurchaseOrder(
        order_no=make_no("PO", payload.order_date),
        order_date=payload.order_date,
        supplier_id=payload.supplier_id,
        handler=payload.handler,
        note=payload.note,
        created_by=user.id,
    )
    replace_purchase_order_lines(db, order, payload)
    db.add(order)
    db.commit()
    db.refresh(order)
    return order


@router.put("/purchase-orders/{order_id}", response_model=PurchaseOrderOut)
def update_purchase_order(order_id: int, payload: PurchaseOrderCreate, user: User = Depends(require("purchase_orders:update")), db: Session = Depends(get_db)):
    order = db.get(PurchaseOrder, order_id)
    if not order:
        raise not_found("采购单")
    if order.status != "DRAFT":
        raise HTTPException(status_code=400, detail="只有草稿采购单可以编辑")
    order.order_date = payload.order_date
    order.supplier_id = payload.supplier_id
    order.handler = payload.handler
    order.note = payload.note
    order.updated_by = user.id
    replace_purchase_order_lines(db, order, payload)
    db.commit()
    db.refresh(order)
    return order


@router.delete("/purchase-orders/{order_id}", status_code=204)
def delete_purchase_order(order_id: int, _: User = Depends(require("purchase_orders:delete")), db: Session = Depends(get_db)):
    order = db.get(PurchaseOrder, order_id)
    if not order:
        raise not_found("采购单")
    if order.status != "DRAFT":
        raise HTTPException(status_code=400, detail="只有草稿采购单可以删除")
    db.delete(order)
    db.commit()


@router.post("/purchase-orders/{order_id}/confirm", response_model=PurchaseOrderOut)
def confirm_purchase_order(order_id: int, user: User = Depends(require("purchase_orders:post")), db: Session = Depends(get_db)):
    order = db.scalar(select(PurchaseOrder).where(PurchaseOrder.id == order_id).with_for_update())
    if not order:
        raise not_found("采购单")
    if order.status != "DRAFT":
        raise HTTPException(status_code=400, detail="采购单已处理，不能重复确认")
    order.status = "CONFIRMED"
    order.confirmed_by = user.id
    order.confirmed_at = datetime.utcnow()
    db.commit()
    db.refresh(order)
    return order


@router.post("/purchase-orders/{order_id}/void", response_model=PurchaseOrderOut)
def void_purchase_order(order_id: int, user: User = Depends(require("purchase_orders:void")), db: Session = Depends(get_db)):
    order = db.scalar(select(PurchaseOrder).where(PurchaseOrder.id == order_id).with_for_update())
    if not order:
        raise not_found("采购单")
    if order.status != "CONFIRMED":
        raise HTTPException(status_code=400, detail="只有已确认采购单可以作废")
    if any(item.status != "VOID" for item in order.inbound_documents):
        raise HTTPException(status_code=400, detail="采购单已有未作废入库单，请先处理对应入库单")
    order.status = "VOID"
    order.voided_by = user.id
    order.voided_at = datetime.utcnow()
    db.commit()
    db.refresh(order)
    return order


@router.put("/purchase-orders/{order_id}/payment", response_model=PurchaseOrderOut)
def update_purchase_order_payment(order_id: int, payload: PurchaseOrderPaymentUpdate, _: User = Depends(require("purchase_orders:payment")), db: Session = Depends(get_db)):
    order = db.get(PurchaseOrder, order_id)
    if not order:
        raise not_found("采购单")
    if order.status == "VOID":
        raise HTTPException(status_code=400, detail="已作废采购单不能修改付款状态")
    order.is_paid = payload.is_paid
    db.commit()
    db.refresh(order)
    return order


def replace_document_lines(db: Session, document: MaterialDocument, payload: MaterialDocumentCreate):
    document.lines.clear()
    seen: set[int] = set()
    for item in payload.lines:
        if item.material_id in seen:
            raise HTTPException(status_code=400, detail="同一张单据不能重复选择原材料")
        seen.add(item.material_id)
        material = db.get(Material, item.material_id)
        if not material or not material.is_active:
            raise HTTPException(status_code=400, detail=f"原材料 {item.material_id} 不存在或已停用")
        if document.document_type == "NORMAL_IN" and document.inbound_mode == "DIRECT":
            expected_type = "SCRAP" if document.inbound_category == "SCRAP" else "MATERIAL"
            if material.material_type != expected_type:
                label = "废料" if expected_type == "SCRAP" else "原料"
                raise HTTPException(status_code=400, detail=f"{label}入库只能选择{label}类型的档案")
        price = item.unit_price if item.unit_price is not None else material.price
        document.lines.append(MaterialDocumentLine(
            material_id=material.id,
            material_name=material.name,
            specification=material.specification,
            unit=material.material_unit,
            quantity=item.quantity,
            unit_price=price,
            amount=(item.quantity * price).quantize(Decimal("0.01")),
        ))


def apply_purchase_order_to_inbound(db: Session, document: MaterialDocument, purchase_order_id: int):
    order = db.scalar(select(PurchaseOrder).where(PurchaseOrder.id == purchase_order_id).with_for_update())
    if not order or order.status != "CONFIRMED":
        raise HTTPException(status_code=400, detail="采购单不存在或尚未确认")
    existing = db.scalar(select(MaterialDocument).where(
        MaterialDocument.purchase_order_id == order.id,
        MaterialDocument.status != "VOID",
        MaterialDocument.id != document.id,
    ))
    if existing:
        raise HTTPException(status_code=400, detail=f"采购单已生成入库单 {existing.document_no}")
    document.purchase_order_id = order.id
    document.supplier_id = order.supplier_id
    document.lines.clear()
    if document.id is not None:
        db.flush()
    for line in order.lines:
        document.lines.append(MaterialDocumentLine(
            material_id=line.material_id,
            material_name=line.material_name,
            specification=line.specification,
            unit=line.unit,
            quantity=line.quantity,
            unit_price=line.unit_price,
            amount=line.amount,
        ))


def check_material_document_permission(user: User, document_type: str, action: str):
    prefix = MATERIAL_DOCUMENT_PERMISSIONS.get(document_type)
    if not prefix or (not user.is_superadmin and f"{prefix}:{action}" not in (user.permissions or [])):
        raise HTTPException(403, "无此单据的操作权限")


@router.get("/material-documents/return-options", response_model=PageOut[MaterialDocumentOut])
def material_return_options(page: int = Query(1, ge=1), page_size: int = Query(100, ge=1, le=100),
    _: User = Depends(require("material_return:view")), db: Session = Depends(get_db)):
    return page_result(db, select(MaterialDocument).where(
        MaterialDocument.document_type == "OUTBOUND", MaterialDocument.status == "POSTED",
        MaterialDocument.outbound_category == "MATERIAL").order_by(MaterialDocument.id.desc()), page, page_size)


@router.get("/material-documents", response_model=PageOut[MaterialDocumentOut])
def list_material_documents(
    document_type: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    material_id: int | None = Query(default=None, ge=1),
    inbound_mode: str | None = None,
    inbound_category: str | None = None,
    outbound_category: str | None = None,
    is_paid: bool | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    if is_paid is not None:
        raise HTTPException(status_code=400, detail="付款状态已迁移到供应商采购单")
    if inbound_mode is not None and inbound_mode not in {"DIRECT", "PURCHASE_ORDER"}:
        raise HTTPException(status_code=400, detail="入库方式无效")
    if inbound_category is not None and inbound_category not in {"MATERIAL", "SCRAP", "PURCHASE"}:
        raise HTTPException(status_code=400, detail="入库类型无效")
    if outbound_category is not None and outbound_category not in {"MATERIAL", "SCRAP"}:
        raise HTTPException(status_code=400, detail="出库类型无效")
    allowed = [kind for kind, prefix in MATERIAL_DOCUMENT_PERMISSIONS.items()
               if user.is_superadmin or f"{prefix}:view" in (user.permissions or [])]
    if not allowed or (document_type and document_type not in allowed):
        raise HTTPException(403, "无此单据的查询权限")
    stmt = select(MaterialDocument).where(MaterialDocument.document_type.in_(allowed)).order_by(MaterialDocument.business_date.desc(), MaterialDocument.id.desc())
    if document_type:
        stmt = stmt.where(MaterialDocument.document_type == document_type)
    if date_from:
        stmt = stmt.where(MaterialDocument.business_date >= date_from)
    if date_to:
        stmt = stmt.where(MaterialDocument.business_date <= date_to)
    if material_id is not None:
        stmt = stmt.where(MaterialDocument.lines.any(MaterialDocumentLine.material_id == material_id))
    if inbound_mode is not None:
        stmt = stmt.where(MaterialDocument.inbound_mode == inbound_mode)
    if inbound_category is not None:
        if inbound_category == "MATERIAL":
            stmt = stmt.where(MaterialDocument.inbound_category.in_({"MATERIAL", "PURCHASE"}))
        else:
            stmt = stmt.where(MaterialDocument.inbound_category == inbound_category)
    if outbound_category is not None:
        stmt = stmt.where(MaterialDocument.outbound_category == outbound_category)
    return page_result(db, stmt, page, page_size)


@router.post("/material-documents", response_model=MaterialDocumentOut, status_code=201)
def create_material_document(payload: MaterialDocumentCreate, user: User = Depends(current_user), db: Session = Depends(get_db)):
    check_material_document_permission(user, payload.document_type, "create")
    prefix = {"NORMAL_IN": "MRI", "OUTBOUND": "MRO", "RETURN_IN": "MRR"}[payload.document_type]
    inbound_category = "MATERIAL" if payload.inbound_category == "PURCHASE" else payload.inbound_category
    document = MaterialDocument(
        document_no=make_no(prefix, payload.business_date),
        document_type=payload.document_type,
        business_date=payload.business_date,
        supplier_id=payload.supplier_id,
        purchase_order_id=payload.purchase_order_id,
        related_outbound_id=payload.related_outbound_id,
        inbound_mode=payload.inbound_mode,
        inbound_category=inbound_category,
        outbound_category=payload.outbound_category,
        is_paid=False,
        handler=payload.handler,
        note=payload.note,
        created_by=user.id,
    )
    if document.document_type == "NORMAL_IN" and payload.inbound_mode == "PURCHASE_ORDER":
        apply_purchase_order_to_inbound(db, document, payload.purchase_order_id)
    else:
        replace_document_lines(db, document, payload)
    if document.document_type == "RETURN_IN":
        validate_return(db, document)
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


@router.put("/material-documents/{document_id}", response_model=MaterialDocumentOut)
def update_material_document(document_id: int, payload: MaterialDocumentCreate, user: User = Depends(current_user), db: Session = Depends(get_db)):
    document = db.get(MaterialDocument, document_id)
    if not document:
        raise not_found("原材料单据")
    check_material_document_permission(user, document.document_type, "update")
    if document.status != "DRAFT":
        raise HTTPException(status_code=400, detail="只有草稿单据可以编辑")
    if payload.document_type != document.document_type:
        raise HTTPException(status_code=400, detail="不能修改单据类型")
    document.business_date = payload.business_date
    inbound_category = "MATERIAL" if payload.inbound_category == "PURCHASE" else payload.inbound_category
    document.supplier_id = payload.supplier_id
    document.purchase_order_id = payload.purchase_order_id
    document.related_outbound_id = payload.related_outbound_id
    document.inbound_mode = payload.inbound_mode
    document.inbound_category = inbound_category
    document.outbound_category = payload.outbound_category
    document.is_paid = False
    document.handler = payload.handler
    document.note = payload.note
    document.updated_by = user.id
    if document.document_type == "NORMAL_IN" and payload.inbound_mode == "PURCHASE_ORDER":
        apply_purchase_order_to_inbound(db, document, payload.purchase_order_id)
    else:
        replace_document_lines(db, document, payload)
    if document.document_type == "RETURN_IN":
        validate_return(db, document)
    db.commit()
    db.refresh(document)
    return document


@router.put("/material-documents/{document_id}/payment", response_model=MaterialDocumentOut)
def update_material_payment(document_id: int, payload: MaterialPaymentUpdate, user: User = Depends(current_user), db: Session = Depends(get_db)):
    document = db.get(MaterialDocument, document_id)
    if not document:
        raise not_found("原材料单据")
    check_material_document_permission(user, document.document_type, "update")
    raise HTTPException(status_code=400, detail="付款状态已迁移到供应商采购单，请在采购单页面维护")


@router.delete("/material-documents/{document_id}", status_code=204)
def delete_material_document(document_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    document = db.get(MaterialDocument, document_id)
    if not document:
        raise not_found("原材料单据")
    check_material_document_permission(user, document.document_type, "delete")
    if document.status != "DRAFT":
        raise HTTPException(status_code=400, detail="只有草稿单据可以删除")
    db.delete(document)
    db.commit()


def validate_return(db: Session, document: MaterialDocument):
    # Serialize validation for the same source outbound, so two concurrent
    # postings cannot both consume the same remaining return quantity.
    outbound = db.scalar(
        select(MaterialDocument)
        .where(MaterialDocument.id == document.related_outbound_id)
        .with_for_update()
    )
    if not outbound or outbound.document_type != "OUTBOUND" or outbound.outbound_category != "MATERIAL" or outbound.status != "POSTED":
        raise HTTPException(status_code=400, detail="绑定的原出库单不存在或尚未生效")
    outbound_qty = {line.material_id: line.quantity for line in outbound.lines}
    previous = db.scalars(select(MaterialDocument).where(
        MaterialDocument.related_outbound_id == outbound.id,
        MaterialDocument.document_type == "RETURN_IN",
        MaterialDocument.status == "POSTED",
        MaterialDocument.id != document.id,
    )).all()
    returned: dict[int, Decimal] = {}
    for item in previous:
        for line in item.lines:
            returned[line.material_id] = returned.get(line.material_id, Decimal(0)) + line.quantity
    for line in document.lines:
        if line.material_id not in outbound_qty:
            raise HTTPException(status_code=400, detail=f"{line.material_name} 不在原出库单中")
        remaining = outbound_qty[line.material_id] - returned.get(line.material_id, Decimal(0))
        if line.quantity > remaining:
            raise HTTPException(
                status_code=400,
                detail=f"{line.material_name} 的退料数量不能超过剩余可退数量 {remaining.quantize(Decimal('0.01'))}",
            )


@router.post("/material-documents/{document_id}/post", response_model=MaterialDocumentOut)
def post_material_document(document_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    document = db.scalar(select(MaterialDocument).where(MaterialDocument.id == document_id).with_for_update())
    if not document:
        raise not_found("原材料单据")
    check_material_document_permission(user, document.document_type, "post")
    if document.status != "DRAFT":
        raise HTTPException(status_code=400, detail="单据已处理，不能重复确认")
    if document.document_type == "RETURN_IN":
        validate_return(db, document)
    direction = Decimal(-1) if document.document_type == "OUTBOUND" else Decimal(1)
    for line in document.lines:
        material = db.scalar(select(Material).where(Material.id == line.material_id).with_for_update())
        new_stock = material.stock + direction * line.quantity
        if new_stock < 0:
            raise HTTPException(status_code=400, detail=f"{material.name} 库存不足")
        material.stock = new_stock
        db.add(MaterialMovement(
            material_id=material.id,
            document_id=document.id,
            movement_type="POST",
            quantity_change=direction * line.quantity,
            balance_after=new_stock,
            created_by=user.id,
        ))
    document.status = "POSTED"
    document.posted_by = user.id
    document.posted_at = datetime.utcnow()
    db.commit()
    db.refresh(document)
    return document


@router.post("/material-documents/{document_id}/void", response_model=MaterialDocumentOut)
def void_material_document(document_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    document = db.scalar(select(MaterialDocument).where(MaterialDocument.id == document_id).with_for_update())
    if not document:
        raise not_found("原材料单据")
    check_material_document_permission(user, document.document_type, "void")
    if document.status != "POSTED":
        raise HTTPException(status_code=400, detail="只有已生效单据可以作废或撤销")
    if document.document_type == "OUTBOUND":
        has_returns = db.scalar(select(func.count()).select_from(MaterialDocument).where(
            MaterialDocument.related_outbound_id == document.id,
            MaterialDocument.status == "POSTED",
        ))
        if has_returns:
            raise HTTPException(status_code=400, detail="该出库单已有生效退料单，请先撤销退料单")
    original_direction = Decimal(-1) if document.document_type == "OUTBOUND" else Decimal(1)
    for line in document.lines:
        material = db.scalar(select(Material).where(Material.id == line.material_id).with_for_update())
        new_stock = material.stock - original_direction * line.quantity
        if new_stock < 0:
            raise HTTPException(status_code=400, detail=f"撤销后 {material.name} 库存将小于零")
        material.stock = new_stock
        db.add(MaterialMovement(
            material_id=material.id,
            document_id=document.id,
            movement_type="VOID",
            quantity_change=-original_direction * line.quantity,
            balance_after=new_stock,
            created_by=user.id,
        ))
    document.status = "VOID"
    document.voided_by = user.id
    document.voided_at = datetime.utcnow()
    db.commit()
    db.refresh(document)
    return document


def parse_inventory_month(value: str, label: str) -> date:
    try:
        return datetime.strptime(value, "%Y-%m").date().replace(day=1)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=f"{label}格式应为 YYYY-MM") from exc


def next_inventory_month(value: date) -> date:
    return date(value.year + (1 if value.month == 12 else 0), 1 if value.month == 12 else value.month + 1, 1)


def inventory_month_range(month_from: date, month_to: date) -> list[date]:
    months: list[date] = []
    current = month_from
    while current <= month_to:
        months.append(current)
        if len(months) > 60:
            raise HTTPException(status_code=400, detail="月度区间最多选择 60 个月")
        current = next_inventory_month(current)
    return months


def pending_inventory_state_changes(db: Session, date_to: date, source_id: int | None = None):
    """Remove resolved bottles from pending; named GOOD outbounds already deduct raw pending.

    Use coding business dates (coding enters GOOD) rather than last status-update timestamps.
    BAD/SCRAP and outbound bottles stay resolved, so later transitions cannot restore pending.
    """
    stmt = (select(Bottle.source_id, Bottle.spec_id, Bottle.zone, Bottle.coding_date,
        Bottle.coded_at, Bottle.inspected_at, Bottle.outbound_origin_zone,
        BottleInbound.business_date, BottleInbound.status,
        BottleOutbound.business_date, BottleOutbound.status, BottleOutbound.source_id)
        .outerjoin(BottleInbound, BottleInbound.id == Bottle.inbound_id)
        .outerjoin(BottleOutbound, BottleOutbound.id == Bottle.outbound_id)
        .where(Bottle.zone.in_(["GOOD", "BAD", "SCRAP", "OUTBOUND"])))
    changes = {}
    def add(source, spec, day, delta):
        if day <= date_to and (source_id is None or source == source_id):
            key = (source, spec, day)
            changes[key] = changes.get(key, 0) + delta
    for source, spec, zone, coding_day, coded_at, inspected_at, origin, inbound_day, inbound_status, outbound_day, outbound_status, outbound_source in db.execute(stmt):
        if inbound_status is not None and inbound_status != "POSTED":
            continue
        if zone == "OUTBOUND" and origin not in {"GOOD", "BAD"}:
            continue
        entered = coding_day or ((inspected_at or coded_at).replace(tzinfo=ZoneInfo("UTC")).astimezone(ZoneInfo("Asia/Shanghai")).date())
        if inbound_day:
            entered = max(entered, inbound_day)
        if outbound_status == "POSTED" and outbound_day:
            entered = min(entered, outbound_day)
        add(source, spec, entered, -1)
        if zone == "OUTBOUND" and origin == "GOOD" and outbound_status == "POSTED" and outbound_source is not None:
            add(outbound_source, spec, outbound_day, 1)
    return changes


def build_monthly_inventory_report(
    db: Session,
    month_from: date,
    month_to: date,
    source_id: int | None = None,
) -> dict:
    months = inventory_month_range(month_from, month_to)
    date_to = next_inventory_month(month_to) - timedelta(days=1)

    source_stmt = select(BottleSource).order_by(BottleSource.is_default.desc(), BottleSource.sort_order, BottleSource.id)
    sources = list(db.scalars(source_stmt).all())
    if source_id is not None and not any(item.id == source_id for item in sources):
        raise not_found("供瓶单位")
    selected_source_ids = [source_id] if source_id is not None else [item.id for item in sources]

    specs = list(db.scalars(select(BottleSpec).order_by(BottleSpec.sort_order, BottleSpec.id)).all())
    spec_ids = [item.id for item in specs]

    inbound_stmt = (
        select(
            BottleInbound.source_id,
            BottleInboundLine.spec_id,
            BottleInbound.business_date,
            func.coalesce(func.sum(BottleInboundLine.pending_quantity), 0),
            func.coalesce(func.sum(BottleInboundLine.temporary_quantity), 0),
            func.coalesce(func.sum(BottleInboundLine.scrap_quantity), 0),
        )
        .join(BottleInboundLine, BottleInboundLine.inbound_id == BottleInbound.id)
        .where(BottleInbound.status == "POSTED", BottleInbound.business_date <= date_to)
        .group_by(BottleInbound.source_id, BottleInboundLine.spec_id, BottleInbound.business_date)
    )
    outbound_stmt = (
        select(
            BottleOutbound.source_id,
            BottleOutboundLine.spec_id,
            BottleOutbound.business_date,
            func.coalesce(func.sum(BottleOutboundLine.good_quantity), 0),
            func.coalesce(func.sum(BottleOutboundLine.temporary_quantity), 0),
            func.coalesce(func.sum(BottleOutboundLine.scrap_quantity), 0),
            func.coalesce(func.sum(BottleOutboundLine.bad_quantity), 0),
        )
        .join(BottleOutboundLine, BottleOutboundLine.outbound_id == BottleOutbound.id)
        .where(BottleOutbound.status == "POSTED", BottleOutbound.source_id.is_not(None), BottleOutbound.business_date <= date_to)
        .group_by(BottleOutbound.source_id, BottleOutboundLine.spec_id, BottleOutbound.business_date)
    )
    override_stmt = select(BottleMonthlyInventoryOverride).where(BottleMonthlyInventoryOverride.month <= month_to)
    if source_id is not None:
        inbound_stmt = inbound_stmt.where(BottleInbound.source_id == source_id)
        outbound_stmt = outbound_stmt.where(BottleOutbound.source_id == source_id)
        override_stmt = override_stmt.where(BottleMonthlyInventoryOverride.source_id == source_id)

    raw_inbound: dict[tuple[int, int, date, str], int] = {}
    raw_outbound: dict[tuple[int, int, date, str], int] = {}
    state_changes = {}
    for (source, spec, day), delta in pending_inventory_state_changes(db, date_to, source_id).items():
        key = (source, spec, day.replace(day=1))
        state_changes[key] = state_changes.get(key, 0) + delta
    event_months = set(months) | {key[2] for key in state_changes}
    for row_source_id, row_spec_id, business_date, pending, temporary, scrap in db.execute(inbound_stmt).all():
        row_month = business_date.replace(day=1)
        base_key = (int(row_source_id), int(row_spec_id), row_month)
        for category, quantity in (("PENDING", pending), ("TEMPORARY", temporary), ("SCRAP", scrap)):
            raw_inbound[(*base_key, category)] = raw_inbound.get((*base_key, category), 0) + int(quantity)
        event_months.add(row_month)
    for row_source_id, row_spec_id, business_date, pending, temporary, scrap, bad in db.execute(outbound_stmt).all():
        row_month = business_date.replace(day=1)
        base_key = (int(row_source_id), int(row_spec_id), row_month)
        for category, quantity in (
            ("PENDING", pending), ("TEMPORARY", temporary), ("SCRAP", scrap), ("BAD", bad),
        ):
            raw_outbound[(*base_key, category)] = raw_outbound.get((*base_key, category), 0) + int(quantity)
        event_months.add(row_month)

    overrides: dict[tuple[int, int, date, str], int] = {}
    for item in db.scalars(override_stmt).all():
        overrides[(item.source_id, item.spec_id, item.month, item.metric)] = item.value
        event_months.add(item.month)

    # Include historical, now-disabled master data if it is still referenced by business data.
    known_source_ids = {item.id for item in sources}
    known_spec_ids = set(spec_ids)
    for key in list(raw_inbound) + list(raw_outbound):
        known_source_ids.add(key[0])
        known_spec_ids.add(key[1])
    for key in overrides:
        known_source_ids.add(key[0])
        known_spec_ids.add(key[1])
    if source_id is None:
        selected_source_ids = sorted(known_source_ids)
    if known_spec_ids - set(spec_ids):
        extra_specs = db.scalars(select(BottleSpec).where(BottleSpec.id.in_(known_spec_ids - set(spec_ids)))).all()
        specs.extend(extra_specs)
        specs.sort(key=lambda item: (item.sort_order, item.id))
        spec_ids = [item.id for item in specs]

    ordered_event_months = sorted(item for item in event_months if item <= month_to)
    table_categories = {
        "INBOUND": (("PENDING", "待检"), ("TEMPORARY", "暂存"), ("SCRAP", "报废")),
        "OUTBOUND": (("PENDING", "待检验"), ("TEMPORARY", "暂存"), ("SCRAP", "报废"), ("BAD", "不良品")),
        "BALANCE": (("PENDING", "待检"), ("TEMPORARY", "暂存"), ("SCRAP", "报废"), ("BAD", "不良品")),
    }
    all_metrics = [f"{table}_{category}" for table, categories in table_categories.items() for category, _ in categories]
    result: dict[str, dict[int, dict[str, int]]] = {
        metric: {spec_id: {month.strftime("%Y-%m"): 0 for month in months} for spec_id in spec_ids}
        for metric in all_metrics
    }
    modified: dict[str, dict[int, set[str]]] = {
        metric: {spec_id: set() for spec_id in spec_ids} for metric in all_metrics
    }

    for row_source_id in selected_source_ids:
        for row_spec_id in spec_ids:
            balances = {category: 0 for category, _ in table_categories["BALANCE"]}
            for row_month in ordered_event_months:
                month_key = row_month.strftime("%Y-%m")
                flow_key = (row_source_id, row_spec_id, row_month)
                for category in balances:
                    inbound_metric = f"INBOUND_{category}"
                    outbound_metric = f"OUTBOUND_{category}"
                    balance_metric = f"BALANCE_{category}"
                    inbound_value = raw_inbound.get((*flow_key, category), 0)
                    outbound_value = raw_outbound.get((*flow_key, category), 0)
                    if (override_value := overrides.get((*flow_key, inbound_metric))) is not None:
                        inbound_value = override_value
                    if (override_value := overrides.get((*flow_key, outbound_metric))) is not None:
                        outbound_value = override_value
                    balance_override = overrides.get((*flow_key, balance_metric))
                    # 不良品没有独立入库来源；保留出库造成的负数，显示真实的累计扣减。
                    calculated_balance = balances[category] + inbound_value - outbound_value
                    if category == "PENDING":
                        calculated_balance += state_changes.get(flow_key, 0)
                    balances[category] = (
                        balance_override if balance_override is not None else
                        calculated_balance if category == "BAD" else max(calculated_balance, 0)
                    )
                    if month_from <= row_month <= month_to:
                        if inbound_metric in result:
                            result[inbound_metric][row_spec_id][month_key] += inbound_value
                        result[outbound_metric][row_spec_id][month_key] += outbound_value
                        result[balance_metric][row_spec_id][month_key] += balances[category]
                        for metric in (inbound_metric, outbound_metric, balance_metric):
                            if metric in modified and (*flow_key, metric) in overrides:
                                modified[metric][row_spec_id].add(month_key)

    def table_payload(table: str) -> dict:
        rows = []
        last_month_key = months[-1].strftime("%Y-%m")
        categories = table_categories[table]
        for category, category_label in categories:
            metric = f"{table}_{category}"
            for index, spec in enumerate(specs):
                values = result[metric][spec.id]
                rows.append({
                    "metric": metric,
                    "category": category,
                    "category_label": category_label,
                    "category_first": index == 0,
                    "category_rowspan": len(specs),
                    "spec_id": spec.id,
                    "spec_name": spec.name,
                    "values": values,
                    "modified_months": sorted(modified[metric][spec.id]),
                    "summary": values[last_month_key] if table == "BALANCE" else sum(values.values()),
                })
        totals = {
            month.strftime("%Y-%m"): sum(
                result[f"{table}_{category}"][spec.id][month.strftime("%Y-%m")]
                for category, _ in categories for spec in specs
            )
            for month in months
        }
        return {
            "rows": rows,
            "totals": totals,
            "summary": totals[last_month_key] if table == "BALANCE" else sum(totals.values()),
        }

    selected_source = next((item for item in sources if item.id == source_id), None)
    return {
        "month_from": month_from.strftime("%Y-%m"),
        "month_to": month_to.strftime("%Y-%m"),
        "months": [{"key": item.strftime("%Y-%m"), "label": f"{item.year}年{item.month}月"} for item in months],
        "source_id": source_id,
        "source_name": selected_source.name if selected_source else "全部公司",
        "editable": source_id is not None,
        "sources": [{"id": item.id, "name": item.name, "is_active": item.is_active} for item in sources],
        "tables": {
            "inbound": table_payload("INBOUND"),
            "outbound": table_payload("OUTBOUND"),
            "balance": table_payload("BALANCE"),
        },
    }


@router.get("/monthly-inventory")
def monthly_inventory(
    month_from: str,
    month_to: str,
    source_id: int | None = None,
    _: User = Depends(require("monthly_inventory:view")),
    db: Session = Depends(get_db),
):
    start = parse_inventory_month(month_from, "开始月份")
    end = parse_inventory_month(month_to, "结束月份")
    if start > end:
        raise HTTPException(status_code=400, detail="开始月份不能晚于结束月份")
    return build_monthly_inventory_report(db, start, end, source_id)


def build_daily_inventory_report(db: Session, date_from: date, date_to: date, source_id: int | None = None) -> dict:
    if date_from > date_to:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    if (date_to - date_from).days > 30:
        raise HTTPException(status_code=400, detail="日度区间最多选择 31 天")
    days = [date_from + timedelta(days=offset) for offset in range((date_to - date_from).days + 1)]
    sources = list(db.scalars(select(BottleSource).order_by(BottleSource.is_default.desc(), BottleSource.sort_order, BottleSource.id)).all())
    if source_id is not None and not any(item.id == source_id for item in sources):
        raise not_found("供瓶单位")
    specs = list(db.scalars(select(BottleSpec).order_by(BottleSpec.sort_order, BottleSpec.id)).all())
    inbound_stmt = (select(BottleInbound.source_id, BottleInboundLine.spec_id, BottleInbound.business_date,
        func.sum(BottleInboundLine.pending_quantity), func.sum(BottleInboundLine.temporary_quantity), func.sum(BottleInboundLine.scrap_quantity))
        .join(BottleInboundLine, BottleInboundLine.inbound_id == BottleInbound.id)
        .where(BottleInbound.status == "POSTED", BottleInbound.business_date <= date_to)
        .group_by(BottleInbound.source_id, BottleInboundLine.spec_id, BottleInbound.business_date))
    outbound_stmt = (select(BottleOutbound.source_id, BottleOutboundLine.spec_id, BottleOutbound.business_date,
        func.sum(BottleOutboundLine.good_quantity), func.sum(BottleOutboundLine.temporary_quantity),
        func.sum(BottleOutboundLine.scrap_quantity), func.sum(BottleOutboundLine.bad_quantity))
        .join(BottleOutboundLine, BottleOutboundLine.outbound_id == BottleOutbound.id)
        .where(BottleOutbound.status == "POSTED", BottleOutbound.source_id.is_not(None), BottleOutbound.business_date <= date_to)
        .group_by(BottleOutbound.source_id, BottleOutboundLine.spec_id, BottleOutbound.business_date))
    daily_stmt = select(BottleDailyInventoryOverride).where(BottleDailyInventoryOverride.business_date <= date_to)
    monthly_stmt = select(BottleMonthlyInventoryOverride).where(BottleMonthlyInventoryOverride.month <= date_to.replace(day=1))
    if source_id is not None:
        inbound_stmt = inbound_stmt.where(BottleInbound.source_id == source_id)
        outbound_stmt = outbound_stmt.where(BottleOutbound.source_id == source_id)
        daily_stmt = daily_stmt.where(BottleDailyInventoryOverride.source_id == source_id)
        monthly_stmt = monthly_stmt.where(BottleMonthlyInventoryOverride.source_id == source_id)
    flows: dict[tuple[int, int, date, str], int] = {}
    for source, spec, day, pending, temporary, scrap in db.execute(inbound_stmt):
        for category, quantity in (("PENDING", pending), ("TEMPORARY", temporary), ("SCRAP", scrap)):
            flows[(source, spec, day, f"INBOUND_{category}")] = int(quantity or 0)
    for source, spec, day, pending, temporary, scrap, bad in db.execute(outbound_stmt):
        for category, quantity in (("PENDING", pending), ("TEMPORARY", temporary), ("SCRAP", scrap), ("BAD", bad)):
            flows[(source, spec, day, f"OUTBOUND_{category}")] = int(quantity or 0)
    daily = {(item.source_id, item.spec_id, item.business_date, item.metric): item.value for item in db.scalars(daily_stmt)}
    monthly = {(item.source_id, item.spec_id, item.month, item.metric): item.value for item in db.scalars(monthly_stmt)}
    known_source_ids = {item.id for item in sources}
    known_spec_ids = {item.id for item in specs}
    for key in list(flows) + list(daily) + list(monthly):
        known_source_ids.add(key[0])
        known_spec_ids.add(key[1])
    if known_spec_ids - {item.id for item in specs}:
        specs.extend(db.scalars(select(BottleSpec).where(BottleSpec.id.in_(known_spec_ids - {item.id for item in specs}))).all())
        specs.sort(key=lambda item: (item.sort_order, item.id))
    selected_sources = [source_id] if source_id is not None else sorted(known_source_ids)
    categories = {
        "INBOUND": (("PENDING", "待检"), ("TEMPORARY", "暂存"), ("SCRAP", "报废")),
        "OUTBOUND": (("PENDING", "待检验"), ("TEMPORARY", "暂存"), ("SCRAP", "报废"), ("BAD", "不良品")),
        "BALANCE": (("PENDING", "待检"), ("TEMPORARY", "暂存"), ("SCRAP", "报废"), ("BAD", "不良品")),
    }
    metrics = [f"{table}_{category}" for table, group in categories.items() for category, _ in group]
    values = {metric: {spec.id: {day.isoformat(): 0 for day in days} for spec in specs} for metric in metrics}
    modified = {metric: {spec.id: set() for spec in specs} for metric in metrics}
    state_changes = pending_inventory_state_changes(db, date_to, source_id)
    event_days = set(days) | {key[2] for key in flows} | {key[2] for key in daily} | {key[2] for key in state_changes}
    for month_key in {key[2] for key in monthly}:
        month_end = next_inventory_month(month_key) - timedelta(days=1)
        if month_end <= date_to:
            event_days.add(month_end)
    for source in selected_sources:
        for spec in specs:
            balance = {category: 0 for category, _ in categories["BALANCE"]}
            month_flow = {(table, category): 0 for table in ("INBOUND", "OUTBOUND") for category in balance}
            active_month = None
            for day in sorted(item for item in event_days if item <= date_to):
                month = day.replace(day=1)
                if month != active_month:
                    month_flow = {(table, category): 0 for table in ("INBOUND", "OUTBOUND") for category in balance}
                    active_month = month
                daily_values = {}
                for category in balance:
                    for table in ("INBOUND", "OUTBOUND"):
                        metric = f"{table}_{category}"
                        raw = flows.get((source, spec.id, day, metric), 0)
                        current = daily.get((source, spec.id, day, metric), raw)
                        daily_values[metric] = current
                        month_flow[(table, category)] += current
                    calculated_balance = balance[category] + daily_values[f"INBOUND_{category}"] - daily_values[f"OUTBOUND_{category}"]
                    if category == "PENDING":
                        calculated_balance += state_changes.get((source, spec.id, day), 0)
                    balance[category] = calculated_balance if category == "BAD" else max(calculated_balance, 0)
                if day == next_inventory_month(month) - timedelta(days=1):
                    for category in balance:
                        for table in ("INBOUND", "OUTBOUND"):
                            metric = f"{table}_{category}"
                            override = monthly.get((source, spec.id, month, metric))
                            if override is not None:
                                delta = override - month_flow[(table, category)]
                                calculated_balance = balance[category] + (delta if table == "INBOUND" else -delta)
                                balance[category] = calculated_balance if category == "BAD" else max(calculated_balance, 0)
                        month_balance = monthly.get((source, spec.id, month, f"BALANCE_{category}"))
                        if month_balance is not None:
                            balance[category] = month_balance
                for category in balance:
                    metric = f"BALANCE_{category}"
                    balance[category] = daily.get((source, spec.id, day, metric), balance[category])
                    daily_values[metric] = balance[category]
                if date_from <= day <= date_to:
                    day_key = day.isoformat()
                    for metric in metrics:
                        values[metric][spec.id][day_key] += daily_values.get(metric, 0)
                        if (source, spec.id, day, metric) in daily:
                            modified[metric][spec.id].add(day_key)

    def table_payload(table: str) -> dict:
        rows = []
        for category, label in categories[table]:
            metric = f"{table}_{category}"
            for index, spec in enumerate(specs):
                row_values = values[metric][spec.id]
                rows.append({"metric": metric, "category": category, "category_label": label,
                    "category_first": index == 0, "category_rowspan": len(specs), "spec_id": spec.id,
                    "spec_name": spec.name, "values": row_values, "modified_months": sorted(modified[metric][spec.id]),
                    "summary": row_values[date_to.isoformat()] if table == "BALANCE" else sum(row_values.values())})
        totals = {day.isoformat(): sum(values[f"{table}_{category}"][spec.id][day.isoformat()]
            for category, _ in categories[table] for spec in specs) for day in days}
        return {"rows": rows, "totals": totals,
            "summary": totals[date_to.isoformat()] if table == "BALANCE" else sum(totals.values())}

    selected_source = next((item for item in sources if item.id == source_id), None)
    return {"date_from": date_from.isoformat(), "date_to": date_to.isoformat(),
        "months": [{"key": day.isoformat(), "label": f"{day.month}月{day.day}日"} for day in days],
        "source_id": source_id, "source_name": selected_source.name if selected_source else "全部公司",
        "editable": source_id is not None,
        "sources": [{"id": item.id, "name": item.name, "is_active": item.is_active} for item in sources],
        "tables": {table.lower(): table_payload(table) for table in categories}}


@router.get("/daily-inventory")
def daily_inventory(date_from: date, date_to: date, source_id: int | None = None,
    _: User = Depends(require("daily_inventory:view")), db: Session = Depends(get_db)):
    return build_daily_inventory_report(db, date_from, date_to, source_id)


@router.put("/daily-inventory/value")
def update_daily_inventory_value(payload: BottleDailyInventoryUpdate,
    user: User = Depends(require("daily_inventory:update")), db: Session = Depends(get_db)):
    if not db.get(BottleSource, payload.source_id):
        raise not_found("供瓶单位")
    if not db.get(BottleSpec, payload.spec_id):
        raise not_found("瓶子规格")
    before = build_daily_inventory_report(db, payload.business_date, payload.business_date, payload.source_id)
    row = next(item for item in before["tables"][payload.metric.split("_", 1)[0].lower()]["rows"]
        if item["spec_id"] == payload.spec_id and item["metric"] == payload.metric)
    old_value = int(row["values"][payload.business_date.isoformat()])
    if old_value == payload.value:
        return {"changed": False, "old_value": old_value, "new_value": payload.value}
    override = db.scalar(select(BottleDailyInventoryOverride).where(
        BottleDailyInventoryOverride.business_date == payload.business_date,
        BottleDailyInventoryOverride.source_id == payload.source_id,
        BottleDailyInventoryOverride.spec_id == payload.spec_id,
        BottleDailyInventoryOverride.metric == payload.metric).with_for_update())
    if override:
        override.value = payload.value
        override.updated_by = user.id
    else:
        db.add(BottleDailyInventoryOverride(business_date=payload.business_date, source_id=payload.source_id,
            spec_id=payload.spec_id, metric=payload.metric, value=payload.value, updated_by=user.id))
    log = BottleDailyInventoryLog(business_date=payload.business_date, source_id=payload.source_id,
        spec_id=payload.spec_id, metric=payload.metric, old_value=old_value, new_value=payload.value,
        reason=(payload.reason or "日度盘点页面直接修改").strip() or "日度盘点页面直接修改", created_by=user.id)
    db.add(log)
    db.commit()
    db.refresh(log)
    return {"changed": True, "old_value": old_value, "new_value": payload.value, "log_id": log.id}


@router.get("/daily-inventory/logs")
def daily_inventory_logs(source_id: int | None = None, date_from: date | None = None, date_to: date | None = None,
    page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require("daily_inventory:view")), db: Session = Depends(get_db)):
    stmt = select(BottleDailyInventoryLog)
    if source_id is not None:
        stmt = stmt.where(BottleDailyInventoryLog.source_id == source_id)
    if date_from:
        stmt = stmt.where(BottleDailyInventoryLog.business_date >= date_from)
    if date_to:
        stmt = stmt.where(BottleDailyInventoryLog.business_date <= date_to)
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    items = db.scalars(stmt.order_by(BottleDailyInventoryLog.created_at.desc(), BottleDailyInventoryLog.id.desc())
        .offset((page - 1) * page_size).limit(page_size)).all()
    return {"items": [{"id": item.id, "month": item.business_date.isoformat(),
        "source_id": item.source_id, "source_name": item.source.name, "spec_id": item.spec_id,
        "spec_name": item.spec.name, "metric": item.metric, "old_value": item.old_value,
        "new_value": item.new_value, "reason": item.reason, "created_by": item.created_by,
        "created_by_name": item.creator.name, "created_at": item.created_at} for item in items], "total": total}


@router.put("/monthly-inventory/value")
def update_monthly_inventory_value(
    payload: BottleMonthlyInventoryUpdate,
    user: User = Depends(require("monthly_inventory:update")),
    db: Session = Depends(get_db),
):
    month = parse_inventory_month(payload.month, "盘点月份")
    source = db.get(BottleSource, payload.source_id)
    spec = db.get(BottleSpec, payload.spec_id)
    if not source:
        raise not_found("供瓶单位")
    if not spec:
        raise not_found("瓶子规格")

    before = build_monthly_inventory_report(db, month, month, payload.source_id)
    table_key = payload.metric.split("_", 1)[0].lower()
    row = next(
        item for item in before["tables"][table_key]["rows"]
        if item["spec_id"] == payload.spec_id and item["metric"] == payload.metric
    )
    old_value = int(row["values"][payload.month])
    if old_value == payload.value:
        return {"changed": False, "old_value": old_value, "new_value": payload.value}

    override = db.scalar(select(BottleMonthlyInventoryOverride).where(
        BottleMonthlyInventoryOverride.month == month,
        BottleMonthlyInventoryOverride.source_id == payload.source_id,
        BottleMonthlyInventoryOverride.spec_id == payload.spec_id,
        BottleMonthlyInventoryOverride.metric == payload.metric,
    ).with_for_update())
    if override:
        override.value = payload.value
        override.updated_by = user.id
    else:
        override = BottleMonthlyInventoryOverride(
            month=month,
            source_id=payload.source_id,
            spec_id=payload.spec_id,
            metric=payload.metric,
            value=payload.value,
            updated_by=user.id,
        )
        db.add(override)
    log = BottleMonthlyInventoryLog(
        month=month,
        source_id=payload.source_id,
        spec_id=payload.spec_id,
        metric=payload.metric,
        old_value=old_value,
        new_value=payload.value,
        reason=(payload.reason or "月度盘点页面直接修改").strip() or "月度盘点页面直接修改",
        created_by=user.id,
    )
    db.add(log)
    db.commit()
    db.refresh(log)
    return {"changed": True, "old_value": old_value, "new_value": payload.value, "log_id": log.id}


@router.get("/monthly-inventory/logs")
def monthly_inventory_logs(
    source_id: int | None = None,
    month_from: str | None = None,
    month_to: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require("monthly_inventory:view")),
    db: Session = Depends(get_db),
):
    stmt = select(BottleMonthlyInventoryLog)
    if source_id is not None:
        stmt = stmt.where(BottleMonthlyInventoryLog.source_id == source_id)
    if month_from:
        stmt = stmt.where(BottleMonthlyInventoryLog.month >= parse_inventory_month(month_from, "开始月份"))
    if month_to:
        stmt = stmt.where(BottleMonthlyInventoryLog.month <= parse_inventory_month(month_to, "结束月份"))
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    items = db.scalars(
        stmt.order_by(BottleMonthlyInventoryLog.created_at.desc(), BottleMonthlyInventoryLog.id.desc())
        .offset((page - 1) * page_size).limit(page_size)
    ).all()
    return {
        "items": [{
            "id": item.id,
            "month": item.month.strftime("%Y-%m"),
            "source_id": item.source_id,
            "source_name": item.source.name,
            "spec_id": item.spec_id,
            "spec_name": item.spec.name,
            "metric": item.metric,
            "old_value": item.old_value,
            "new_value": item.new_value,
            "reason": item.reason,
            "created_by": item.created_by,
            "created_by_name": item.creator.name,
            "created_at": item.created_at,
        } for item in items],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@router.get("/bottle-inbounds", response_model=PageOut[BottleInboundOut])
def list_bottle_inbounds(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    date_from: date | None = None,
    date_to: date | None = None,
    source_id: int | None = None,
    document_no: str | None = None,
    handler: str | None = None,
    _: User = Depends(require("bottle_inbound:view")),
    db: Session = Depends(get_db),
):
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    stmt = select(BottleInbound)
    if date_from is not None:
        stmt = stmt.where(BottleInbound.business_date >= date_from)
    if date_to is not None:
        stmt = stmt.where(BottleInbound.business_date <= date_to)
    if source_id is not None:
        stmt = stmt.where(BottleInbound.source_id == source_id)
    if document_no and document_no.strip():
        stmt = stmt.where(BottleInbound.document_no.contains(document_no.strip()))
    if handler and handler.strip():
        stmt = stmt.where(BottleInbound.handler.contains(handler.strip()))
    return page_result(db, stmt.order_by(BottleInbound.business_date.desc(), BottleInbound.id.desc()), page, page_size)


@router.get("/bottle-inbounds/summary")
def bottle_inbound_summary(business_date: date = Query(default_factory=china_today), _: User = Depends(require("bottle_inbound:view")), db: Session = Depends(get_db)):
    rows = db.execute(select(
        func.count(BottleInbound.id), func.coalesce(func.sum(BottleInbound.pending_quantity), 0),
        func.coalesce(func.sum(BottleInbound.temporary_quantity), 0), func.coalesce(func.sum(BottleInbound.scrap_quantity), 0),
    ).where(BottleInbound.status == "POSTED", BottleInbound.business_date == business_date)).one()
    return {"batch_count": rows[0], "pending": rows[1], "temporary": rows[2], "scrap": rows[3], "total": rows[1] + rows[2] + rows[3]}


@router.get("/bottle-inbounds/report", response_model=list[BottleInboundOut])
def bottle_inbound_report(business_date: date = Query(default_factory=china_today), _: User = Depends(require("bottle_inbound:view")), db: Session = Depends(get_db)):
    return db.scalars(select(BottleInbound).where(BottleInbound.business_date == business_date).order_by(BottleInbound.id)).all()


def replace_bottle_inbound_lines(db: Session, obj: BottleInbound, payload: BottleInboundCreate):
    lines = payload.lines
    if not lines:
        default_spec = db.scalar(select(BottleSpec).where(BottleSpec.is_active.is_(True)).order_by(BottleSpec.is_default.desc(), BottleSpec.sort_order, BottleSpec.id))
        if not default_spec:
            raise HTTPException(status_code=400, detail="请先配置一个启用的钢瓶规格")
        lines = [{
            "spec_id": default_spec.id,
            "pending_quantity": payload.pending_quantity,
            "temporary_quantity": payload.temporary_quantity,
            "scrap_quantity": payload.scrap_quantity,
        }]
    obj.lines.clear()
    if obj.id is not None:
        db.flush()
    for item in lines:
        values = item.model_dump() if hasattr(item, "model_dump") else item
        spec = db.get(BottleSpec, values["spec_id"])
        if not spec or not spec.is_active:
            raise HTTPException(status_code=400, detail="钢瓶规格不存在或已停用")
        obj.lines.append(BottleInboundLine(**values))
    obj.pending_quantity = sum(line.pending_quantity for line in obj.lines)
    obj.temporary_quantity = sum(line.temporary_quantity for line in obj.lines)
    obj.scrap_quantity = sum(line.scrap_quantity for line in obj.lines)


@router.post("/bottle-inbounds", response_model=BottleInboundOut, status_code=201)
def create_bottle_inbound(payload: BottleInboundCreate, user: User = Depends(require("bottle_inbound:create")), db: Session = Depends(get_db)):
    source = db.get(BottleSource, payload.source_id)
    if not source or not source.is_active:
        raise HTTPException(status_code=400, detail="供瓶单位不存在或已停用")
    obj = BottleInbound(
        document_no=make_no("CYI", payload.business_date),
        created_by=user.id,
        **payload.model_dump(exclude={"lines"}),
    )
    replace_bottle_inbound_lines(db, obj, payload)
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@router.put("/bottle-inbounds/{object_id}", response_model=BottleInboundOut)
def update_bottle_inbound(object_id: int, payload: BottleInboundCreate, _: User = Depends(require("bottle_inbound:update")), db: Session = Depends(get_db)):
    obj = db.get(BottleInbound, object_id)
    if not obj:
        raise not_found("钢瓶入库记录")
    apply_master_values(obj, payload, exclude={"lines"})
    replace_bottle_inbound_lines(db, obj, payload)
    db.commit()
    db.refresh(obj)
    return obj


@router.post("/quick-inbound", response_model=BottleInboundOut, status_code=201)
def quick_bottle_inbound(payload: BottleInboundCreate, user: User = Depends(require("quick_inbound:create")), db: Session = Depends(get_db)):
    source = db.get(BottleSource, payload.source_id)
    if not source or not source.is_active:
        raise HTTPException(status_code=400, detail="供瓶单位不存在或已停用")
    obj = BottleInbound(
        document_no=make_no("CYI", payload.business_date),
        created_by=user.id,
        status="POSTED",
        posted_at=datetime.utcnow(),
        **payload.model_dump(exclude={"lines"}),
    )
    replace_bottle_inbound_lines(db, obj, payload)
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


def replace_bottle_outbound_lines(db: Session, obj: BottleOutbound, payload: BottleOutboundCreate):
    obj.lines.clear()
    if obj.id is not None:
        db.flush()
    for item in payload.lines:
        spec = db.get(BottleSpec, item.spec_id)
        if not spec or not spec.is_active:
            raise HTTPException(status_code=400, detail="钢瓶规格不存在或已停用")
        obj.lines.append(BottleOutboundLine(**item.model_dump(exclude={"bad_bottle_ids", "scrap_bottle_ids"})))


def reconcile_edited_posted_outbound(db: Session, obj: BottleOutbound):
    """Keep linked bottles aligned when a posted historical outbound is edited directly."""
    linked_bottles = list(db.scalars(
        select(Bottle).where(Bottle.outbound_id == obj.id).order_by(Bottle.id).with_for_update()
    ).all())
    desired: dict[tuple[int, str], int] = {}
    for line in obj.lines:
        desired[(line.spec_id, "GOOD")] = line.good_quantity
        desired[(line.spec_id, "BAD")] = line.bad_quantity

    retained: dict[tuple[int, str], int] = {}
    for bottle in linked_bottles:
        origin = bottle.outbound_origin_zone or "GOOD"
        key = (bottle.spec_id, origin)
        if retained.get(key, 0) < desired.get(key, 0):
            retained[key] = retained.get(key, 0) + 1
            continue
        bottle.zone = origin
        bottle.outbound_id = None
        bottle.outbound_origin_zone = None
        bottle.good_entry_date = china_today() if origin == "GOOD" else None

    for line in obj.lines:
        line.good_adjustment_quantity = max(line.good_quantity - retained.get((line.spec_id, "GOOD"), 0), 0)
        line.bad_adjustment_quantity = max(line.bad_quantity - retained.get((line.spec_id, "BAD"), 0), 0)


def bottle_outbound_availability(db: Session, source_id: int | None = None) -> dict[str, dict[str, int]]:
    stock: dict[str, dict[str, int]] = {}
    def empty_stock():
        return {"good": 0, "bad": 0, "temporary": 0, "scrap": 0}

    bottle_stmt = select(Bottle.spec_id, Bottle.zone, func.count(Bottle.id)).where(
        Bottle.zone.in_(["GOOD", "BAD"]), Bottle.outbound_id.is_(None)
    )
    if source_id is not None:
        bottle_stmt = bottle_stmt.where(func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == source_id)
    for spec_id, zone, quantity in db.execute(
        bottle_stmt
        .group_by(Bottle.spec_id, Bottle.zone)
    ).all():
        stock.setdefault(str(spec_id), empty_stock())[zone.lower()] = int(quantity)
    inbound_stmt = (select(
            BottleInboundLine.spec_id,
            func.coalesce(func.sum(BottleInboundLine.temporary_quantity), 0),
            func.coalesce(func.sum(BottleInboundLine.scrap_quantity), 0),
        )
        .join(BottleInbound, BottleInbound.id == BottleInboundLine.inbound_id)
        .where(BottleInbound.status == "POSTED"))
    if source_id is not None:
        inbound_stmt = inbound_stmt.where(BottleInbound.source_id == source_id)
    for spec_id, temporary, scrap in db.execute(inbound_stmt.group_by(BottleInboundLine.spec_id)).all():
        item = stock.setdefault(str(spec_id), empty_stock())
        item["temporary"] = int(temporary)
        item["scrap"] = int(scrap)
    adjustment_stmt = select(
            BottleInventoryAdjustment.spec_id,
            BottleInventoryAdjustment.zone,
            func.coalesce(func.sum(BottleInventoryAdjustment.quantity), 0),
        )
    if source_id is not None:
        adjustment_stmt = adjustment_stmt.where(or_(
            BottleInventoryAdjustment.source_id == source_id,
            and_(BottleInventoryAdjustment.source_id.is_(None), BottleInventoryAdjustment.zone == "GOOD"),
        ))
    for spec_id, zone, quantity in db.execute(
        adjustment_stmt.group_by(BottleInventoryAdjustment.spec_id, BottleInventoryAdjustment.zone)
    ).all():
        if zone == "BAD":
            continue  # Added below after subtracting posted bad outbounds.
        item = stock.setdefault(str(spec_id), empty_stock())
        item[{"GOOD": "good", "TEMPORARY": "temporary", "SCRAP": "scrap"}[zone]] += int(quantity)
    outbound_quantity_stmt = (select(
            BottleOutboundLine.spec_id,
            func.coalesce(func.sum(BottleOutboundLine.temporary_quantity), 0),
            func.coalesce(func.sum(BottleOutboundLine.scrap_quantity), 0),
        )
        .join(BottleOutbound, BottleOutbound.id == BottleOutboundLine.outbound_id)
        .where(BottleOutbound.status == "POSTED"))
    if source_id is not None:
        outbound_quantity_stmt = outbound_quantity_stmt.where(BottleOutbound.source_id == source_id)
    for spec_id, temporary, scrap in db.execute(outbound_quantity_stmt.group_by(BottleOutboundLine.spec_id)).all():
        item = stock.setdefault(str(spec_id), empty_stock())
        item["temporary"] = max(item["temporary"] - int(temporary), 0)
        item["scrap"] = max(item["scrap"] - int(scrap), 0)
    outbound_good_stmt = (select(BottleOutboundLine.spec_id, func.coalesce(func.sum(BottleOutboundLine.good_adjustment_quantity), 0))
        .join(BottleOutbound, BottleOutbound.id == BottleOutboundLine.outbound_id)
        .where(BottleOutbound.status == "POSTED"))
    # Anonymous good adjustments predate customer ownership and remain one shared
    # legacy pool, so their posted consumption must also be subtracted globally.
    for spec_id, quantity in db.execute(outbound_good_stmt.group_by(BottleOutboundLine.spec_id)).all():
        item = stock.setdefault(str(spec_id), empty_stock())
        item["good"] = max(item["good"] - int(quantity), 0)
    for spec_id, quantity in manual_bad_stock(db, source_id).items():
        stock.setdefault(str(spec_id), empty_stock())["bad"] += quantity
    return stock


def active_inspection_exists():
    return select(InspectionDocumentLine.id).join(
        InspectionDocument, InspectionDocument.id == InspectionDocumentLine.document_id
    ).where(
        InspectionDocumentLine.bottle_id == Bottle.id,
        InspectionDocument.status == "ACTIVE",
    ).exists()


def lock_bottle_stock(db: Session, spec_ids: list[int]) -> None:
    if db.bind and db.bind.dialect.name == "postgresql":
        for spec_id in sorted(set(spec_ids)):
            db.execute(text("SELECT pg_advisory_xact_lock(902620, :spec_id)"), {"spec_id": spec_id})


def oldest_verified_good_bottle(db: Session, spec_id: int, source_id: int | None, excluded_id: int):
    filters = [
        Bottle.spec_id == spec_id, Bottle.zone == "GOOD", Bottle.outbound_id.is_(None),
        Bottle.id != excluded_id, active_inspection_exists(),
    ]
    if source_id is not None:
        filters.append(Bottle.source_id == source_id)
    return select(Bottle).where(*filters).order_by(
        Bottle.coding_date.is_(None), Bottle.coding_date, Bottle.id
    )


def quick_outbound_stock(db: Session, source_id: int | None = None) -> dict[str, dict[str, int]]:
    stock = bottle_outbound_availability(db, source_id)
    verified_stmt = select(Bottle.spec_id, Bottle.zone, func.count(Bottle.id)).where(
        Bottle.zone.in_(["GOOD", "BAD"]), Bottle.outbound_id.is_(None), active_inspection_exists()
    )
    named_stmt = select(Bottle.spec_id, Bottle.zone, func.count(Bottle.id)).where(
        Bottle.zone.in_(["GOOD", "BAD", "SCRAP"]), Bottle.outbound_id.is_(None)
    )
    if source_id is not None:
        source_filter = func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == source_id
        verified_stmt = verified_stmt.where(source_filter)
        named_stmt = named_stmt.where(source_filter)
    verified = {(spec_id, zone): int(quantity) for spec_id, zone, quantity in db.execute(
        verified_stmt
        .group_by(Bottle.spec_id, Bottle.zone)
    ).all()}
    named = {(spec_id, zone): int(quantity) for spec_id, zone, quantity in db.execute(
        named_stmt
        .group_by(Bottle.spec_id, Bottle.zone)
    ).all()}
    for key, item in stock.items():
        spec_id = int(key)
        for zone in ("GOOD", "BAD"):
            label = zone.lower()
            available_total = item[label]
            anonymous = max(available_total - named.get((spec_id, zone), 0), 0)
            eligible_named = verified.get((spec_id, zone), 0) if zone == "GOOD" else named.get((spec_id, zone), 0)
            eligible = min(eligible_named, max(available_total - anonymous, 0))
            item[f"anonymous_{label}"] = anonymous
            item[f"verified_{label}"] = eligible
            item[label] = anonymous + eligible
        item["selectable_bad"] = min(item["bad"], named.get((spec_id, "BAD"), 0))
        item["selectable_scrap"] = min(item["scrap"], named.get((spec_id, "SCRAP"), 0))
    return stock


def manual_bad_stock(db: Session, source_id: int | None = None) -> dict[int, int]:
    adjustment_stmt = (select(
        BottleInventoryAdjustment.spec_id,
        func.coalesce(func.sum(BottleInventoryAdjustment.quantity), 0),
    ).where(BottleInventoryAdjustment.zone == "BAD"))
    if source_id is not None:
        adjustment_stmt = adjustment_stmt.where(BottleInventoryAdjustment.source_id == source_id)
    stock = {
        spec_id: int(quantity)
        for spec_id, quantity in db.execute(
            adjustment_stmt.group_by(BottleInventoryAdjustment.spec_id)
        ).all()
    }
    outbound_stmt = (select(BottleOutboundLine.spec_id, func.coalesce(func.sum(BottleOutboundLine.bad_adjustment_quantity), 0))
        .join(BottleOutbound, BottleOutbound.id == BottleOutboundLine.outbound_id)
        .where(BottleOutbound.status == "POSTED"))
    if source_id is not None:
        outbound_stmt = outbound_stmt.where(BottleOutbound.source_id == source_id)
    for spec_id, quantity in db.execute(outbound_stmt.group_by(BottleOutboundLine.spec_id)).all():
        stock[spec_id] = stock.get(spec_id, 0) - int(quantity)
    return stock


def allocate_bottle_outbound(
    db: Session, obj: BottleOutbound, user: User, *, quick: bool = False,
    payload: BottleOutboundCreate | None = None,
):
    allocations: list[tuple[Bottle, str]] = []
    selected_by_spec = {item.spec_id: item for item in payload.lines} if payload else {}
    if quick:
        lock_bottle_stock(db, [line.spec_id for line in obj.lines])
    quantity_stock = quick_outbound_stock(db, obj.source_id) if quick else bottle_outbound_availability(db, obj.source_id)
    for line in obj.lines:
        requested = selected_by_spec.get(line.spec_id)
        spec_stock = quantity_stock.get(str(line.spec_id), {"good": 0, "bad": 0, "temporary": 0, "scrap": 0})
        if line.good_quantity > spec_stock["good"]:
            raise HTTPException(status_code=400, detail=f"{line.spec_name} 良品库存不足，需要 {line.good_quantity} 只，当前可用 {spec_stock['good']} 只")
        good_named_needed = max(line.good_quantity - spec_stock.get("anonymous_good", 0), 0) if quick else line.good_quantity
        good_filters = [
            Bottle.zone == "GOOD",
            Bottle.spec_id == line.spec_id,
            Bottle.outbound_id.is_(None),
            func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == obj.source_id,
        ]
        if quick:
            good_filters.append(active_inspection_exists())
        good_order = (Bottle.coding_date, Bottle.id) if quick else (
            case((active_inspection_exists(), 0), else_=1), Bottle.coding_date, Bottle.id
        )
        available = db.scalars(select(Bottle).where(*good_filters)
            .order_by(*good_order).with_for_update().limit(good_named_needed)).all()
        if quick and len(available) != good_named_needed:
            raise HTTPException(status_code=409, detail=f"{line.spec_name} 可出库的已检验良品数量已变化，请刷新重试")
        line.good_adjustment_quantity = line.good_quantity - len(available)
        if line.bad_quantity > spec_stock["bad"]:
            raise HTTPException(status_code=400, detail=f"{line.spec_name} 不良品库存不足，需要 {line.bad_quantity} 只，当前可用 {spec_stock['bad']} 只")
        selected_bad_ids = requested.bad_bottle_ids if requested else []
        bad_named_needed = len(selected_bad_ids) if quick else line.bad_quantity
        bad_filters = [
            Bottle.zone == "BAD",
            Bottle.spec_id == line.spec_id,
            Bottle.outbound_id.is_(None),
            func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == obj.source_id,
        ]
        if quick and selected_bad_ids:
            bad_filters.append(Bottle.id.in_(selected_bad_ids))
        bad_bottles = db.scalars(select(Bottle).where(*bad_filters)
            .order_by(Bottle.coding_date, Bottle.id).with_for_update().limit(bad_named_needed)).all()
        if quick and len(bad_bottles) != bad_named_needed:
            raise HTTPException(status_code=409, detail=f"{line.spec_name} 已选择的不良品状态已变化，请重新选择")
        line.bad_adjustment_quantity = line.bad_quantity - len(bad_bottles)
        if line.temporary_quantity > spec_stock["temporary"]:
            raise HTTPException(status_code=400, detail=f"{line.spec_name} 暂存库存不足，需要 {line.temporary_quantity} 只，当前可用 {spec_stock['temporary']} 只")
        if line.scrap_quantity > spec_stock["scrap"]:
            raise HTTPException(status_code=400, detail=f"{line.spec_name} 报废库存不足，需要 {line.scrap_quantity} 只，当前可用 {spec_stock['scrap']} 只")
        selected_scrap_ids = requested.scrap_bottle_ids if requested else []
        scrap_bottles: list[Bottle] = []
        if quick and selected_scrap_ids:
            scrap_bottles = list(db.scalars(select(Bottle).where(
                Bottle.id.in_(selected_scrap_ids),
                Bottle.zone == "SCRAP",
                Bottle.spec_id == line.spec_id,
                Bottle.outbound_id.is_(None),
                func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == obj.source_id,
            ).order_by(Bottle.coding_date, Bottle.id).with_for_update()).all())
            if len(scrap_bottles) != len(selected_scrap_ids):
                raise HTTPException(status_code=409, detail=f"{line.spec_name} 已选择的报废瓶状态已变化，请重新选择")
        allocations.extend((bottle, "GOOD") for bottle in available)
        allocations.extend((bottle, "BAD") for bottle in bad_bottles)
        allocations.extend((bottle, "SCRAP") for bottle in scrap_bottles)
        for bottle in [*bad_bottles, *scrap_bottles]:
            line.allocations.append(BottleOutboundAllocation(
                outbound_id=obj.id, bottle_id=bottle.id, body_code=bottle.body_code,
                origin_zone=bottle.zone, source_id=obj.source_id, spec_id=bottle.spec_id,
                allocation_method="SELECTED" if quick else "AUTO",
            ))
    for bottle, original_zone in allocations:
        bottle.outbound_origin_zone = original_zone
        bottle.zone = "OUTBOUND"
        bottle.outbound_id = obj.id
    obj.status = "POSTED"
    obj.posted_by = user.id
    obj.posted_at = datetime.utcnow()


@router.get("/quick-outbound/availability")
def quick_outbound_availability(source_id: int | None = None, _: User = Depends(require("quick_outbound:view")), db: Session = Depends(get_db)):
    return quick_outbound_stock(db, source_id)


@router.get("/quick-outbound/bottle-candidates", response_model=PageOut[BottleOut])
def quick_outbound_bottle_candidates(
    source_id: int, spec_id: int, zone: str,
    body_code: str | None = None, date_from: date | None = None, date_to: date | None = None,
    page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require("quick_outbound:view")), db: Session = Depends(get_db),
):
    if zone not in {"BAD", "SCRAP"}:
        raise HTTPException(status_code=400, detail="只能选择不良品或报废瓶")
    available = quick_outbound_stock(db, source_id).get(str(spec_id), {}).get(zone.lower(), 0)
    if available <= 0:
        return {"items": [], "total": 0, "page": page, "page_size": page_size}
    pool_ids = select(Bottle.id).where(
        Bottle.zone == zone, Bottle.spec_id == spec_id, Bottle.outbound_id.is_(None),
        func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == source_id,
    ).order_by(Bottle.coding_date, Bottle.id).limit(available)
    stmt = select(Bottle).where(Bottle.id.in_(pool_ids)).order_by(Bottle.coding_date, Bottle.id)
    if body_code:
        stmt = stmt.where(Bottle.body_code.ilike(f"%{body_code.strip()}%"))
    if date_from:
        stmt = stmt.where(Bottle.coding_date >= date_from)
    if date_to:
        stmt = stmt.where(Bottle.coding_date <= date_to)
    return page_result(db, stmt, page, page_size)


@router.post("/quick-outbound", response_model=BottleOutboundOut, status_code=201)
def quick_bottle_outbound(payload: BottleOutboundCreate, user: User = Depends(require("quick_outbound:create")), db: Session = Depends(get_db)):
    source = db.get(BottleSource, payload.source_id)
    if not source or not source.is_active:
        raise HTTPException(status_code=400, detail="供瓶单位不存在或已停用")
    obj = BottleOutbound(
        document_no=make_no("CYO", payload.business_date),
        business_date=payload.business_date,
        source_id=source.id,
        customer=source.name,
        handler=payload.handler,
        note=payload.note,
        created_by=user.id,
    )
    replace_bottle_outbound_lines(db, obj, payload)
    db.add(obj)
    db.flush()
    allocate_bottle_outbound(db, obj, user, quick=True, payload=payload)
    db.commit()
    db.refresh(obj)
    return obj


@router.get("/bottle-outbounds", response_model=PageOut[BottleOutboundOut])
def list_bottle_outbounds(
    date_from: date | None = None, date_to: date | None = None, source_id: int | None = None,
    page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require("bottle_outbound:view")), db: Session = Depends(get_db),
):
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    stmt = select(BottleOutbound).order_by(BottleOutbound.business_date.desc(), BottleOutbound.id.desc())
    if date_from:
        stmt = stmt.where(BottleOutbound.business_date >= date_from)
    if date_to:
        stmt = stmt.where(BottleOutbound.business_date <= date_to)
    if source_id is not None:
        stmt = stmt.where(BottleOutbound.source_id == source_id)
    return page_result(db, stmt, page, page_size)


@router.post("/bottle-outbounds", response_model=BottleOutboundOut, status_code=201)
def create_bottle_outbound(payload: BottleOutboundCreate, user: User = Depends(require("bottle_outbound:create")), db: Session = Depends(get_db)):
    source = db.get(BottleSource, payload.source_id)
    if not source or not source.is_active:
        raise HTTPException(status_code=400, detail="供瓶单位不存在或已停用")
    obj = BottleOutbound(document_no=make_no("CYO", payload.business_date), business_date=payload.business_date, source_id=source.id, customer=source.name, handler=payload.handler, note=payload.note, created_by=user.id)
    replace_bottle_outbound_lines(db, obj, payload)
    db.add(obj); db.commit(); db.refresh(obj)
    return obj


@router.put("/bottle-outbounds/{object_id}", response_model=BottleOutboundOut)
def update_bottle_outbound(object_id: int, payload: BottleOutboundCreate, user: User = Depends(require("bottle_outbound:update")), db: Session = Depends(get_db)):
    obj = db.get(BottleOutbound, object_id)
    if not obj:
        raise not_found("钢瓶出库单")
    if obj.status == "POSTED" and any(line.allocations for line in obj.lines):
        raise HTTPException(status_code=400, detail="该出库单已绑定具体瓶号，请先撤销后再调整")
    source = db.get(BottleSource, payload.source_id)
    if not source or not source.is_active:
        raise HTTPException(status_code=400, detail="供瓶单位不存在或已停用")
    obj.business_date = payload.business_date; obj.source_id = source.id; obj.customer = source.name; obj.handler = payload.handler; obj.note = payload.note; obj.updated_by = user.id
    replace_bottle_outbound_lines(db, obj, payload)
    if obj.status == "POSTED":
        reconcile_edited_posted_outbound(db, obj)
    db.commit(); db.refresh(obj)
    return obj


@router.delete("/bottle-outbounds/{object_id}", status_code=204)
def delete_bottle_outbound(object_id: int, user: User = Depends(require("bottle_outbound:delete")), db: Session = Depends(get_db)):
    obj = db.get(BottleOutbound, object_id)
    if not obj:
        raise not_found("钢瓶出库单")
    if obj.status == "VOID":
        if not user.is_superadmin:
            raise HTTPException(status_code=403, detail="只有超级管理员可以彻底删除已作废记录")
    elif obj.status != "DRAFT":
        raise HTTPException(status_code=400, detail="只有草稿或已作废出库单可以删除")
    db.delete(obj); db.commit()


@router.post("/bottle-outbounds/{object_id}/post", response_model=BottleOutboundOut)
def post_bottle_outbound(object_id: int, user: User = Depends(require("bottle_outbound:post")), db: Session = Depends(get_db)):
    obj = db.scalar(select(BottleOutbound).where(BottleOutbound.id == object_id).with_for_update())
    if not obj:
        raise not_found("钢瓶出库单")
    if obj.status != "DRAFT":
        raise HTTPException(status_code=400, detail="出库单已处理")
    allocate_bottle_outbound(db, obj, user)
    db.commit(); db.refresh(obj)
    return obj


@router.post("/bottle-outbounds/{object_id}/void", response_model=BottleOutboundOut)
def void_bottle_outbound(object_id: int, user: User = Depends(require("bottle_outbound:void")), db: Session = Depends(get_db)):
    obj = db.scalar(select(BottleOutbound).where(BottleOutbound.id == object_id).with_for_update())
    if not obj:
        raise not_found("钢瓶出库单")
    if obj.status != "POSTED":
        raise HTTPException(status_code=400, detail="只有已审批出库单可以撤销")
    bottles = db.scalars(select(Bottle).where(Bottle.outbound_id == obj.id).with_for_update()).all()
    for bottle in bottles:
        original_zone = bottle.outbound_origin_zone or "GOOD"
        bottle.zone = original_zone
        bottle.outbound_id = None
        bottle.outbound_origin_zone = None
        bottle.good_entry_date = china_today() if original_zone == "GOOD" else None
    for line in obj.lines:
        for allocation in line.allocations:
            allocation.is_active = False
    obj.status = "VOID"; obj.voided_by = user.id; obj.voided_at = datetime.utcnow()
    db.commit(); db.refresh(obj)
    return obj


@router.delete("/bottle-inbounds/{object_id}", status_code=204)
def delete_bottle_inbound(object_id: int, user: User = Depends(require("bottle_inbound:delete")), db: Session = Depends(get_db)):
    obj = db.get(BottleInbound, object_id)
    if not obj:
        raise not_found("钢瓶入库记录")
    if obj.status == "VOID":
        if not user.is_superadmin:
            raise HTTPException(status_code=403, detail="只有超级管理员可以彻底删除已作废记录")
    elif obj.status != "DRAFT":
        raise HTTPException(status_code=400, detail="只有草稿或已作废记录可以删除")
    db.delete(obj)
    db.commit()


@router.post("/bottle-inbounds/{object_id}/post", response_model=BottleInboundOut)
def post_bottle_inbound(object_id: int, _: User = Depends(require("bottle_inbound:post")), db: Session = Depends(get_db)):
    obj = db.get(BottleInbound, object_id)
    if not obj:
        raise not_found("钢瓶入库记录")
    if obj.status != "DRAFT":
        raise HTTPException(status_code=400, detail="入库记录已处理")
    obj.status = "POSTED"
    obj.posted_at = datetime.utcnow()
    db.commit()
    db.refresh(obj)
    return obj


@router.post("/bottle-inbounds/{object_id}/void", response_model=BottleInboundOut)
def void_bottle_inbound(object_id: int, _: User = Depends(require("bottle_inbound:void")), db: Session = Depends(get_db)):
    obj = db.get(BottleInbound, object_id)
    if not obj:
        raise not_found("钢瓶入库记录")
    if obj.status != "POSTED":
        raise HTTPException(status_code=400, detail="只有已生效入库记录可以作废")
    obj.status = "VOID"
    obj.voided_at = datetime.utcnow()
    db.commit()
    db.refresh(obj)
    return obj


@router.get("/bottles", response_model=PageOut[BottleOut])
def list_bottles(
    zone: str | None = None,
    process_status: str | None = None,
    returned_to_customer: bool | None = None,
    body_code: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    coded_by: int | None = None,
    coding_customer_id: int | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require_any("coding:view", "inspection:view")),
    db: Session = Depends(get_db),
):
    stmt = select(Bottle).order_by(Bottle.id.desc())
    if returned_to_customer is not None:
        stmt = stmt.where(Bottle.returned_to_customer.is_(returned_to_customer))
    if zone:
        stmt = stmt.where(Bottle.zone == zone)
    if process_status:
        stmt = stmt.where(Bottle.process_status == process_status)
    if body_code:
        stmt = stmt.where(Bottle.body_code.ilike(f"%{body_code.strip()}%"))
    if date_from:
        stmt = stmt.where(Bottle.coding_date >= date_from)
    if date_to:
        stmt = stmt.where(Bottle.coding_date <= date_to)
    if coding_customer_id:
        stmt = stmt.where(func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == coding_customer_id)
    if coded_by:
        stmt = stmt.where(Bottle.coded_by == coded_by)
    return page_result(db, stmt, page, page_size)


@router.get("/bottles/lifecycle")
def bottle_lifecycle(
    body_code: str = Query(min_length=1),
    _: User = Depends(require("bottle_lifecycle:view")),
    db: Session = Depends(get_db),
):
    code = body_code.strip()
    if not code:
        raise HTTPException(status_code=400, detail="请输入瓶身编号")
    bottle = db.scalar(select(Bottle).where(func.lower(Bottle.body_code) == code.lower()))
    if not bottle:
        raise not_found("瓶身编号")
    source = db.get(BottleSource, bottle.source_id)
    spec = db.get(BottleSpec, bottle.spec_id)
    coder = db.get(User, bottle.coded_by)
    inspector = db.get(User, bottle.inspected_by) if bottle.inspected_by else None
    inbound = db.get(BottleInbound, bottle.inbound_id) if bottle.inbound_id else None
    outbound = db.get(BottleOutbound, bottle.outbound_id) if bottle.outbound_id else None
    inspection_lines = db.scalars(
        select(InspectionDocumentLine).where(InspectionDocumentLine.bottle_id == bottle.id)
        .order_by(InspectionDocumentLine.id)
    ).all()
    unverified = db.scalars(select(UnverifiedBottle).where(UnverifiedBottle.bottle_id == bottle.id)
        .order_by(UnverifiedBottle.id)).all()
    corrections = db.scalars(select(BottleOutboundCorrection).where(or_(
        BottleOutboundCorrection.original_bottle_id == bottle.id,
        BottleOutboundCorrection.replacement_bottle_id == bottle.id,
    )).order_by(BottleOutboundCorrection.id.desc())).all()
    correction_option = None
    if outbound and outbound.status == "POSTED" and bottle.zone == "OUTBOUND" and bottle.outbound_origin_zone == "GOOD":
        anonymous = quick_outbound_stock(db).get(str(bottle.spec_id), {}).get("anonymous_good", 0)
        candidate = db.scalar(oldest_verified_good_bottle(db, bottle.spec_id, outbound.source_id, bottle.id).limit(1))
        correction_option = {
            "anonymous_good": anonymous,
            "oldest_replacement_body_code": candidate.body_code if candidate else None,
            "oldest_replacement_coding_date": candidate.coding_date if candidate else None,
        }
    inspections = []
    for line in inspection_lines:
        document = line.document
        inspections.append({
            "document_no": document.document_no, "document_status": document.status,
            "inspection_date": document.inspection_date, "customer_name": document.customer_name,
            "inspector_name": line.inspector_name or document.inspector_name,
            "sequence": line.sequence, "result": line.inspection_result,
            "inspection_time": line.inspection_time, "inspected_at": line.inspected_at,
            "next_inspection_date": line.next_inspection_date,
            "manufacturer": line.manufacturer, "manufacture_date": line.manufacture_date,
            "volume": line.volume, "is_archived": document.is_archived,
            "is_retry": document.is_retry, "source_line_id": line.source_line_id,
        })
    return {
        "bottle": {
            "id": bottle.id, "body_code": bottle.body_code,
            "manufacture_date": bottle.manufacture_date, "manufacturer": bottle.manufacturer,
            "source_name": source.name if source else None,
            "spec_name": spec.name if spec else None, "inspection_price": bottle.inspection_price,
            "zone": bottle.zone, "process_status": bottle.process_status,
            "inspection_result": bottle.inspection_result,
            "coding_date": bottle.coding_date, "coding_sequence": bottle.coding_sequence,
            "coding_note": bottle.coding_note, "coded_by_name": coder.name if coder else None,
            "coded_at": bottle.coded_at, "good_entry_date": bottle.good_entry_date,
            "inspected_by_name": inspector.name if inspector else None,
            "inspected_at": bottle.inspected_at, "inspection_note": bottle.inspection_note,
            "outbound_origin_zone": bottle.outbound_origin_zone,
            "created_at": bottle.created_at, "updated_at": bottle.updated_at,
        },
        "inbound": ({
            "document_no": inbound.document_no, "business_date": inbound.business_date,
            "status": inbound.status, "handler": inbound.handler, "note": inbound.note,
            "posted_at": inbound.posted_at, "created_at": inbound.created_at,
        } if inbound else None),
        "inspections": inspections,
        "unverified": [{
            "document_no": item.document_no, "sequence": item.sequence,
            "status": item.status, "error_message": item.error_message,
            "created_at": item.created_at,
        } for item in unverified],
        "outbound": ({
            "id": outbound.id, "document_no": outbound.document_no, "business_date": outbound.business_date,
            "status": outbound.status, "customer": outbound.customer,
            "handler": outbound.handler, "note": outbound.note,
            "posted_at": outbound.posted_at, "created_at": outbound.created_at,
            "origin_zone": bottle.outbound_origin_zone,
        } if outbound else None),
        "correction_option": correction_option,
        "corrections": [{
            "outbound_document_no": item.outbound_document_no,
            "original_body_code": item.original_body_code,
            "replacement_body_code": item.replacement_body_code,
            "replacement_type": item.replacement_type,
            "created_by_name": item.creator.name,
            "created_at": item.created_at,
        } for item in corrections],
    }


@router.post("/bottles/lifecycle/restore-good")
def restore_outbound_bottle_to_good(
    payload: BottleOutboundCorrectionCreate,
    user: User = Depends(require("bottle_lifecycle:restore")),
    db: Session = Depends(get_db),
):
    outbound = db.scalar(select(BottleOutbound).where(BottleOutbound.id == payload.outbound_id).with_for_update())
    bottle = db.scalar(select(Bottle).where(func.lower(Bottle.body_code) == payload.body_code.strip().lower()).with_for_update())
    if not bottle:
        raise not_found("瓶身编号")
    if not outbound or outbound.status != "POSTED" or bottle.outbound_id != outbound.id or bottle.zone != "OUTBOUND":
        raise HTTPException(status_code=409, detail="出库关联已变化，请刷新后重试")
    if bottle.outbound_origin_zone != "GOOD":
        raise HTTPException(status_code=400, detail="只有从良品库出库的瓶子可以恢复为良品")
    line = db.scalar(select(BottleOutboundLine).where(
        BottleOutboundLine.outbound_id == outbound.id,
        BottleOutboundLine.spec_id == bottle.spec_id,
    ).with_for_update())
    if not line or line.good_quantity <= 0 or line.good_adjustment_quantity >= line.good_quantity:
        raise HTTPException(status_code=409, detail="出库单良品明细与钢瓶关联不一致，请先核对单据")
    lock_bottle_stock(db, [bottle.spec_id])
    anonymous = quick_outbound_stock(db).get(str(bottle.spec_id), {}).get("anonymous_good", 0)
    replacement = None
    if anonymous >= 1:
        line.good_adjustment_quantity += 1
        replacement_type = "ANONYMOUS"
    else:
        replacement = db.scalar(oldest_verified_good_bottle(
            db, bottle.spec_id, outbound.source_id, bottle.id
        ).with_for_update().limit(1))
        if not replacement:
            raise HTTPException(status_code=400, detail="没有可扣减的库存数量，也没有同供瓶单位、同规格且有有效检验单的良品瓶可替换")
        replacement.zone = "OUTBOUND"
        replacement.outbound_id = outbound.id
        replacement.outbound_origin_zone = "GOOD"
        replacement_type = "BOTTLE"
    bottle.zone = "GOOD"
    bottle.outbound_id = None
    bottle.outbound_origin_zone = None
    bottle.good_entry_date = bottle.good_entry_date or bottle.coding_date or china_today()
    db.add(BottleOutboundCorrection(
        outbound_id=outbound.id, outbound_document_no=outbound.document_no,
        original_bottle_id=bottle.id, original_body_code=bottle.body_code,
        replacement_bottle_id=replacement.id if replacement else None,
        replacement_body_code=replacement.body_code if replacement else None,
        replacement_type=replacement_type, created_by=user.id,
    ))
    db.commit()
    return {
        "body_code": bottle.body_code, "zone": bottle.zone,
        "outbound_document_no": outbound.document_no,
        "replacement_type": replacement_type,
        "replacement_body_code": replacement.body_code if replacement else None,
    }


@router.get("/coding/summary")
def coding_summary(
    date_from: date = Query(default_factory=china_today),
    date_to: date = Query(default_factory=china_today),
    coded_by: int | None = None,
    body_code: str | None = None,
    coding_customer_id: int | None = None,
    _: User = Depends(require_any("coding:view", "quick_entry:view")),
    db: Session = Depends(get_db),
):
    filters = [Bottle.coding_date >= date_from, Bottle.coding_date <= date_to]
    if coding_customer_id:
        filters.append(func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == coding_customer_id)
    if coded_by:
        filters.append(Bottle.coded_by == coded_by)
    if body_code:
        filters.append(Bottle.body_code.ilike(f"%{body_code.strip()}%"))
    rows = db.execute(
        select(User.id, User.name, BottleSpec.id, BottleSpec.name, func.count(Bottle.id), BottleSource.id, BottleSource.name)
        .join(User, Bottle.coded_by == User.id)
        .join(BottleSpec, Bottle.spec_id == BottleSpec.id)
        .join(BottleSource, func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == BottleSource.id)
        .where(*filters)
        .group_by(User.id, User.name, BottleSpec.id, BottleSpec.name, BottleSource.id, BottleSource.name)
        .order_by(User.name, BottleSpec.sort_order, BottleSpec.id)
    ).all()
    items = [{"user_id": row[0], "user_name": row[1], "spec_id": row[2], "spec_name": row[3], "quantity": row[4], "customer_id": row[5], "customer_name": row[6]} for row in rows]
    return {"date_from": date_from, "date_to": date_to, "total": sum(item["quantity"] for item in items), "items": items}


def coding_suspect_condition():
    return or_(
        func.length(Bottle.body_code) > 13,
        func.coalesce(func.trim(Bottle.manufacturer), "") == "",
        Bottle.manufacture_date.is_(None),
    )


@router.get("/coding/review-bottles", response_model=CodingReviewPageOut)
def coding_review_bottles(
    date_from: date = Query(default_factory=china_today),
    date_to: date = Query(default_factory=china_today),
    body_code: str | None = None,
    coding_customer_id: int | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require("coding:view")),
    db: Session = Depends(get_db),
):
    if date_from > date_to:
        raise HTTPException(status_code=422, detail="开始日期不能晚于结束日期")
    filters = [Bottle.coding_date >= date_from, Bottle.coding_date <= date_to]
    if body_code and body_code.strip():
        filters.append(Bottle.body_code.ilike(f"%{body_code.strip()}%"))
    if coding_customer_id:
        filters.append(func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == coding_customer_id)
    suspect = coding_suspect_condition()
    stmt = select(Bottle).where(*filters).order_by(
        case((suspect, 0), else_=1), Bottle.coding_date, Bottle.coding_sequence, Bottle.id,
    )
    result = page_result(db, stmt, page, page_size)
    result["suspect_count"] = db.scalar(select(func.count(Bottle.id)).where(*filters, suspect)) or 0
    return result


@router.get("/coding/today-bottles", response_model=list[BottleOut])
def coding_today_bottles(
    _: User = Depends(require("coding:view")),
    db: Session = Depends(get_db),
):
    target_date = china_today()
    suspect = coding_suspect_condition()
    return db.scalars(
        select(Bottle)
        .where(Bottle.coding_date == target_date)
        .order_by(case((suspect, 0), else_=1), Bottle.coding_sequence, Bottle.id)
    ).all()


@router.get("/coding/inbounds", response_model=PageOut[BottleInboundOut])
def coding_inbounds(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("coding:view")), db: Session = Depends(get_db)):
    return page_result(db, select(BottleInbound).where(BottleInbound.status == "POSTED").order_by(BottleInbound.business_date.desc(), BottleInbound.id.desc()), page, page_size)


def remaining_for_line(db: Session, line: BottleInboundLine) -> int:
    coded = db.scalar(select(func.count()).select_from(Bottle).where(
        Bottle.inbound_id == line.inbound_id,
        Bottle.spec_id == line.spec_id,
    )) or 0
    return max(line.pending_quantity - coded, 0)


def next_coding_sequence(db: Session, coding_date: date) -> int:
    if db.bind and db.bind.dialect.name == "postgresql":
        db.execute(text("SELECT pg_advisory_xact_lock(902610)"))
    return (db.scalar(select(func.max(Bottle.coding_sequence)).where(Bottle.coding_date == coding_date)) or 0) + 1


@router.get("/coding/pool")
def coding_pool(
    date_from: date | None = None,
    date_to: date | None = None,
    daily_page: int = Query(1, ge=1),
    coding_page: int = Query(1, ge=1),
    batch_page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    include_all_batches: bool = False,
    _: User = Depends(require_any("coding_pool:view", "quick_entry:view")), db: Session = Depends(get_db),
):
    date_to = date_to or china_today()
    date_from = date_from or (date_to - timedelta(days=9))
    if date_from > date_to:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    inbounds = db.scalars(select(BottleInbound).where(BottleInbound.status == "POSTED").order_by(BottleInbound.business_date, BottleInbound.id)).all()
    batches = []
    by_date: dict[str, dict] = {}
    total = 0
    for inbound in inbounds:
        spec_rows = []
        batch_remaining = 0
        for line in inbound.lines:
            remaining = remaining_for_line(db, line)
            if line.pending_quantity:
                spec_rows.append({"spec_id": line.spec_id, "spec_name": line.spec_name, "pending_quantity": line.pending_quantity, "remaining": remaining})
            batch_remaining += remaining
        if inbound.pending_quantity and (include_all_batches or date_from <= inbound.business_date <= date_to):
            batches.append({
                "id": inbound.id, "document_no": inbound.document_no, "business_date": inbound.business_date,
                "source_id": inbound.source_id, "source_name": inbound.source.name,
                "pending_quantity": inbound.pending_quantity, "remaining": batch_remaining, "specs": spec_rows,
            })
        if date_from <= inbound.business_date <= date_to:
            key = inbound.business_date.isoformat()
            row = by_date.setdefault(key, {"business_date": inbound.business_date, "pending_quantity": 0, "coded_quantity": 0, "remaining": 0})
            row["pending_quantity"] += inbound.pending_quantity
            row["remaining"] += batch_remaining
            row["coded_quantity"] = row["pending_quantity"] - row["remaining"]
        total += batch_remaining
    coding_daily = [
        {"coding_date": row[0], "coded_by": row[1], "coded_by_name": row[2], "coded_quantity": row[3]}
        for row in db.execute(
            select(Bottle.coding_date, Bottle.coded_by, User.name, func.count(Bottle.id))
            .join(User, User.id == Bottle.coded_by)
            .where(Bottle.coding_date >= date_from, Bottle.coding_date <= date_to)
            .group_by(Bottle.coding_date, Bottle.coded_by, User.name)
            .order_by(Bottle.coding_date.desc(), User.name)
        ).all()
    ]
    def section(items: list, page: int):
        return {"items": items[(page - 1) * page_size:page * page_size], "total": len(items), "page": page, "page_size": page_size}
    return {
        "date_from": date_from, "date_to": date_to, "total_remaining": total,
        "daily": section(sorted(by_date.values(), key=lambda item: item["business_date"], reverse=True), daily_page),
        "coding_daily": section(coding_daily, coding_page),
        "batches": section(list(reversed(batches)), batch_page),
    }


@router.get("/coding/good-bottles", response_model=list[GoodBottleExportOut])
def coding_good_bottles(date: date = Query(...), _: User = Depends(require_any("quick_entry:view", "inspection:view")), db: Session = Depends(get_db)):
    """Bottles that entered the good-stock zone on the given business date (Asia/Shanghai)."""
    rows = db.execute(
        select(Bottle, BottleSpec.volume)
        .join(BottleSpec, Bottle.spec_id == BottleSpec.id)
        .where(Bottle.zone == "GOOD", Bottle.good_entry_date == date)
        .order_by(Bottle.manufacture_date.is_(None), Bottle.manufacture_date.desc(), Bottle.id)
    ).all()
    return [{"id": bottle.id, "body_code": bottle.body_code, "manufacture_date": bottle.manufacture_date, "volume": volume or Decimal(0), "manufacturer": bottle.manufacturer} for bottle, volume in rows]


def inspection_document_candidates(db: Session, inspection_date: date | list[date], source_id: int | None = None):
    stmt = (
        select(Bottle, BottleSpec.volume, BottleSource.name)
        .join(BottleSpec, Bottle.spec_id == BottleSpec.id)
        .join(BottleSource, func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == BottleSource.id)
        .where(Bottle.zone == "GOOD", Bottle.good_entry_date.in_(inspection_date if isinstance(inspection_date, list) else [inspection_date]))
        .order_by(
            Bottle.manufacture_date.is_(None),
            Bottle.manufacture_date.desc(),
            Bottle.inspected_at.is_(None),
            Bottle.inspected_at,
            Bottle.id,
        )
        .with_for_update()
    )
    if source_id is not None:
        stmt = stmt.where(func.coalesce(Bottle.coding_customer_id, Bottle.source_id) == source_id)
    return db.execute(stmt).all()


def resequence_inspection_document(document: InspectionDocument) -> None:
    ordered = sorted(document.lines, key=lambda line: (
        line.manufacture_date is None,
        -line.manufacture_date.toordinal() if line.manufacture_date else 0,
        line.inspected_at is None,
        line.inspected_at or datetime.max,
        line.sequence,
        line.id or 0,
    ))
    for sequence, line in enumerate(ordered, start=1):
        line.sequence = sequence
        if line.unverified_record:
            line.unverified_record.sequence = sequence
            line.unverified_record.body_code = line.body_code
            line.unverified_record.manufacture_date = line.manufacture_date
            line.unverified_record.manufacturer = line.manufacturer
    document.lines[:] = ordered


def add_inspection_document_lines(document: InspectionDocument, rows) -> int:
    existing_ids = {line.bottle_id for line in document.lines}
    added = 0
    for bottle, volume, _ in rows:
        if bottle.id in existing_ids:
            continue
        document.lines.append(InspectionDocumentLine(
            bottle_id=bottle.id,
            spec_id=bottle.spec_id,
            sequence=len(document.lines) + 1,
            body_code=bottle.body_code,
            manufacture_date=bottle.manufacture_date,
            manufacturer=bottle.manufacturer or "",
            volume=volume or Decimal(0),
            inspected_at=bottle.inspected_at,
            inspection_result=bottle.inspection_result or "GOOD",
        ))
        added += 1
    resequence_inspection_document(document)
    return added


def report_spec_for_inspection_line(
    db: Session,
    line: InspectionDocumentLine,
    specs_by_volume: dict[Decimal, list[BottleSpec]] | None = None,
) -> BottleSpec:
    if line.spec_id:
        spec = db.get(BottleSpec, line.spec_id)
        if spec:
            return spec
    if line.bottle:
        return line.bottle.spec
    matches = specs_by_volume.get(line.volume, []) if specs_by_volume is not None else db.scalars(select(BottleSpec).where(BottleSpec.volume == line.volume)).all()
    if len(matches) == 1:
        return matches[0]
    raise HTTPException(status_code=400, detail=f"{line.body_code} 缺少明确的钢瓶规格，请先编辑该检验单明细后再导出单瓶报告")


def regular_inspection_documents(db: Session, active_only: bool = True) -> list[InspectionDocument]:
    stmt = select(InspectionDocument).options(lazyload(InspectionDocument.lines)).where(InspectionDocument.is_retry.is_(False))
    if active_only:
        stmt = stmt.where(InspectionDocument.status == "ACTIVE")
    return list(db.scalars(stmt.order_by(InspectionDocument.id)).all())


def claimed_inspection_bottle_ids(db: Session) -> set[int]:
    return set(db.scalars(select(InspectionDocumentLine.bottle_id).join(InspectionDocument).where(
        InspectionDocument.status == "ACTIVE", InspectionDocument.is_retry.is_(False),
        InspectionDocumentLine.bottle_id.is_not(None),
    )).all())


@router.get("/inspection-documents/generation-dates")
def inspection_generation_dates(
    month: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    _: User = Depends(require("inspection_documents:create")), db: Session = Depends(get_db),
):
    try:
        first_day = date.fromisoformat(f"{month}-01")
    except ValueError:
        raise HTTPException(status_code=422, detail="月份格式无效")
    next_month = (first_day.replace(day=28) + timedelta(days=4)).replace(day=1)
    generated_dates = {day for document in regular_inspection_documents(db) for day in document.source_dates}
    claimed_ids = claimed_inspection_bottle_ids(db)
    rows = db.execute(select(Bottle.good_entry_date, func.coalesce(Bottle.coding_customer_id, Bottle.source_id), Bottle.id).where(
        Bottle.zone == "GOOD", Bottle.good_entry_date >= first_day, Bottle.good_entry_date < next_month,
    )).all()
    by_date: dict[date, list] = {}
    for day, source_id, bottle_id in rows:
        if bottle_id not in claimed_ids:
            by_date.setdefault(day, []).append(source_id)
    days = []
    day = first_day
    while day < next_month:
        sources = by_date.get(day, [])
        generated = day in generated_dates
        days.append({"date": day, "quantity": len(sources), "customer_count": len(set(sources)),
                     "generated": generated, "selectable": bool(sources) and not generated})
        day += timedelta(days=1)
    return days


@router.get("/inspection-documents", response_model=PageOut[InspectionDocumentSummaryOut])
def list_inspection_documents(
    date_from: date | None = None,
    date_to: date | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require("inspection_documents:view")),
    db: Session = Depends(get_db),
):
    stmt = select(InspectionDocument).order_by(InspectionDocument.inspection_date.desc(), InspectionDocument.id.desc())
    if date_from:
        stmt = stmt.where(InspectionDocument.inspection_date >= date_from)
    if date_to:
        stmt = stmt.where(InspectionDocument.inspection_date <= date_to)
    return page_result(db, stmt, page, page_size)


@router.get("/inspection-documents/{object_id}", response_model=InspectionDocumentOut)
def get_inspection_document(object_id: int, _: User = Depends(require("inspection_documents:view")), db: Session = Depends(get_db)):
    lines = selectinload(InspectionDocument.lines)
    document = db.scalar(select(InspectionDocument).options(
        lines.selectinload(InspectionDocumentLine.bottle),
        lines.selectinload(InspectionDocumentLine.source_line)
            .selectinload(InspectionDocumentLine.bottle),
    ).where(InspectionDocument.id == object_id))
    if not document:
        raise not_found("检验单")
    return document


def remember_inspection_export(
    db: Session, document: InspectionDocument, kind: str, selected_date: date,
    lines: list[InspectionDocumentLine],
) -> None:
    record = db.scalar(select(InspectionDocumentExportRecord).where(
        InspectionDocumentExportRecord.document_id == document.id,
        InspectionDocumentExportRecord.kind == kind,
    ).with_for_update())
    if record is None:
        db.add(InspectionDocumentExportRecord(
            document_id=document.id, kind=kind, export_date=selected_date,
        ))
    else:
        record.export_date = selected_date
    for line in lines:
        line.inspection_time = selected_date
    document.updated_at = datetime.utcnow()
    db.commit()


@router.post("/inspection-documents/{object_id}/export-xls")
def export_inspection_document_xls(
    object_id: int,
    kind: str = Query(..., pattern="^(full|verified|unverified)$"),
    export_date: date | None = None,
    _: User = Depends(require("inspection_documents:export")),
    db: Session = Depends(get_db),
):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能导出")
    if not document.customer_name.strip():
        raise HTTPException(status_code=400, detail="请先统一设置客户/充装单位")
    if not document.inspector_name.strip():
        raise HTTPException(status_code=400, detail="请先统一设置检验人员")
    if kind == "unverified" and not any(line.is_unverified for line in document.lines):
        raise HTTPException(status_code=400, detail="当前检验单没有标黄的未经系统验证钢瓶")
    matching_lines = [line for line in document.lines if kind == "full" or line.is_unverified == (kind == "unverified")]
    effective_export_date = export_date or document.export_dates.get(f"xls:{kind}") or next(
        (line.inspection_time for line in matching_lines if line.inspection_time), china_today(),
    )
    try:
        content, quantity = render_inspection_xls(document, kind, effective_export_date)
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=500, detail=f"XLS 检验单生成失败：{exc}") from exc
    affected_lines = [line for line in document.lines if line.is_unverified] if kind == "unverified" else list(document.lines)
    remember_inspection_export(db, document, f"xls:{kind}", effective_export_date, affected_lines)
    safe_customer = re.sub(r'[\\/:*?"<>|]', "_", document.customer_name.strip())
    prefix, category = {
        "full": ("JY-", "全量检测单"),
        "verified": ("JY-", "成功上传"),
        "unverified": ("JY-", "未经系统验证"),
    }[kind]
    filename = f"{prefix}{safe_customer}{effective_export_date:%Y%m%d}{category}[{quantity}].xls"
    return Response(
        content=content,
        media_type="application/vnd.ms-excel",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/inspection-documents/{object_id}/export-receipt-word")
def export_inspection_document_receipt_word(
    object_id: int,
    _: User = Depends(require("inspection_documents:export")), db: Session = Depends(get_db),
):
    document = db.get(InspectionDocument, object_id)
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能导出")
    if not document.customer_name.strip():
        raise HTTPException(status_code=400, detail="请先统一设置客户/充装单位")
    if not document.inspector_name.strip():
        raise HTTPException(status_code=400, detail="请先统一设置检验人员")
    source = db.scalar(select(BottleSource).where(BottleSource.name == document.customer_name.strip()))
    try:
        content = render_inspection_receipt_word(
            document, contact=source.contact if source else "", phone=source.phone if source else "",
        )
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=500, detail=f"来瓶登记 Word 生成失败：{exc}") from exc
    safe_customer = re.sub(r'[\\/:*?"<>|]', "_", document.customer_name.strip())
    filename = f"{safe_customer}{document.inspection_date:%Y%m%d}来瓶登记[{len(document.lines)}].docx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def lock_inspection_word_counter(db: Session, report_date: date) -> InspectionWordDailyCounter:
    db.execute(text(
        "INSERT INTO inspection_word_daily_counters (report_date, last_number) "
        "VALUES (:report_date, 0) ON CONFLICT (report_date) DO NOTHING"
    ), {"report_date": report_date})
    counter = db.scalar(select(InspectionWordDailyCounter).where(
        InspectionWordDailyCounter.report_date == report_date,
    ).with_for_update())
    if counter is None:
        raise HTTPException(status_code=500, detail="报告序号分配失败")
    return counter


def inspection_word_code(report_date: date, number: int) -> str:
    return f"WFJA-{report_date:%y}-{report_date:%m%d}{number:03d}"


def inspection_word_assignment_out(assignment: InspectionWordExportAssignment) -> dict:
    return {
        "id": assignment.id,
        "report_date": assignment.report_date,
        "first_number": assignment.first_number,
        "regular_code": inspection_word_code(assignment.report_date, assignment.first_number),
        "assessment_code": inspection_word_code(assignment.report_date, assignment.first_number + 1),
    }


@router.get("/inspection-word-code-ledger")
def get_inspection_word_code_ledger(
    user: User = Depends(current_user), db: Session = Depends(get_db),
):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="仅超级管理员可管理 Word 报告编号")
    counter_by_date = dict(db.execute(select(
        InspectionWordDailyCounter.report_date, InspectionWordDailyCounter.last_number,
    )).all())
    released_by_date: dict[date, list[int]] = {}
    for released_date, first_number in db.execute(select(
        InspectionWordReleasedPair.report_date, InspectionWordReleasedPair.first_number,
    )).all():
        released_by_date.setdefault(released_date, []).append(first_number)
    assignment_rows = db.execute(select(
        InspectionWordExportAssignment.id,
        InspectionWordExportAssignment.document_id,
        InspectionWordExportAssignment.report_date,
        InspectionWordExportAssignment.first_number,
        InspectionDocument.document_no,
        InspectionDocument.inspection_date,
        InspectionDocument.customer_name,
        InspectionDocument.status,
        InspectionDocumentExportRecord.export_date.label("default_report_date"),
    ).join(InspectionDocument, InspectionDocument.id == InspectionWordExportAssignment.document_id).outerjoin(
        InspectionDocumentExportRecord,
        (InspectionDocumentExportRecord.document_id == InspectionDocument.id)
        & (InspectionDocumentExportRecord.kind == "word:all"),
    ).order_by(InspectionWordExportAssignment.report_date.desc(), InspectionWordExportAssignment.first_number)).all()
    assigned_by_date: dict[date, list] = {}
    for item in assignment_rows:
        assigned_by_date.setdefault(item.report_date, []).append(item)
    days = []
    for report_date in sorted(counter_by_date.keys() | assigned_by_date.keys(), reverse=True):
        day_assignments = assigned_by_date.get(report_date, [])
        required_number = max((item.first_number + 1 for item in day_assignments), default=0)
        last_number = counter_by_date.get(report_date, 0)
        effective_number = max(last_number, required_number)
        assigned_numbers = {
            number
            for item in day_assignments
            for number in (item.first_number, item.first_number + 1)
        }
        released_pairs = sorted(released_by_date.get(report_date, []))
        released_numbers = {number for first in released_pairs for number in (first, first + 1)}
        unassigned_numbers = [number for number in range(1, effective_number + 1) if number not in assigned_numbers and number not in released_numbers]
        unassigned_set = set(unassigned_numbers)
        days.append({
            "report_date": report_date,
            "last_number": last_number,
            "required_number": required_number,
            "unassigned_numbers": unassigned_numbers,
            "unassigned_pairs": [first for first in range(1, effective_number, 2) if first in unassigned_set and first + 1 in unassigned_set],
            "released_pairs": released_pairs,
            "next_first_number": released_pairs[0] if released_pairs else effective_number + 1 if effective_number + 2 <= 999 else None,
            "assignments": [{
                "id": item.id,
                "document_id": item.document_id,
                "report_date": item.report_date,
                "first_number": item.first_number,
                "regular_code": inspection_word_code(item.report_date, item.first_number),
                "assessment_code": inspection_word_code(item.report_date, item.first_number + 1),
                "document_no": item.document_no,
                "inspection_date": item.inspection_date,
                "customer_name": item.customer_name,
                "document_status": item.status,
                "default_report_date": item.default_report_date,
            } for item in day_assignments],
        })
    assigned_keys = {(item.document_id, item.report_date) for item in assignment_rows}
    legacy_records = db.execute(select(
        InspectionDocumentExportRecord.document_id,
        InspectionDocumentExportRecord.export_date,
        InspectionDocument.document_no,
        InspectionDocument.inspection_date,
        InspectionDocument.customer_name,
    ).join(InspectionDocument, InspectionDocument.id == InspectionDocumentExportRecord.document_id).where(
        InspectionDocumentExportRecord.kind == "word:all",
    ).order_by(
        InspectionDocumentExportRecord.export_date.desc(), InspectionDocumentExportRecord.id.desc(),
    )).all()
    return {"days": days, "without_saved_code": [{
        "document_id": item.document_id,
        "document_no": item.document_no,
        "inspection_date": item.inspection_date,
        "customer_name": item.customer_name,
        "report_date": item.export_date,
    } for item in legacy_records if (item.document_id, item.export_date) not in assigned_keys]}


@router.delete("/inspection-word-code-ledger/dates/{report_date}/unassigned-pairs/{first_number}", status_code=204)
def release_unassigned_inspection_word_pair(
    report_date: date, first_number: int, user: User = Depends(current_user), db: Session = Depends(get_db),
):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="仅超级管理员可管理 Word 报告编号")
    if first_number < 1 or first_number % 2 != 1 or first_number + 1 > 999:
        raise HTTPException(status_code=400, detail="只能释放完整的 001/002、003/004 等编号组")
    counter = lock_inspection_word_counter(db, report_date)
    highest_first = db.scalar(select(func.max(InspectionWordExportAssignment.first_number)).where(
        InspectionWordExportAssignment.report_date == report_date,
    ))
    effective_number = max(counter.last_number, highest_first + 1 if highest_first is not None else 0)
    if first_number + 1 > effective_number:
        raise HTTPException(status_code=409, detail="这组编号不在当前日期计数范围内，请刷新页面")
    existing = db.scalar(select(InspectionWordExportAssignment.id).where(
        InspectionWordExportAssignment.report_date == report_date,
        InspectionWordExportAssignment.first_number <= first_number + 1,
        InspectionWordExportAssignment.first_number + 1 >= first_number,
    ))
    if existing:
        raise HTTPException(status_code=409, detail="这组编号已有检验单分配，不能作为未关联编号释放")
    if db.get(InspectionWordReleasedPair, (report_date, first_number)):
        raise HTTPException(status_code=409, detail="这组编号已释放，请刷新页面")
    db.add(InspectionWordReleasedPair(report_date=report_date, first_number=first_number))
    db.commit()
    return Response(status_code=204)


@router.delete("/inspection-word-code-ledger/assignments/{assignment_id}", status_code=204)
def delete_inspection_word_assignment(
    assignment_id: int, user: User = Depends(current_user), db: Session = Depends(get_db),
):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="仅超级管理员可管理 Word 报告编号")
    assignment = db.get(InspectionWordExportAssignment, assignment_id)
    if not assignment:
        raise not_found("Word 报告编号")
    lock_inspection_word_counter(db, assignment.report_date)
    db.delete(assignment)
    db.commit()
    return Response(status_code=204)


@router.post("/inspection-word-code-ledger/dates/{report_date}/reset")
def reset_inspection_word_counter(
    report_date: date, user: User = Depends(current_user), db: Session = Depends(get_db),
):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="仅超级管理员可管理 Word 报告编号")
    counter = lock_inspection_word_counter(db, report_date)
    highest_first = db.scalar(select(func.max(InspectionWordExportAssignment.first_number)).where(
        InspectionWordExportAssignment.report_date == report_date,
    ))
    counter.last_number = highest_first + 1 if highest_first is not None else 0
    for released in db.scalars(select(InspectionWordReleasedPair).where(
        InspectionWordReleasedPair.report_date == report_date,
        InspectionWordReleasedPair.first_number > counter.last_number,
    )).all():
        db.delete(released)
    db.commit()
    return {
        "report_date": report_date,
        "last_number": counter.last_number,
        "next_first_number": counter.last_number + 1 if counter.last_number + 2 <= 999 else None,
    }


@router.get("/inspection-documents/{object_id}/word-codes")
def get_inspection_word_codes(
    object_id: int, _: User = Depends(require("inspection_documents:view")), db: Session = Depends(get_db),
):
    if not db.get(InspectionDocument, object_id):
        raise not_found("检验单")
    assignments = db.scalars(select(InspectionWordExportAssignment).where(
        InspectionWordExportAssignment.document_id == object_id,
    ).order_by(InspectionWordExportAssignment.report_date.desc(), InspectionWordExportAssignment.id.desc())).all()
    return [inspection_word_assignment_out(item) for item in assignments]


@router.get("/inspection-documents/{object_id}/word-codes/{assignment_id}/check")
def check_inspection_word_code(
    object_id: int, assignment_id: int, report_date: date, first_number: int = Query(ge=1, le=998),
    user: User = Depends(current_user), db: Session = Depends(get_db),
):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="仅超级管理员可修改 Word 报告编码")
    assignment = db.scalar(select(InspectionWordExportAssignment).where(
        InspectionWordExportAssignment.id == assignment_id,
        InspectionWordExportAssignment.document_id == object_id,
    ))
    if not assignment:
        raise not_found("Word 报告编码")
    others = db.scalars(select(InspectionWordExportAssignment).where(
        InspectionWordExportAssignment.report_date == report_date,
        InspectionWordExportAssignment.id != assignment_id,
    )).all()
    conflicting = [item for item in others if item.first_number <= first_number + 1 and item.first_number + 1 >= first_number]
    same_document = any(item.document_id == object_id for item in others)
    return {
        "existing_count": len(others),
        "conflict": bool(conflicting or same_document),
        "conflict_reason": "本检验单在目标日期已有另一组编码" if same_document else "目标编号已被其他检验单使用" if conflicting else None,
        "regular_code": inspection_word_code(report_date, first_number),
        "assessment_code": inspection_word_code(report_date, first_number + 1),
    }


@router.put("/inspection-documents/{object_id}/word-codes/{assignment_id}")
def update_inspection_word_code(
    object_id: int, assignment_id: int, payload: InspectionWordCodeUpdate,
    user: User = Depends(current_user), db: Session = Depends(get_db),
):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="仅超级管理员可修改 Word 报告编码")
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能修改报告编码")
    counter = lock_inspection_word_counter(db, payload.report_date)
    assignment = db.scalar(select(InspectionWordExportAssignment).where(
        InspectionWordExportAssignment.id == assignment_id,
        InspectionWordExportAssignment.document_id == object_id,
    ).with_for_update())
    if not assignment:
        raise not_found("Word 报告编码")
    others = db.scalars(select(InspectionWordExportAssignment).where(
        InspectionWordExportAssignment.report_date == payload.report_date,
        InspectionWordExportAssignment.id != assignment_id,
    )).all()
    if any(item.document_id == object_id for item in others):
        raise HTTPException(status_code=409, detail="本检验单在目标日期已有另一组编码")
    if any(item.first_number <= payload.first_number + 1 and item.first_number + 1 >= payload.first_number for item in others):
        raise HTTPException(status_code=409, detail="目标编号已被其他检验单使用，请换一个序号")
    assignment.report_date = payload.report_date
    assignment.first_number = payload.first_number
    for released in db.scalars(select(InspectionWordReleasedPair).where(
        InspectionWordReleasedPair.report_date == payload.report_date,
        InspectionWordReleasedPair.first_number <= payload.first_number + 1,
        InspectionWordReleasedPair.first_number + 1 >= payload.first_number,
    )).all():
        db.delete(released)
    counter.last_number = max(counter.last_number, payload.first_number + 1)
    record = db.scalar(select(InspectionDocumentExportRecord).where(
        InspectionDocumentExportRecord.document_id == object_id,
        InspectionDocumentExportRecord.kind == "word:all",
    ).with_for_update())
    if record is None:
        db.add(InspectionDocumentExportRecord(document_id=object_id, kind="word:all", export_date=payload.report_date))
    else:
        record.export_date = payload.report_date
    document.updated_at = datetime.utcnow()
    save(db, "目标日期或报告编码已被占用，请刷新后重试")
    return inspection_word_assignment_out(assignment)


def inspection_word_numbers_for_document(
    db: Session, document_id: int, report_date: date, *, allocate: bool,
) -> tuple[int, int] | None:
    counter = lock_inspection_word_counter(db, report_date) if allocate else None
    assignment = db.scalar(select(InspectionWordExportAssignment).where(
        InspectionWordExportAssignment.document_id == document_id,
        InspectionWordExportAssignment.report_date == report_date,
    ).with_for_update())
    if assignment:
        return assignment.first_number, assignment.first_number + 1
    if not allocate:
        return None
    highest_first = db.scalar(select(func.max(InspectionWordExportAssignment.first_number)).where(
        InspectionWordExportAssignment.report_date == report_date,
    ))
    if highest_first is not None:
        counter.last_number = max(counter.last_number, highest_first + 1)
    released = db.scalar(select(InspectionWordReleasedPair).where(
        InspectionWordReleasedPair.report_date == report_date,
    ).order_by(InspectionWordReleasedPair.first_number).with_for_update())
    if released:
        first_number = released.first_number
        conflict = db.scalar(select(InspectionWordExportAssignment.id).where(
            InspectionWordExportAssignment.report_date == report_date,
            InspectionWordExportAssignment.first_number <= first_number + 1,
            InspectionWordExportAssignment.first_number + 1 >= first_number,
        ))
        if conflict:
            raise HTTPException(status_code=409, detail="可重用编号已被占用，请刷新编号管理页面")
        db.delete(released)
    else:
        if counter.last_number + 2 > 999:
            raise HTTPException(status_code=409, detail="该日期的 Word 报告序号已超过 999，请核对报告日期")
        first_number = counter.last_number + 1
        counter.last_number += 2
    db.add(InspectionWordExportAssignment(
        document_id=document_id, report_date=report_date, first_number=first_number,
    ))
    db.flush()
    return first_number, first_number + 1


def inspection_word_template_for_customer(db: Session, customer_name: str) -> str:
    choice = db.scalar(select(InspectionWordTemplateChoice).where(
        InspectionWordTemplateChoice.customer_name == customer_name.strip(),
    ))
    return choice.template_key if choice and choice.template_key in TEMPLATE_LIBRARY else "default"


@router.get("/inspection-word-templates")
def list_inspection_word_templates(
    customer_name: str = "",
    _: User = Depends(require("inspection_documents:export")),
    db: Session = Depends(get_db),
):
    return {
        "templates": [
            {"key": key, "name": entry["name"], "description": entry["description"]}
            for key, entry in TEMPLATE_LIBRARY.items()
        ],
        "selected_key": inspection_word_template_for_customer(db, customer_name) if customer_name else "default",
    }


@router.put("/inspection-word-templates/choice")
def save_inspection_word_template_choice(
    payload: InspectionWordTemplateChoiceUpdate,
    customer_name: str,
    _: User = Depends(require("inspection_documents:export")),
    db: Session = Depends(get_db),
):
    customer_name = customer_name.strip()
    if not customer_name:
        raise HTTPException(status_code=400, detail="请先设置客户/充装单位")
    if payload.template_key not in TEMPLATE_LIBRARY:
        raise HTTPException(status_code=400, detail="未知的 Word 模板")
    choice = db.scalar(select(InspectionWordTemplateChoice).where(
        InspectionWordTemplateChoice.customer_name == customer_name,
    ).with_for_update())
    if choice is None:
        choice = InspectionWordTemplateChoice(customer_name=customer_name, template_key=payload.template_key)
        db.add(choice)
    else:
        choice.template_key = payload.template_key
    db.commit()
    return {"customer_name": customer_name, "template_key": payload.template_key}


@router.post("/inspection-documents/{object_id}/export-word-pair")
def export_inspection_document_word_pair(
    object_id: int,
    report_date: date | None = None,
    _: User = Depends(require("inspection_documents:export")),
    db: Session = Depends(get_db),
):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能导出")
    if not document.customer_name.strip():
        raise HTTPException(status_code=400, detail="请先统一设置客户/充装单位")
    selected_date = report_date or document.export_dates.get("word:all") or next(
        (line.inspection_time for line in document.lines if line.inspection_time), document.inspection_date,
    )
    template_key = inspection_word_template_for_customer(db, document.customer_name)
    regular_number, assessment_number = inspection_word_numbers_for_document(
        db, document.id, selected_date, allocate=True,
    )
    older, within = split_inspection_lines_by_age(document, selected_date)
    safe_customer = re.sub(r'[\\/:*?"<>|]', "_", document.customer_name.strip())
    reports = []
    try:
        for kind, lines, number in (
            ("regular", within if template_key == "default" else older, regular_number),
            ("assessment", older if template_key == "default" else within, assessment_number),
        ):
            label = ("评定" if kind == "regular" else "检验") if template_key == "photo" else ("定检" if kind == "regular" else "评定")
            content = render_inspection_word(
                kind, document.customer_name.strip(), selected_date, len(lines),
                inspection_age_table_rows(document, lines, selected_date), report_number=number,
                template_key=template_key,
            )
            reports.append({
                "kind": kind,
                "filename": f"{safe_customer}{selected_date:%m%d}{label}[{len(lines)}].docx",
                "report_number": number,
                "content_base64": base64.b64encode(content).decode("ascii"),
            })
    except (OSError, ValueError) as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Word 模板生成失败：{exc}") from exc
    remember_inspection_export(db, document, "word:all", selected_date, within + older)
    return {"report_date": selected_date, "reports": reports}


@router.post("/inspection-documents/{object_id}/export-word")
@router.get("/inspection-documents/{object_id}/export-word")
def export_inspection_document_word(
    object_id: int,
    request: Request,
    kind: str = Query(..., pattern="^(regular|assessment)$"),
    report_date: date | None = None,
    _: User = Depends(require("inspection_documents:export")),
    db: Session = Depends(get_db),
):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update()) if request.method == "POST" else db.get(InspectionDocument, object_id)
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能导出")
    if not document.customer_name.strip():
        raise HTTPException(status_code=400, detail="请先统一设置客户/充装单位")

    effective_report_date = report_date or document.export_dates.get("word:all") or next(
        (line.inspection_time for line in document.lines if line.inspection_time), document.inspection_date,
    )
    older, within = split_inspection_lines_by_age(document, effective_report_date)
    template_key = inspection_word_template_for_customer(db, document.customer_name)
    # Default templates use regular for inspection; photo templates use assessment.
    inspection_kind = "regular" if template_key == "default" else "assessment"
    lines = within if kind == inspection_kind else older
    quantity = len(lines)
    label = ("评定" if kind == "regular" else "检验") if template_key == "photo" else ("定检" if kind == "regular" else "评定")
    # Legacy GET is read-only; once assigned, both routes render the same number.
    pair = inspection_word_numbers_for_document(
        db, document.id, effective_report_date, allocate=request.method == "POST",
    )
    report_number = (pair[0] if kind == "regular" else pair[1]) if pair else None
    try:
        content = render_inspection_word(
            kind, document.customer_name.strip(), effective_report_date, quantity,
            inspection_age_table_rows(document, lines, effective_report_date),
            report_number=report_number,
            template_key=template_key,
        )
    except (OSError, ValueError) as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Word 模板生成失败：{exc}") from exc
    if request.method == "POST":
        remember_inspection_export(db, document, "word:all", effective_report_date, lines)
    safe_customer = re.sub(r'[\\/:*?"<>|]', "_", document.customer_name.strip())
    filename = f"{safe_customer}{effective_report_date:%m%d}{label}[{quantity}].docx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/inspection-documents/{object_id}/export-bottle-reports")
@router.get("/inspection-documents/{object_id}/export-bottle-reports")
def export_inspection_document_bottle_reports(
    object_id: int,
    request: Request,
    report_date: date | None = None,
    _: User = Depends(require("inspection_documents:export")),
    db: Session = Depends(get_db),
):
    document = db.get(InspectionDocument, object_id)
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能导出")
    if not document.customer_name.strip():
        raise HTTPException(status_code=400, detail="请先统一设置客户/充装单位")
    eligible_lines = [line for line in document.lines if not line.is_unverified]
    if not eligible_lines:
        raise HTTPException(status_code=400, detail="检验单中没有未被上游驳回的钢瓶")
    if len(eligible_lines) > 9999:
        raise HTTPException(status_code=400, detail="单张检验单最多导出 9999 份单瓶报告")
    specs_by_volume: dict[Decimal, list[BottleSpec]] = {}
    for spec in db.scalars(select(BottleSpec)).all():
        specs_by_volume.setdefault(spec.volume, []).append(spec)
    report_lines = [
        (line, report_spec_for_inspection_line(db, line, specs_by_volume))
        for line in sorted(eligible_lines, key=lambda item: item.sequence)
    ]
    effective_report_date = report_date or document.export_dates.get("bottle:all") or next(
        (line.inspection_time for line in eligible_lines if line.inspection_time), document.inspection_date,
    )

    try:
        template_bytes = BOTTLE_REPORT_TEMPLATE_PATH.read_bytes()
        archive = SpooledTemporaryFile(max_size=8 * 1024 * 1024)
        export_date = china_today()
        used_filenames: set[str] = set()
        with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zipped:
            for line, spec in report_lines:
                content = render_bottle_report_word(
                    body_code=line.body_code,
                    sequence=line.sequence,
                    customer_name=document.customer_name.strip(),
                    manufacturer=line.manufacturer or "",
                    spec_name=spec.name,
                    manufacture_date=line.manufacture_date,
                    volume=line.volume,
                    wall_thickness=spec.wall_thickness,
                    inspection_date=effective_report_date,
                    template_bytes=template_bytes,
                )
                safe_code = re.sub(r'[\\/:*?"<>|]', "_", line.body_code)
                filename = f"{safe_code}[{export_date:%Y%m%d}].docx"
                if filename in used_filenames:
                    filename = f"{safe_code}[{export_date:%Y%m%d}]-{line.sequence:04d}.docx"
                used_filenames.add(filename)
                zipped.writestr(filename, content)
        archive.seek(0)
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise HTTPException(status_code=500, detail=f"单瓶报告生成失败：{exc}") from exc
    if request.method == "POST":
        remember_inspection_export(db, document, "bottle:all", effective_report_date, eligible_lines)

    def chunks():
        try:
            while part := archive.read(1024 * 1024):
                yield part
        finally:
            archive.close()

    filename = f"{document.document_no}-单瓶检测报告[{len(eligible_lines)}].zip"
    return StreamingResponse(
        chunks(),
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/inspection-documents/{object_id}/lines/{line_id}/export-bottle-report")
@router.get("/inspection-documents/{object_id}/lines/{line_id}/export-bottle-report")
def export_inspection_document_bottle_report(
    object_id: int,
    line_id: int,
    request: Request,
    report_date: date | None = None,
    _: User = Depends(require("inspection_documents:export")),
    db: Session = Depends(get_db),
):
    document = db.get(InspectionDocument, object_id)
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能导出")
    if not document.customer_name.strip():
        raise HTTPException(status_code=400, detail="请先统一设置客户/充装单位")
    line = next((item for item in document.lines if item.id == line_id), None)
    if not line:
        raise not_found("检验单钢瓶")
    if line.is_unverified:
        raise HTTPException(status_code=400, detail="被上游驳回的钢瓶不能导出单瓶检测报告")
    if line.sequence > 9999:
        raise HTTPException(status_code=400, detail="单瓶报告序号不能超过 9999")
    spec = report_spec_for_inspection_line(db, line)
    effective_report_date = report_date or document.export_dates.get(f"bottle:{line.id}") or line.inspection_time or document.inspection_date
    try:
        content = render_bottle_report_word(
            body_code=line.body_code,
            sequence=line.sequence,
            customer_name=document.customer_name.strip(),
            manufacturer=line.manufacturer or "",
            spec_name=spec.name,
            manufacture_date=line.manufacture_date,
            volume=line.volume,
            wall_thickness=spec.wall_thickness,
            inspection_date=effective_report_date,
        )
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise HTTPException(status_code=500, detail=f"单瓶报告生成失败：{exc}") from exc
    if request.method == "POST":
        remember_inspection_export(db, document, f"bottle:{line.id}", effective_report_date, [line])
    safe_code = re.sub(r'[\\/:*?"<>|]', "_", line.body_code)
    filename = f"{safe_code}[{china_today():%Y%m%d}].docx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/inspection-documents/import", response_model=InspectionDocumentOut, status_code=201)
def import_inspection_document(
    inspection_date: date = Form(...),
    file: UploadFile = File(...),
    user: User = Depends(require("inspection_documents:create")),
    db: Session = Depends(get_db),
):
    try:
        if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
            raise HTTPException(status_code=400, detail="请上传导出全量检验单格式的 .xlsx 或 .xls 文件")
        content = file.file.read(10 * 1024 * 1024 + 1)
        if len(content) > 10 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="Excel 文件不能超过 10 MB")
        try:
            customer_name, inspector_name, rows = parse_full_inspection_excel(content, inspection_date, file.filename)
        except (ValueError, InvalidFileException, XLRDError, zipfile.BadZipFile, KeyError) as exc:
            raise HTTPException(status_code=400, detail=f"检验单 Excel 导入失败：{exc}") from exc
    finally:
        # UploadFile may spool large requests to a temporary file; closing it
        # removes that temporary copy before the database transaction starts.
        file.file.close()
    if db.bind and db.bind.dialect.name == "postgresql":
        db.execute(text("SELECT pg_advisory_xact_lock(902612)"))
    existing = next((document for document in regular_inspection_documents(db, active_only=False)
                     if inspection_date in document.source_dates), None)
    if existing:
        detail = f"所选日期已有检验单 {existing.document_no}，历史归档不能重复导入"
        if existing.status == "ACTIVE":
            detail += "；如需更新，请打开该检验单使用“回导修改后的检验单”"
        raise HTTPException(status_code=409, detail=detail)
    document = InspectionDocument(
        document_no=f"JY{inspection_date:%Y%m%d}",
        inspection_date=inspection_date,
        customer_name=customer_name,
        inspector_name=inspector_name,
        is_archived=True,
        status="ACTIVE",
        created_by=user.id,
    )
    specs_by_volume: dict[Decimal, list[int]] = {}
    for spec in db.scalars(select(BottleSpec)).all():
        specs_by_volume.setdefault(spec.volume, []).append(spec.id)
    for row in rows:
        matches = specs_by_volume.get(row["volume"], [])
        row["spec_id"] = matches[0] if len(matches) == 1 else None
    document.lines = [InspectionDocumentLine(bottle_id=None, **row) for row in rows]
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


def generate_inspection_documents(db: Session, inspection_date: date, user: User) -> list[InspectionDocument]:
    if db.bind and db.bind.dialect.name == "postgresql":
        db.execute(text("SELECT pg_advisory_xact_lock(902612)"))
    existing = [document for document in regular_inspection_documents(db) if inspection_date in document.source_dates]
    if any(len(document.source_dates) > 1 for document in existing):
        return existing
    active_documents = [document for document in existing if not document.is_archived]
    if not active_documents and existing and existing[0].is_archived:
        return [existing[0]]
    rows = inspection_document_candidates(db, inspection_date)
    if not rows and not active_documents:
        raise HTTPException(status_code=400, detail="该日期没有可生成检验单的良品钢瓶")
    by_source: dict[int, list] = {}
    for row in rows:
        by_source.setdefault(row[0].effective_coding_customer_id, []).append(row)
    existing_by_source = {document.source_id: document for document in active_documents if document.source_id is not None}
    claimed_bottle_ids = claimed_inspection_bottle_ids(db)
    base_no = f"JY{inspection_date:%Y%m%d}"
    used_numbers = set(db.scalars(select(InspectionDocument.document_no).where(
        InspectionDocument.document_no.like(f"{base_no}%"),
    )).all())
    for source_id in sorted(by_source):
        source_rows = by_source[source_id]
        document = existing_by_source.get(source_id)
        if document is None:
            source_rows = [row for row in source_rows if row[0].id not in claimed_bottle_ids]
            if not source_rows:
                continue
            version = 1
            while True:
                document_no = base_no if version == 1 else f"{base_no}-{version:02d}"
                if document_no not in used_numbers:
                    used_numbers.add(document_no)
                    break
                version += 1
            document = InspectionDocument(
                document_no=document_no,
                inspection_date=inspection_date,
                source_id=source_id,
                customer_name=source_rows[0][2],
                status="ACTIVE",
                created_by=user.id,
            )
            db.add(document)
            active_documents.append(document)
            existing_by_source[source_id] = document
        else:
            own_bottle_ids = {line.bottle_id for line in document.lines}
            source_rows = [row for row in source_rows if row[0].id not in claimed_bottle_ids or row[0].id in own_bottle_ids]
        if add_inspection_document_lines(document, source_rows):
            document.updated_at = datetime.utcnow()
        claimed_bottle_ids.update(row[0].id for row in source_rows)
    db.commit()
    for document in active_documents:
        db.refresh(document)
    return active_documents


def generate_merged_inspection_documents(
    db: Session, dates: list[date], inspection_date: date, user: User,
) -> list[InspectionDocument]:
    if db.bind and db.bind.dialect.name == "postgresql":
        db.execute(text("SELECT pg_advisory_xact_lock(902612)"))
    dates = sorted(set(dates))
    generated_dates = {day for document in regular_inspection_documents(db) for day in document.source_dates}
    conflicts = generated_dates.intersection(dates)
    if conflicts:
        raise HTTPException(status_code=409, detail="以下日期已有检验单，请刷新后重新选择：" + "、".join(day.isoformat() for day in sorted(conflicts)))
    claimed_ids = claimed_inspection_bottle_ids(db)
    rows = [row for row in inspection_document_candidates(db, dates) if row[0].id not in claimed_ids]
    available_dates = {row[0].good_entry_date for row in rows}
    if set(dates) - available_dates:
        raise HTTPException(status_code=400, detail="以下日期没有可生成检验单的良品钢瓶：" + "、".join(day.isoformat() for day in sorted(set(dates) - available_dates)))
    by_source: dict[int, list] = {}
    for row in rows:
        by_source.setdefault(row[0].effective_coding_customer_id, []).append(row)
    base_no = f"JY{inspection_date:%Y%m%d}"
    used_numbers = set(db.scalars(select(InspectionDocument.document_no).where(
        InspectionDocument.document_no.like(f"{base_no}%"),
    )).all())
    documents = []
    for source_id in sorted(by_source):
        version = 1
        document_no = base_no
        while document_no in used_numbers:
            version += 1
            document_no = f"{base_no}-{version:02d}"
        used_numbers.add(document_no)
        document = InspectionDocument(
            document_no=document_no, inspection_date=inspection_date,
            generation_dates=[day.isoformat() for day in dates], source_id=source_id,
            customer_name=by_source[source_id][0][2], status="ACTIVE", created_by=user.id,
        )
        add_inspection_document_lines(document, by_source[source_id])
        db.add(document)
        documents.append(document)
    db.commit()
    for document in documents:
        db.refresh(document)
    return documents


@router.post("/inspection-documents/generate", response_model=list[InspectionDocumentOut], status_code=201)
def create_inspection_documents_by_source(
    payload: InspectionDocumentGenerate, user: User = Depends(require("inspection_documents:create")), db: Session = Depends(get_db),
):
    if payload.inspection_dates is not None:
        return generate_merged_inspection_documents(db, payload.inspection_dates, payload.document_date, user)
    return generate_inspection_documents(db, payload.inspection_date, user)


@router.post("/inspection-documents", response_model=InspectionDocumentOut, status_code=201)
def create_inspection_document(payload: InspectionDocumentCreate, user: User = Depends(require("inspection_documents:create")), db: Session = Depends(get_db)):
    return generate_inspection_documents(db, payload.inspection_date, user)[0]


@router.post("/inspection-documents/{object_id}/sync", response_model=InspectionDocumentOut)
def sync_inspection_document(object_id: int, _: User = Depends(require("inspection_documents:create")), db: Session = Depends(get_db)):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能同步数据")
    if document.is_archived:
        raise HTTPException(status_code=400, detail="历史归档检验单不能同步良品库存数据")
    if document.source_id is None:
        raise HTTPException(status_code=400, detail="旧版混合来源检验单不能直接同步，请先核对供瓶单位")
    own_ids = {line.bottle_id for line in document.lines}
    claimed_ids = claimed_inspection_bottle_ids(db) - own_ids
    candidates = inspection_document_candidates(db, document.source_dates, document.source_id)
    added = add_inspection_document_lines(document, [row for row in candidates if row[0].id not in claimed_ids])
    if added:
        document.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(document)
    return document


def inspection_split_context(db: Session, object_id: int, payload: InspectionDocumentSplitPreview, *, lock: bool = False):
    stmt = select(InspectionDocument).where(InspectionDocument.id == object_id)
    document = db.scalar(stmt.with_for_update() if lock else stmt)
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE" or document.is_retry or document.is_archived:
        raise HTTPException(status_code=400, detail="只有有效的系统常规检验单可以拆单")
    customer = db.get(BottleSource, payload.customer_id)
    if not customer or not customer.is_active:
        raise HTTPException(status_code=400, detail="请选择启用的客户")
    named_source_customer = db.scalar(select(BottleSource.id).where(BottleSource.name == document.customer_name))
    if (named_source_customer or document.source_id) == customer.id:
        raise HTTPException(status_code=400, detail="目标客户与原单客户相同，无需拆单")
    lines = sorted((line for line in document.lines if payload.start_sequence <= line.sequence <= payload.end_sequence), key=lambda line: line.sequence)
    if not lines or len(lines) != payload.end_sequence - payload.start_sequence + 1:
        raise HTTPException(status_code=400, detail="所选序号区间不存在或不连续，请刷新检验单后重新选择")
    if any(line.unverified_record and line.unverified_record.status == "RETRYING" for line in lines):
        raise HTTPException(status_code=409, detail="所选钢瓶正在重检，请先处理关联的重检单")
    # Older whole-document edits changed the displayed customer name without changing source_id.
    named_customer_id = select(BottleSource.id).where(
        BottleSource.name == InspectionDocument.customer_name,
    ).correlate(InspectionDocument).scalar_subquery()
    targets_stmt = select(InspectionDocument).where(
        InspectionDocument.inspection_date == document.inspection_date,
        func.coalesce(named_customer_id, InspectionDocument.source_id) == customer.id,
        InspectionDocument.status == "ACTIVE",
        InspectionDocument.is_archived.is_(False), InspectionDocument.is_retry.is_(False),
        InspectionDocument.id != document.id,
    ).order_by(InspectionDocument.id)
    targets = list(db.scalars(targets_stmt.with_for_update() if lock else targets_stmt).all())
    target = next((item for item in targets if item.id == payload.target_document_id), None) if payload.target_document_id else (targets[0] if len(targets) == 1 else None)
    if payload.target_document_id and target is None:
        raise HTTPException(status_code=409, detail="所选目标检验单已变化，请重新预览")
    return document, customer, lines, targets, target


@router.post("/inspection-documents/{object_id}/split-preview")
def preview_inspection_split(object_id: int, payload: InspectionDocumentSplitPreview,
                             _: User = Depends(require("inspection_documents:update")), db: Session = Depends(get_db)):
    document, customer, lines, targets, target = inspection_split_context(db, object_id, payload)
    return {
        "source_updated_at": document.updated_at,
        "line_ids": [line.id for line in lines], "quantity": len(lines),
        "lines": [InspectionDocumentLineOut.model_validate(line).model_dump(mode="json") for line in lines],
        "customer_name": customer.name, "source_quantity_after": len(document.lines) - len(lines),
        "target_candidates": [InspectionDocumentSummaryOut.model_validate(item) for item in targets],
        "target": InspectionDocumentSummaryOut.model_validate(target) if target else None,
        "requires_target_selection": len(targets) > 1 and target is None,
        "will_create": not targets,
    }


@router.get("/inspection-documents/{object_id}/split-history", response_model=PageOut[InspectionDocumentSplitHistoryOut])
def inspection_split_history(object_id: int, page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100),
                             _: User = Depends(require("inspection_documents:view")), db: Session = Depends(get_db)):
    if not db.get(InspectionDocument, object_id):
        raise not_found("检验单")
    stmt = select(InspectionDocumentSplitHistory).where(or_(
        InspectionDocumentSplitHistory.source_document_id == object_id,
        InspectionDocumentSplitHistory.target_document_id == object_id,
    )).order_by(InspectionDocumentSplitHistory.id.desc())
    return page_result(db, stmt, page, page_size)


@router.post("/inspection-documents/{object_id}/split")
def split_inspection_document(object_id: int, payload: InspectionDocumentSplitCreate,
                              user: User = Depends(require("inspection_documents:update")), db: Session = Depends(get_db)):
    if db.bind and db.bind.dialect.name == "postgresql":
        db.execute(text("SELECT pg_advisory_xact_lock(902612)"))
    document, customer, lines, targets, target = inspection_split_context(db, object_id, payload, lock=True)
    if document.updated_at != payload.expected_source_updated_at or [line.id for line in lines] != payload.expected_line_ids:
        raise HTTPException(status_code=409, detail="原单明细或序号已变化，请重新预览后再拆单")
    if len(targets) > 1 and target is None:
        raise HTTPException(status_code=409, detail="该客户同日有多张检验单，请选择合并目标")
    if (target.id if target else None) != payload.expected_target_id or (target.updated_at if target else None) != payload.expected_target_updated_at:
        raise HTTPException(status_code=409, detail="目标检验单已变化，请重新预览后再拆单")
    target_created = target is None
    if target_created:
        if not payload.confirm_create:
            raise HTTPException(status_code=409, detail="当天该客户没有检验单，请确认创建新单后再拆单")
        if not user.is_superadmin and "inspection_documents:create" not in user.permissions:
            raise HTTPException(status_code=403, detail="创建目标检验单需要检验单创建权限")
        base_no = f"JY{document.inspection_date:%Y%m%d}"
        used_numbers = set(db.scalars(select(InspectionDocument.document_no).where(InspectionDocument.document_no.like(f"{base_no}%"))).all())
        version, document_no = 1, base_no
        while document_no in used_numbers:
            version += 1
            document_no = f"{base_no}-{version:02d}"
        target = InspectionDocument(document_no=document_no, inspection_date=document.inspection_date,
                                    source_id=customer.id, customer_name=customer.name, generation_dates=[],
                                    inspector_id=document.inspector_id, inspector_name=document.inspector_name,
                                    created_by=user.id, status="ACTIVE")
        db.add(target)
        db.flush()
    target_bottle_ids = {item.bottle_id for item in target.lines if item.bottle_id is not None}
    target_body_codes = {item.body_code.casefold() for item in target.lines}
    if any(line.bottle_id in target_bottle_ids or line.body_code.casefold() in target_body_codes for line in lines):
        raise HTTPException(status_code=409, detail="目标检验单已有相同钢瓶，不能重复合并")
    source_before, target_before = len(document.lines), len(target.lines)
    snapshot = [dict(InspectionDocumentLineOut.model_validate(line).model_dump(mode="json"),
                     original_source_id=line.bottle.source_id if line.bottle else None,
                     original_coding_customer_id=line.bottle.effective_coding_customer_id if line.bottle else None)
                for line in lines]
    # Customer assignment is independent of the original stock/inbound/outbound source.
    source_customer_id = db.scalar(select(BottleSource.id).where(BottleSource.name == document.customer_name)) or document.source_id
    moving_bottle_ids = {line.bottle_id for line in lines}
    source_bottle_ids = {line.bottle_id for line in document.lines if line.bottle_id is not None}
    bottles = list(db.scalars(select(Bottle).where(Bottle.id.in_(source_bottle_ids | target_bottle_ids)).with_for_update()).all())
    for bottle in bottles:
        assigned_customer = customer.id if bottle.id in moving_bottle_ids or bottle.id in target_bottle_ids else source_customer_id
        if bottle.effective_coding_customer_id != assigned_customer:
            bottle.coding_customer_id = assigned_customer
    document.source_id = source_customer_id
    target.source_id = customer.id
    target.customer_name = customer.name
    source_days = set(document.source_dates)
    target_days = set(target.source_dates) if not target_created else set()
    target.generation_dates = [day.isoformat() for day in sorted(source_days | target_days)]
    for line in lines:
        if not line.inspector_name:
            line.inspector_name = document.inspector_name or None
        target.lines.append(line)
        if line.unverified_record:
            line.unverified_record.inspection_document_id = target.id
    moving_kinds = [f"bottle:{line.id}" for line in lines]
    for record in db.scalars(select(InspectionDocumentExportRecord).where(
        InspectionDocumentExportRecord.document_id == document.id,
        InspectionDocumentExportRecord.kind.in_(moving_kinds),
    )).all():
        record.document_id = target.id
    resequence_inspection_document(document)
    resequence_inspection_document(target)
    source_emptied = not document.lines
    if source_emptied:
        document.status = "SPLIT_EMPTY"
    document.updated_at = target.updated_at = datetime.utcnow()
    history = InspectionDocumentSplitHistory(
        source_document_id=document.id, target_document_id=target.id,
        source_document_no=document.document_no, target_document_no=target.document_no,
        source_customer_name=document.customer_name, target_customer_name=customer.name,
        inspection_date=document.inspection_date, start_sequence=payload.start_sequence, end_sequence=payload.end_sequence,
        quantity=len(lines), target_created=target_created, source_emptied=source_emptied,
        source_quantity_before=source_before, target_quantity_before=target_before, snapshot=snapshot, created_by=user.id,
    )
    db.add(history)
    db.commit()
    return {"source": InspectionDocumentOut.model_validate(document), "target": InspectionDocumentOut.model_validate(target),
            "history": InspectionDocumentSplitHistoryOut.model_validate(history)}


@router.put("/inspection-documents/{object_id}", response_model=InspectionDocumentOut)
def update_inspection_document(object_id: int, payload: InspectionDocumentUpdate, _: User = Depends(require("inspection_documents:update")), db: Session = Depends(get_db)):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能修改")
    inspector = None
    if payload.inspector_id is not None:
        inspector = db.get(User, payload.inspector_id)
        if not inspector or not inspector.is_active or inspector.is_superadmin:
            raise HTTPException(status_code=400, detail="请选择一名启用的普通员工作为检验人员")
    customer_name = payload.customer_name.strip()
    if customer_name and not document.is_archived and not document.is_retry:
        customer = db.scalar(select(BottleSource).where(BottleSource.name == customer_name))
        if customer and (document.source_id != customer.id or document.customer_name != customer_name):
            document.source_id = customer.id
            for line in document.lines:
                if line.bottle:
                    line.bottle.coding_customer_id = customer.id
    document.customer_name = customer_name
    if inspector or not document.is_archived:
        document.inspector_id = inspector.id if inspector else None
        document.inspector_name = inspector.name if inspector else ""
    if document.is_archived and inspector:
        for line in document.lines:
            line.inspector_name = document.inspector_name
    db.commit()
    db.refresh(document)
    return document


@router.put("/inspection-documents/{object_id}/lines/{line_id}", response_model=InspectionDocumentOut)
def update_inspection_document_line(
    object_id: int,
    line_id: int,
    payload: InspectionDocumentLineUpdate,
    _: User = Depends(require("inspection_documents:update")),
    db: Session = Depends(get_db),
):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能修改")
    line = next((item for item in document.lines if item.id == line_id), None)
    if not line:
        raise not_found("检验单钢瓶")
    bottle = None
    if not document.is_archived:
        bottle = db.scalar(select(Bottle).where(Bottle.id == line.bottle_id).with_for_update())
        if not bottle:
            raise HTTPException(status_code=400, detail="检验单关联的库存钢瓶不存在")
    fields = payload.model_fields_set
    if not fields:
        raise HTTPException(status_code=400, detail="请填写需要修改的明细信息")
    if "body_code" in fields:
        body_code = (payload.body_code or "").strip()
        if not body_code:
            raise HTTPException(status_code=400, detail="产品编号不能为空")
        if any(item.id != line.id and item.body_code.casefold() == body_code.casefold() for item in document.lines):
            raise HTTPException(status_code=400, detail="同一检验单中的产品编号不能重复")
        if bottle and db.scalar(select(Bottle.id).where(func.lower(Bottle.body_code) == body_code.lower(), Bottle.id != bottle.id)):
            raise HTTPException(status_code=400, detail="产品编号已被其他库存钢瓶使用")
        line.body_code = body_code
        if bottle:
            bottle.body_code = body_code
    if "manufacture_date" in fields:
        if payload.manufacture_date and payload.manufacture_date > document.inspection_date:
            raise HTTPException(status_code=400, detail="生产日期不能晚于检验日期")
        line.manufacture_date = payload.manufacture_date
        if bottle:
            bottle.manufacture_date = payload.manufacture_date
    if "manufacturer" in fields:
        line.manufacturer = (payload.manufacturer or "").strip()
        if bottle:
            bottle.manufacturer = line.manufacturer or None
    if "volume" in fields:
        if payload.volume is None:
            raise HTTPException(status_code=400, detail="容积不能为空")
        if bottle and payload.volume != line.volume:
            raise HTTPException(status_code=400, detail="系统生成的检验单容积由钢瓶规格决定，不能单独修改")
        line.volume = payload.volume
    if "next_inspection_date" in fields:
        if bottle and payload.next_inspection_date != line.next_inspection_date:
            raise HTTPException(status_code=400, detail="系统生成的检验单不能单独修改下次检验日期")
        line.next_inspection_date = payload.next_inspection_date
    if "spec_id" in fields:
        if not document.is_archived:
            if payload.spec_id != bottle.spec_id:
                raise HTTPException(status_code=400, detail="系统生成的检验单不能修改钢瓶规格")
        elif payload.spec_id is not None and not db.get(BottleSpec, payload.spec_id):
            raise HTTPException(status_code=400, detail="钢瓶规格不存在")
        else:
            line.spec_id = payload.spec_id
    resequence_inspection_document(document)
    document.updated_at = datetime.utcnow()
    save(db, "产品编号已被其他库存钢瓶使用")
    db.refresh(document)
    return document


@router.post("/inspection-documents/{object_id}/reimport", response_model=InspectionDocumentOut)
def reimport_inspection_document(
    object_id: int,
    file: UploadFile = File(...),
    _: User = Depends(require("inspection_documents:update")),
    db: Session = Depends(get_db),
):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="已作废检验单不能回导修改")
    try:
        if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
            raise HTTPException(status_code=400, detail="请上传导出全量检验单格式的 .xlsx 或 .xls 文件")
        content = file.file.read(10 * 1024 * 1024 + 1)
        if len(content) > 10 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="Excel 文件不能超过 10 MB")
        try:
            customer_name, _, rows = parse_full_inspection_excel(
                content, document.inspection_date, file.filename,
                match_by_sequence=True, require_matching_date=False,
            )
        except (ValueError, InvalidFileException, XLRDError, zipfile.BadZipFile, KeyError) as exc:
            raise HTTPException(status_code=400, detail=f"检验单 Excel 回导失败：{exc}") from exc
    finally:
        file.file.close()

    if customer_name != document.customer_name:
        raise HTTPException(status_code=400, detail="Excel 的客户/充装单位与当前检验单不一致；如需修改客户，请使用整单批量设置")
    lines_by_sequence = {line.sequence: line for line in document.lines}
    if set(row["sequence"] for row in rows) != set(lines_by_sequence):
        raise HTTPException(status_code=400, detail="Excel 序号与当前检验单不一致，请重新导出全量检验单后再修改回导")

    codes = [row["body_code"].casefold() for row in rows]
    if len(codes) != len(set(codes)):
        raise HTTPException(status_code=400, detail="同一检验单中的产品编号不能重复")
    bottles_by_line: dict[int, Bottle] = {}
    if not document.is_archived and not document.is_retry:
        bottle_ids = [line.bottle_id for line in document.lines]
        bottles = db.scalars(select(Bottle).where(Bottle.id.in_(bottle_ids)).with_for_update()).all()
        bottles_by_id = {bottle.id: bottle for bottle in bottles}
        if any(line.bottle_id not in bottles_by_id for line in document.lines):
            raise HTTPException(status_code=400, detail="检验单关联的库存钢瓶不存在")
        bottles_by_line = {line.id: bottles_by_id[line.bottle_id] for line in document.lines}
        changed_codes = {row["body_code"].lower() for row in rows
                         if row["body_code"] != bottles_by_line[lines_by_sequence[row["sequence"]].id].body_code}
        if changed_codes and db.scalar(select(Bottle.id).where(func.lower(Bottle.body_code).in_(changed_codes))):
            raise HTTPException(status_code=400, detail="Excel 中的产品编号已被其他库存钢瓶使用")

    for row in rows:
        line = lines_by_sequence[row["sequence"]]
        manufactured = row["manufacture_date"]
        if manufactured and manufactured > document.inspection_date:
            raise HTTPException(status_code=400, detail=f"序号 {line.sequence} 的生产日期不能晚于检验日期")
        if line.manufacture_date and manufactured and manufactured.year == line.manufacture_date.year and manufactured.month == line.manufacture_date.month:
            manufactured = line.manufacture_date
        row["manufacture_date"] = manufactured
        if line.id in bottles_by_line:
            if row["volume"] != line.volume:
                raise HTTPException(status_code=400, detail=f"序号 {line.sequence} 的容积由钢瓶规格决定，不能通过 Excel 单独修改")
            expected_next_date = default_next_inspection_date(line.inspection_time or document.inspection_date)
            if row["next_inspection_date"] != expected_next_date:
                raise HTTPException(status_code=400, detail=f"序号 {line.sequence} 的下次检验日期不能通过 Excel 单独修改")

    for row in rows:
        line = lines_by_sequence[row["sequence"]]
        line.body_code = row["body_code"]
        line.manufacture_date = row["manufacture_date"]
        line.manufacturer = row["manufacturer"]
        line.device_category = row["device_category"]
        line.inspection_result = row["inspection_result"]
        line.inspector_name = row["inspector_name"]
        if line.id in bottles_by_line:
            bottle = bottles_by_line[line.id]
            bottle.body_code = line.body_code
            bottle.manufacture_date = line.manufacture_date
            bottle.manufacturer = line.manufacturer or None
        else:
            line.volume = row["volume"]
            line.next_inspection_date = row["next_inspection_date"]
    resequence_inspection_document(document)
    document.updated_at = datetime.utcnow()
    save(db, "Excel 中的产品编号已被其他库存钢瓶使用")
    db.refresh(document)
    return document


@router.post("/inspection-documents/{object_id}/mark-unverified", response_model=UnverifiedBottleImportResult)
def mark_unverified_bottles(
    object_id: int,
    payload: UnverifiedBottleImport,
    user: User = Depends(require("unverified_bottles:create")),
    db: Session = Depends(get_db),
):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="只能标注有效检验单")
    matches = list(re.finditer(r"气瓶列表\s*第\s*(\d+)\s*行", payload.raw_text, flags=re.IGNORECASE))
    if not matches:
        raise HTTPException(status_code=400, detail="没有识别到“气瓶列表第N行”，请检查粘贴内容")
    fragments: dict[int, str] = {}
    for index, match in enumerate(matches):
        sequence = int(match.group(1))
        end = matches[index + 1].start() if index + 1 < len(matches) else len(payload.raw_text)
        fragment = payload.raw_text[match.start():end].strip()
        fragments.setdefault(sequence, fragment[:2000])
    lines_by_sequence = {line.sequence: line for line in document.lines}
    matched_sequences: list[int] = []
    unmatched_sequences: list[int] = []
    for sequence, fragment in fragments.items():
        line = lines_by_sequence.get(sequence)
        if not line:
            unmatched_sequences.append(sequence)
            continue
        message_text = re.sub(r"^气瓶列表\s*第\s*\d+\s*行", "", fragment, count=1, flags=re.IGNORECASE)
        error_message = re.sub(r"\s+", " ", message_text).strip(" ，,。；;：【】[]")
        if not error_message:
            error_message = "气瓶信息未通过上游系统验证"
        record = db.scalar(select(UnverifiedBottle).where(UnverifiedBottle.inspection_line_id == line.id).with_for_update())
        if record and record.status == "RETRYING":
            raise HTTPException(status_code=400, detail=f"第 {sequence} 行已转入重检单，不能再次按原检验单录入驳回")
        values = {
            "inspection_document_id": document.id,
            "bottle_id": line.bottle_id,
            "sequence": line.sequence,
            "body_code": line.body_code,
            "manufacture_date": line.manufacture_date,
            "manufacturer": line.manufacturer,
            "error_message": error_message[:1000],
            "raw_fragment": fragment,
            "status": "PENDING",
            "created_by": user.id,
        }
        if record:
            for key, value in values.items():
                setattr(record, key, value)
        else:
            db.add(UnverifiedBottle(inspection_line_id=line.id, **values))
        matched_sequences.append(sequence)
    db.commit()
    return {
        "matched_sequences": matched_sequences,
        "unmatched_sequences": unmatched_sequences,
        "total_marked": len(matched_sequences),
    }


@router.post("/inspection-documents/{object_id}/void", response_model=InspectionDocumentOut)
def void_inspection_document(object_id: int, user: User = Depends(require("inspection_documents:void")), db: Session = Depends(get_db)):
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "ACTIVE":
        raise HTTPException(status_code=400, detail="检验单已经作废")
    if any(line.bottle and line.bottle.zone == "OUTBOUND" for line in document.lines):
        raise HTTPException(status_code=400, detail="检验单中已有钢瓶出库，不能作废")
    if document.is_retry:
        source_line_ids = [line.source_line_id for line in document.lines if line.source_line_id is not None]
        if source_line_ids:
            for record in db.scalars(select(UnverifiedBottle).where(
                UnverifiedBottle.inspection_line_id.in_(source_line_ids),
                UnverifiedBottle.status == "RETRYING",
            )).all():
                record.status = "PENDING"
    document.status = "VOID"
    document.voided_by = user.id
    document.voided_at = datetime.utcnow()
    db.commit()
    db.refresh(document)
    return document


@router.delete("/inspection-documents/{object_id}", status_code=204)
def delete_void_inspection_document(
    object_id: int,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="只有超级管理员可以彻底删除作废检验单")
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == object_id).with_for_update())
    if not document:
        raise not_found("检验单")
    if document.status != "VOID":
        raise HTTPException(status_code=400, detail="只能彻底删除已经作废的检验单")
    if db.scalar(select(InspectionDocumentSplitHistory.id).where(or_(
        InspectionDocumentSplitHistory.source_document_id == document.id,
        InspectionDocumentSplitHistory.target_document_id == document.id,
    )).limit(1)):
        raise HTTPException(status_code=409, detail="该检验单有拆单历史，需保留原单号与追溯记录，不能彻底删除")
    line_ids = db.scalars(select(InspectionDocumentLine.id).where(
        InspectionDocumentLine.document_id == document.id,
    ).with_for_update()).all()
    if line_ids and db.scalar(select(InspectionDocumentLine.id).where(
        InspectionDocumentLine.source_line_id.in_(line_ids),
    )):
        raise HTTPException(status_code=409, detail="该检验单的明细已被重检单引用，请先处理关联的重检单")
    for record in db.scalars(select(UnverifiedBottle).where(
        UnverifiedBottle.inspection_document_id == document.id,
    ).with_for_update()).all():
        db.delete(record)
    db.flush()
    db.delete(document)
    save(db, "检验单仍被其他业务数据引用，不能删除")


def filtered_unverified_bottles(
    document_no: str | None,
    body_code: str | None,
    date_from: date | None,
    date_to: date | None,
):
    stmt = select(UnverifiedBottle).join(InspectionDocument).where(
        UnverifiedBottle.status == "PENDING",
        InspectionDocument.status == "ACTIVE",
    )
    if document_no:
        stmt = stmt.where(InspectionDocument.document_no.ilike(f"%{document_no.strip()}%"))
    if body_code:
        stmt = stmt.where(UnverifiedBottle.body_code.ilike(f"%{body_code.strip()}%"))
    if date_from:
        stmt = stmt.where(InspectionDocument.inspection_date >= date_from)
    if date_to:
        stmt = stmt.where(InspectionDocument.inspection_date <= date_to)
    stmt = stmt.order_by(InspectionDocument.inspection_date.desc(), InspectionDocument.id.desc(), UnverifiedBottle.sequence, UnverifiedBottle.id)
    return stmt


@router.get("/unverified-bottles", response_model=PageOut[UnverifiedBottleOut])
def list_unverified_bottles(
    document_no: str | None = None,
    body_code: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require("unverified_bottles:view")),
    db: Session = Depends(get_db),
):
    stmt = filtered_unverified_bottles(document_no, body_code, date_from, date_to)
    return page_result(db, stmt, page, page_size)


@router.get("/unverified-bottles/export-xls")
def export_unverified_bottles_xls(
    _: User = Depends(require("unverified_bottles:export")),
    db: Session = Depends(get_db),
):
    records = db.scalars(filtered_unverified_bottles(None, None, None, None)).all()
    content = render_unverified_bottles_xls(records)
    filename = f"未经系统验证瓶子池{china_today():%Y%m%d}全量[{len(records)}].xls"
    return Response(
        content=content,
        media_type="application/vnd.ms-excel",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/unverified-bottles/create-retry-document", response_model=list[InspectionDocumentSummaryOut], status_code=201)
def create_retry_inspection_document(
    payload: UnverifiedBottleBatchCreate,
    user: User = Depends(require("unverified_bottles:batch")),
    _: User = Depends(require("inspection_documents:create")),
    db: Session = Depends(get_db),
):
    if payload.date_from and payload.date_to and payload.date_from > payload.date_to:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    if db.bind and db.bind.dialect.name == "postgresql":
        # Serialize customer assignment/splitting with retry snapshot generation.
        db.execute(text("SELECT pg_advisory_xact_lock(902612)"))
        db.execute(text("SELECT pg_advisory_xact_lock(902613)"))
    stmt = filtered_unverified_bottles(payload.document_no, payload.body_code, payload.date_from, payload.date_to)
    source_ids = list(db.scalars(stmt.with_only_columns(UnverifiedBottle.inspection_document_id).order_by(None).distinct()).all())
    if not source_ids:
        raise HTTPException(status_code=400, detail="当前筛选条件下没有待处理的未验证钢瓶")
    # Lock documents before rejection records, matching edits/voids and preserving customer ownership.
    db.scalars(select(InspectionDocument).where(InspectionDocument.id.in_(source_ids))
               .order_by(InspectionDocument.id).with_for_update()).all()
    selected_records = db.scalars(stmt.where(UnverifiedBottle.inspection_document_id.in_(source_ids))
                                 .with_for_update(of=UnverifiedBottle)).all()
    if not selected_records:
        raise HTTPException(status_code=409, detail="待处理池数据已变化，请刷新后重试")
    missing_customers = sorted({record.document_no for record in selected_records if not record.customer_name.strip()})
    if missing_customers:
        raise HTTPException(status_code=400, detail="以下原检验单未设置客户，请先补齐客户后再生成重检单：" + "、".join(missing_customers))
    customers = list(db.scalars(select(BottleSource)).all())
    customers_by_name = {item.name: item for item in customers}
    groups = {}
    for record in selected_records:
        source_document = record.inspection_document
        name = source_document.customer_name.strip()
        # Current document customer takes precedence over the original inventory source.
        customer = customers_by_name.get(name)
        key = ("id", customer.id) if customer else ("name", name)
        if key not in groups:
            groups[key] = {"source_id": customer.id if customer else None,
                           "customer_name": customer.name if customer else name, "records": []}
        groups[key]["records"].append(record)
    inspection_date = payload.inspection_date or china_today()
    base_no = f"JY{inspection_date:%Y%m%d}-R"
    used_numbers = set(db.scalars(select(InspectionDocument.document_no).where(
        InspectionDocument.document_no.like(f"{base_no}%"),
    )).all())
    version = max((int(number[len(base_no):]) for number in used_numbers if number[len(base_no):].isdigit()), default=0) + 1
    documents = []
    for group in sorted(groups.values(), key=lambda item: (item["customer_name"], item["source_id"] or 0)):
        inspectors = {record.inspector_name for record in group["records"]}
        document_no = f"{base_no}{version:02d}"
        used_numbers.add(document_no)
        version += 1
        document = InspectionDocument(
            document_no=document_no, inspection_date=inspection_date,
            source_id=group["source_id"], customer_name=group["customer_name"],
            inspector_name=next(iter(inspectors)) if len(inspectors) == 1 else "",
            is_archived=True, is_retry=True, status="ACTIVE", created_by=user.id,
        )
        for sequence, record in enumerate(group["records"], start=1):
            source_line = record.inspection_line
            document.lines.append(InspectionDocumentLine(
                bottle_id=None, source_line_id=source_line.id,
                spec_id=source_line.spec_id or (record.bottle.spec_id if record.bottle else None),
                sequence=sequence, device_category=source_line.device_category,
                body_code=record.body_code, manufacture_date=record.manufacture_date,
                manufacturer=record.manufacturer, volume=source_line.volume,
                next_inspection_date=source_line.next_inspection_date,
                inspector_name=source_line.inspector_name or record.inspector_name,
                inspected_at=source_line.inspected_at, inspection_result=source_line.inspection_result,
            ))
            record.status = "RETRYING"
        resequence_inspection_document(document)
        db.add(document)
        documents.append(document)
    save(db, "所选日期的重检单编号已被使用，请刷新后重试")
    for document in documents:
        db.refresh(document)
    return documents


@router.put("/unverified-bottles/{object_id}", response_model=UnverifiedBottleOut)
def update_unverified_bottle(
    object_id: int,
    payload: InspectionDocumentLineUpdate,
    user: User = Depends(require("unverified_bottles:update")),
    db: Session = Depends(get_db),
):
    record = db.scalar(select(UnverifiedBottle).where(UnverifiedBottle.id == object_id).with_for_update())
    if not record or record.status != "PENDING":
        raise not_found("未经系统验证钢瓶")
    if not payload.model_fields_set or payload.model_fields_set - {"body_code", "manufacture_date", "manufacturer"}:
        raise HTTPException(status_code=400, detail="这里只能更正瓶身编号、生产日期和生产厂家")
    update_inspection_document_line(record.inspection_document_id, record.inspection_line_id, payload, user, db)
    db.refresh(record)
    return record


@router.delete("/unverified-bottles/{object_id}", status_code=204)
def delete_unverified_bottle(
    object_id: int,
    user: User = Depends(require("unverified_bottles:delete")),
    db: Session = Depends(get_db),
):
    document_id = db.scalar(select(UnverifiedBottle.inspection_document_id).where(UnverifiedBottle.id == object_id))
    if document_id is None:
        raise not_found("待处理的未经系统验证钢瓶")
    document = db.scalar(select(InspectionDocument).where(InspectionDocument.id == document_id).with_for_update())
    record = db.scalar(select(UnverifiedBottle).where(UnverifiedBottle.id == object_id).with_for_update())
    if not document or not record or record.inspection_document_id != document.id or record.status != "PENDING" or document.status != "ACTIVE":
        raise not_found("待处理的未经系统验证钢瓶")
    line = db.scalar(select(InspectionDocumentLine).where(InspectionDocumentLine.id == record.inspection_line_id).with_for_update())
    if not line or line.document_id != record.inspection_document_id or line.bottle_id != record.bottle_id:
        raise HTTPException(status_code=409, detail="驳回记录与检验单明细关联不一致，请先核对数据")
    if line.source_line_id is not None or db.scalar(select(InspectionDocumentLine.id).where(
        InspectionDocumentLine.source_line_id == line.id,
    )):
        raise HTTPException(status_code=409, detail="该钢瓶已关联重检单，不能直接删除，请先处理关联的重检记录")

    bottle = None
    if record.bottle_id is not None:
        # The outbound order is quantity based, but posting assigns concrete bottles
        # internally. Lock the order before its bottle, matching the void path.
        bottle = db.get(Bottle, record.bottle_id)
        if not bottle:
            raise HTTPException(status_code=409, detail="关联的库存钢瓶不存在，请先核对数据")
        outbound = db.scalar(select(BottleOutbound).where(BottleOutbound.id == bottle.outbound_id).with_for_update()) if bottle.outbound_id else None
        bottle = db.scalar(select(Bottle).where(Bottle.id == record.bottle_id).with_for_update().execution_options(populate_existing=True))
        if not bottle:
            raise HTTPException(status_code=409, detail="关联的库存钢瓶不存在，请先核对数据")
        if db.scalar(select(InspectionDocumentLine.id).where(
            InspectionDocumentLine.bottle_id == bottle.id,
            InspectionDocumentLine.id != line.id,
        )) or db.scalar(select(UnverifiedBottle.id).where(
            UnverifiedBottle.bottle_id == bottle.id,
            UnverifiedBottle.id != record.id,
        )):
            raise HTTPException(status_code=409, detail="该钢瓶仍被其他检验单或驳回记录引用，不能直接删除")
        if bottle.outbound_id is not None or bottle.zone == "OUTBOUND":
            if not outbound or outbound.id != bottle.outbound_id or outbound.status != "POSTED" or bottle.zone != "OUTBOUND":
                raise HTTPException(status_code=409, detail="钢瓶出库状态已变化，请刷新后重试")
            origin = bottle.outbound_origin_zone
            if origin not in {"GOOD", "BAD"}:
                raise HTTPException(status_code=409, detail="钢瓶出库来源不明，不能调整出库数量")
            outbound_line = db.scalar(select(BottleOutboundLine).where(
                BottleOutboundLine.outbound_id == outbound.id,
                BottleOutboundLine.spec_id == bottle.spec_id,
            ).with_for_update())
            if not outbound_line:
                raise HTTPException(status_code=409, detail="出库单缺少对应规格明细，不能删除钢瓶")
            quantity_field = "good_quantity" if origin == "GOOD" else "bad_quantity"
            adjustment_field = "good_adjustment_quantity" if origin == "GOOD" else "bad_adjustment_quantity"
            assigned = db.scalar(select(func.count(Bottle.id)).where(
                Bottle.outbound_id == outbound.id,
                Bottle.spec_id == bottle.spec_id,
                Bottle.outbound_origin_zone == origin,
            )) or 0
            if assigned + getattr(outbound_line, adjustment_field) != getattr(outbound_line, quantity_field):
                raise HTTPException(status_code=409, detail="出库单分配数量不一致，请先核对数据")
            # Replace this concrete allocation with an equal anonymous source and
            # consumption. Available stock stays unchanged; voiding returns one
            # anonymous unit instead of resurrecting the deleted bottle.
            setattr(outbound_line, adjustment_field, getattr(outbound_line, adjustment_field) + 1)
            db.add(BottleInventoryAdjustment(
                business_date=china_today(), source_id=bottle.effective_coding_customer_id,
                spec_id=bottle.spec_id, zone=origin,
                quantity=1, created_by=user.id,
                note=f"删除未验证钢瓶后转为匿名库存：{document.document_no} 第{line.sequence}行，出库单 {outbound.document_no}",
            ))

    db.delete(record)
    document.lines.remove(line)
    db.flush()
    if bottle:
        for allocation in db.scalars(select(BottleOutboundAllocation).where(
            BottleOutboundAllocation.bottle_id == bottle.id
        )).all():
            allocation.bottle_id = None
        db.delete(bottle)
    resequence_inspection_document(document)
    document.updated_at = datetime.utcnow()
    save(db, "钢瓶仍被其他业务数据引用，不能彻底删除")


@router.get("/inspection/bottles", response_model=PageOut[BottleOut])
def inspection_bottles(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), _: User = Depends(require("inspection:view")), db: Session = Depends(get_db)):
    return page_result(db, select(Bottle).order_by(Bottle.process_status, Bottle.id.desc()), page, page_size)


@router.get("/bottle-zones")
def bottle_zones(
    zone: str | None = None,
    body_code: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(require("bottle_zones:view")), db: Session = Depends(get_db),
):
    active_inbounds = db.scalars(select(BottleInbound).where(BottleInbound.status == "POSTED")).all()
    all_bottles = db.scalars(select(Bottle).order_by(Bottle.id.desc())).all()
    inspected_count = sum(1 for item in all_bottles if item.process_status == "INSPECTED")
    coded_count = sum(1 for item in all_bottles if item.process_status == "CODED")
    pending_total = sum(item.pending_quantity for item in active_inbounds)
    outbound_stock = bottle_outbound_availability(db)
    manual_bad_by_spec = manual_bad_stock(db)
    adjustment_good = db.scalar(select(func.coalesce(func.sum(BottleInventoryAdjustment.quantity), 0)).where(BottleInventoryAdjustment.zone == "GOOD")) or 0
    used_adjustment_good = db.scalar(select(func.coalesce(func.sum(BottleOutboundLine.good_adjustment_quantity), 0)).join(BottleOutbound, BottleOutbound.id == BottleOutboundLine.outbound_id).where(BottleOutbound.status == "POSTED")) or 0
    stock_by_spec = []
    for spec in db.scalars(select(BottleSpec).order_by(BottleSpec.sort_order, BottleSpec.id)).all():
        quantities = outbound_stock.get(str(spec.id), {"good": 0, "bad": 0, "temporary": 0, "scrap": 0})
        bad_quantity = quantities["bad"]
        total_quantity = quantities["good"] + quantities["temporary"] + quantities["scrap"] + bad_quantity
        if total_quantity > 0 or spec.is_active:
            stock_by_spec.append({
                "spec_id": spec.id,
                "spec_name": spec.name,
                "good": quantities["good"],
                "temporary": quantities["temporary"],
                "scrap": quantities["scrap"],
                "bad": bad_quantity,
                "total": total_quantity,
            })
    stmt = select(Bottle).order_by(Bottle.id.desc())
    if zone and zone not in {"TEMPORARY", "SCRAP"}:
        stmt = stmt.where(Bottle.zone == zone)
    if body_code:
        stmt = stmt.where(Bottle.body_code.ilike(f"%{body_code.strip()}%"))
    if date_from:
        stmt = stmt.where(Bottle.coding_date >= date_from)
    if date_to:
        stmt = stmt.where(Bottle.coding_date <= date_to)
    bottle_page = page_result(db, stmt, page, page_size)
    return {
        "zones": {
            "pending": pending_total - inspected_count,
            "temporary": sum(item["temporary"] for item in outbound_stock.values()),
            "scrap": sum(item["scrap"] for item in outbound_stock.values()),
            "good": sum(item["good"] for item in outbound_stock.values()),
            "bad": sum(item["bad"] for item in outbound_stock.values()),
        },
        "manual_good": max(int(adjustment_good) - int(used_adjustment_good), 0),
        "manual_bad": sum(manual_bad_by_spec.values()),
        "stock_by_spec": stock_by_spec,
        "adjustments": [BottleInventoryAdjustmentOut.model_validate(item) for item in db.scalars(select(BottleInventoryAdjustment).order_by(BottleInventoryAdjustment.business_date.desc(), BottleInventoryAdjustment.id.desc()).limit(50)).all()],
        "uncoded": max(pending_total - inspected_count - coded_count, 0),
        "bottles": {**bottle_page, "items": [BottleOut.model_validate(item) for item in bottle_page["items"]]},
    }


@router.post("/bottle-inventory-adjustments", response_model=BottleInventoryAdjustmentOut, status_code=201)
def create_bottle_inventory_adjustment(payload: BottleInventoryAdjustmentCreate, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="只有超级管理员可以调整库存")
    spec = db.get(BottleSpec, payload.spec_id)
    if not spec or not spec.is_active:
        raise HTTPException(status_code=400, detail="钢瓶规格不存在或已停用")
    adjustment_data = payload.model_dump()
    if payload.source_id is None and payload.zone != "GOOD":
        default_source = db.scalar(select(BottleSource).where(BottleSource.is_active.is_(True)).order_by(
            BottleSource.is_default.desc(), BottleSource.sort_order, BottleSource.id
        ).limit(1))
        if default_source:
            adjustment_data["source_id"] = default_source.id
    if adjustment_data["source_id"] is not None and not db.get(BottleSource, adjustment_data["source_id"]):
        raise HTTPException(status_code=400, detail="供瓶单位不存在")
    if payload.quantity < 0:
        if payload.zone == "BAD":
            current_quantity = manual_bad_stock(db, adjustment_data["source_id"]).get(payload.spec_id, 0)
        else:
            stock_key = {"GOOD": "good", "TEMPORARY": "temporary", "SCRAP": "scrap"}[payload.zone]
            current_quantity = bottle_outbound_availability(db, adjustment_data["source_id"]).get(str(payload.spec_id), {}).get(stock_key, 0)
        deduction = abs(payload.quantity)
        if deduction > current_quantity:
            zone_name = {"GOOD": "良品", "TEMPORARY": "暂存", "SCRAP": "报废", "BAD": "不良品补录"}[payload.zone]
            raise HTTPException(
                status_code=400,
                detail=f"{spec.name} {zone_name}库存不足，当前可用 {current_quantity} 只，不能扣减 {deduction} 只",
            )
    adjustment = BottleInventoryAdjustment(created_by=user.id, **adjustment_data)
    db.add(adjustment)
    db.commit()
    db.refresh(adjustment)
    return adjustment


@router.delete("/bottle-inventory-adjustments/{adjustment_id}", status_code=204)
def delete_bottle_inventory_adjustment(adjustment_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not user.is_superadmin:
        raise HTTPException(status_code=403, detail="只有超级管理员可以删除库存调整记录")
    adjustment = db.get(BottleInventoryAdjustment, adjustment_id)
    if not adjustment:
        raise HTTPException(status_code=404, detail="库存调整记录不存在")
    if adjustment.quantity > 0:
        if adjustment.zone == "BAD":
            current_quantity = manual_bad_stock(db, adjustment.source_id).get(adjustment.spec_id, 0)
        else:
            stock_key = {"GOOD": "good", "TEMPORARY": "temporary", "SCRAP": "scrap"}[adjustment.zone]
            current_quantity = bottle_outbound_availability(db, adjustment.source_id).get(str(adjustment.spec_id), {}).get(stock_key, 0)
        if adjustment.quantity > current_quantity:
            zone_name = {"GOOD": "良品", "TEMPORARY": "暂存", "SCRAP": "报废", "BAD": "不良品补录"}[adjustment.zone]
            raise HTTPException(
                status_code=400,
                detail=f"该补录库存已有部分被使用，当前 {zone_name}可用 {current_quantity} 只，不能删除此记录",
            )
    db.delete(adjustment)
    db.commit()


@router.post("/bottles/coding", response_model=BottleOut, status_code=201)
def create_coding(payload: BottleCodingCreate, user: User = Depends(require("coding:create")), db: Session = Depends(get_db)):
    inbound = db.scalar(select(BottleInbound).where(BottleInbound.id == payload.inbound_id).with_for_update())
    if not inbound or inbound.status != "POSTED":
        raise HTTPException(status_code=400, detail="入库批次不存在或尚未确认")
    line = next((item for item in inbound.lines if item.spec_id == payload.spec_id), None)
    if not line or remaining_for_line(db, line) <= 0:
        raise HTTPException(status_code=400, detail="该批次对应规格的待检测数量已全部抄码")
    source = db.get(BottleSource, payload.source_id)
    spec = db.get(BottleSpec, payload.spec_id)
    if not source or not source.is_active or not spec or not spec.is_active:
        raise HTTPException(status_code=400, detail="供瓶单位或瓶子规格不存在或已停用")
    coding_date = china_today()
    now = datetime.utcnow()
    obj = Bottle(
        **payload.model_dump(),
        inspection_price=spec.inspection_price,
        zone="GOOD",
        process_status="INSPECTED",
        inspection_result="GOOD",
        inspected_by=user.id,
        inspected_at=now,
        coding_date=coding_date,
        coding_sequence=next_coding_sequence(db, coding_date),
        coded_by=user.id,
        coded_at=now,
    )
    db.add(obj)
    save(db, "瓶身编号已存在")
    db.refresh(obj)
    return obj


@router.post("/quick-entry", response_model=BottleOut, status_code=201)
def quick_bottle_entry(payload: QuickEntryCreate, user: User = Depends(require("quick_entry:create")), db: Session = Depends(get_db)):
    spec = db.get(BottleSpec, payload.spec_id) if payload.spec_id else None
    if not spec or not spec.is_active:
        raise HTTPException(status_code=400, detail="请选择一个启用的瓶子规格")
    source = db.get(BottleSource, payload.source_id) if payload.source_id is not None else None
    if payload.source_id is not None and (not source or not source.is_active):
        raise HTTPException(status_code=400, detail="请选择一个启用的供瓶单位")
    if db.scalar(select(Bottle.id).where(Bottle.body_code == payload.body_code.strip())):
        raise HTTPException(status_code=409, detail="瓶身编号已存在")
    if source and payload.inbound_id is None:
        inbound = None
    elif payload.inbound_id:
        inbound = db.scalar(select(BottleInbound).where(BottleInbound.id == payload.inbound_id).with_for_update())
        if inbound and payload.source_id is not None and inbound.source_id != payload.source_id:
            raise HTTPException(status_code=400, detail="所选入库单与供瓶单位不一致")
        candidates = [inbound] if inbound else []
    else:
        inbound_query = select(BottleInbound).where(BottleInbound.status == "POSTED")
        if payload.source_id is not None:
            inbound_query = inbound_query.where(BottleInbound.source_id == payload.source_id)
        candidates = db.scalars(inbound_query.order_by(BottleInbound.business_date, BottleInbound.id).with_for_update()).all()
    if not (source and payload.inbound_id is None):
        inbound = next((item for item in candidates if item and item.status == "POSTED" and any(
            line.spec_id == spec.id and remaining_for_line(db, line) > 0 for line in item.lines
        )), None)
        if not inbound:
            raise HTTPException(status_code=400, detail="所选供瓶单位的待抄码池中没有该规格的剩余数量" if payload.source_id is not None else "待抄码池中没有该规格的剩余数量")
    coding_date = china_today()
    sequence = next_coding_sequence(db, coding_date)
    now = datetime.utcnow()
    bottle = Bottle(
        inbound_id=inbound.id if inbound else None,
        body_code=payload.body_code.strip(),
        manufacture_date=payload.manufacture_date,
        source_id=source.id if source else inbound.source_id,
        spec_id=spec.id,
        inspection_price=spec.inspection_price,
        zone="GOOD",
        process_status="INSPECTED",
        inspection_result="GOOD",
        inspection_note=payload.note,
        inspected_by=user.id,
        inspected_at=now,
        coding_date=coding_date,
        coding_sequence=sequence,
        coding_note=payload.note,
        good_entry_date=coding_date,
        manufacturer=(payload.manufacturer or "").strip() or None,
        coded_by=user.id,
        coded_at=now,
    )
    db.add(bottle)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        detail = str(exc.orig).lower()
        if "body_code" in detail or "ix_bottles_body_code" in detail:
            raise HTTPException(status_code=409, detail="瓶身编号已存在")
        raise HTTPException(status_code=500, detail="抄码保存失败，请刷新页面后重试")
    db.refresh(bottle)
    return bottle


@router.put("/bottles/{object_id}/coding", response_model=BottleOut)
def update_coding(object_id: int, payload: BottleCodingUpdate, _: User = Depends(require("coding:update")), db: Session = Depends(get_db)):
    obj = db.scalar(select(Bottle).where(Bottle.id == object_id).with_for_update())
    if not obj:
        raise not_found("钢瓶")
    if obj.zone == "BAD":
        raise HTTPException(status_code=400, detail="不良品钢瓶请先恢复为良品再修改抄码信息")
    spec = db.get(BottleSpec, payload.spec_id)
    source = db.get(BottleSource, payload.source_id)
    if not spec or not spec.is_active or not source or not source.is_active:
        raise HTTPException(status_code=400, detail="供瓶单位或瓶子规格不存在或已停用")
    body_code = payload.body_code.strip()
    manufacturer = (payload.manufacturer or "").strip()
    active_lines = db.scalars(
        select(InspectionDocumentLine)
        .join(InspectionDocument, InspectionDocument.id == InspectionDocumentLine.document_id)
        .where(InspectionDocumentLine.bottle_id == obj.id, InspectionDocument.status == "ACTIVE")
        .order_by(InspectionDocumentLine.id)
        .with_for_update()
    ).all()
    if active_lines and (payload.source_id != obj.source_id or payload.spec_id != obj.spec_id):
        raise HTTPException(status_code=400, detail="该钢瓶已进入检验单，只能在此更正编号、生产日期和生产厂家")
    for line in active_lines:
        document = line.document
        if payload.manufacture_date and payload.manufacture_date > document.inspection_date:
            raise HTTPException(status_code=400, detail=f"生产日期不能晚于检验单 {document.document_no} 的检验日期")
        if any(item.id != line.id and item.body_code.casefold() == body_code.casefold() for item in document.lines):
            raise HTTPException(status_code=400, detail=f"检验单 {document.document_no} 中已存在相同产品编号")
        if line.unverified_record and line.unverified_record.status == "RETRYING":
            raise HTTPException(status_code=409, detail="该钢瓶已转入重检单，请在重检单中更正")
    obj.body_code = body_code
    obj.manufacture_date = payload.manufacture_date
    obj.manufacturer = manufacturer or None
    obj.source_id = payload.source_id
    obj.spec_id = payload.spec_id
    obj.inspection_price = spec.inspection_price
    documents = {}
    for line in active_lines:
        line.body_code = body_code
        line.manufacture_date = payload.manufacture_date
        line.manufacturer = manufacturer
        documents[line.document_id] = line.document
    for document in documents.values():
        resequence_inspection_document(document)
        document.updated_at = datetime.utcnow()
    save(db, "瓶身编号已存在")
    db.refresh(obj)
    return obj


@router.delete("/bottles/{object_id}/coding", status_code=204)
def delete_coding(object_id: int, _: User = Depends(require("coding:delete")), db: Session = Depends(get_db)):
    obj = db.get(Bottle, object_id)
    if not obj:
        raise not_found("钢瓶")
    if obj.outbound_id is not None or obj.zone == "OUTBOUND":
        raise HTTPException(status_code=400, detail="该钢瓶已经出库，请先撤销对应出库单")
    if obj.inspection_document_status == "INSPECTED":
        raise HTTPException(status_code=400, detail="该钢瓶已进入检验单，不能撤销抄码")
    if obj.zone in {"BAD", "SCRAP"}:
        raise HTTPException(status_code=400, detail="不良品或已报废钢瓶不能撤销抄码")
    db.delete(obj)
    db.commit()


@router.post("/bottles/{object_id}/inspection", response_model=BottleOut)
def inspect_bottle(object_id: int, payload: InspectionCreate, user: User = Depends(require("inspection:create")), db: Session = Depends(get_db)):
    obj = db.scalar(select(Bottle).where(Bottle.id == object_id).with_for_update())
    if not obj:
        raise not_found("钢瓶")
    if obj.inspection_document_status == "INSPECTED":
        raise HTTPException(status_code=400, detail="该钢瓶已进入检验单，不能重新判定")
    if obj.process_status == "INSPECTED" and obj.zone != "GOOD":
        raise HTTPException(status_code=400, detail="只有待检测或良品钢瓶可以重新判定")
    if obj.process_status not in {"CODED", "INSPECTED"}:
        raise HTTPException(status_code=400, detail="钢瓶当前状态不能进行检测判定")
    obj.inspection_result = payload.result
    obj.inspection_note = payload.note
    obj.inspected_by = user.id
    obj.inspected_at = datetime.utcnow()
    obj.process_status = "INSPECTED"
    obj.zone = payload.result
    obj.good_entry_date = china_today() if payload.result == "GOOD" else None
    db.commit()
    db.refresh(obj)
    return obj


@router.put("/bottles/{object_id}/status", response_model=BottleOut)
def update_bottle_status(object_id: int, payload: BottleStatusUpdate, user: User = Depends(require_any("inspection:create", "bottle_lifecycle:update")), db: Session = Depends(get_db)):
    obj = db.scalar(select(Bottle).where(Bottle.id == object_id).with_for_update())
    if not obj:
        raise not_found("钢瓶")
    if obj.inspection_document_status == "INSPECTED":
        raise HTTPException(status_code=400, detail="该钢瓶已进入检验单，不能调整质量状态")
    if obj.zone not in {"GOOD", "BAD"}:
        raise HTTPException(status_code=400, detail="只有良品库和不良品库中的钢瓶可以调整状态")
    if payload.zone == "SCRAP":
        note_parts = [f"{('良品' if obj.zone == 'GOOD' else '不良品')}转报废：{obj.body_code}"]
        if payload.note:
            note_parts.append(payload.note)
        db.add(BottleInventoryAdjustment(
            business_date=china_today(),
            source_id=obj.effective_coding_customer_id,
            spec_id=obj.spec_id,
            zone="SCRAP",
            quantity=1,
            note="；".join(note_parts),
            created_by=user.id,
        ))
    if payload.returned_to_customer is True and payload.zone != "BAD":
        raise HTTPException(status_code=400, detail="仅转为不良品时可标记退还客户")
    if payload.zone == "BAD":
        obj.returned_to_customer = payload.returned_to_customer if payload.returned_to_customer is not None else (
            obj.returned_to_customer if obj.zone == "BAD" and obj.returned_to_customer is not None else True)
    elif payload.zone == "GOOD":
        obj.returned_to_customer = None
    obj.zone = payload.zone
    obj.process_status = "INSPECTED"
    obj.inspection_result = "BAD" if payload.zone == "SCRAP" else payload.zone
    obj.good_entry_date = china_today() if payload.zone == "GOOD" else None
    obj.inspection_note = payload.note
    obj.inspected_by = user.id
    obj.inspected_at = datetime.utcnow()
    db.commit()
    db.refresh(obj)
    return obj


@router.delete("/bottles/{object_id}/inspection", response_model=BottleOut)
def undo_inspection(object_id: int, _: User = Depends(require("inspection:delete")), db: Session = Depends(get_db)):
    obj = db.get(Bottle, object_id)
    if not obj:
        raise not_found("钢瓶")
    if obj.inspection_document_status == "INSPECTED":
        raise HTTPException(status_code=400, detail="该钢瓶已进入检验单，不能撤销检验")
    if obj.process_status != "INSPECTED":
        raise HTTPException(status_code=400, detail="该钢瓶尚未完成检测")
    if obj.zone == "SCRAP":
        raise HTTPException(status_code=400, detail="已报废钢瓶不能撤销检测状态")
    if obj.zone == "BAD":
        obj.inspection_result = "GOOD"
        obj.zone = "GOOD"
        obj.inspection_note = None
        obj.inspected_at = datetime.utcnow()
        obj.good_entry_date = china_today()
    else:
        obj.inspection_result = None
        obj.inspection_note = None
        obj.inspected_by = None
        obj.inspected_at = None
        obj.process_status = "CODED"
        obj.zone = "PENDING"
        obj.good_entry_date = None
    db.commit()
    db.refresh(obj)
    return obj


@router.get("/dashboard")
def dashboard(
    target_date: date = Query(default_factory=china_today),
    daily_from: date | None = None,
    daily_to: date | None = None,
    monthly_from: date | None = None,
    monthly_to: date | None = None,
    yearly_from: date | None = None,
    yearly_to: date | None = None,
    _: User = Depends(require("dashboard:view")),
    db: Session = Depends(get_db),
):
    active_inbounds = db.scalars(select(BottleInbound).where(BottleInbound.status == "POSTED")).all()
    def period_summary(date_from: date, date_to: date):
        period_inbounds = [item for item in active_inbounds if date_from <= item.business_date <= date_to]
        inbound_ids = [item.id for item in period_inbounds]
        pending = sum(item.pending_quantity for item in period_inbounds)
        linked_inspected_count = 0
        if inbound_ids:
            linked_inspected_count = db.scalar(
                select(func.count()).select_from(Bottle).where(
                    Bottle.inbound_id.in_(inbound_ids),
                    Bottle.process_status == "INSPECTED",
                )
            ) or 0
        direct_inspected_count = db.scalar(select(func.count()).select_from(Bottle).where(
            Bottle.inbound_id.is_(None),
            Bottle.coding_date >= date_from,
            Bottle.coding_date <= date_to,
            Bottle.process_status == "INSPECTED",
        )) or 0
        outbound = db.scalar(
            select(func.coalesce(func.sum(BottleOutboundLine.quantity), 0))
            .join(BottleOutbound, BottleOutbound.id == BottleOutboundLine.outbound_id)
            .where(
                BottleOutbound.status == "POSTED",
                BottleOutbound.business_date >= date_from,
                BottleOutbound.business_date <= date_to,
            )
        ) or 0
        return {
            "date_from": date_from,
            "date_to": date_to,
            "batch_count": len(period_inbounds),
            "inbound": sum(item.total_quantity for item in period_inbounds),
            "temporary": sum(item.temporary_quantity for item in period_inbounds),
            "scrap": sum(item.scrap_quantity for item in period_inbounds),
            "outbound": outbound,
            "inspected": linked_inspected_count + direct_inspected_count,
            "uninspected": max(pending - linked_inspected_count, 0),
        }

    month_from = target_date.replace(day=1)
    month_to = (date(target_date.year + (target_date.month == 12), target_date.month % 12 + 1, 1) - timedelta(days=1))
    year_from = date(target_date.year, 1, 1)
    year_to = date(target_date.year, 12, 31)
    daily_from = daily_from or target_date
    daily_to = daily_to or target_date
    monthly_from = monthly_from or month_from
    monthly_to = monthly_to or month_to
    yearly_from = yearly_from or year_from
    yearly_to = yearly_to or year_to
    for label, date_from, date_to in (
        ("日汇总", daily_from, daily_to),
        ("月汇总", monthly_from, monthly_to),
        ("年汇总", yearly_from, yearly_to),
    ):
        if date_from > date_to:
            raise HTTPException(status_code=400, detail=f"{label}开始时间不能晚于结束时间")
    today_summary = period_summary(target_date, target_date)
    daily_summary = period_summary(daily_from, daily_to)
    monthly_summary = period_summary(monthly_from, monthly_to)
    yearly_summary = period_summary(yearly_from, yearly_to)
    total_pending = sum(item.pending_quantity for item in active_inbounds)
    linked_inspected = db.scalar(select(func.count()).select_from(Bottle).where(
        Bottle.inbound_id.is_not(None), Bottle.process_status == "INSPECTED",
    )) or 0
    coded = db.scalar(select(func.count()).select_from(Bottle).where(Bottle.process_status == "CODED")) or 0
    linked_coded = db.scalar(select(func.count()).select_from(Bottle).where(
        Bottle.inbound_id.is_not(None), Bottle.process_status == "CODED",
    )) or 0
    outbound_stock = bottle_outbound_availability(db)
    low_materials = db.scalars(select(Material).where(
        Material.warning_enabled.is_(True),
        Material.safety_threshold.is_not(None),
        case(
            (Material.conversion_enabled.is_(True), Material.stock * Material.conversion_ratio),
            else_=Material.stock,
        ) < Material.safety_threshold,
        Material.is_active.is_(True),
    ).order_by(Material.stock)).all()
    return {
        "target_date": target_date,
        "today_inbound_total": today_summary["inbound"],
        "today_batch_count": today_summary["batch_count"],
        "daily_summary": daily_summary,
        "monthly_summary": monthly_summary,
        "yearly_summary": yearly_summary,
        "zones": {
            "pending": max(total_pending - linked_inspected, 0),
            "temporary": sum(item["temporary"] for item in outbound_stock.values()),
            "scrap": sum(item["scrap"] for item in outbound_stock.values()),
            "good": sum(item["good"] for item in outbound_stock.values()),
            "bad": sum(item["bad"] for item in outbound_stock.values()),
        },
        "process": {"coded_waiting": coded, "uncoded": max(total_pending - linked_inspected - linked_coded, 0)},
        "material_warning_count": len(low_materials),
        "low_materials": [MaterialOut.model_validate(item) for item in low_materials[:8]],
    }


@router.get("/profit-analysis")
def profit_analysis(
    date_from: date,
    date_to: date,
    _: User = Depends(require("profit:view")),
    db: Session = Depends(get_db),
):
    if date_from > date_to:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")

    document_period = (
        MaterialDocument.status == "POSTED",
        MaterialDocument.business_date >= date_from,
        MaterialDocument.business_date <= date_to,
    )

    purchase_rows = db.execute(
        select(
            PurchaseOrder.is_paid,
            func.count(func.distinct(PurchaseOrder.id)),
            func.coalesce(func.sum(PurchaseOrderLine.amount), 0),
        )
        .join(PurchaseOrderLine, PurchaseOrderLine.purchase_order_id == PurchaseOrder.id)
        .where(
            PurchaseOrder.status == "CONFIRMED",
            PurchaseOrder.order_date >= date_from,
            PurchaseOrder.order_date <= date_to,
        )
        .group_by(PurchaseOrder.is_paid)
    ).all()
    purchase = {
        "paid_count": 0, "paid_amount": Decimal("0"),
        "unpaid_count": 0, "unpaid_amount": Decimal("0"),
    }
    for is_paid, count, amount in purchase_rows:
        prefix = "paid" if is_paid else "unpaid"
        purchase[f"{prefix}_count"] = count
        purchase[f"{prefix}_amount"] = amount
    purchase_orders = db.scalars(
        select(PurchaseOrder).where(
            PurchaseOrder.status == "CONFIRMED",
            PurchaseOrder.order_date >= date_from,
            PurchaseOrder.order_date <= date_to,
        ).order_by(PurchaseOrder.order_date, PurchaseOrder.id)
    ).all()
    purchase["items"] = [
        {
            "order_no": item.order_no,
            "order_date": item.order_date,
            "supplier_name": item.supplier_name,
            "is_paid": item.is_paid,
            "amount": item.total_amount,
        }
        for item in purchase_orders
    ]
    # Preserve the historical financial result for legacy purchase inbounds
    # created before payment tracking moved to purchase orders.
    legacy_purchase_rows = db.execute(
        select(
            MaterialDocument.is_paid,
            func.count(func.distinct(MaterialDocument.id)),
            func.coalesce(func.sum(MaterialDocumentLine.amount), 0),
        )
        .join(MaterialDocumentLine, MaterialDocumentLine.document_id == MaterialDocument.id)
        .where(
            *document_period,
            MaterialDocument.document_type == "NORMAL_IN",
            MaterialDocument.inbound_category == "PURCHASE",
            MaterialDocument.purchase_order_id.is_(None),
        )
        .group_by(MaterialDocument.is_paid)
    ).all()
    for is_paid, count, amount in legacy_purchase_rows:
        prefix = "paid" if is_paid else "unpaid"
        purchase[f"{prefix}_count"] += count
        purchase[f"{prefix}_amount"] += amount

    def grouped_material_amount(document_type: str, category_field=None, category_value: str | None = None):
        filters = [*document_period, MaterialDocument.document_type == document_type]
        if category_field is not None and category_value is not None:
            filters.append(category_field == category_value)
        return db.execute(
            select(
                MaterialDocumentLine.material_id,
                MaterialDocumentLine.material_name,
                func.coalesce(func.sum(MaterialDocumentLine.amount), 0),
            )
            .join(MaterialDocument, MaterialDocument.id == MaterialDocumentLine.document_id)
            .where(*filters)
            .group_by(MaterialDocumentLine.material_id, MaterialDocumentLine.material_name)
            .order_by(MaterialDocumentLine.material_name)
        ).all()

    material_out_rows = grouped_material_amount("OUTBOUND", MaterialDocument.outbound_category, "MATERIAL")
    return_rows = grouped_material_amount("RETURN_IN")
    def grouped_scrap(document_type: str, category_field, category_value: str):
        return db.execute(
            select(
                MaterialDocumentLine.material_id,
                MaterialDocumentLine.material_name,
                MaterialDocumentLine.unit,
                func.coalesce(func.sum(MaterialDocumentLine.quantity), 0),
                func.coalesce(func.sum(MaterialDocumentLine.amount), 0),
            )
            .join(MaterialDocument, MaterialDocument.id == MaterialDocumentLine.document_id)
            .where(
                *document_period,
                MaterialDocument.document_type == document_type,
                category_field == category_value,
            )
            .group_by(
                MaterialDocumentLine.material_id,
                MaterialDocumentLine.material_name,
                MaterialDocumentLine.unit,
            )
            .order_by(MaterialDocumentLine.material_name)
        ).all()

    scrap_inbound_rows = grouped_scrap("NORMAL_IN", MaterialDocument.inbound_category, "SCRAP")
    scrap_outbound_rows = grouped_scrap("OUTBOUND", MaterialDocument.outbound_category, "SCRAP")

    material_map: dict[int, dict] = {}
    for material_id, material_name, amount in material_out_rows:
        material_map[material_id] = {
            "material_id": material_id, "material_name": material_name,
            "outbound_cost": amount, "return_amount": Decimal("0"),
        }
    for material_id, material_name, amount in return_rows:
        row = material_map.setdefault(material_id, {
            "material_id": material_id, "material_name": material_name,
            "outbound_cost": Decimal("0"), "return_amount": Decimal("0"),
        })
        row["return_amount"] = amount
    material_breakdown = []
    for row in sorted(material_map.values(), key=lambda item: item["material_name"]):
        row["net_cost"] = row["outbound_cost"] - row["return_amount"]
        material_breakdown.append(row)

    scrap_inbound_breakdown = [
        {"material_id": material_id, "material_name": material_name, "unit": unit, "quantity": quantity, "cost": amount}
        for material_id, material_name, unit, quantity, amount in scrap_inbound_rows
    ]
    scrap_outbound_breakdown = [
        {"material_id": material_id, "material_name": material_name, "unit": unit, "quantity": quantity, "revenue": amount}
        for material_id, material_name, unit, quantity, amount in scrap_outbound_rows
    ]
    bottle_rows = db.execute(
        select(
            BottleOutboundLine.spec_id,
            BottleSpec.name,
            BottleSpec.inspection_price,
            func.coalesce(func.sum(BottleOutboundLine.good_quantity), 0),
            func.coalesce(func.sum(BottleOutboundLine.good_quantity * BottleSpec.inspection_price), 0),
        )
        .join(BottleOutbound, BottleOutbound.id == BottleOutboundLine.outbound_id)
        .join(BottleSpec, BottleSpec.id == BottleOutboundLine.spec_id)
        .where(
            BottleOutbound.status == "POSTED",
            BottleOutbound.business_date >= date_from,
            BottleOutbound.business_date <= date_to,
            BottleOutboundLine.good_quantity > 0,
        )
        .group_by(BottleOutboundLine.spec_id, BottleSpec.name, BottleSpec.inspection_price)
        .order_by(BottleSpec.name, BottleSpec.inspection_price)
    ).all()
    bottle_breakdown = [
        {
            "spec_id": spec_id, "spec_name": spec_name, "inspection_price": inspection_price,
            "quantity": quantity, "revenue": revenue,
        }
        for spec_id, spec_name, inspection_price, quantity, revenue in bottle_rows
    ]

    purchase_cost = purchase["paid_amount"]
    material_outbound_cost = sum((row["outbound_cost"] for row in material_breakdown), Decimal("0"))
    material_return_amount = sum((row["return_amount"] for row in material_breakdown), Decimal("0"))
    net_material_cost = material_outbound_cost - material_return_amount
    scrap_inbound_amount = sum((row["cost"] for row in scrap_inbound_breakdown), Decimal("0"))
    scrap_revenue = sum((row["revenue"] for row in scrap_outbound_breakdown), Decimal("0"))
    inspection_revenue = sum((row["revenue"] for row in bottle_breakdown), Decimal("0"))
    total_cost = purchase_cost + net_material_cost
    total_revenue = scrap_revenue + inspection_revenue
    return {
        "date_from": date_from,
        "date_to": date_to,
        "purchase": purchase,
        "materials": {
            "outbound_cost": material_outbound_cost,
            "return_amount": material_return_amount,
            "net_cost": net_material_cost,
            "items": material_breakdown,
        },
        "scrap": {
            "inbound_amount": scrap_inbound_amount,
            # Legacy field retained for older clients; no longer counted as a cost.
            "inbound_cost": scrap_inbound_amount,
            "outbound_revenue": scrap_revenue,
            "profit": scrap_revenue,
            "inbound_items": scrap_inbound_breakdown,
            "outbound_items": scrap_outbound_breakdown,
            # Kept for clients using the previous response shape.
            "revenue": scrap_revenue,
            "items": scrap_outbound_breakdown,
        },
        "inspection": {
            "quantity": sum(row["quantity"] for row in bottle_breakdown),
            "revenue": inspection_revenue,
            "items": bottle_breakdown,
        },
        "totals": {
            "cost": total_cost,
            "revenue": total_revenue,
            "profit": total_revenue - total_cost,
        },
        "calculation": {
            "cost": "总成本 = 已付款采购单金额 + 原材料净耗用成本（废料入库不计成本）",
            "revenue": "总营收 = 废料出库收入 + 良品钢瓶出库检测费",
            "profit": "区间盈利 = 总营收 - 总成本",
            "scrap": "废料收益 = 废料出库明细的数量 × 出库单价之和",
        },
    }
