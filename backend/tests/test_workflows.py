import os
import io
import base64
import zipfile
import xlrd
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import unquote
from openpyxl import Workbook
import pytest


TEST_DB = Path(__file__).with_name("test_ailitewms.db")
if TEST_DB.exists():
    TEST_DB.unlink()
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB}"
os.environ["JWT_SECRET"] = "test-secret"

from fastapi.testclient import TestClient

from backend.app.main import app


def request(client: TestClient, method: str, path: str, token: str, **kwargs):
    headers = kwargs.pop("headers", {})
    headers["Authorization"] = f"Bearer {token}"
    return client.request(method, path, headers=headers, **kwargs)


def test_first_phase_workflows():
    with TestClient(app) as client:
        login = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"})
        assert login.status_code == 200
        token = login.json()["access_token"]
        users_page = request(client, "GET", "/api/users?page=1&page_size=20", token).json()
        assert users_page["total"] == 0
        assert users_page["items"] == []
        backup_login = client.post("/api/auth/login", json={"username": "admin_backup", "password": "BackupAdmin@123456"})
        assert backup_login.status_code == 200

        material = request(client, "POST", "/api/materials", token, json={
            "name": "银灰粉末", "specification": "银灰色", "material_unit": "瓶", "stock_unit": "kg",
            "conversion_enabled": True, "conversion_ratio": 2.5, "price": 11.9,
            "warning_enabled": True, "safety_threshold": 100, "is_active": True,
        })
        assert material.status_code == 201
        material_id = material.json()["id"]
        scrap_material = request(client, "POST", "/api/materials", token, json={
            "name": "回收废料", "material_type": "SCRAP", "material_unit": "kg", "price": 0,
        })
        assert scrap_material.status_code == 201
        scrap_material_id = scrap_material.json()["id"]

        supplier = request(client, "POST", "/api/suppliers", token, json={
            "name": "潍坊材料供应商", "is_active": True,
            "material_prices": [{"material_id": material_id, "purchase_price": 11.9}],
        })
        assert supplier.status_code == 201
        assert float(supplier.json()["material_prices"][0]["purchase_price"]) == 11.9
        source = request(client, "POST", "/api/bottle-sources", token, json={"name": "安丘送检单位", "is_default": True, "sort_order": 1, "is_active": True})
        assert source.status_code == 201
        source_id = source.json()["id"]
        spec = request(client, "POST", "/api/bottle-specs", token, json={"name": "15kg", "inspection_price": 30, "volume": 35.5, "wall_thickness": 2.45, "is_default": True, "sort_order": 1, "is_active": True})
        assert spec.status_code == 201
        assert float(spec.json()["volume"]) == 35.5
        assert float(spec.json()["wall_thickness"]) == 2.45
        spec_id = spec.json()["id"]

        purchase_order = request(client, "POST", "/api/purchase-orders", token, json={
            "order_date": "2026-09-09", "supplier_id": supplier.json()["id"], "handler": "采购员",
            "lines": [{"material_id": material_id, "quantity": 120}],
        })
        assert purchase_order.status_code == 201
        assert float(purchase_order.json()["total_amount"]) == 1428
        purchase_order_id = purchase_order.json()["id"]
        purchase_order = request(client, "POST", f"/api/purchase-orders/{purchase_order_id}/confirm", token)
        assert purchase_order.status_code == 200
        assert purchase_order.json()["status"] == "CONFIRMED"
        paid_order = request(client, "PUT", f"/api/purchase-orders/{purchase_order_id}/payment", token, json={"is_paid": True})
        assert paid_order.status_code == 200
        assert paid_order.json()["is_paid"] is True

        normal_in = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "NORMAL_IN", "business_date": "2026-09-09", "purchase_order_id": purchase_order_id,
            "inbound_mode": "PURCHASE_ORDER", "inbound_category": "MATERIAL", "handler": "仓库员",
        })
        assert normal_in.status_code == 201
        assert normal_in.json()["inbound_mode"] == "PURCHASE_ORDER"
        assert normal_in.json()["inbound_category"] == "MATERIAL"
        assert normal_in.json()["purchase_order_id"] == purchase_order_id
        assert normal_in.json()["supplier_id"] == supplier.json()["id"]
        assert float(normal_in.json()["lines"][0]["amount"]) == 1428
        assert normal_in.json()["is_paid"] is False
        assert normal_in.json()["created_by_name"] == "超级管理员"
        normal_in_id = normal_in.json()["id"]
        normal_in = request(client, "PUT", f"/api/material-documents/{normal_in_id}", token, json={
            "document_type": "NORMAL_IN", "business_date": "2026-09-09", "purchase_order_id": purchase_order_id,
            "inbound_mode": "PURCHASE_ORDER", "inbound_category": "MATERIAL", "handler": "仓库员", "note": "已复核",
        })
        assert normal_in.status_code == 200
        assert normal_in.json()["updated_by_name"] == "超级管理员"
        assert request(client, "GET", "/api/material-documents", token, params={"document_type": "NORMAL_IN", "is_paid": False}).status_code == 400
        assert request(client, "GET", "/api/purchase-orders", token, params={"is_paid": True}).json()["total"] == 1
        assert request(client, "GET", "/api/material-documents", token, params={
            "document_type": "NORMAL_IN", "date_from": "2026-10-01", "date_to": "2026-10-31",
        }).json()["total"] == 0
        assert request(client, "GET", "/api/material-documents", token, params={
            "document_type": "NORMAL_IN", "date_from": "2026-10-01", "date_to": "2026-09-01",
        }).status_code == 400
        normal_in = request(client, "POST", f"/api/material-documents/{normal_in_id}/post", token)
        assert normal_in.status_code == 200
        assert normal_in.json()["posted_by_name"] == "超级管理员"
        assert request(client, "PUT", f"/api/material-documents/{normal_in_id}/payment", token, json={"is_paid": True}).status_code == 400
        duplicate_inbound = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "NORMAL_IN", "business_date": "2026-09-09", "purchase_order_id": purchase_order_id,
            "inbound_mode": "PURCHASE_ORDER", "inbound_category": "MATERIAL", "handler": "仓库员",
        })
        assert duplicate_inbound.status_code == 400
        scrap_inbound = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "NORMAL_IN", "business_date": "2026-09-10", "inbound_category": "SCRAP",
            "handler": "废料库管", "lines": [{"material_id": scrap_material_id, "quantity": 2, "unit_price": 0}],
        })
        assert scrap_inbound.status_code == 201
        assert scrap_inbound.json()["inbound_mode"] == "DIRECT"
        assert scrap_inbound.json()["inbound_category"] == "SCRAP"
        assert scrap_inbound.json()["is_paid"] is False
        assert request(client, "GET", "/api/material-documents", token, params={
            "document_type": "NORMAL_IN", "inbound_category": "SCRAP",
        }).json()["total"] == 1
        invalid_purchase_scrap = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "NORMAL_IN", "business_date": "2026-09-10", "purchase_order_id": purchase_order_id,
            "inbound_mode": "PURCHASE_ORDER", "inbound_category": "SCRAP", "handler": "仓库员",
        })
        assert invalid_purchase_scrap.status_code == 422
        invalid_scrap_material = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "NORMAL_IN", "business_date": "2026-09-10", "inbound_category": "SCRAP",
            "handler": "废料库管", "lines": [{"material_id": material_id, "quantity": 1, "unit_price": 0}],
        })
        assert invalid_scrap_material.status_code == 400
        assert request(client, "GET", "/api/material-documents", token, params={
            "document_type": "NORMAL_IN", "is_paid": False,
        }).status_code == 400
        assert request(client, "PUT", f"/api/material-documents/{scrap_inbound.json()['id']}/payment", token, json={"is_paid": True}).status_code == 400
        assert request(client, "DELETE", f"/api/material-documents/{scrap_inbound.json()['id']}", token).status_code == 204

        outbound = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "OUTBOUND", "business_date": "2026-09-09", "handler": "领料员",
            "lines": [{"material_id": material_id, "quantity": 30, "unit_price": 11.9}],
        })
        assert outbound.status_code == 201
        assert outbound.json()["outbound_category"] == "MATERIAL"
        outbound_id = outbound.json()["id"]
        assert request(client, "POST", f"/api/material-documents/{outbound_id}/post", token).status_code == 200

        scrap_outbound = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "OUTBOUND", "outbound_category": "SCRAP", "business_date": "2026-09-09", "handler": "废料销售员",
            "lines": [{"material_id": material_id, "quantity": 1, "unit_price": 6.5}],
        })
        assert scrap_outbound.status_code == 201
        assert scrap_outbound.json()["outbound_category"] == "SCRAP"
        scrap_outbounds = request(client, "GET", "/api/material-documents", token, params={
            "document_type": "OUTBOUND", "outbound_category": "SCRAP",
            "date_from": "2026-09-01", "date_to": "2026-09-30",
        })
        assert scrap_outbounds.status_code == 200
        assert scrap_outbounds.json()["total"] == 1
        assert request(client, "GET", "/api/material-documents", token, params={
            "document_type": "OUTBOUND", "date_from": "2026-10-01", "date_to": "2026-10-31",
        }).json()["total"] == 0
        assert request(client, "DELETE", f"/api/material-documents/{scrap_outbound.json()['id']}", token).status_code == 204

        exceeds_original = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "RETURN_IN", "business_date": "2026-09-09", "related_outbound_id": outbound_id,
            "handler": "领料员", "lines": [{"material_id": material_id, "quantity": 30.01, "unit_price": 11.9}],
        })
        assert exceeds_original.status_code == 400
        assert "剩余可退数量 30.00" in exceeds_original.json()["detail"]

        returned = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "RETURN_IN", "business_date": "2026-09-09", "related_outbound_id": outbound_id,
            "handler": "领料员", "lines": [{"material_id": material_id, "quantity": 10, "unit_price": 11.9}],
        })
        assert returned.status_code == 201
        returned_in_range = request(client, "GET", "/api/material-documents", token, params={
            "document_type": "RETURN_IN", "date_from": "2026-09-01", "date_to": "2026-09-30",
        })
        assert returned_in_range.status_code == 200
        assert returned_in_range.json()["total"] == 1
        assert request(client, "GET", "/api/material-documents", token, params={
            "document_type": "RETURN_IN", "date_from": "2026-10-01", "date_to": "2026-10-31",
        }).json()["total"] == 0
        assert request(client, "POST", f"/api/material-documents/{returned.json()['id']}/post", token).status_code == 200
        materials_page = request(client, "GET", "/api/materials?page=1&page_size=20", token).json()
        assert materials_page["total"] == 2
        material_row = next(item for item in materials_page["items"] if item["id"] == material_id)
        assert material_row["material_type"] == "MATERIAL"
        assert float(material_row["stock"]) == 100
        assert float(material_row["converted_stock"]) == 250
        assert material_row["material_unit"] == "瓶"
        assert material_row["stock_unit"] == "kg"

        exported = request(client, "GET", "/api/materials/export", token)
        assert exported.status_code == 200
        assert "类型" in exported.text
        assert "物料单位" in exported.text
        imported = request(client, "POST", "/api/materials/import", token, files={
            "file": ("materials.csv", "名称,规格,物料单位,库存单位,换算比例,参考价格,启用预警,安全库存阈值,档案状态,备注\n清洗液,20L桶,桶,L,20,80,是,100,启用,测试导入\n", "text/csv"),
        })
        assert imported.status_code == 200
        assert imported.json()["created"] == 1

        plain = request(client, "POST", "/api/materials", token, json={
            "name": "普通包装材料", "material_unit": "箱", "price": 20,
        })
        assert plain.status_code == 201
        assert plain.json()["conversion_enabled"] is False
        assert plain.json()["converted_stock"] is None
        assert plain.json()["stock_unit"] == "箱"

        excessive_return = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "RETURN_IN", "business_date": "2026-09-09", "related_outbound_id": outbound_id,
            "handler": "领料员", "lines": [{"material_id": material_id, "quantity": 21, "unit_price": 11.9}],
        })
        assert excessive_return.status_code == 400
        assert "剩余可退数量 20.00" in excessive_return.json()["detail"]
        assert request(client, "POST", f"/api/material-documents/{outbound_id}/void", token).status_code == 400
        returned = request(client, "POST", f"/api/material-documents/{returned.json()['id']}/void", token)
        assert returned.status_code == 200
        assert returned.json()["voided_by_name"] == "超级管理员"

        inbound = request(client, "POST", "/api/bottle-inbounds", token, json={
            "business_date": "2026-09-09", "source_id": source_id, "handler": "接瓶员",
            "lines": [{"spec_id": spec_id, "pending_quantity": 4, "temporary_quantity": 3, "scrap_quantity": 1}],
        })
        assert inbound.status_code == 201
        assert inbound.json()["lines"][0]["spec_name"] == "15kg"
        inbound_id = inbound.json()["id"]
        inbound_no = inbound.json()["document_no"]
        for query in (
            {"date_from": "2026-09-09", "date_to": "2026-09-09"},
            {"source_id": source_id},
            {"document_no": inbound_no[-6:]},
            {"handler": "接瓶"},
        ):
            filtered_inbounds = request(client, "GET", "/api/bottle-inbounds", token, params=query)
            assert filtered_inbounds.status_code == 200
            assert filtered_inbounds.json()["total"] == 1
            assert filtered_inbounds.json()["items"][0]["id"] == inbound_id
        no_inbounds = request(client, "GET", "/api/bottle-inbounds", token, params={"handler": "不存在的经办人"})
        assert no_inbounds.status_code == 200
        assert no_inbounds.json()["total"] == 0
        before_range = request(client, "GET", "/api/bottle-inbounds", token, params={"date_to": "2026-09-08"})
        assert before_range.status_code == 200
        assert before_range.json()["total"] == 0
        invalid_range = request(client, "GET", "/api/bottle-inbounds", token, params={"date_from": "2026-09-10", "date_to": "2026-09-09"})
        assert invalid_range.status_code == 400

        purge_candidate = request(client, "POST", "/api/bottle-inbounds", token, json={
            "business_date": "2026-09-08", "source_id": source_id, "handler": "纠错测试",
            "lines": [{"spec_id": spec_id, "pending_quantity": 1, "temporary_quantity": 0, "scrap_quantity": 0}],
        })
        purge_id = purge_candidate.json()["id"]
        assert request(client, "POST", f"/api/bottle-inbounds/{purge_id}/post", token).status_code == 200
        assert request(client, "POST", f"/api/bottle-inbounds/{purge_id}/void", token).status_code == 200
        inbound_operator = request(client, "POST", "/api/users", token, json={
            "username": "inbound_operator", "name": "入库管理员", "role_name": "入库管理员", "password": "Inbound@123",
            "permissions": ["bottle_inbound:view", "bottle_inbound:delete"],
        })
        assert inbound_operator.status_code == 201
        inbound_operator_login = client.post("/api/auth/login", json={"username": "inbound_operator", "password": "Inbound@123"})
        inbound_operator_token = inbound_operator_login.json()["access_token"]
        forbidden_purge = request(client, "DELETE", f"/api/bottle-inbounds/{purge_id}", inbound_operator_token)
        assert forbidden_purge.status_code == 403
        assert request(client, "DELETE", f"/api/bottle-inbounds/{purge_id}", token).status_code == 204
        assert request(client, "GET", "/api/bottle-inbounds", token, params={"date_from": "2026-09-08", "date_to": "2026-09-08"}).json()["total"] == 0
        posted_inbound = request(client, "POST", f"/api/bottle-inbounds/{inbound_id}/post", token)
        assert posted_inbound.status_code == 200
        inbound_posted_at = posted_inbound.json()["posted_at"]

        bottle_ids = []
        for index in range(2):
            coded = request(client, "POST", "/api/bottles/coding", token, json={
                "inbound_id": inbound_id, "body_code": f"JN-2026-{index + 1:04d}",
                "manufacture_date": "2020-06-01", "source_id": source_id, "spec_id": spec_id,
            })
            assert coded.status_code == 201
            bottle_ids.append(coded.json()["id"])
        inspected = request(client, "POST", f"/api/bottles/{bottle_ids[0]}/inspection", token, json={"result": "GOOD", "note": "检测合格"})
        assert inspected.status_code == 200
        assert inspected.json()["zone"] == "GOOD"
        zones = request(client, "GET", "/api/bottle-zones", token).json()
        assert zones["zones"] == {"pending": 2, "temporary": 3, "scrap": 1, "good": 2, "bad": 0}
        inbound_count = request(client, "GET", "/api/bottle-inbounds", token).json()["total"]
        manufacturer = request(client, "POST", "/api/bottle-manufacturers", token, json={
            "name": "济南钢瓶厂", "code_prefix": "quick-, qk；quick-", "sort_order": 1, "is_active": True,
        })
        assert manufacturer.status_code == 201
        assert manufacturer.json()["code_prefix"] == "QUICK-,QK"
        assert request(client, "POST", "/api/bottle-manufacturers", token, json={"name": "济南钢瓶厂"}).status_code == 409
        updated_manufacturer = request(client, "PUT", f"/api/bottle-manufacturers/{manufacturer.json()['id']}", token, json={
            "name": "济南钢瓶厂", "code_prefix": "QUICK- QK JN-", "sort_order": 2, "is_active": True,
        })
        assert updated_manufacturer.status_code == 200
        manufacturer_option = request(client, "GET", "/api/options", token).json()["manufacturers"][0]
        assert manufacturer_option["name"] == "济南钢瓶厂"
        assert manufacturer_option["code_prefix"] == "QUICK-,QK,JN-"
        inspector_options = request(client, "GET", "/api/options", token).json()["users"]
        assert [item["name"] for item in inspector_options] == ["入库管理员"]
        assert inspector_options[0]["is_active"] is True
        assert all(not item["is_superadmin"] for item in inspector_options)
        for unusual_length_code in ("12345678901", "1234567890123"):
            unusual_entry = request(client, "POST", "/api/quick-entry", token, json={
                "source_id": source_id, "spec_id": spec_id, "body_code": unusual_length_code,
            })
            assert unusual_entry.status_code == 201
            assert request(client, "DELETE", f"/api/bottles/{unusual_entry.json()['id']}/coding", token).status_code == 204
        quick = request(client, "POST", "/api/quick-entry", token, json={
            "inbound_id": inbound_id, "source_id": source_id, "spec_id": spec_id, "body_code": "QUICK-2026-0001",
            "manufacture_date": "2021-08-15", "note": "移动端连续录入", "manufacturer": "济南钢瓶厂",
        })
        assert quick.status_code == 201
        assert quick.json()["source_id"] == source_id
        assert quick.json()["spec_id"] == spec_id
        assert quick.json()["manufacturer"] == "济南钢瓶厂"
        assert quick.json()["zone"] == "GOOD"
        assert quick.json()["process_status"] == "INSPECTED"
        assert quick.json()["coding_sequence"] == 3
        suspect_candidate = request(client, "POST", "/api/quick-entry", token, json={
            "source_id": source_id, "spec_id": spec_id, "body_code": "SUSPECT-TOO-LONG-001",
        })
        assert suspect_candidate.status_code == 201
        suspects = request(client, "GET", "/api/coding/today-bottles", token)
        assert suspects.status_code == 200
        assert suspect_candidate.json()["id"] in {item["id"] for item in suspects.json()}
        suspect_flags = [
            len(item["body_code"].strip()) > 13
            or not (item["manufacturer"] or "").strip()
            or item["manufacture_date"] is None
            for item in suspects.json()
        ]
        assert suspect_flags == sorted(suspect_flags, reverse=True)
        review_items = []
        for review_page in range(1, len(suspects.json()) + 1):
            review = request(client, "GET", "/api/coding/review-bottles", token, params={"page": review_page, "page_size": 1})
            assert review.status_code == 200, review.text
            assert review.json()["total"] == len(suspects.json())
            assert review.json()["suspect_count"] == sum(suspect_flags)
            review_items.extend(review.json()["items"])
        assert [item["id"] for item in review_items] == [item["id"] for item in suspects.json()]
        review_date = suspect_candidate.json()["coding_date"]
        filtered_review = request(client, "GET", "/api/coding/review-bottles", token, params={
            "date_from": review_date, "date_to": review_date, "body_code": "  suspect-too-long  ",
        }).json()
        assert filtered_review["total"] == filtered_review["suspect_count"] == 1
        assert filtered_review["items"][0]["id"] == suspect_candidate.json()["id"]
        customer_review = request(client, "GET", "/api/coding/review-bottles", token, params={"coding_customer_id": source_id}).json()
        assert customer_review["total"] == len(suspects.json())
        assert all(item["effective_coding_customer_id"] == source_id for item in customer_review["items"])
        assert request(client, "GET", "/api/coding/review-bottles", token, params={"coding_customer_id": 999999}).json()["total"] == 0
        assert request(client, "GET", "/api/coding/review-bottles", token, params={"body_code": "NO-MATCH-REVIEW"}).json()["total"] == 0
        assert request(client, "GET", "/api/coding/review-bottles", token, params={"date_from": "2000-01-01", "date_to": "2000-01-02"}).json()["total"] == 0
        assert request(client, "GET", "/api/coding/review-bottles", token, params={"date_from": "2026-02-01", "date_to": "2026-01-01"}).status_code == 422
        corrected_suspect = request(client, "PUT", f"/api/bottles/{suspect_candidate.json()['id']}/coding", token, json={
            "body_code": "SUSPECT-001", "manufacture_date": "2020-01-01", "manufacturer": "济南钢瓶厂",
            "source_id": source_id, "spec_id": spec_id,
        })
        assert corrected_suspect.status_code == 200, corrected_suspect.text
        assert corrected_suspect.json()["manufacturer"] == "济南钢瓶厂"
        corrected_today_rows = request(client, "GET", "/api/coding/today-bottles", token).json()
        corrected_today = next(item for item in corrected_today_rows if item["id"] == suspect_candidate.json()["id"])
        assert corrected_today["body_code"] == "SUSPECT-001" and corrected_today["manufacturer"] == "济南钢瓶厂"
        assert corrected_today["manufacture_date"] == "2020-01-01"
        assert request(client, "DELETE", f"/api/bottles/{suspect_candidate.json()['id']}/coding", token).status_code == 204
        other_source = request(client, "POST", "/api/bottle-sources", token, json={
            "name": "其他供瓶单位", "is_active": True,
        })
        assert other_source.status_code == 201
        assert request(client, "POST", "/api/quick-entry", token, json={
            "inbound_id": inbound_id, "source_id": other_source.json()["id"], "spec_id": spec_id, "body_code": "OTHER-SOURCE-BATCH",
        }).json()["detail"] == "所选入库单与供瓶单位不一致"
        assert request(client, "POST", "/api/quick-entry", token, json={
            "source_id": 99999, "spec_id": spec_id, "body_code": "OTHER-SOURCE-POOL",
        }).status_code == 400
        assert request(client, "GET", "/api/bottle-inbounds", token).json()["total"] == inbound_count
        duplicate = request(client, "POST", "/api/quick-entry", token, json={
            "spec_id": spec_id, "body_code": "QUICK-2026-0001", "manufacture_date": "2021-08-15",
        })
        assert duplicate.status_code == 409

        without_manufacture_date = request(client, "POST", "/api/quick-entry", token, json={
            "spec_id": spec_id, "body_code": "QUICK-2026-0002", "note": "生产日期后补",
        })
        assert without_manufacture_date.status_code == 201
        assert without_manufacture_date.json()["manufacture_date"] is None
        assert without_manufacture_date.json()["coding_sequence"] == 4
        assert request(client, "DELETE", f"/api/bottles/{without_manufacture_date.json()['id']}/coding", token).status_code == 204
        restored_pool = request(client, "GET", "/api/coding/pool", token)
        assert restored_pool.status_code == 200
        assert restored_pool.json()["total_remaining"] == 1
        without_manufacture_date = request(client, "POST", "/api/quick-entry", token, json={
            "spec_id": spec_id, "body_code": "QUICK-2026-0002", "note": "纠错后重新录入",
        })
        assert without_manufacture_date.status_code == 201
        assert without_manufacture_date.json()["coding_sequence"] == 4
        overflow = request(client, "POST", "/api/bottles/coding", token, json={
            "inbound_id": inbound_id, "body_code": "JN-2026-9999", "manufacture_date": "2020-06-01",
            "source_id": source_id, "spec_id": spec_id,
        })
        assert overflow.status_code == 400
        pool = request(client, "GET", "/api/coding/pool", token)
        assert pool.status_code == 200
        assert pool.json()["total_remaining"] == 0
        coding_summary = request(client, "GET", "/api/coding/summary?coded_by=1", token)
        assert coding_summary.status_code == 200
        assert coding_summary.json()["total"] == 4
        inbound_report = request(client, "GET", "/api/bottle-inbounds/report?business_date=2026-09-09", token)
        assert inbound_report.status_code == 200
        assert inbound_report.json()[0]["document_no"] == inbound.json()["document_no"]
        shanghai_today = (datetime.utcnow() + timedelta(hours=8)).date().isoformat()
        good_bottles = request(client, "GET", "/api/coding/good-bottles", token, params={"date": shanghai_today})
        assert good_bottles.status_code == 200
        quick_export = next(row for row in good_bottles.json() if row["body_code"] == "QUICK-2026-0001")
        assert quick_export["manufacturer"] == "济南钢瓶厂"
        plain_export = next(row for row in good_bottles.json() if row["body_code"] == "QUICK-2026-0002")
        assert plain_export["manufacturer"] is None

        marked_bad = request(client, "POST", f"/api/bottles/{quick.json()['id']}/inspection", token, json={"result": "BAD", "note": "复判不良"})
        assert marked_bad.status_code == 200
        assert marked_bad.json()["zone"] == "BAD"
        restored = request(client, "DELETE", f"/api/bottles/{quick.json()['id']}/inspection", token)
        assert restored.status_code == 200
        assert restored.json()["zone"] == "GOOD"
        marked_bad_again = request(client, "PUT", f"/api/bottles/{quick.json()['id']}/status", token, json={"zone": "BAD", "note": "确认不良"})
        assert marked_bad_again.status_code == 200
        scrapped_bottle = request(client, "PUT", f"/api/bottles/{quick.json()['id']}/status", token, json={"zone": "SCRAP", "note": "无法修复"})
        assert scrapped_bottle.status_code == 200
        assert scrapped_bottle.json()["zone"] == "SCRAP"
        assert scrapped_bottle.json()["inspection_result"] == "BAD"
        assert request(client, "PUT", f"/api/bottles/{quick.json()['id']}/status", token, json={"zone": "SCRAP"}).status_code == 400
        zones_after_scrap = request(client, "GET", "/api/bottle-zones", token).json()
        assert zones_after_scrap["zones"]["scrap"] == 2
        assert zones_after_scrap["adjustments"][0]["zone"] == "SCRAP"
        assert zones_after_scrap["adjustments"][0]["quantity"] == 1

        inspection_document = request(client, "POST", "/api/inspection-documents", token, json={"inspection_date": shanghai_today})
        assert inspection_document.status_code == 201
        inspection_document_data = inspection_document.json()
        assert inspection_document_data["status"] == "ACTIVE"
        assert inspection_document_data["total_quantity"] == 2
        assert [line["sequence"] for line in inspection_document_data["lines"]] == [1, 2]
        assert inspection_document_data["lines"][0]["manufacture_date"] == "2020-06-01"
        assert inspection_document_data["lines"][1]["manufacture_date"] is None
        assert "QUICK-2026-0002" in [line["body_code"] for line in inspection_document_data["lines"]]
        coded_times = {item["body_code"]: item["coded_at"] for item in request(client, "GET", "/api/bottles", token).json()["items"]}
        assert all(line["inbound_at"] == coded_times[line["body_code"]] for line in inspection_document_data["lines"])
        assert all(line["inbound_at"] != inbound_posted_at for line in inspection_document_data["lines"])
        assert all(line["inspection_time"] is None for line in inspection_document_data["lines"])
        assert all("inspected_at" not in line for line in inspection_document_data["lines"])
        assert request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-xls?kind=full", token).status_code == 405
        same_daily_document = request(client, "POST", "/api/inspection-documents", token, json={"inspection_date": shanghai_today})
        assert same_daily_document.json()["id"] == inspection_document_data["id"]
        updated_inspection_document = request(client, "PUT", f"/api/inspection-documents/{inspection_document_data['id']}", token, json={
            "customer_name": "安丘送检单位", "inspector_id": inbound_operator.json()["id"],
        })
        assert updated_inspection_document.status_code == 200
        assert updated_inspection_document.json()["customer_name"] == "安丘送检单位"
        assert updated_inspection_document.json()["inspector_name"] == "入库管理员"
        receipt_word = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-receipt-word", token)
        assert receipt_word.status_code == 200
        assert receipt_word.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert "来瓶登记[2].docx" in unquote(receipt_word.headers["content-disposition"])
        from docx import Document
        receipt = Document(io.BytesIO(receipt_word.content))
        assert receipt.paragraphs[0].text == "液化石油气钢瓶接收记录表"
        assert receipt.tables[0].cell(0, 0).text == "检验单位名称：安丘送检单位"
        assert receipt.tables[1].cell(1, 1).text == "锦安特种设备检测有限公司"
        from docx.oxml.ns import qn
        assert receipt.tables[1]._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")) == receipt.tables[2]._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))
        assert len(receipt.tables[2].rows) == 3
        assert receipt.tables[2].cell(1, 1).text == updated_inspection_document.json()["lines"][0]["body_code"]
        assert receipt.tables[2].cell(1, 4).text == receipt.tables[2].cell(1, 6).text == "2020.06"
        assert receipt.tables[2].cell(2, 6).text == ""
        assert receipt.tables[2].cell(1, 7).text == "√裂纹  √凹陷  √鼓包  √腐蚀  √变形\n√钢印清晰  √检验标志有效  √使用登记证"
        detail_after_word = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
        assert "xls:full" not in detail_after_word["export_dates"]
        assert all(line["inspection_time"] is None for line in detail_after_word["lines"])
        for kind, quantity in (("full", 2), ("verified", 2)):
            exported = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-xls?kind={kind}&export_date=2031-10-05", token)
            assert exported.status_code == 200
            assert exported.headers["content-type"].startswith("application/vnd.ms-excel")
            assert exported.content.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
            assert f"[{quantity}].xls" in unquote(exported.headers["content-disposition"])
            sheet = xlrd.open_workbook(file_contents=exported.content).sheet_by_index(0)
            assert sheet.nrows == quantity + 2
            assert [sheet.cell_value(row, 6) for row in range(2, sheet.nrows)] == ["2031.10.05"] * quantity
            assert [sheet.cell_value(row, 7) for row in range(2, sheet.nrows)] == ["2035.10.04"] * quantity
            assert "20311005" in unquote(exported.headers["content-disposition"])
            if kind == "full":
                full_sheet = sheet
            exported_detail = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
            assert all(line["inspection_time"] == "2031-10-05" for line in exported_detail["lines"])
        removed_age_xls = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-xls?kind=age-split", token)
        assert removed_age_xls.status_code == 422
        empty_unverified_xls = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-xls?kind=unverified", token)
        assert empty_unverified_xls.status_code == 400
        import xlwt
        revised_book = xlwt.Workbook()
        revised_sheet = revised_book.add_sheet("检验气瓶列表")
        for row_index in range(full_sheet.nrows):
            for column_index, value in enumerate(full_sheet.row_values(row_index)):
                if row_index == 2 and column_index == 4:
                    value = "回导修正厂家"
                revised_sheet.write(row_index, column_index, value)
        revised_content = io.BytesIO()
        revised_book.save(revised_content)
        reimported = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/reimport", token,
            files={"file": ("revised.xls", revised_content.getvalue(), "application/vnd.ms-excel")})
        assert reimported.status_code == 200, reimported.text
        assert reimported.json()["lines"][0]["manufacturer"] == "回导修正厂家"
        first_bottle = next(item for item in request(client, "GET", "/api/bottles", token).json()["items"]
                            if item["id"] == reimported.json()["lines"][0]["bottle_id"])
        assert first_bottle["manufacturer"] == "回导修正厂家"
        regular_word = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-word?kind=regular&report_date=2031-10-05", token)
        assessment_word = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-word?kind=assessment&report_date=2031-10-05", token)
        assert regular_word.status_code == 200
        assert assessment_word.status_code == 200
        assert "定检[0].docx" in unquote(regular_word.headers["content-disposition"])
        assert "评定[2].docx" in unquote(assessment_word.headers["content-disposition"])
        assert regular_word.content[:2] == b"PK"
        assert assessment_word.content[:2] == b"PK"
        for report, expected_rows, appendix_title in (
            (regular_word, 1, "附表:定期检验合格气瓶一览表"),
            (assessment_word, 3, "附表:安全评定合格气瓶一览表"),
        ):
            with zipfile.ZipFile(io.BytesIO(report.content)) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
                assert "WFJA-31-1005" in xml
                assert xml.count("<w:tr>") == expected_rows + 1
                assert appendix_title in xml
                if expected_rows > 1:
                    assert xml.count("2035.10.04") == expected_rows - 1
        templates = request(client, "GET", "/api/inspection-word-templates",
            token, params={"customer_name": "安丘送检单位"})
        assert templates.json()["selected_key"] == "default"
        assert {item["key"] for item in templates.json()["templates"]} == {"default", "photo"}
        chosen = request(client, "PUT", "/api/inspection-word-templates/choice",
            token, params={"customer_name": "安丘送检单位"}, json={"template_key": "photo"})
        assert chosen.status_code == 200, chosen.text
        photo_word = request(client, "GET",
            f"/api/inspection-documents/{inspection_document_data['id']}/export-word?kind=regular&report_date=2031-10-05", token)
        assert photo_word.status_code == 200
        assert "评定[2].docx" in unquote(photo_word.headers["content-disposition"])
        with zipfile.ZipFile(io.BytesIO(photo_word.content)) as archive:
            photo_xml = archive.read("word/document.xml").decode("utf-8")
        assert "委托单位：" in photo_xml
        assert "安丘送检单位" in photo_xml
        photo_pair = request(client, "POST",
            f"/api/inspection-documents/{inspection_document_data['id']}/export-word-pair?report_date=2031-10-05", token)
        assert photo_pair.status_code == 200
        assert [item["filename"] for item in photo_pair.json()["reports"]] == [
            "安丘送检单位1005评定[2].docx", "安丘送检单位1005检验[0].docx",
        ]
        assert request(client, "GET", "/api/inspection-word-templates",
            token, params={"customer_name": "安丘送检单位"}).json()["selected_key"] == "photo"
        assert request(client, "PUT", "/api/inspection-word-templates/choice",
            token, params={"customer_name": "安丘送检单位"}, json={"template_key": "default"}).status_code == 200
        removed_bundle = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-word-bundle?report_date=2031-10-05", token)
        assert removed_bundle.status_code == 404
        bottle_reports = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-bottle-reports?report_date=2026-09-14", token)
        assert bottle_reports.status_code == 200
        with zipfile.ZipFile(io.BytesIO(bottle_reports.content)) as reports:
            assert len(reports.namelist()) == inspection_document_data["total_quantity"]
            assert all(name.endswith(".docx") for name in reports.namelist())
            assert all(reports.read(name)[:2] == b"PK" for name in reports.namelist())
            for name in reports.namelist():
                with zipfile.ZipFile(io.BytesIO(reports.read(name))) as report:
                    assert "2030-09-13" in report.read("word/document.xml").decode("utf-8")
        first_line_id = inspection_document_data["lines"][0]["id"]
        second_line_id = inspection_document_data["lines"][1]["id"]
        one_bottle_report = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/lines/{first_line_id}/export-bottle-report?report_date=2031-10-05", token)
        assert one_bottle_report.status_code == 200
        assert one_bottle_report.content[:2] == b"PK"
        with zipfile.ZipFile(io.BytesIO(one_bottle_report.content)) as report:
            assert "2035-10-04" in report.read("word/document.xml").decode("utf-8")
        assert ".docx" in unquote(one_bottle_report.headers["content-disposition"])
        first_rejection = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/mark-unverified", token, json={"raw_text": "气瓶列表第2行【气瓶信息不存在】"})
        assert first_rejection.status_code == 200
        unverified_xls = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-xls?kind=unverified&export_date=2032-02-03", token)
        assert unverified_xls.status_code == 200
        assert "未经系统验证[1].xls" in unquote(unverified_xls.headers["content-disposition"])
        unverified_sheet = xlrd.open_workbook(file_contents=unverified_xls.content).sheet_by_index(0)
        assert unverified_sheet.nrows == 3
        assert unverified_sheet.cell_value(2, 0) == "2"
        assert unverified_sheet.cell_value(2, 2) == inspection_document_data["lines"][1]["body_code"]
        assert unverified_sheet.cell_value(2, 6) == "2032.02.03"
        partial_export_detail = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
        assert [line["inspection_time"] for line in partial_export_detail["lines"]] == ["2031-10-05", "2032-02-03"]
        assert partial_export_detail["export_dates"]["xls:unverified"] == "2032-02-03"
        verified_xls = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-xls?kind=verified&export_date=2032-01-02", token)
        assert verified_xls.status_code == 200
        assert "[1].xls" in unquote(verified_xls.headers["content-disposition"])
        assert xlrd.open_workbook(file_contents=verified_xls.content).sheet_by_index(0).nrows == 3
        after_verified_export = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
        assert all(line["inspection_time"] == "2032-01-02" for line in after_verified_export["lines"])
        assert after_verified_export["export_dates"]["xls:verified"] == "2032-01-02"
        reused_xls = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-xls?kind=verified", token)
        assert reused_xls.status_code == 200
        assert "20320102" in unquote(reused_xls.headers["content-disposition"])
        for selected_date in ("2032-03-04", "2032-04-05"):
            posted_word = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-word?kind=assessment&report_date={selected_date}", token)
            assert posted_word.status_code == 200, posted_word.text
            report_detail = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
            assert report_detail["export_dates"]["word:all"] == selected_date
            assert [line["inspection_time"] for line in report_detail["lines"]] == [selected_date, "2032-01-02"]
        reused_word = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-word?kind=assessment", token)
        assert reused_word.status_code == 200
        with zipfile.ZipFile(io.BytesIO(reused_word.content)) as report:
            assert "2036.04.04" in report.read("word/document.xml").decode("utf-8")
        for selected_date in ("2032-05-06", "2032-06-07"):
            posted_reports = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/export-bottle-reports?report_date={selected_date}", token)
            assert posted_reports.status_code == 200, posted_reports.text
            report_detail = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
            assert report_detail["export_dates"]["bottle:all"] == selected_date
            assert [line["inspection_time"] for line in report_detail["lines"]] == [selected_date, "2032-01-02"]
        for selected_date in ("2032-07-08", "2032-08-09"):
            posted_single = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/lines/{first_line_id}/export-bottle-report?report_date={selected_date}", token)
            assert posted_single.status_code == 200, posted_single.text
            report_detail = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
            assert report_detail["export_dates"][f"bottle:{first_line_id}"] == selected_date
            assert [line["inspection_time"] for line in report_detail["lines"]] == [selected_date, "2032-01-02"]
        assert report_detail["export_dates"]["bottle:all"] == "2032-06-07"
        assert report_detail["export_dates"]["xls:unverified"] == "2032-02-03"
        rejected_bottle_report = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/lines/{second_line_id}/export-bottle-report", token)
        assert rejected_bottle_report.status_code == 400
        remaining_bottle_report = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/lines/{first_line_id}/export-bottle-report", token)
        assert remaining_bottle_report.status_code == 200
        filtered_bottle_reports = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-bottle-reports", token)
        assert filtered_bottle_reports.status_code == 200
        with zipfile.ZipFile(io.BytesIO(filtered_bottle_reports.content)) as reports:
            assert len(reports.namelist()) == 1
        rejection_text = """气瓶列表第2行【气瓶信息不存在，请检查】
气瓶列表第1行【气瓶信息不存在，
请检杳]
气瓶列表第29行
气瓶信息不存在，请检查"""
        marked_unverified = request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/mark-unverified", token, json={"raw_text": rejection_text})
        assert marked_unverified.status_code == 200
        assert marked_unverified.json()["matched_sequences"] == [2, 1]
        assert marked_unverified.json()["unmatched_sequences"] == [29]
        marked_document = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
        assert all(line["is_unverified"] for line in marked_document["lines"])
        all_bottle_reports = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-bottle-reports", token)
        assert all_bottle_reports.status_code == 400
        filtered_assessment_word = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}/export-word?kind=assessment", token)
        assert filtered_assessment_word.status_code == 200
        assert "评定[0].docx" in unquote(filtered_assessment_word.headers["content-disposition"])
        unverified_pool = request(client, "GET", "/api/unverified-bottles", token)
        assert unverified_pool.status_code == 200
        assert unverified_pool.json()["total"] == 2
        assert [item["sequence"] for item in unverified_pool.json()["items"]] == [1, 2]
        assert len(request(client, "GET", "/api/unverified-bottles?page_size=1", token).json()["items"]) == 1
        exported_pool = request(client, "GET", "/api/unverified-bottles/export-xls", token)
        assert exported_pool.status_code == 200
        assert exported_pool.headers["content-type"].startswith("application/vnd.ms-excel")
        assert "全量[2].xls" in unquote(exported_pool.headers["content-disposition"])
        workbook = xlrd.open_workbook(file_contents=exported_pool.content, formatting_info=True)
        sheet = workbook.sheet_by_index(0)
        assert sheet.nrows == 4
        assert sheet.cell_value(0, 0) == "未经系统验证瓶子池"
        assert sheet.row_values(1) == [
            "检验单号", "检验日期", "单据序号", "瓶身编号", "生产日期", "生产厂家",
            "钢瓶规格", "容积（L）", "客户", "检验人员", "上游驳回信息", "状态",
        ]
        assert [sheet.cell_value(row, 3) for row in (2, 3)] == [item["body_code"] for item in unverified_pool.json()["items"]]
        assert [sheet.cell_value(row, 10) for row in (2, 3)] == [item["error_message"] for item in unverified_pool.json()["items"]]
        assert all(sheet.cell_value(row, 11) == "未经系统验证" for row in (2, 3))
        border = workbook.xf_list[sheet.cell(2, 3).xf_index].border
        assert (border.left_line_style, border.right_line_style, border.top_line_style, border.bottom_line_style) == (1, 1, 1, 1)
        assert unverified_pool.json()["items"][0]["document_no"] == inspection_document_data["document_no"]
        corrected_live = request(client, "PUT", f"/api/inspection-documents/{inspection_document_data['id']}/lines/{second_line_id}", token, json={
            "body_code": "QUICK-2026-0002-CORRECTED", "manufacture_date": "2026-01-01", "manufacturer": "修正制造厂",
        })
        assert corrected_live.status_code == 200, corrected_live.text
        assert corrected_live.json()["lines"][0]["id"] == second_line_id
        assert corrected_live.json()["lines"][0]["sequence"] == 1
        corrected_live_pool = request(client, "GET", "/api/unverified-bottles?body_code=QUICK-2026-0002-CORRECTED", token).json()["items"]
        assert len(corrected_live_pool) == 1 and corrected_live_pool[0]["sequence"] == 1
        original_bottle = next(item for item in request(client, "GET", "/api/bottles", token).json()["items"] if item["id"] == without_manufacture_date.json()["id"])
        assert original_bottle["body_code"] == "QUICK-2026-0002-CORRECTED"
        assert original_bottle["manufacture_date"] == "2026-01-01"
        assert original_bottle["manufacturer"] == "修正制造厂"
        assert corrected_live_pool[0]["spec_name"] == "15kg"
        assert float(corrected_live_pool[0]["wall_thickness"]) == 2.45
        edited_from_pool = request(client, "PUT", f"/api/unverified-bottles/{corrected_live_pool[0]['id']}", token, json={
            "body_code": "QUICK-2026-0002-POOL", "manufacture_date": "2026-02-01", "manufacturer": "池中更正厂家",
        })
        assert edited_from_pool.status_code == 200, edited_from_pool.text
        assert edited_from_pool.json()["body_code"] == "QUICK-2026-0002-POOL"
        assert edited_from_pool.json()["sequence"] == 1
        pool_bottle = next(item for item in request(client, "GET", "/api/bottles", token).json()["items"] if item["id"] == without_manufacture_date.json()["id"])
        assert (pool_bottle["body_code"], pool_bottle["manufacture_date"], pool_bottle["manufacturer"]) == (
            "QUICK-2026-0002-POOL", "2026-02-01", "池中更正厂家",
        )
        corrected_from_coding = request(client, "PUT", f"/api/bottles/{without_manufacture_date.json()['id']}/coding", token, json={
            "body_code": "QUICK-2026-DIRECT", "manufacture_date": "2026-03-01", "manufacturer": "抄码表更正厂家",
            "source_id": source_id, "spec_id": spec_id,
        })
        assert corrected_from_coding.status_code == 200, corrected_from_coding.text
        assert (corrected_from_coding.json()["body_code"], corrected_from_coding.json()["manufacturer"]) == (
            "QUICK-2026-DIRECT", "抄码表更正厂家",
        )
        coding_corrected_document = request(client, "GET", f"/api/inspection-documents/{inspection_document_data['id']}", token).json()
        coding_corrected_line = next(line for line in coding_corrected_document["lines"] if line["bottle_id"] == without_manufacture_date.json()["id"])
        assert (coding_corrected_line["body_code"], coding_corrected_line["manufacture_date"], coding_corrected_line["manufacturer"]) == (
            "QUICK-2026-DIRECT", "2026-03-01", "抄码表更正厂家",
        )
        coding_corrected_pool = request(client, "GET", "/api/unverified-bottles?body_code=QUICK-2026-DIRECT", token).json()["items"]
        assert len(coding_corrected_pool) == 1 and coding_corrected_pool[0]["manufacturer"] == "抄码表更正厂家"
        assert request(client, "DELETE", f"/api/bottles/{without_manufacture_date.json()['id']}/coding", token).status_code == 400
        assert request(client, "PUT", f"/api/bottles/{without_manufacture_date.json()['id']}/status", token, json={"zone": "BAD"}).status_code == 400
        inspection_zones = request(client, "GET", "/api/bottle-zones", token).json()
        assert inspection_zones["zones"]["good"] == 3

        daily_dashboard = request(client, "GET", "/api/dashboard?target_date=2026-09-09", token)
        assert daily_dashboard.status_code == 200
        assert daily_dashboard.json()["daily_summary"]["inbound"] == 8
        assert daily_dashboard.json()["daily_summary"]["outbound"] == 0
        assert daily_dashboard.json()["daily_summary"]["inspected"] + daily_dashboard.json()["daily_summary"]["uninspected"] == 4
        assert daily_dashboard.json()["monthly_summary"]["inbound"] == 8
        assert daily_dashboard.json()["yearly_summary"]["inbound"] == 8

        quick_config = request(client, "POST", "/api/quick-configs", token, json={
            "category": "CODE_PREFIX", "name": "济南", "value": "JN-", "sort_order": 1,
        })
        assert quick_config.status_code == 201
        options = request(client, "GET", "/api/options", token).json()
        assert "inspection_price" not in options["specs"][0]
        assert options["code_prefixes"][0]["value"] == "JN-"

        availability = request(client, "GET", "/api/quick-outbound/availability", token)
        assert availability.status_code == 200
        assert availability.json()[str(spec_id)]["good"] >= 1
        assert availability.json()[str(spec_id)]["temporary"] == 3
        assert availability.json()[str(spec_id)]["scrap"] == 2
        forbidden_adjustment = request(client, "POST", "/api/bottle-inventory-adjustments", inbound_operator_token, json={
            "business_date": "2026-09-10", "spec_id": spec_id, "zone": "GOOD", "quantity": 1,
        })
        assert forbidden_adjustment.status_code == 403
        bad_before = request(client, "GET", "/api/bottle-zones?zone=BAD", token).json()
        bad_supply = request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
            "business_date": "2026-09-10", "spec_id": spec_id, "zone": "BAD", "quantity": 3, "note": "不良品期初补录",
        })
        assert bad_supply.status_code == 201
        assert bad_supply.json()["zone"] == "BAD"
        bad_stock = request(client, "GET", "/api/bottle-zones?zone=BAD", token).json()
        assert bad_stock["zones"]["bad"] == bad_before["zones"]["bad"] + 3
        assert bad_stock["manual_bad"] == 3
        assert bad_stock["bottles"]["total"] == bad_before["bottles"]["total"]
        assert next(row for row in bad_stock["stock_by_spec"] if row["spec_id"] == spec_id)["bad"] == next(row for row in bad_before["stock_by_spec"] if row["spec_id"] == spec_id)["bad"] + 3
        assert request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec_id)]["bad"] == 3
        bad_deduction = request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
            "business_date": "2026-09-10", "spec_id": spec_id, "zone": "BAD", "quantity": -2,
        })
        assert bad_deduction.status_code == 201
        assert request(client, "GET", "/api/bottle-zones?zone=BAD", token).json()["manual_bad"] == 1
        assert request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
            "business_date": "2026-09-10", "spec_id": spec_id, "zone": "BAD", "quantity": -2,
        }).status_code == 400
        assert request(client, "DELETE", f"/api/bottle-inventory-adjustments/{bad_supply.json()['id']}", token).status_code == 400
        assert request(client, "DELETE", f"/api/bottle-inventory-adjustments/{bad_deduction.json()['id']}", token).status_code == 204
        assert request(client, "DELETE", f"/api/bottle-inventory-adjustments/{bad_supply.json()['id']}", token).status_code == 204
        assert request(client, "GET", "/api/bottle-zones?zone=BAD", token).json()["zones"]["bad"] == bad_before["zones"]["bad"]
        assert request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
            "business_date": "2026-09-10", "spec_id": spec_id, "zone": "BAD", "quantity": -1,
        }).status_code == 400
        for adjustment_zone, adjustment_quantity in (("GOOD", 2), ("TEMPORARY", 2), ("SCRAP", 1)):
            adjustment = request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
                "business_date": "2026-09-10", "spec_id": spec_id, "zone": adjustment_zone,
                "quantity": adjustment_quantity, "note": "期初库存补录",
            })
            assert adjustment.status_code == 201
            assert adjustment.json()["created_by_name"] == "超级管理员"
            if adjustment_zone == "GOOD":
                good_adjustment_id = adjustment.json()["id"]
        adjusted_availability = request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec_id)]
        assert adjusted_availability["good"] >= 3
        assert adjusted_availability["temporary"] == 5
        assert adjusted_availability["scrap"] == 3
        deduction = request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
            "business_date": "2026-09-10", "spec_id": spec_id, "zone": "TEMPORARY",
            "quantity": -1, "note": "盘点扣减",
        })
        assert deduction.status_code == 201
        assert deduction.json()["quantity"] == -1
        assert request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec_id)]["temporary"] == 4
        excessive_deduction = request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
            "business_date": "2026-09-10", "spec_id": spec_id, "zone": "TEMPORARY", "quantity": -99,
        })
        assert excessive_deduction.status_code == 400
        assert request(client, "DELETE", f"/api/bottle-inventory-adjustments/{deduction.json()['id']}", inbound_operator_token).status_code == 403
        assert request(client, "DELETE", f"/api/bottle-inventory-adjustments/{deduction.json()['id']}", token).status_code == 204
        assert request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec_id)]["temporary"] == 5
        assert request(client, "DELETE", f"/api/bottle-inventory-adjustments/{good_adjustment_id}", token).status_code == 204
        quick_outbound = request(client, "POST", "/api/quick-outbound", token, json={
            "business_date": "2026-09-10", "source_id": source_id, "handler": "快速发货员",
            "lines": [{"spec_id": spec_id, "quantity": 1}],
        })
        assert quick_outbound.status_code == 201
        assert quick_outbound.json()["status"] == "POSTED"
        assert quick_outbound.json()["posted_by_name"] == "超级管理员"
        outbound_bottle = request(client, "GET", "/api/bottles?zone=OUTBOUND", token).json()["items"][0]
        assert outbound_bottle["inspection_document_status"] == "INSPECTED"
        assert outbound_bottle["inspection_document_no"] == inspection_document_data["document_no"]
        assert request(client, "POST", f"/api/inspection-documents/{inspection_document_data['id']}/void", token).status_code == 400
        blocked_undo = request(client, "DELETE", f"/api/bottles/{outbound_bottle['id']}/coding", token)
        assert blocked_undo.status_code == 400
        assert "先撤销对应出库单" in blocked_undo.json()["detail"]
        dashboard_with_outbound = request(client, "GET", "/api/dashboard?target_date=2026-09-10", token)
        assert dashboard_with_outbound.json()["daily_summary"]["outbound"] == 1
        assert dashboard_with_outbound.json()["monthly_summary"]["outbound"] == 1
        assert dashboard_with_outbound.json()["yearly_summary"]["outbound"] == 1
        ranged_dashboard = request(client, "GET", "/api/dashboard", token, params={
            "daily_from": "2026-09-09", "daily_to": "2026-09-10",
            "monthly_from": "2025-01-01", "monthly_to": "2025-12-31",
            "yearly_from": "2026-01-01", "yearly_to": "2026-12-31",
        })
        assert ranged_dashboard.status_code == 200
        assert ranged_dashboard.json()["daily_summary"]["inbound"] == 8
        assert ranged_dashboard.json()["daily_summary"]["outbound"] == 1
        assert ranged_dashboard.json()["monthly_summary"]["inbound"] == 0
        assert ranged_dashboard.json()["monthly_summary"]["outbound"] == 0
        assert ranged_dashboard.json()["yearly_summary"]["inbound"] == 8
        assert ranged_dashboard.json()["yearly_summary"]["outbound"] == 1
        invalid_dashboard_range = request(client, "GET", "/api/dashboard", token, params={
            "yearly_from": "2027-01-01", "yearly_to": "2026-12-31",
        })
        assert invalid_dashboard_range.status_code == 400
        assert request(client, "POST", f"/api/bottle-outbounds/{quick_outbound.json()['id']}/void", token).status_code == 200

        bottle_outbound = request(client, "POST", "/api/bottle-outbounds", token, json={
            "business_date": "2026-09-10", "source_id": source_id, "handler": "发货员",
            "lines": [{"spec_id": spec_id, "quantity": 2}],
        })
        assert bottle_outbound.status_code == 201
        assert bottle_outbound.json()["source_id"] == source_id
        assert bottle_outbound.json()["customer"] == "安丘送检单位"
        assert bottle_outbound.json()["created_by_name"] == "超级管理员"
        bottle_outbound = request(client, "PUT", f"/api/bottle-outbounds/{bottle_outbound.json()['id']}", token, json={
            "business_date": "2026-09-10", "source_id": source_id, "handler": "发货员", "note": "已复核",
            "lines": [{"spec_id": spec_id, "quantity": 2}],
        })
        assert bottle_outbound.json()["updated_by_name"] == "超级管理员"
        bottle_outbound = request(client, "POST", f"/api/bottle-outbounds/{bottle_outbound.json()['id']}/post", token)
        assert bottle_outbound.status_code == 200
        assert bottle_outbound.json()["posted_by_name"] == "超级管理员"
        assert request(client, "GET", "/api/bottles?zone=OUTBOUND", token).json()["total"] == 2
        bottle_outbound = request(client, "POST", f"/api/bottle-outbounds/{bottle_outbound.json()['id']}/void", token)
        assert bottle_outbound.status_code == 200
        assert bottle_outbound.json()["voided_by_name"] == "超级管理员"

        quantity_outbound = request(client, "POST", "/api/quick-outbound", token, json={
            "business_date": "2026-09-09", "source_id": source_id, "handler": "补录出库员",
            "lines": [{"spec_id": spec_id, "temporary_quantity": 2, "scrap_quantity": 1}],
        })
        assert quantity_outbound.status_code == 201
        assert quantity_outbound.json()["total_quantity"] == 3
        assert quantity_outbound.json()["lines"][0]["temporary_quantity"] == 2
        assert quantity_outbound.json()["lines"][0]["scrap_quantity"] == 1
        quantity_zones = request(client, "GET", "/api/bottle-zones", token).json()["zones"]
        assert quantity_zones["temporary"] == 3
        assert quantity_zones["scrap"] == 2
        restored_quantity_outbound = request(client, "POST", f"/api/bottle-outbounds/{quantity_outbound.json()['id']}/void", token)
        assert restored_quantity_outbound.status_code == 200
        restored_quantity_zones = request(client, "GET", "/api/bottle-zones", token).json()["zones"]
        assert restored_quantity_zones["temporary"] == 5
        assert restored_quantity_zones["scrap"] == 3

        bad_replenishment = request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
            "business_date": "2026-09-10", "spec_id": spec_id, "zone": "BAD", "quantity": 2,
        })
        assert bad_replenishment.status_code == 201
        assert request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec_id)]["bad"] == 2
        bad_outbound = request(client, "POST", "/api/quick-outbound", token, json={
            "business_date": "2026-09-10", "source_id": source_id, "handler": "不良品发货员",
            "lines": [{"spec_id": spec_id, "bad_quantity": 2}],
        })
        assert bad_outbound.status_code == 201
        assert bad_outbound.json()["total_quantity"] == 2
        assert bad_outbound.json()["lines"][0]["bad_quantity"] == 2
        assert request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec_id)]["bad"] == 0
        assert request(client, "GET", "/api/bottle-zones?zone=BAD", token).json()["manual_bad"] == 0
        assert request(client, "POST", "/api/quick-outbound", token, json={
            "business_date": "2026-09-10", "source_id": source_id, "handler": "不良品发货员",
            "lines": [{"spec_id": spec_id, "bad_quantity": 1}],
        }).status_code == 400
        assert request(client, "POST", f"/api/bottle-outbounds/{bad_outbound.json()['id']}/void", token).status_code == 200
        assert request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec_id)]["bad"] == 2
        assert request(client, "DELETE", f"/api/bottle-inventory-adjustments/{bad_replenishment.json()['id']}", token).status_code == 204

        unpaid_order = request(client, "POST", "/api/purchase-orders", token, json={
            "order_date": "2026-09-10", "supplier_id": supplier.json()["id"], "handler": "采购员",
            "lines": [{"material_id": material_id, "quantity": 2}],
        })
        assert unpaid_order.status_code == 201
        unpaid_order = request(client, "POST", f"/api/purchase-orders/{unpaid_order.json()['id']}/confirm", token)
        assert unpaid_order.status_code == 200
        unpaid_purchase = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "NORMAL_IN", "business_date": "2026-09-10", "purchase_order_id": unpaid_order.json()["id"],
            "inbound_mode": "PURCHASE_ORDER", "inbound_category": "MATERIAL", "handler": "仓库员",
        })
        assert unpaid_purchase.status_code == 201
        assert request(client, "POST", f"/api/material-documents/{unpaid_purchase.json()['id']}/post", token).status_code == 200
        scrap_cost_inbound = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "NORMAL_IN", "business_date": "2026-09-09", "inbound_category": "SCRAP",
            "handler": "废料库管", "lines": [{"material_id": scrap_material_id, "quantity": 2, "unit_price": 1.5}],
        })
        assert scrap_cost_inbound.status_code == 201
        assert request(client, "POST", f"/api/material-documents/{scrap_cost_inbound.json()['id']}/post", token).status_code == 200
        material_filtered = request(client, "GET", "/api/material-documents", token, params={
            "document_type": "NORMAL_IN", "material_id": scrap_material_id,
        })
        assert material_filtered.status_code == 200
        assert [item["id"] for item in material_filtered.json()["items"]] == [scrap_cost_inbound.json()["id"]]
        assert request(client, "GET", "/api/material-documents", token, params={
            "document_type": "NORMAL_IN", "material_id": material_id,
        }).json()["total"] == 2
        scrap_sale = request(client, "POST", "/api/material-documents", token, json={
            "document_type": "OUTBOUND", "outbound_category": "SCRAP", "business_date": "2026-09-10", "handler": "废料销售员",
            "lines": [{"material_id": scrap_material_id, "quantity": 1, "unit_price": 5}],
        })
        assert scrap_sale.status_code == 201
        assert request(client, "POST", f"/api/material-documents/{scrap_sale.json()['id']}/post", token).status_code == 200
        revenue_outbound = request(client, "POST", "/api/quick-outbound", token, json={
            "business_date": "2026-09-10", "source_id": source_id, "handler": "营收验证员",
            "lines": [{"spec_id": spec_id, "good_quantity": 1, "temporary_quantity": 1, "scrap_quantity": 1}],
        })
        assert revenue_outbound.status_code == 201
        assert revenue_outbound.json()["total_quantity"] == 3
        profit = request(client, "GET", "/api/profit-analysis", token, params={"date_from": "2026-09-09", "date_to": "2026-09-10"})
        assert profit.status_code == 200
        assert profit.json()["purchase"]["paid_count"] == 1
        assert float(profit.json()["purchase"]["paid_amount"]) == 1428
        assert profit.json()["purchase"]["unpaid_count"] == 1
        assert float(profit.json()["purchase"]["unpaid_amount"]) == 23.8
        assert float(profit.json()["materials"]["outbound_cost"]) == 357
        assert float(profit.json()["materials"]["return_amount"]) == 0
        assert float(profit.json()["materials"]["net_cost"]) == 357
        assert float(profit.json()["scrap"]["inbound_amount"]) == 3
        assert float(profit.json()["scrap"]["outbound_revenue"]) == 5
        assert float(profit.json()["scrap"]["profit"]) == 5
        assert float(profit.json()["scrap"]["revenue"]) == 5
        assert float(profit.json()["scrap"]["inbound_items"][0]["quantity"]) == 2
        assert float(profit.json()["scrap"]["outbound_items"][0]["quantity"]) == 1
        assert len(profit.json()["purchase"]["items"]) == 2
        assert profit.json()["inspection"]["quantity"] == 1
        assert float(profit.json()["inspection"]["revenue"]) == 30
        assert float(profit.json()["totals"]["cost"]) == 1785
        assert float(profit.json()["totals"]["revenue"]) == 35
        assert float(profit.json()["totals"]["profit"]) == -1750
        inbound_only = request(client, "GET", "/api/profit-analysis", token, params={"date_from": "2026-09-09", "date_to": "2026-09-09"})
        assert float(inbound_only.json()["scrap"]["inbound_amount"]) == 3
        assert float(inbound_only.json()["scrap"]["outbound_revenue"]) == 0
        assert float(inbound_only.json()["scrap"]["profit"]) == 0
        assert float(inbound_only.json()["totals"]["cost"]) == float(inbound_only.json()["purchase"]["paid_amount"]) + float(inbound_only.json()["materials"]["net_cost"])
        assert request(client, "GET", "/api/profit-analysis", token, params={"date_from": "2026-09-11", "date_to": "2026-09-10"}).status_code == 400

        limited = request(client, "POST", "/api/users", token, json={
            "username": "viewer", "name": "库存查看员", "role_name": "材料库管员", "password": "Viewer@123",
            "permissions": ["dashboard:view", "materials:view"],
        })
        assert limited.status_code == 201
        assert limited.json()["role_name"] == "材料库管员"
        viewer_login = client.post("/api/auth/login", json={"username": "viewer", "password": "Viewer@123"})
        viewer_token = viewer_login.json()["access_token"]
        assert request(client, "GET", "/api/materials", viewer_token).status_code == 200
        assert request(client, "POST", "/api/materials", viewer_token, json={"name": "无权限物料"}).status_code == 403
        assert request(client, "GET", "/api/profit-analysis", viewer_token, params={"date_from": "2026-09-01", "date_to": "2026-09-30"}).status_code == 403

        quality_user = request(client, "POST", "/api/users", token, json={
            "username": "quality", "name": "质检人员", "role_name": "质检员", "password": "Quality@123",
            "permissions": ["inspection:view"],
        })
        assert quality_user.status_code == 201
        quality_login = client.post("/api/auth/login", json={"username": "quality", "password": "Quality@123"})
        quality_token = quality_login.json()["access_token"]
        assert quality_login.json()["user"]["role_name"] == "质检员"
        assert request(client, "GET", "/api/bottles?zone=GOOD", quality_token).status_code == 200
        assert request(client, "GET", "/api/coding/inbounds", quality_token).status_code == 403


def test_word_pair_numbers_continue_across_documents_and_restart_for_new_date():
    from backend.app.database import SessionLocal
    from backend.app.models import InspectionDocument, InspectionWordDailyCounter, InspectionWordExportAssignment, User

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        with SessionLocal() as db:
            user = db.query(User).filter_by(username="admin").one()
            documents = [InspectionDocument(
                document_no=f"JY20370306-WORD-{index}", inspection_date=date(2037, 3, 6),
                customer_name=f"报告客户{index}", inspector_name="检验员", created_by=user.id,
            ) for index in (1, 2, 3)]
            db.add_all(documents)
            db.commit()
            document_ids = [document.id for document in documents]

        for document_id, selected_date, first_number in (
            (document_ids[0], "2037-03-06", 1),
            (document_ids[0], "2037-03-06", 1),  # Refresh and re-export keeps the original pair.
            (document_ids[1], "2037-03-06", 3),
            (document_ids[0], "2037-03-07", 1),
            (document_ids[1], "2037-03-07", 3),
            (document_ids[0], "2037-03-06", 1),
            (document_ids[0], "2037-03-07", 1),
            (document_ids[0], "2038-03-06", 1),
            (document_ids[0], "2037-03-06", 1),
            (document_ids[2], "2037-03-06", 5),
        ):
            response = request(client, "POST", f"/api/inspection-documents/{document_id}/export-word-pair?report_date={selected_date}", token)
            assert response.status_code == 200, response.text
            data = response.json()
            assert data["report_date"] == selected_date
            assert [(item["kind"], item["report_number"]) for item in data["reports"]] == [
                ("regular", first_number), ("assessment", first_number + 1),
            ]
            for item in data["reports"]:
                content = base64.b64decode(item["content_base64"])
                with zipfile.ZipFile(io.BytesIO(content)) as report:
                    assert report.testzip() is None
                    xml = report.read("word/document.xml").decode("utf-8")
                prefix = f"WFJA-{selected_date[2:4]}-{selected_date[5:7]}{selected_date[8:10]}"
                assert f"{prefix}{item['report_number']:03d}" in xml
                assert item["filename"].endswith(".docx")
            detail = request(client, "GET", f"/api/inspection-documents/{document_id}", token).json()
            assert detail["export_dates"]["word:all"] == selected_date

        legacy = request(client, "POST", f"/api/inspection-documents/{document_ids[0]}/export-word?kind=assessment&report_date=2037-03-06", token)
        assert legacy.status_code == 200
        with zipfile.ZipFile(io.BytesIO(legacy.content)) as report:
            assert "WFJA-37-0306002" in report.read("word/document.xml").decode("utf-8")
        reopened = request(client, "POST", f"/api/inspection-documents/{document_ids[0]}/export-word-pair", token)
        assert [item["report_number"] for item in reopened.json()["reports"]] == [1, 2]

        with SessionLocal() as db:
            assert db.get(InspectionWordDailyCounter, date(2037, 3, 6)).last_number == 6
            assert db.get(InspectionWordDailyCounter, date(2037, 3, 7)).last_number == 4
            assert db.get(InspectionWordDailyCounter, date(2038, 3, 6)).last_number == 2
            assert db.query(InspectionWordExportAssignment).filter_by(report_date=date(2037, 3, 6)).count() == 3
            assert db.query(InspectionWordExportAssignment).filter_by(report_date=date(2037, 3, 7)).count() == 2
            db.add(InspectionWordDailyCounter(report_date=date(2037, 3, 8), last_number=998))
            db.commit()
        exhausted = request(client, "POST", f"/api/inspection-documents/{document_ids[0]}/export-word-pair?report_date=2037-03-08", token)
        assert exhausted.status_code == 409
        with SessionLocal() as db:
            assert db.get(InspectionWordDailyCounter, date(2037, 3, 8)).last_number == 998


def test_word_codes_can_be_viewed_and_superadmin_can_change_pair_date_and_number():
    from backend.app.database import SessionLocal
    from backend.app.models import InspectionDocument, InspectionWordDailyCounter, User

    with TestClient(app) as client:
        admin = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        with SessionLocal() as db:
            user = db.query(User).filter_by(username="admin").one()
            documents = [InspectionDocument(
                document_no=f"JY20390410-CODE-{index}", inspection_date=date(2039, 4, 10),
                customer_name=f"编码客户{index}", inspector_name="检验员", created_by=user.id,
            ) for index in (1, 2)]
            db.add_all(documents)
            db.commit()
            first, second = [item.id for item in documents]

        for document_id in (first, second):
            response = request(client, "POST", f"/api/inspection-documents/{document_id}/export-word-pair?report_date=2039-04-10", admin)
            assert response.status_code == 200, response.text
        codes = request(client, "GET", f"/api/inspection-documents/{first}/word-codes", admin).json()
        assert codes[0]["regular_code"] == "WFJA-39-0410001"
        assignment_id = codes[0]["id"]
        path = f"/api/inspection-documents/{first}/word-codes/{assignment_id}"

        check = request(client, "GET", f"{path}/check?report_date=2039-04-10&first_number=2", admin).json()
        assert check["conflict"] is True  # 002 overlaps the second document's 003 pair.
        assert check["existing_count"] == 1
        free_slot = request(client, "GET", f"{path}/check?report_date=2039-04-10&first_number=7", admin).json()
        assert free_slot["conflict"] is False
        assert free_slot["existing_count"] == 1  # The UI warns before using an occupied date.
        assert request(client, "PUT", path, admin, json={"report_date": "2039-04-10", "first_number": 2}).status_code == 409

        changed = request(client, "PUT", path, admin, json={"report_date": "2039-04-11", "first_number": 21})
        assert changed.status_code == 200, changed.text
        assert changed.json()["regular_code"] == "WFJA-39-0411021"
        assert changed.json()["assessment_code"] == "WFJA-39-0411022"
        assert request(client, "GET", f"/api/inspection-documents/{first}", admin).json()["export_dates"]["word:all"] == "2039-04-11"
        reopened = request(client, "POST", f"/api/inspection-documents/{first}/export-word-pair", admin)
        assert reopened.status_code == 200
        assert reopened.json()["report_date"] == "2039-04-11"
        assert [item["report_number"] for item in reopened.json()["reports"]] == [21, 22]
        with SessionLocal() as db:
            assert db.get(InspectionWordDailyCounter, date(2039, 4, 11)).last_number == 22
        next_pair = request(client, "POST", f"/api/inspection-documents/{second}/export-word-pair?report_date=2039-04-11", admin)
        assert [item["report_number"] for item in next_pair.json()["reports"]] == [23, 24]

        staff = request(client, "POST", "/api/users", admin, json={
            "username": "word_code_viewer", "name": "编码查看员", "password": "Viewer@1234",
            "permissions": ["inspection_documents:view"],
        })
        assert staff.status_code == 201
        viewer = client.post("/api/auth/login", json={"username": "word_code_viewer", "password": "Viewer@1234"}).json()["access_token"]
        assert request(client, "GET", f"/api/inspection-documents/{first}/word-codes", viewer).status_code == 200
        assert request(client, "GET", f"{path}/check?report_date=2039-04-12&first_number=1", viewer).status_code == 403
        assert request(client, "PUT", path, viewer, json={"report_date": "2039-04-12", "first_number": 1}).status_code == 403


def test_word_code_ledger_deletes_assignments_and_resets_only_to_retained_numbers():
    from backend.app.database import SessionLocal
    from backend.app.models import (
        InspectionDocument, InspectionDocumentExportRecord, InspectionWordDailyCounter, User,
    )

    with TestClient(app) as client:
        admin = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        with SessionLocal() as db:
            user = db.query(User).filter_by(username="admin").one()
            documents = [InspectionDocument(
                document_no=f"JY20400501-LEDGER-{index}", inspection_date=date(2040, 5, 1),
                customer_name=f"编号台账客户{index}", inspector_name="检验员", created_by=user.id,
            ) for index in (1, 2, 3, 4)]
            db.add_all(documents)
            db.flush()
            document_ids = [item.id for item in documents]
            db.add(InspectionDocumentExportRecord(document_id=document_ids[2], kind="word:all", export_date=date(2040, 5, 2)))
            db.add(InspectionWordDailyCounter(report_date=date(2040, 5, 3), last_number=998))
            db.commit()

        for document_id in document_ids[:2]:
            exported = request(client, "POST", f"/api/inspection-documents/{document_id}/export-word-pair?report_date=2040-05-01", admin)
            assert exported.status_code == 200
        ledger = request(client, "GET", "/api/inspection-word-code-ledger", admin)
        assert ledger.status_code == 200
        by_date = {item["report_date"]: item for item in ledger.json()["days"]}
        assert by_date["2040-05-01"]["last_number"] == 4
        assert by_date["2040-05-01"]["required_number"] == 4
        assert by_date["2040-05-01"]["unassigned_numbers"] == []
        assert [item["regular_code"] for item in by_date["2040-05-01"]["assignments"]] == ["WFJA-40-0501001", "WFJA-40-0501003"]
        assert by_date["2040-05-03"]["last_number"] == 998
        assert by_date["2040-05-03"]["assignments"] == []
        assert by_date["2040-05-03"]["unassigned_numbers"] == list(range(1, 999))
        assert any(item["document_id"] == document_ids[2] for item in ledger.json()["without_saved_code"])

        second_assignment = by_date["2040-05-01"]["assignments"][1]
        deleted = request(client, "DELETE", f"/api/inspection-word-code-ledger/assignments/{second_assignment['id']}", admin)
        assert deleted.status_code == 204
        before_reset = request(client, "GET", "/api/inspection-word-code-ledger", admin).json()
        day = next(item for item in before_reset["days"] if item["report_date"] == "2040-05-01")
        assert (day["last_number"], day["required_number"]) == (4, 2)
        assert day["unassigned_numbers"] == [3, 4]
        reset = request(client, "POST", "/api/inspection-word-code-ledger/dates/2040-05-01/reset", admin)
        assert reset.status_code == 200
        assert (reset.json()["last_number"], reset.json()["next_first_number"]) == (2, 3)
        next_pair = request(client, "POST", f"/api/inspection-documents/{document_ids[2]}/export-word-pair?report_date=2040-05-01", admin)
        assert [item["report_number"] for item in next_pair.json()["reports"]] == [3, 4]
        first_assignment = by_date["2040-05-01"]["assignments"][0]
        assert request(client, "DELETE", f"/api/inspection-word-code-ledger/assignments/{first_assignment['id']}", admin).status_code == 204
        gap_day = next(item for item in request(client, "GET", "/api/inspection-word-code-ledger", admin).json()["days"] if item["report_date"] == "2040-05-01")
        assert (gap_day["last_number"], gap_day["required_number"]) == (4, 4)
        assert gap_day["unassigned_numbers"] == [1, 2]
        assert gap_day["unassigned_pairs"] == [1]
        release_path = "/api/inspection-word-code-ledger/dates/2040-05-01/unassigned-pairs/1"
        assert request(client, "DELETE", release_path, admin).status_code == 204
        released_day = next(item for item in request(client, "GET", "/api/inspection-word-code-ledger", admin).json()["days"] if item["report_date"] == "2040-05-01")
        assert released_day["unassigned_numbers"] == []
        assert released_day["released_pairs"] == [1]
        assert released_day["next_first_number"] == 1
        assert request(client, "DELETE", release_path, admin).status_code == 409
        reused_pair = request(client, "POST", f"/api/inspection-documents/{document_ids[0]}/export-word-pair?report_date=2040-05-01", admin)
        assert [item["report_number"] for item in reused_pair.json()["reports"]] == [1, 2]
        reused_day = next(item for item in request(client, "GET", "/api/inspection-word-code-ledger", admin).json()["days"] if item["report_date"] == "2040-05-01")
        assert reused_day["released_pairs"] == []
        assert reused_day["last_number"] == 4
        with SessionLocal() as db:
            db.get(InspectionWordDailyCounter, date(2040, 5, 1)).last_number = 0
            db.commit()
        recovered_pair = request(client, "POST", f"/api/inspection-documents/{document_ids[3]}/export-word-pair?report_date=2040-05-01", admin)
        assert [item["report_number"] for item in recovered_pair.json()["reports"]] == [5, 6]
        empty_reset = request(client, "POST", "/api/inspection-word-code-ledger/dates/2040-05-03/reset", admin)
        assert (empty_reset.json()["last_number"], empty_reset.json()["next_first_number"]) == (0, 1)

        staff = request(client, "POST", "/api/users", admin, json={
            "username": "word_ledger_viewer", "name": "台账查看员", "password": "Viewer@1234",
            "permissions": ["inspection_documents:view"],
        })
        assert staff.status_code == 201
        viewer = client.post("/api/auth/login", json={"username": "word_ledger_viewer", "password": "Viewer@1234"}).json()["access_token"]
        assert request(client, "GET", "/api/inspection-word-code-ledger", viewer).status_code == 403
        assert request(client, "DELETE", f"/api/inspection-word-code-ledger/assignments/{second_assignment['id']}", viewer).status_code == 403
        assert request(client, "DELETE", release_path, viewer).status_code == 403
        assert request(client, "POST", "/api/inspection-word-code-ledger/dates/2040-05-01/reset", viewer).status_code == 403


def test_import_historical_inspection_excel_without_inventory_changes():
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        before_bottles = request(client, "GET", "/api/bottles", token).json()["total"]
        archived_spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "5kg", "inspection_price": 0, "volume": 8.8, "wall_thickness": 2.0,
        })
        assert archived_spec.status_code == 201
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "检验气瓶列表"
        sheet.append(["充装单位名称", "", "历史客户"])
        sheet.append(["序号", "设备品种", "产品编号", "生产日期", "生产厂家", "容积", "检验日期", "下次检验日期", "检验结果", "检验人员"])
        sheet.append(["1", "焊接气瓶", "0001001", "2019.01", "历史制造厂", 35.5, "2020.01.02", "2024.01.01", "合格", "历史检验员"])
        sheet.append(["2", "焊接气瓶", "0001002", "2018.02", "历史制造厂", 8.8, "2020.01.02", "2024.01.01", "合格", "历史检验员"])
        content = io.BytesIO()
        workbook.save(content)
        upload = lambda selected_date: request(client, "POST", "/api/inspection-documents/import", token,
            data={"inspection_date": selected_date},
            files={"file": ("history.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert upload("2020-01-03").status_code == 400
        imported = upload("2020-01-02")
        assert imported.status_code == 201, imported.text
        document = imported.json()
        assert document["is_archived"] is True
        assert document["customer_name"] == "历史客户"
        assert document["inspector_name"] == "历史检验员"
        assert [line["sequence"] for line in document["lines"]] == [1, 2]
        assert [line["body_code"] for line in document["lines"]] == ["0001001", "0001002"]
        assert all(line["bottle_id"] is None for line in document["lines"])
        assert all(line["inbound_at"] == document["created_at"] for line in document["lines"])
        assert all(line["spec_id"] is not None for line in document["lines"])
        assert all(line["inspection_time"] is None for line in document["lines"])
        assert document["lines"][0]["next_inspection_date"] == "2024-01-01"
        archive_word = request(client, "GET", f"/api/inspection-documents/{document['id']}/export-word?kind=assessment", token)
        assert archive_word.status_code == 200
        archived_reports = request(client, "GET", f"/api/inspection-documents/{document['id']}/export-bottle-reports", token)
        assert archived_reports.status_code == 200
        with zipfile.ZipFile(io.BytesIO(archived_reports.content)) as reports:
            assert len(reports.namelist()) == 2
        archived_first_line_id = document["lines"][0]["id"]
        archived_second_line_id = document["lines"][1]["id"]
        assert request(client, "GET", f"/api/inspection-documents/{document['id']}/lines/{archived_first_line_id}/export-bottle-report", token).status_code == 200
        assert request(client, "GET", "/api/bottles", token).json()["total"] == before_bottles
        assert upload("2020-01-02").status_code == 409
        assert request(client, "POST", "/api/inspection-documents", token, json={"inspection_date": "2020-01-02"}).json()["total_quantity"] == 2
        assert request(client, "POST", f"/api/inspection-documents/{document['id']}/sync", token).status_code == 400
        marked = request(client, "POST", f"/api/inspection-documents/{document['id']}/mark-unverified", token,
            json={"raw_text": "气瓶列表第2行【气瓶信息不存在】"})
        assert marked.status_code == 200, marked.text
        assert marked.json()["matched_sequences"] == [2]
        pool = request(client, "GET", "/api/unverified-bottles?body_code=0001002", token).json()["items"]
        assert len(pool) == 1 and pool[0]["bottle_id"] is None
        corrected = request(client, "PUT", f"/api/inspection-documents/{document['id']}/lines/{archived_second_line_id}", token, json={
            "body_code": "0001002X", "manufacture_date": "2019-12-01",
        })
        assert corrected.status_code == 200, corrected.text
        assert [line["body_code"] for line in corrected.json()["lines"]] == ["0001002X", "0001001"]
        assert [line["sequence"] for line in corrected.json()["lines"]] == [1, 2]
        assert corrected.json()["lines"][0]["is_unverified"] is True
        corrected_pool = request(client, "GET", "/api/unverified-bottles?body_code=0001002X", token).json()["items"]
        assert len(corrected_pool) == 1 and corrected_pool[0]["sequence"] == 1
        assert request(client, "PUT", f"/api/inspection-documents/{document['id']}/lines/{archived_second_line_id}", token, json={
            "body_code": "0001001",
        }).status_code == 400
        assert request(client, "PUT", f"/api/inspection-documents/{document['id']}/lines/{archived_second_line_id}", token, json={
            "manufacture_date": "2021-01-01",
        }).status_code == 400
        assert request(client, "GET", f"/api/inspection-documents/{document['id']}/lines/{archived_second_line_id}/export-bottle-report", token).status_code == 400
        filtered_reports = request(client, "GET", f"/api/inspection-documents/{document['id']}/export-bottle-reports", token)
        assert filtered_reports.status_code == 200
        with zipfile.ZipFile(io.BytesIO(filtered_reports.content)) as reports:
            assert len(reports.namelist()) == 1
        assert request(client, "POST", f"/api/inspection-documents/{document['id']}/void", token).status_code == 200


def test_import_historical_inspection_xls():
    xlwt = pytest.importorskip("xlwt")
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("检验气瓶列表")
    rows = [
        ["充装单位名称", "", "旧版客户"],
        ["序号", "设备品种", "产品编号", "生产日期", "生产厂家", "容积", "检验日期", "下次检验日期", "检验结果", "检验人员"],
        [1, "焊接气瓶", "0003001", "2018.03", "旧版制造厂", 35.5, datetime(2020, 1, 4), "2024.01.03", "合格", "旧版检验员"],
        [2, "焊接气瓶", "0003002", "2018.04", "旧版制造厂", 8.8, "2020.01", "2024.01", "合格", "旧版检验员"],
    ]
    for row_number, row in enumerate(rows):
        for column, value in enumerate(row):
            if isinstance(value, datetime):
                sheet.write(row_number, column, value, xlwt.easyxf(num_format_str="YYYY.MM.DD"))
            else:
                sheet.write(row_number, column, value)
    content = io.BytesIO()
    workbook.save(content)
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        imported = request(client, "POST", "/api/inspection-documents/import", token,
            data={"inspection_date": "2020-01-04"},
            files={"file": ("history.xls", content.getvalue(), "application/vnd.ms-excel")})
        assert imported.status_code == 201, imported.text
        document = imported.json()
        assert document["customer_name"] == "旧版客户"
        assert document["lines"][0]["body_code"] == "0003001"
        assert document["lines"][0]["volume"] == "35.50"
        assert document["lines"][0]["inspection_result"] == "GOOD"
        assert document["lines"][1]["body_code"] == "0003002"
        assert document["lines"][1]["next_inspection_date"] == "2024-01-01"


def test_reimport_historical_document_matches_original_sequences_and_keeps_rejections():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "检验气瓶列表"
    sheet.append(["充装单位名称", "", "回导客户"])
    sheet.append(["序号", "设备品种", "产品编号", "生产日期", "生产厂家", "容积", "检验日期", "下次检验日期", "检验结果", "检验人员"])
    sheet.append([1, "焊接气瓶", "REIMPORT-1", "2018.01", "原厂家", "35.50", "2020.01.07", "2024.01.06", "合格", "检验员"])
    sheet.append([2, "焊接气瓶", "REIMPORT-2", "2017.01", "原厂家", "35.50", "2020.01.07", "2024.01.06", "合格", "检验员"])
    initial = io.BytesIO()
    workbook.save(initial)
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        imported = request(client, "POST", "/api/inspection-documents/import", token,
            data={"inspection_date": "2020-01-07"}, files={"file": ("initial.xlsx", initial.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert imported.status_code == 201, imported.text
        document_id = imported.json()["id"]
        rejected_id = imported.json()["lines"][1]["id"]
        assert request(client, "POST", f"/api/inspection-documents/{document_id}/mark-unverified", token,
            json={"raw_text": "气瓶列表第2行【气瓶信息不存在】"}).status_code == 200

        sheet.delete_rows(3, 2)
        sheet.append([2, "焊接气瓶", "REIMPORT-2-NEW", "2019.12", "新厂家", "35.50", "2020.01.07", "2024.01.06", "合格", "检验员"])
        sheet.append([1, "焊接气瓶", "REIMPORT-1", "2018.01", "原厂家", "35.50", "2020.01.07", "2024.01.06", "合格", "检验员"])
        revised = io.BytesIO()
        workbook.save(revised)
        updated = request(client, "POST", f"/api/inspection-documents/{document_id}/reimport", token,
            files={"file": ("revised.xlsx", revised.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert updated.status_code == 200, updated.text
        assert [line["body_code"] for line in updated.json()["lines"]] == ["REIMPORT-2-NEW", "REIMPORT-1"]
        assert updated.json()["lines"][0]["id"] == rejected_id
        assert updated.json()["lines"][0]["sequence"] == 1
        assert updated.json()["lines"][0]["is_unverified"] is True
        assert updated.json()["lines"][0]["manufacturer"] == "新厂家"
        assert request(client, "GET", "/api/unverified-bottles?body_code=REIMPORT-2-NEW", token).json()["items"][0]["sequence"] == 1

        sheet.delete_rows(4)
        incomplete = io.BytesIO()
        workbook.save(incomplete)
        failed = request(client, "POST", f"/api/inspection-documents/{document_id}/reimport", token,
            files={"file": ("incomplete.xlsx", incomplete.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert failed.status_code == 400
        after_failure = request(client, "GET", f"/api/inspection-documents/{document_id}", token).json()
        assert [line["body_code"] for line in after_failure["lines"]] == ["REIMPORT-2-NEW", "REIMPORT-1"]


def test_import_historical_month_only_dates():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "检验气瓶列表"
    sheet.append(["充装单位名称", "", "月份客户"])
    sheet.append(["序号", "设备品种", "产品编号", "生产日期", "生产厂家", "容积", "检验日期", "下次检验日期", "检验结果", "检验人员"])
    sheet.append([1, "焊接气瓶", "0004001", "2022.01", "月份制造厂", 35.5, "2026.09", "2030.09", "合格", "月份检验员"])
    sheet.append([2, "焊接气瓶", "0004002", 2022.10, "月份制造厂", 8.8, 2026.09, 2030.10, "合格", "月份检验员"])
    content = io.BytesIO()
    workbook.save(content)
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        def upload(selected_date):
            return request(client, "POST", "/api/inspection-documents/import", token,
                data={"inspection_date": selected_date},
                files={"file": ("month.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert upload("2026-08-12").status_code == 400
        imported = upload("2026-09-12")
        assert imported.status_code == 201, imported.text
        assert imported.json()["inspection_date"] == "2026-09-12"
        assert imported.json()["lines"][0]["manufacture_date"] == "2022-01-01"
        assert imported.json()["lines"][0]["next_inspection_date"] == "2030-09-01"
        assert imported.json()["lines"][1]["manufacture_date"] == "2022-10-01"
        assert imported.json()["lines"][1]["next_inspection_date"] == "2030-10-01"
        corrected_xls = request(client, "POST", f"/api/inspection-documents/{imported.json()['id']}/export-xls?kind=full&export_date=2026-09-14", token)
        assert corrected_xls.status_code == 200
        corrected_sheet = xlrd.open_workbook(file_contents=corrected_xls.content).sheet_by_index(0)
        assert [corrected_sheet.cell_value(row, 7) for row in range(2, corrected_sheet.nrows)] == ["2030.09.13"] * 2
        corrected_word = request(client, "GET", f"/api/inspection-documents/{imported.json()['id']}/export-word?kind=assessment&report_date=2026-09-14", token)
        assert corrected_word.status_code == 200
        with zipfile.ZipFile(io.BytesIO(corrected_word.content)) as report:
            assert "2030.09.13" in report.read("word/document.xml").decode("utf-8")


def test_only_superadmin_can_permanently_delete_void_inspection_document():
    from backend.app.database import SessionLocal
    from backend.app.models import InspectionDocumentLine, UnverifiedBottle

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "检验气瓶列表"
    sheet.append(["充装单位名称", "", "待删除客户"])
    sheet.append(["序号", "设备品种", "产品编号", "生产日期", "生产厂家", "容积", "检验日期", "下次检验日期", "检验结果", "检验人员"])
    sheet.append([1, "焊接气瓶", "DELETE-TEST-1", "2018.01", "测试厂家", "35.50", "2020.01.08", "2024.01.07", "合格", "测试员"])
    content = io.BytesIO()
    workbook.save(content)
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        regular_user = request(client, "POST", "/api/users", token, json={
            "username": "void_delete_operator", "name": "删除权限测试员", "password": "Test@12345678",
            "permissions": ["inspection_documents:view", "inspection_documents:void"],
        })
        assert regular_user.status_code == 201, regular_user.text
        regular_token = client.post("/api/auth/login", json={
            "username": "void_delete_operator", "password": "Test@12345678",
        }).json()["access_token"]
        bottles_before = request(client, "GET", "/api/bottles", token).json()["total"]
        imported = request(client, "POST", "/api/inspection-documents/import", token,
            data={"inspection_date": "2020-01-08"},
            files={"file": ("delete-test.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert imported.status_code == 201, imported.text
        document_id = imported.json()["id"]
        assert request(client, "DELETE", f"/api/inspection-documents/{document_id}", token).status_code == 400
        assert request(client, "POST", f"/api/inspection-documents/{document_id}/mark-unverified", token,
            json={"raw_text": "气瓶列表第1行【气瓶信息不存在】"}).status_code == 200
        assert request(client, "POST", f"/api/inspection-documents/{document_id}/void", token).status_code == 200
        assert request(client, "DELETE", f"/api/inspection-documents/{document_id}", regular_token).status_code == 403
        assert request(client, "GET", f"/api/inspection-documents/{document_id}", token).status_code == 200
        deleted = request(client, "DELETE", f"/api/inspection-documents/{document_id}", token)
        assert deleted.status_code == 204, deleted.text
        assert request(client, "GET", f"/api/inspection-documents/{document_id}", token).status_code == 404
        assert request(client, "GET", "/api/bottles", token).json()["total"] == bottles_before
        with SessionLocal() as db:
            assert db.query(InspectionDocumentLine).filter_by(document_id=document_id).count() == 0
            assert db.query(UnverifiedBottle).filter_by(inspection_document_id=document_id).count() == 0


def test_unverified_pool_retry_any_quantity_preserves_source_history():
    from backend.app.database import SessionLocal
    from backend.app.models import BottleSpec, InspectionDocument, InspectionDocumentLine, UnverifiedBottle, User

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        with SessionLocal() as db:
            user_id = db.query(User).filter(User.username == "admin").one().id
            spec = db.query(BottleSpec).filter(BottleSpec.name == "15kg").one()
            document = InspectionDocument(
                document_no="JY20260115-POOLTEST", inspection_date=date(2026, 1, 15),
                customer_name="批量客户", inspector_name="批量检验员", is_archived=True,
                status="ACTIVE", created_by=user_id,
            )
            db.add(document)
            db.flush()
            lines = [InspectionDocumentLine(
                document_id=document.id, sequence=index, body_code=f"POOL-TEST-{index:04d}",
                manufacture_date=date(2022, 1, 1), manufacturer="原厂家", spec_id=spec.id,
                volume=spec.volume, inspection_result="GOOD",
            ) for index in range(1, 502)]
            db.add_all(lines)
            db.flush()
            db.add_all(UnverifiedBottle(
                inspection_document_id=document.id, inspection_line_id=line.id,
                sequence=line.sequence, body_code=line.body_code,
                manufacture_date=line.manufacture_date, manufacturer=line.manufacturer,
                error_message="气瓶信息不存在", raw_fragment=f"气瓶列表第{line.sequence}行", status="PENDING", created_by=user_id,
            ) for line in lines)
            db.commit()
            source_document_id = document.id

        filtered = request(client, "GET", "/api/unverified-bottles", token, params={"document_no": "POOLTEST"}).json()
        assert filtered["total"] == 501
        assert filtered["items"][0]["spec_name"] == "15kg"
        assert float(filtered["items"][0]["volume"]) == 35.5
        first_id = filtered["items"][0]["id"]
        corrected = request(client, "PUT", f"/api/unverified-bottles/{first_id}", token, json={
            "body_code": "POOL-TEST-CORRECTED", "manufacture_date": "2023-01-01", "manufacturer": "新厂家",
        })
        assert corrected.status_code == 200, corrected.text
        assert corrected.json()["body_code"] == "POOL-TEST-CORRECTED"
        assert request(client, "PUT", f"/api/unverified-bottles/{first_id}", token, json={"spec_id": 1}).status_code == 400

        created = request(client, "POST", "/api/unverified-bottles/create-retry-document", token, json={"document_no": "POOLTEST"})
        assert created.status_code == 201, created.text
        assert len(created.json()) == 1
        retry = created.json()[0]
        assert retry["is_retry"] is True and retry["is_archived"] is True
        assert retry["document_no"].endswith("-R01")
        assert retry["total_quantity"] == 501
        retry_detail = request(client, "GET", f"/api/inspection-documents/{retry['id']}", token).json()
        assert len(retry_detail["lines"]) == 501
        assert all(line["source_line_id"] for line in retry_detail["lines"])
        assert all(line["inbound_at"] for line in retry_detail["lines"])
        assert all(line["inspection_time"] is None for line in retry_detail["lines"])
        assert any(line["body_code"] == "POOL-TEST-CORRECTED" for line in retry_detail["lines"])
        assert request(client, "GET", "/api/unverified-bottles", token, params={"document_no": "POOLTEST"}).json()["total"] == 0
        assert request(client, "POST", "/api/unverified-bottles/create-retry-document", token, json={"document_no": "POOLTEST"}).status_code == 400
        original_detail = request(client, "GET", f"/api/inspection-documents/{source_document_id}", token).json()
        assert sum(line["is_unverified"] for line in original_detail["lines"]) == 501
        assert request(client, "POST", f"/api/inspection-documents/{retry['id']}/void", token).status_code == 200
        assert request(client, "GET", "/api/unverified-bottles", token, params={"document_no": "POOLTEST"}).json()["total"] == 501
        referenced_record = request(client, "GET", "/api/unverified-bottles", token, params={"body_code": "POOL-TEST-CORRECTED"}).json()["items"][0]
        assert request(client, "DELETE", f"/api/unverified-bottles/{referenced_record['id']}", token).status_code == 409
        single_retry = request(client, "POST", "/api/unverified-bottles/create-retry-document", token, json={"body_code": "POOL-TEST-CORRECTED"})
        assert single_retry.status_code == 201, single_retry.text
        assert len(single_retry.json()) == 1
        assert single_retry.json()[0]["total_quantity"] == 1
        assert request(client, "GET", "/api/unverified-bottles", token, params={"body_code": "POOL-TEST-CORRECTED"}).json()["total"] == 0
        assert request(client, "POST", f"/api/inspection-documents/{source_document_id}/void", token).status_code == 200
        blocked_delete = request(client, "DELETE", f"/api/inspection-documents/{source_document_id}", token)
        assert blocked_delete.status_code == 409
        assert request(client, "GET", f"/api/inspection-documents/{retry['id']}", token).status_code == 200


def test_unverified_pool_permanent_delete_removes_exact_bottle_and_snapshot():
    from backend.app.database import SessionLocal
    from backend.app.models import Bottle, BottleInbound, BottleSpec, InspectionDocument, InspectionDocumentLine, UnverifiedBottle, User

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        viewer = request(client, "POST", "/api/users", token, json={
            "username": "pool_delete_viewer", "name": "未验证池只读员", "password": "Viewer@12345678",
            "permissions": ["unverified_bottles:view"],
        })
        assert viewer.status_code == 201, viewer.text
        viewer_token = client.post("/api/auth/login", json={
            "username": "pool_delete_viewer", "password": "Viewer@12345678",
        }).json()["access_token"]
        with SessionLocal() as db:
            user = db.query(User).filter_by(username="admin").one()
            spec = db.query(BottleSpec).filter_by(name="15kg").one()
            inbound = db.query(BottleInbound).filter_by(status="POSTED").first()
            assert inbound is not None
            bottle = Bottle(
                inbound_id=inbound.id, source_id=inbound.source_id, spec_id=spec.id,
                body_code="POOL-PERMANENT-LIVE", manufacturer="测试厂家", zone="GOOD",
                process_status="INSPECTED", inspection_price=spec.inspection_price,
                coding_date=date(2026, 2, 11), good_entry_date=date(2026, 2, 11),
                coded_by=user.id, coded_at=datetime.utcnow(),
            )
            db.add(bottle)
            live_doc = InspectionDocument(
                document_no="JY20260211-DELETE-TEST", inspection_date=date(2026, 2, 11),
                customer_name="删除核验客户", inspector_name="核验员", status="ACTIVE", created_by=user.id,
            )
            archive_doc = InspectionDocument(
                document_no="JY20260212-DELETE-TEST", inspection_date=date(2026, 2, 12),
                customer_name="删除核验客户", inspector_name="核验员", is_archived=True,
                status="ACTIVE", created_by=user.id,
            )
            db.add_all((live_doc, archive_doc))
            db.flush()
            live_line = InspectionDocumentLine(
                document_id=live_doc.id, bottle_id=bottle.id, sequence=1,
                body_code=bottle.body_code, manufacture_date=date(2022, 1, 1),
                manufacturer="测试厂家", spec_id=spec.id, volume=spec.volume,
            )
            archived_lines = [InspectionDocumentLine(
                document_id=archive_doc.id, sequence=index, body_code=f"POOL-PERMANENT-ARCHIVE-{index}",
                manufacture_date=date(2021, index, 1), manufacturer="归档厂家",
                spec_id=spec.id, volume=spec.volume,
            ) for index in (1, 2)]
            db.add_all((live_line, *archived_lines))
            db.flush()
            records = [UnverifiedBottle(
                inspection_document_id=line.document_id, inspection_line_id=line.id,
                bottle_id=line.bottle_id, sequence=line.sequence, body_code=line.body_code,
                manufacture_date=line.manufacture_date, manufacturer=line.manufacturer,
                error_message="上游驳回", raw_fragment=f"气瓶列表第{line.sequence}行",
                status="PENDING", created_by=user.id,
            ) for line in (live_line, *archived_lines)]
            db.add_all(records)
            db.commit()
            live_id, archive_id = live_doc.id, archive_doc.id
            live_bottle_id, live_record_id = bottle.id, records[0].id
            archived_record_id = records[1].id

        assert request(client, "DELETE", f"/api/unverified-bottles/{live_record_id}", viewer_token).status_code == 403
        with SessionLocal() as db:
            db.get(Bottle, live_bottle_id).zone = "OUTBOUND"
            db.commit()
        blocked = request(client, "DELETE", f"/api/unverified-bottles/{live_record_id}", token)
        assert blocked.status_code == 409
        with SessionLocal() as db:
            assert db.get(UnverifiedBottle, live_record_id) is not None
            db.get(Bottle, live_bottle_id).zone = "GOOD"
            db.commit()

        before_bottles = request(client, "GET", "/api/bottles", token).json()["total"]
        deleted = request(client, "DELETE", f"/api/unverified-bottles/{live_record_id}", token)
        assert deleted.status_code == 204, deleted.text
        assert request(client, "GET", "/api/bottles", token).json()["total"] == before_bottles - 1
        assert request(client, "GET", f"/api/inspection-documents/{live_id}", token).json()["lines"] == []
        assert request(client, "GET", "/api/unverified-bottles", token, params={"body_code": "POOL-PERMANENT-LIVE"}).json()["total"] == 0
        assert request(client, "DELETE", f"/api/unverified-bottles/{live_record_id}", token).status_code == 404

        assert request(client, "DELETE", f"/api/unverified-bottles/{archived_record_id}", token).status_code == 204
        archive_lines = request(client, "GET", f"/api/inspection-documents/{archive_id}", token).json()["lines"]
        assert len(archive_lines) == 1
        assert archive_lines[0]["body_code"] == "POOL-PERMANENT-ARCHIVE-2"
        assert archive_lines[0]["sequence"] == 1
        with SessionLocal() as db:
            assert db.get(Bottle, live_bottle_id) is None
            assert db.get(UnverifiedBottle, archived_record_id) is None


@pytest.mark.parametrize("origin", ["GOOD", "BAD"])
def test_unverified_pool_can_edit_and_delete_allocated_outbound_bottle_without_changing_outbound_quantity(origin):
    from backend.app.database import SessionLocal
    from backend.app.models import (
        Bottle, BottleInbound, BottleInventoryAdjustment, BottleOutbound,
        BottleOutboundLine, BottleSpec, InspectionDocument, InspectionDocumentLine,
        UnverifiedBottle, User,
    )

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        with SessionLocal() as db:
            user = db.query(User).filter_by(username="admin").one()
            spec = db.query(BottleSpec).filter_by(name="15kg").one()
            inbound = db.query(BottleInbound).filter_by(status="POSTED").first()
            outbound = BottleOutbound(
                document_no=f"CYO-POOL-DELETE-{origin}", business_date=date(2026, 9, 14),
                source_id=inbound.source_id, customer="删除核验客户", handler="出库员",
                status="POSTED", created_by=user.id, posted_by=user.id, posted_at=datetime.utcnow(),
            )
            outbound.lines.append(BottleOutboundLine(
                spec_id=spec.id, quantity=1,
                good_quantity=1 if origin == "GOOD" else 0,
                bad_quantity=1 if origin == "BAD" else 0,
            ))
            db.add(outbound)
            db.flush()
            bottle = Bottle(
                inbound_id=inbound.id, source_id=inbound.source_id, spec_id=spec.id,
                body_code=f"POOL-OUTBOUND-DELETE-{origin}", zone="OUTBOUND",
                outbound_id=outbound.id, outbound_origin_zone=origin,
                process_status="INSPECTED", inspection_price=spec.inspection_price,
                coding_date=date(2026, 9, 14), good_entry_date=date(2026, 9, 14),
                coded_by=user.id, coded_at=datetime.utcnow(),
            )
            document = InspectionDocument(
                document_no=f"JY20260914-POOL-DELETE-{origin}", inspection_date=date(2026, 9, 14),
                customer_name="删除核验客户", inspector_name="核验员", status="ACTIVE", created_by=user.id,
            )
            db.add_all((bottle, document))
            db.flush()
            line = InspectionDocumentLine(
                document_id=document.id, bottle_id=bottle.id, sequence=1,
                body_code=bottle.body_code, spec_id=spec.id, volume=spec.volume,
            )
            db.add(line)
            db.flush()
            record = UnverifiedBottle(
                inspection_document_id=document.id, inspection_line_id=line.id,
                bottle_id=bottle.id, sequence=1, body_code=bottle.body_code,
                error_message="上游驳回", raw_fragment="气瓶列表第1行", status="PENDING", created_by=user.id,
            )
            db.add(record)
            db.commit()
            bottle_id, record_id, document_id, outbound_id = bottle.id, record.id, document.id, outbound.id

        corrected_code = f"POOL-OUTBOUND-CORRECTED-{origin}"
        edited = request(client, "PUT", f"/api/unverified-bottles/{record_id}", token, json={"body_code": corrected_code})
        assert edited.status_code == 200, edited.text
        assert edited.json()["body_code"] == corrected_code
        with SessionLocal() as db:
            assert db.get(Bottle, bottle_id).body_code == corrected_code
        before_stock = request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec.id)][origin.lower()]
        before_zone_total = next(item["total"] for item in request(client, "GET", "/api/bottle-zones", token).json()["stock_by_spec"] if item["spec_id"] == spec.id)
        deleted = request(client, "DELETE", f"/api/unverified-bottles/{record_id}", token)
        assert deleted.status_code == 204, deleted.text
        assert request(client, "GET", f"/api/inspection-documents/{document_id}", token).json()["lines"] == []
        assert request(client, "GET", "/api/unverified-bottles", token, params={"body_code": corrected_code}).json()["total"] == 0
        assert request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec.id)][origin.lower()] == before_stock
        assert next(item["total"] for item in request(client, "GET", "/api/bottle-zones", token).json()["stock_by_spec"] if item["spec_id"] == spec.id) == before_zone_total
        with SessionLocal() as db:
            assert db.get(Bottle, bottle_id) is None
            assert db.get(UnverifiedBottle, record_id) is None
            order = db.get(BottleOutbound, outbound_id)
            assert order.status == "POSTED" and order.total_quantity == 1
            outbound_line = order.lines[0]
            assert getattr(outbound_line, "good_adjustment_quantity" if origin == "GOOD" else "bad_adjustment_quantity") == 1
            adjustment = db.query(BottleInventoryAdjustment).filter_by(zone=origin, quantity=1).order_by(BottleInventoryAdjustment.id.desc()).first()
            assert adjustment is not None and order.document_no in adjustment.note
        voided = request(client, "POST", f"/api/bottle-outbounds/{outbound_id}/void", token)
        assert voided.status_code == 200, voided.text
        assert request(client, "GET", "/api/quick-outbound/availability", token).json()[str(spec.id)][origin.lower()] == before_stock + 1


def test_inbound_time_uses_coding_then_creation_fallback():
    from backend.app.models import Bottle, InspectionDocument, InspectionDocumentLine

    coded_at = datetime(2026, 9, 14, 2, 30)
    created_at = datetime(2026, 9, 14, 2, 29)
    document_created_at = datetime(2026, 9, 14, 3, 0)
    document = InspectionDocument(created_at=document_created_at)
    bottle = Bottle(coded_at=coded_at, created_at=created_at)
    line = InspectionDocumentLine(bottle=bottle, document=document)

    assert line.inbound_at == coded_at
    bottle.coded_at = None
    assert line.inbound_at == created_at
    bottle.created_at = None
    assert line.inbound_at == document_created_at

    archived = InspectionDocumentLine(document=InspectionDocument(created_at=created_at))
    retry = InspectionDocumentLine(document=document, source_line=archived)
    assert archived.inbound_at == created_at
    assert retry.inbound_at == created_at


def test_quick_entry_can_code_selected_source_without_using_inbound_quantity():
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        sources = []
        for name in ("快捷抄码甲单位", "快捷抄码乙单位"):
            response = request(client, "POST", "/api/bottle-sources", token, json={"name": name, "is_active": True})
            assert response.status_code == 201
            sources.append(response.json()["id"])
        spec_response = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "快捷抄码测试规格", "inspection_price": 10, "volume": 99.5,
            "wall_thickness": 2.5, "is_active": True,
        })
        assert spec_response.status_code == 201
        spec_id = spec_response.json()["id"]
        inbound_ids = []
        for source_id, business_date in zip(sources, ("2026-09-08", "2026-09-09")):
            inbound = request(client, "POST", "/api/bottle-inbounds", token, json={
                "business_date": business_date, "source_id": source_id, "handler": "测试员",
                "lines": [{"spec_id": spec_id, "pending_quantity": 1, "temporary_quantity": 0, "scrap_quantity": 0}],
            })
            assert inbound.status_code == 201
            inbound_id = inbound.json()["id"]
            assert request(client, "POST", f"/api/bottle-inbounds/{inbound_id}/post", token).status_code == 200
            inbound_ids.append(inbound_id)
        coded = request(client, "POST", "/api/quick-entry", token, json={
            "source_id": sources[1], "spec_id": spec_id, "body_code": "QUICK-SOURCE-B-0001",
        })
        assert coded.status_code == 201, coded.text
        assert (coded.json()["source_id"], coded.json()["inbound_id"]) == (sources[1], None)
        dashboard = request(client, "GET", "/api/dashboard", token, params={
            "target_date": coded.json()["coding_date"],
        })
        assert dashboard.status_code == 200
        assert dashboard.json()["daily_summary"]["inspected"] >= 1
        assert dashboard.json()["zones"]["pending"] >= 0
        pool = request(client, "GET", "/api/coding/pool?include_all_batches=true", token).json()
        assert {item["id"]: item["remaining"] for item in pool["batches"]["items"] if item["id"] in inbound_ids} == {
            inbound_ids[0]: 1, inbound_ids[1]: 1,
        }


def test_inspection_generation_splits_sources_and_restarts_each_document_sequence():
    from backend.app.database import SessionLocal
    from backend.app.models import Bottle

    inspection_date = date(2045, 1, 7)
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        source_ids = []
        for name in ("分单测试甲单位", "分单测试乙单位"):
            response = request(client, "POST", "/api/bottle-sources", token, json={"name": name, "is_active": True})
            assert response.status_code == 201
            source_ids.append(response.json()["id"])
        spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "分单测试规格", "inspection_price": 10, "volume": 35.5, "wall_thickness": 2.5, "is_active": True,
        })
        assert spec.status_code == 201
        spec_id = spec.json()["id"]

        def code(source_id: int, body_code: str, manufacture_date: str) -> None:
            created = request(client, "POST", "/api/quick-entry", token, json={
                "source_id": source_id, "spec_id": spec_id,
                "body_code": body_code, "manufacture_date": manufacture_date,
            })
            assert created.status_code == 201, created.text
            assert created.json()["inbound_id"] is None
            with SessionLocal() as db:
                db.get(Bottle, created.json()["id"]).good_entry_date = inspection_date
                db.commit()

        code(source_ids[0], "SPLIT-SOURCE-A-0001", "2022-01-01")
        code(source_ids[1], "SPLIT-SOURCE-B-0001", "2023-01-01")
        generated = request(client, "POST", "/api/inspection-documents/generate", token, json={
            "inspection_date": inspection_date.isoformat(),
        })
        assert generated.status_code == 201, generated.text
        documents = generated.json()
        assert [item["document_no"] for item in documents] == ["JY20450107", "JY20450107-02"]
        assert [item["source_id"] for item in documents] == source_ids
        assert [item["customer_name"] for item in documents] == ["分单测试甲单位", "分单测试乙单位"]
        assert [[line["sequence"] for line in item["lines"]] for item in documents] == [[1], [1]]
        assert [[line["body_code"] for line in item["lines"]] for item in documents] == [
            ["SPLIT-SOURCE-A-0001"], ["SPLIT-SOURCE-B-0001"],
        ]

        code(source_ids[0], "SPLIT-SOURCE-A-0002", "2024-01-01")
        regenerated = request(client, "POST", "/api/inspection-documents/generate", token, json={
            "inspection_date": inspection_date.isoformat(),
        })
        assert regenerated.status_code == 201, regenerated.text
        assert [item["id"] for item in regenerated.json()] == [item["id"] for item in documents]
        assert [[line["sequence"] for line in item["lines"]] for item in regenerated.json()] == [[1, 2], [1]]
        assert [line["body_code"] for line in regenerated.json()[0]["lines"]] == [
            "SPLIT-SOURCE-A-0002", "SPLIT-SOURCE-A-0001",
        ]
        synced = request(client, "POST", f"/api/inspection-documents/{documents[1]['id']}/sync", token)
        assert synced.status_code == 200
        assert [line["body_code"] for line in synced.json()["lines"]] == ["SPLIT-SOURCE-B-0001"]


def test_monthly_inventory_keeps_each_stock_status_separate():
    from backend.app.database import SessionLocal
    from backend.app.models import BottleOutbound, BottleOutboundLine, User

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        source = request(client, "POST", "/api/bottle-sources", token, json={
            "name": "月度盘点状态拆分单位", "is_active": True,
        })
        assert source.status_code == 201
        source_id = source.json()["id"]
        spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "月度盘点状态拆分规格", "inspection_price": 10, "volume": 88.8,
            "wall_thickness": 2.5, "is_active": True,
        })
        assert spec.status_code == 201
        spec_id = spec.json()["id"]
        inbound = request(client, "POST", "/api/bottle-inbounds", token, json={
            "business_date": "2046-03-05", "source_id": source_id, "handler": "盘点测试员",
            "lines": [{
                "spec_id": spec_id, "pending_quantity": 10,
                "temporary_quantity": 6, "scrap_quantity": 3,
            }],
        })
        assert inbound.status_code == 201
        assert request(client, "POST", f"/api/bottle-inbounds/{inbound.json()['id']}/post", token).status_code == 200

        with SessionLocal() as db:
            admin_id = db.query(User.id).filter(User.username == "admin").scalar()
            outbound = BottleOutbound(
                document_no="CK204603STATUS", business_date=date(2046, 3, 18), source_id=source_id,
                customer="月度盘点状态拆分单位", handler="盘点测试员", status="POSTED",
                created_by=admin_id, posted_by=admin_id, posted_at=datetime(2046, 3, 18, 12, 0),
            )
            outbound.lines = [BottleOutboundLine(
                spec_id=spec_id, quantity=9, good_quantity=4, temporary_quantity=2,
                scrap_quantity=1, bad_quantity=2,
            )]
            db.add(outbound)
            db.commit()

        report_response = request(client, "GET", "/api/monthly-inventory", token, params={
            "month_from": "2046-03", "month_to": "2046-03", "source_id": source_id,
        })
        assert report_response.status_code == 200, report_response.text
        report = report_response.json()

        def values(table: str) -> dict[str, int]:
            return {
                row["metric"]: row["values"]["2046-03"]
                for row in report["tables"][table]["rows"] if row["spec_id"] == spec_id
            }

        assert values("inbound") == {
            "INBOUND_PENDING": 10, "INBOUND_TEMPORARY": 6, "INBOUND_SCRAP": 3,
        }
        assert values("outbound") == {
            "OUTBOUND_PENDING": 4, "OUTBOUND_TEMPORARY": 2,
            "OUTBOUND_SCRAP": 1, "OUTBOUND_BAD": 2,
        }
        assert values("balance") == {
            "BALANCE_PENDING": 6, "BALANCE_TEMPORARY": 4,
            "BALANCE_SCRAP": 2, "BALANCE_BAD": -2,
        }

        changed = request(client, "PUT", "/api/monthly-inventory/value", token, json={
            "month": "2046-03", "source_id": source_id, "spec_id": spec_id,
            "metric": "BALANCE_BAD", "value": 5, "reason": "不良品月末实盘",
        })
        assert changed.status_code == 200, changed.text
        assert changed.json()["old_value"] == -2
        logs = request(client, "GET", "/api/monthly-inventory/logs", token, params={
            "source_id": source_id, "month_from": "2046-03", "month_to": "2046-03",
        }).json()
        assert logs["items"][0]["metric"] == "BALANCE_BAD"
        assert logs["items"][0]["reason"] == "不良品月末实盘"


def test_daily_inventory_uses_business_dates_and_carries_over_month_end():
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        source = request(client, "POST", "/api/bottle-sources", token, json={
            "name": "日度盘点跨月单位", "is_active": True,
        }).json()
        spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "日度盘点跨月规格", "inspection_price": 10, "volume": 88.8,
            "wall_thickness": 2.5, "is_active": True,
        }).json()
        for day, quantity in (("2047-01-30", 5), ("2047-02-01", 3)):
            inbound = request(client, "POST", "/api/bottle-inbounds", token, json={
                "business_date": day, "source_id": source["id"], "handler": "盘点员",
                "lines": [{"spec_id": spec["id"], "pending_quantity": quantity,
                    "temporary_quantity": 0, "scrap_quantity": 0}],
            })
            assert inbound.status_code == 201, inbound.text
            assert request(client, "POST", f"/api/bottle-inbounds/{inbound.json()['id']}/post", token).status_code == 200

        assert request(client, "PUT", "/api/monthly-inventory/value", token, json={
            "month": "2047-01", "source_id": source["id"], "spec_id": spec["id"],
            "metric": "BALANCE_PENDING", "value": 7,
        }).status_code == 200
        report = request(client, "GET", "/api/daily-inventory", token, params={
            "date_from": "2047-01-30", "date_to": "2047-02-02", "source_id": source["id"],
        })
        assert report.status_code == 200, report.text
        rows = report.json()["tables"]
        pending = lambda table: next(row for row in rows[table]["rows"]
            if row["metric"] == f"{table.upper()}_PENDING" and row["spec_id"] == spec["id"])
        assert pending("inbound")["values"] == {
            "2047-01-30": 5, "2047-01-31": 0, "2047-02-01": 3, "2047-02-02": 0,
        }
        assert pending("balance")["values"] == {
            "2047-01-30": 5, "2047-01-31": 7, "2047-02-01": 10, "2047-02-02": 10,
        }
        changed = request(client, "PUT", "/api/daily-inventory/value", token, json={
            "business_date": "2047-02-01", "source_id": source["id"], "spec_id": spec["id"],
            "metric": "BALANCE_PENDING", "value": 12, "reason": "当日实盘",
        })
        assert changed.status_code == 200, changed.text
        assert changed.json()["old_value"] == 10
        updated = request(client, "GET", "/api/daily-inventory", token, params={
            "date_from": "2047-02-01", "date_to": "2047-02-02", "source_id": source["id"],
        }).json()
        balance_row = next(row for row in updated["tables"]["balance"]["rows"]
            if row["metric"] == "BALANCE_PENDING" and row["spec_id"] == spec["id"])
        assert balance_row["values"] == {"2047-02-01": 12, "2047-02-02": 12}
        assert balance_row["modified_months"] == ["2047-02-01"]
        logs = request(client, "GET", "/api/daily-inventory/logs", token, params={
            "source_id": source["id"], "date_from": "2047-02-01", "date_to": "2047-02-01",
        }).json()
        assert logs["items"][0]["reason"] == "当日实盘"
        assert request(client, "GET", "/api/daily-inventory", token, params={
            "date_from": "2047-01-01", "date_to": "2047-02-02",
        }).status_code == 400


def test_bad_outbound_deducts_by_date_source_and_spec_without_opening_balance():
    from backend.app.database import SessionLocal
    from backend.app.models import BottleOutbound, BottleOutboundLine, User

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        source_ids = []
        for name in ("不良品扣减甲公司", "不良品扣减乙公司"):
            result = request(client, "POST", "/api/bottle-sources", token, json={"name": name, "is_active": True})
            assert result.status_code == 201, result.text
            source_ids.append(result.json()["id"])
        spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "不良品扣减测试规格", "inspection_price": 10, "volume": 88.8,
            "wall_thickness": 2.5, "is_active": True,
        }).json()
        with SessionLocal() as db:
            admin_id = db.query(User.id).filter(User.username == "admin").scalar()
            for number, (day, source_id, quantity) in enumerate((
                (date(2049, 5, 2), source_ids[0], 2),
                (date(2049, 5, 4), source_ids[0], 3),
                (date(2049, 5, 4), source_ids[1], 1),
            ), start=1):
                outbound = BottleOutbound(
                    document_no=f"CK204905BAD{number:03d}", business_date=day, source_id=source_id,
                    customer="不良品扣减测试", handler="测试员", status="POSTED",
                    created_by=admin_id, posted_by=admin_id, posted_at=datetime(2049, 5, 4, 12, 0),
                )
                outbound.lines = [BottleOutboundLine(spec_id=spec["id"], quantity=quantity, bad_quantity=quantity)]
                db.add(outbound)
            db.commit()

        def bad_values(report, table):
            return next(row["values"] for row in report["tables"][table]["rows"]
                if row["spec_id"] == spec["id"] and row["metric"] == f"{table.upper()}_BAD")

        first = request(client, "GET", "/api/daily-inventory", token, params={
            "date_from": "2049-05-01", "date_to": "2049-05-05", "source_id": source_ids[0],
        })
        assert first.status_code == 200, first.text
        assert bad_values(first.json(), "outbound") == {
            "2049-05-01": 0, "2049-05-02": 2, "2049-05-03": 0, "2049-05-04": 3, "2049-05-05": 0,
        }
        assert bad_values(first.json(), "balance") == {
            "2049-05-01": 0, "2049-05-02": -2, "2049-05-03": -2,
            "2049-05-04": -5, "2049-05-05": -5,
        }
        second = request(client, "GET", "/api/daily-inventory", token, params={
            "date_from": "2049-05-04", "date_to": "2049-05-04", "source_id": source_ids[1],
        }).json()
        assert bad_values(second, "outbound") == {"2049-05-04": 1}
        assert bad_values(second, "balance") == {"2049-05-04": -1}
        monthly = request(client, "GET", "/api/monthly-inventory", token, params={
            "month_from": "2049-05", "month_to": "2049-05", "source_id": source_ids[0],
        }).json()
        assert bad_values(monthly, "outbound") == {"2049-05": 5}
        assert bad_values(monthly, "balance") == {"2049-05": -5}


def test_posted_inbound_can_be_voided_without_removing_coding_records():
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        source = request(client, "POST", "/api/bottle-sources", token, json={
            "name": "作废保留抄码测试单位", "is_active": True,
        })
        assert source.status_code == 201
        spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "作废保留抄码测试规格", "inspection_price": 10, "volume": 66.6,
            "wall_thickness": 2.5, "is_active": True,
        })
        assert spec.status_code == 201
        inbound = request(client, "POST", "/api/bottle-inbounds", token, json={
            "business_date": "2047-04-08", "source_id": source.json()["id"], "handler": "作废测试员",
            "lines": [{
                "spec_id": spec.json()["id"], "pending_quantity": 1,
                "temporary_quantity": 0, "scrap_quantity": 0,
            }],
        })
        assert inbound.status_code == 201
        inbound_id = inbound.json()["id"]
        assert request(client, "POST", f"/api/bottle-inbounds/{inbound_id}/post", token).status_code == 200
        coded = request(client, "POST", "/api/bottles/coding", token, json={
            "inbound_id": inbound_id, "body_code": "VOID-INBOUND-KEEP-CODING-001",
            "manufacture_date": "2040-01-01",
            "source_id": source.json()["id"], "spec_id": spec.json()["id"],
        })
        assert coded.status_code == 201, coded.text

        voided = request(client, "POST", f"/api/bottle-inbounds/{inbound_id}/void", token)
        assert voided.status_code == 200, voided.text
        assert voided.json()["status"] == "VOID"
        bottles = request(client, "GET", "/api/bottles", token, params={
            "body_code": "VOID-INBOUND-KEEP-CODING-001",
        }).json()
        assert bottles["total"] == 1
        assert bottles["items"][0]["inbound_id"] == inbound_id

        report = request(client, "GET", "/api/monthly-inventory", token, params={
            "month_from": "2047-04", "month_to": "2047-04", "source_id": source.json()["id"],
        }).json()
        pending_row = next(row for row in report["tables"]["inbound"]["rows"] if row["metric"] == "INBOUND_PENDING" and row["spec_id"] == spec.json()["id"])
        assert pending_row["values"]["2047-04"] == 0


def test_history_records_can_be_edited_and_refresh_monthly_inventory():
    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        source = request(client, "POST", "/api/bottle-sources", token, json={
            "name": "历史记录修改测试单位", "is_active": True,
        }).json()
        spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "历史记录修改测试规格", "inspection_price": 10, "volume": 77.7,
            "wall_thickness": 2.5, "is_active": True,
        }).json()
        inbound_payload = {
            "business_date": "2048-05-06", "source_id": source["id"], "handler": "原经办人",
            "lines": [{
                "spec_id": spec["id"], "pending_quantity": 2,
                "temporary_quantity": 1, "scrap_quantity": 0,
            }],
        }
        inbound = request(client, "POST", "/api/bottle-inbounds", token, json=inbound_payload).json()
        assert request(client, "POST", f"/api/bottle-inbounds/{inbound['id']}/post", token).status_code == 200
        inbound_payload.update({
            "handler": "修改后经办人",
            "lines": [{
                "spec_id": spec["id"], "pending_quantity": 4,
                "temporary_quantity": 2, "scrap_quantity": 1,
            }],
        })
        edited_inbound = request(client, "PUT", f"/api/bottle-inbounds/{inbound['id']}", token, json=inbound_payload)
        assert edited_inbound.status_code == 200, edited_inbound.text
        assert edited_inbound.json()["status"] == "POSTED"
        assert edited_inbound.json()["handler"] == "修改后经办人"
        assert edited_inbound.json()["total_quantity"] == 7

        bottle_codes = []
        for index in range(2):
            code = f"HISTORY-EDIT-{index + 1:03d}"
            bottle_codes.append(code)
            coded = request(client, "POST", "/api/bottles/coding", token, json={
                "inbound_id": inbound["id"], "body_code": code, "manufacture_date": "2040-01-01",
                "source_id": source["id"], "spec_id": spec["id"],
            })
            assert coded.status_code == 201, coded.text
        outbound = request(client, "POST", "/api/bottle-outbounds", token, json={
            "business_date": "2048-05-18", "source_id": source["id"], "handler": "原出库人",
            "lines": [{"spec_id": spec["id"], "good_quantity": 2, "temporary_quantity": 1}],
        })
        assert outbound.status_code == 201, outbound.text
        outbound = request(client, "POST", f"/api/bottle-outbounds/{outbound.json()['id']}/post", token)
        assert outbound.status_code == 200, outbound.text
        edited_outbound = request(client, "PUT", f"/api/bottle-outbounds/{outbound.json()['id']}", token, json={
            "business_date": "2048-05-18", "source_id": source["id"], "handler": "修改后出库人",
            "lines": [{
                "spec_id": spec["id"], "good_quantity": 1,
                "temporary_quantity": 2, "scrap_quantity": 1, "bad_quantity": 0,
            }],
        })
        assert edited_outbound.status_code == 200, edited_outbound.text
        assert edited_outbound.json()["status"] == "POSTED"
        assert edited_outbound.json()["total_quantity"] == 4
        bottle_zones = [
            request(client, "GET", "/api/bottles", token, params={"body_code": code}).json()["items"][0]["zone"]
            for code in bottle_codes
        ]
        assert sorted(bottle_zones) == ["GOOD", "OUTBOUND"]

        report = request(client, "GET", "/api/monthly-inventory", token, params={
            "month_from": "2048-05", "month_to": "2048-05", "source_id": source["id"],
        }).json()
        inbound_values = {
            row["metric"]: row["values"]["2048-05"] for row in report["tables"]["inbound"]["rows"]
            if row["spec_id"] == spec["id"]
        }
        outbound_values = {
            row["metric"]: row["values"]["2048-05"] for row in report["tables"]["outbound"]["rows"]
            if row["spec_id"] == spec["id"]
        }
        assert inbound_values == {
            "INBOUND_PENDING": 4, "INBOUND_TEMPORARY": 2, "INBOUND_SCRAP": 1,
        }
        assert outbound_values == {
            "OUTBOUND_PENDING": 1, "OUTBOUND_TEMPORARY": 2,
            "OUTBOUND_SCRAP": 1, "OUTBOUND_BAD": 0,
        }

        assert request(client, "POST", f"/api/bottle-outbounds/{outbound.json()['id']}/void", token).status_code == 200
        edited_void = request(client, "PUT", f"/api/bottle-outbounds/{outbound.json()['id']}", token, json={
            "business_date": "2048-05-18", "source_id": source["id"], "handler": "作废后备注修改",
            "lines": [{"spec_id": spec["id"], "scrap_quantity": 5}],
        })
        assert edited_void.status_code == 200, edited_void.text
        assert edited_void.json()["status"] == "VOID"
        operator = request(client, "POST", "/api/users", token, json={
            "username": "outbound_history_deleter", "name": "出库记录删除员", "role_name": "出库管理员",
            "password": "Outbound@123", "permissions": ["bottle_outbound:delete"],
        })
        assert operator.status_code == 201, operator.text
        operator_token = client.post("/api/auth/login", json={
            "username": "outbound_history_deleter", "password": "Outbound@123",
        }).json()["access_token"]
        forbidden = request(client, "DELETE", f"/api/bottle-outbounds/{outbound.json()['id']}", operator_token)
        assert forbidden.status_code == 403
        deleted = request(client, "DELETE", f"/api/bottle-outbounds/{outbound.json()['id']}", token)
        assert deleted.status_code == 204, deleted.text


def test_quick_outbound_uses_anonymous_stock_then_oldest_inspected_bottle_and_lifecycle():
    from backend.app.database import SessionLocal
    from backend.app.models import InspectionDocument, InspectionDocumentLine, User

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        source = request(client, "POST", "/api/bottle-sources", token, json={
            "name": "快速出库优先级测试单位", "is_active": True,
        }).json()
        spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "快速出库优先级测试规格", "inspection_price": 10, "is_active": True,
        }).json()
        codes = ["PRIORITY-OLD-001", "PRIORITY-OLD-002", "PRIORITY-NODOC-003"]
        bottles = []
        for code in codes:
            response = request(client, "POST", "/api/quick-entry", token, json={
                "source_id": source["id"], "spec_id": spec["id"],
                "body_code": code, "manufacture_date": "2020-01-01",
            })
            assert response.status_code == 201, response.text
            bottles.append(response.json())
        with SessionLocal() as db:
            user = db.query(User).filter_by(username="admin").one()
            document = InspectionDocument(
                document_no="PRIORITY-INSPECTION-001", inspection_date=date.today(),
                source_id=source["id"], customer_name=source["name"],
                created_by=user.id, status="ACTIVE",
            )
            db.add(document)
            db.flush()
            for sequence, bottle in enumerate(bottles[:2], 1):
                db.add(InspectionDocumentLine(
                    document_id=document.id, bottle_id=bottle["id"],
                    spec_id=spec["id"], sequence=sequence,
                    body_code=bottle["body_code"], manufacturer="测试厂家", volume=35,
                ))
            db.commit()
        path = "/api/quick-outbound"
        payload = {
            "business_date": "2026-09-24", "source_id": source["id"],
            "handler": "测试出库员", "lines": [{"spec_id": spec["id"], "good_quantity": 1}],
        }
        stock = request(client, "GET", f"{path}/availability", token).json()[str(spec["id"])]
        assert (stock["good"], stock["anonymous_good"], stock["verified_good"]) == (2, 0, 2)
        assert request(client, "POST", path, token, json={**payload, "lines": [{"spec_id": spec["id"], "good_quantity": 3}]}).status_code == 400
        adjustment = request(client, "POST", "/api/bottle-inventory-adjustments", token, json={
            "business_date": "2026-09-24", "spec_id": spec["id"], "zone": "GOOD",
            "quantity": 1, "note": "未绑定瓶码库存",
        })
        assert adjustment.status_code == 201, adjustment.text
        stock = request(client, "GET", f"{path}/availability", token).json()[str(spec["id"])]
        assert (stock["good"], stock["anonymous_good"], stock["verified_good"]) == (3, 1, 2)
        first = request(client, "POST", path, token, json=payload)
        assert first.status_code == 201, first.text
        assert first.json()["lines"][0]["good_adjustment_quantity"] == 1
        assert all(request(client, "GET", "/api/bottles", token, params={"body_code": code}).json()["items"][0]["zone"] == "GOOD" for code in codes)
        second = request(client, "POST", path, token, json=payload)
        assert second.status_code == 201, second.text
        assert second.json()["lines"][0]["good_adjustment_quantity"] == 0
        oldest = request(client, "GET", "/api/bottles", token, params={"body_code": codes[0]}).json()["items"][0]
        assert oldest["zone"] == "OUTBOUND"
        assert request(client, "GET", "/api/bottles", token, params={"body_code": codes[1]}).json()["items"][0]["zone"] == "GOOD"
        lifecycle = request(client, "GET", "/api/bottles/lifecycle", token, params={"body_code": codes[0]})
        assert lifecycle.status_code == 200, lifecycle.text
        assert lifecycle.json()["inspections"][0]["document_no"] == "PRIORITY-INSPECTION-001"
        assert lifecycle.json()["outbound"]["document_no"] == second.json()["document_no"]
        assert lifecycle.json()["outbound"]["posted_at"]
        assert request(client, "POST", path, token, json=payload).status_code == 201
        assert request(client, "POST", path, token, json=payload).status_code == 400
        uninspected = request(client, "GET", "/api/bottles/lifecycle", token, params={"body_code": codes[2]}).json()
        assert uninspected["inspections"] == []
        assert uninspected["outbound"] is None
        assert uninspected["bottle"]["zone"] == "GOOD"


def test_lifecycle_restores_outbound_bottle_with_stock_then_oldest_verified_replacement():
    from backend.app.database import SessionLocal
    from backend.app.models import (
        Bottle, BottleInventoryAdjustment, BottleOutbound, BottleOutboundCorrection,
        BottleOutboundLine, InspectionDocument, InspectionDocumentLine, User,
    )

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        source = request(client, "POST", "/api/bottle-sources", token, json={
            "name": "生命周期纠正测试单位", "is_active": True,
        }).json()
        spec = request(client, "POST", "/api/bottle-specs", token, json={
            "name": "生命周期纠正测试规格", "inspection_price": 10, "is_active": True,
        }).json()
        with SessionLocal() as db:
            user = db.query(User).filter_by(username="admin").one()
            outbound_ids = []
            for index in (1, 2):
                outbound = BottleOutbound(
                    document_no=f"CYO-LIFECYCLE-FIX-{index}", business_date=date(2026, 1, index),
                    source_id=source["id"], customer=source["name"], handler="出库员",
                    status="POSTED", created_by=user.id, posted_by=user.id, posted_at=datetime.utcnow(),
                )
                outbound.lines.append(BottleOutboundLine(spec_id=spec["id"], quantity=1, good_quantity=1))
                db.add(outbound)
                db.flush()
                outbound_ids.append(outbound.id)
                db.add(Bottle(
                    body_code=f"LIFECYCLE-RETURN-{index}", source_id=source["id"], spec_id=spec["id"],
                    zone="OUTBOUND", process_status="INSPECTED", outbound_id=outbound.id,
                    outbound_origin_zone="GOOD", coding_date=date(2026, 3, index),
                    good_entry_date=date(2026, 3, index), coded_by=user.id, coded_at=datetime.utcnow(),
                ))
            document = InspectionDocument(
                document_no="LIFECYCLE-VERIFIED-001", inspection_date=date(2026, 3, 1),
                source_id=source["id"], customer_name=source["name"], status="ACTIVE", created_by=user.id,
            )
            db.add(document)
            db.flush()
            for index, coding_date in ((1, date(2026, 1, 1)), (2, date(2026, 2, 1))):
                candidate = Bottle(
                    body_code=f"LIFECYCLE-CANDIDATE-{index}", source_id=source["id"], spec_id=spec["id"],
                    zone="GOOD", process_status="INSPECTED", coding_date=coding_date,
                    good_entry_date=coding_date, coded_by=user.id, coded_at=datetime.utcnow(),
                )
                db.add(candidate)
                db.flush()
                db.add(InspectionDocumentLine(
                    document_id=document.id, bottle_id=candidate.id, sequence=index,
                    body_code=candidate.body_code, spec_id=spec["id"], volume=35,
                ))
            db.add(Bottle(
                body_code="LIFECYCLE-UNVERIFIED-OLDER", source_id=source["id"], spec_id=spec["id"],
                zone="GOOD", process_status="INSPECTED", coding_date=date(2025, 1, 1),
                good_entry_date=date(2025, 1, 1), coded_by=user.id, coded_at=datetime.utcnow(),
            ))
            db.add(BottleInventoryAdjustment(
                business_date=date(2026, 3, 1), spec_id=spec["id"], zone="GOOD",
                quantity=1, note="未绑定瓶码库存", created_by=user.id,
            ))
            db.commit()

        endpoint = "/api/bottles/lifecycle"
        first_detail = request(client, "GET", endpoint, token, params={"body_code": "LIFECYCLE-RETURN-1"}).json()
        assert first_detail["correction_option"]["anonymous_good"] == 1
        assert first_detail["correction_option"]["oldest_replacement_body_code"] == "LIFECYCLE-CANDIDATE-1"
        first = request(client, "POST", f"{endpoint}/restore-good", token, json={
            "body_code": "LIFECYCLE-RETURN-1", "outbound_id": outbound_ids[0],
        })
        assert first.status_code == 200, first.text
        assert first.json()["replacement_type"] == "ANONYMOUS"
        second_detail = request(client, "GET", endpoint, token, params={"body_code": "LIFECYCLE-RETURN-2"}).json()
        assert second_detail["correction_option"]["anonymous_good"] == 0
        second = request(client, "POST", f"{endpoint}/restore-good", token, json={
            "body_code": "LIFECYCLE-RETURN-2", "outbound_id": outbound_ids[1],
        })
        assert second.status_code == 200, second.text
        assert second.json()["replacement_type"] == "BOTTLE"
        assert second.json()["replacement_body_code"] == "LIFECYCLE-CANDIDATE-1"
        assert request(client, "POST", f"{endpoint}/restore-good", token, json={
            "body_code": "LIFECYCLE-RETURN-2", "outbound_id": outbound_ids[1],
        }).status_code == 409
        with SessionLocal() as db:
            for index, outbound_id in enumerate(outbound_ids, 1):
                outbound = db.get(BottleOutbound, outbound_id)
                assert outbound.status == "POSTED" and outbound.total_quantity == 1
                assert outbound.lines[0].good_adjustment_quantity == (1 if index == 1 else 0)
                assert db.query(Bottle).filter_by(body_code=f"LIFECYCLE-RETURN-{index}").one().zone == "GOOD"
            replacement = db.query(Bottle).filter_by(body_code="LIFECYCLE-CANDIDATE-1").one()
            assert replacement.zone == "OUTBOUND" and replacement.outbound_id == outbound_ids[1]
            assert db.query(Bottle).filter_by(body_code="LIFECYCLE-UNVERIFIED-OLDER").one().zone == "GOOD"
            assert db.query(BottleOutboundCorrection).count() >= 2
        corrected = request(client, "GET", endpoint, token, params={"body_code": "LIFECYCLE-RETURN-2"}).json()
        assert corrected["outbound"] is None
        assert corrected["corrections"][0]["replacement_body_code"] == "LIFECYCLE-CANDIDATE-1"
        replacement_detail = request(client, "GET", endpoint, token, params={"body_code": "LIFECYCLE-CANDIDATE-1"}).json()
        assert replacement_detail["corrections"][0]["original_body_code"] == "LIFECYCLE-RETURN-2"


def teardown_module():
    if TEST_DB.exists():
        TEST_DB.unlink()


def test_word_templates_export_same_age_groups_for_single_and_pair():
    from backend.app.database import SessionLocal
    from backend.app.models import InspectionDocument, InspectionDocumentLine, User
    from backend.tests.test_inspection_word import document_text

    with TestClient(app) as client:
        token = client.post("/api/auth/login", json={"username": "admin", "password": "Admin@123456"}).json()["access_token"]
        with SessionLocal() as db:
            user = db.query(User).filter_by(username="admin").one()
            document = InspectionDocument(
                document_no="JY20450912-AGE-GROUPS", inspection_date=date(2045, 9, 12),
                customer_name="四年分组回归客户", created_by=user.id,
                lines=[InspectionDocumentLine(sequence=index, body_code=code, manufacture_date=manufactured)
                       for index, (code, manufactured) in enumerate([
                           ("AGE-OLDER", date(2041, 8, 31)),
                           ("AGE-BOUNDARY", date(2041, 9, 1)),
                           ("AGE-NEWER", date(2044, 1, 1)),
                           ("AGE-MISSING", None),
                       ], start=1)],
            )
            db.add(document)
            db.commit()
            document_id = document.id

        for template_key in ("default", "photo"):
            choice = request(client, "PUT", "/api/inspection-word-templates/choice", token,
                             params={"customer_name": "四年分组回归客户"}, json={"template_key": template_key})
            assert choice.status_code == 200
            pair = request(client, "POST", f"/api/inspection-documents/{document_id}/export-word-pair?report_date=2045-09-12", token)
            assert pair.status_code == 200, pair.text
            for report in pair.json()["reports"]:
                kind = report["kind"]
                is_inspection = kind == ("regular" if template_key == "default" else "assessment")
                expected = {"AGE-BOUNDARY", "AGE-NEWER"} if is_inspection else {"AGE-OLDER", "AGE-MISSING"}
                title = "定期检验" if is_inspection else "安全评定"
                label = ("定检" if template_key == "default" else "检验") if is_inspection else "评定"
                assert report["filename"].endswith(f"{label}[2].docx")
                contents = [base64.b64decode(report["content_base64"])]
                for method in ("GET", "POST"):
                    single = request(client, method, f"/api/inspection-documents/{document_id}/export-word?kind={kind}&report_date=2045-09-12", token)
                    assert single.status_code == 200, single.text
                    assert f"{label}[2].docx" in unquote(single.headers["content-disposition"])
                    contents.append(single.content)
                for content in contents:
                    text = document_text(content)
                    assert f"附表:{title}合格气瓶一览表" in text
                    assert f"WFJA-45-0912{report['report_number']:03d}" in text
                    for code in ("AGE-OLDER", "AGE-BOUNDARY", "AGE-NEWER", "AGE-MISSING"):
                        assert (code in text) == (code in expected)


def test_inspection_generation_merges_month_dates_by_customer_and_prevents_duplicates():
    from backend.app.database import SessionLocal
    from backend.app.models import Bottle

    with TestClient(app) as client:
        token = client.post('/api/auth/login', json={'username': 'admin', 'password': 'Admin@123456'}).json()['access_token']
        source_ids = []
        for name in ('多日合并甲客户', '多日合并乙客户'):
            source_ids.append(request(client, 'POST', '/api/bottle-sources', token,
                                      json={'name': name, 'is_active': True}).json()['id'])
        spec_id = request(client, 'POST', '/api/bottle-specs', token, json={
            'name': '多日合并规格', 'inspection_price': 10, 'volume': 35.5, 'is_active': True,
        }).json()['id']

        def code(source_index, number, day, manufactured):
            response = request(client, 'POST', '/api/quick-entry', token, json={
                'source_id': source_ids[source_index], 'spec_id': spec_id,
                'body_code': f'MERGE-DAYS-{number}', 'manufacture_date': manufactured,
            })
            assert response.status_code == 201, response.text
            with SessionLocal() as db:
                db.get(Bottle, response.json()['id']).good_entry_date = date.fromisoformat(day)
                db.commit()
            return response.json()['id']

        first_id = code(0, 1, '2055-02-01', '2022-01-01')
        second_id = code(0, 2, '2055-02-03', '2024-01-01')
        code(1, 3, '2055-02-03', '2023-01-01')
        code(0, 4, '2055-02-05', '2023-01-01')
        code(1, 5, '2055-02-07', '2023-01-01')
        assert request(client, 'POST', '/api/inspection-documents/generate', token,
                       json={'inspection_date': '2055-02-07'}).status_code == 201
        path = '/api/inspection-documents/generation-dates?month=2055-02'
        days = request(client, 'GET', path, token).json()
        assert len(days) == 28
        assert days[0]['selectable'] and days[0]['quantity'] == 1
        assert days[2]['quantity'] == 2 and days[2]['customer_count'] == 2
        assert not days[1]['selectable'] and not days[1]['generated']
        assert days[6]['generated'] and not days[6]['selectable']
        for payload in (
            {},
            {'inspection_dates': []},
            {'inspection_dates': ['2055-02-01']},
            {'inspection_dates': ['2055-02-01', '2055-03-01'], 'document_date': '2055-02-09'},
            {'inspection_date': '2055-02-01', 'inspection_dates': ['2055-02-01'], 'document_date': '2055-02-09'},
            {'document_date': '2055-02-09'},
        ):
            assert request(client, 'POST', '/api/inspection-documents/generate', token, json=payload).status_code == 422
        assert request(client, 'GET', '/api/inspection-documents/generation-dates?month=2055-13', token).status_code == 422
        # Conflict and empty-date validation must not partially generate the other dates.
        assert request(client, 'POST', '/api/inspection-documents/generate', token,
                       json={'inspection_dates': ['2055-02-01', '2055-02-07'], 'document_date': '2055-02-09'}).status_code == 409
        assert request(client, 'POST', '/api/inspection-documents/generate', token,
                       json={'inspection_dates': ['2055-02-01', '2055-02-02'], 'document_date': '2055-02-09'}).status_code == 400
        assert request(client, 'GET', path, token).json()[0]['selectable']
        payload = {
            'inspection_dates': ['2055-02-05', '2055-02-01', '2055-02-03', '2055-02-01'],
            'document_date': '2055-02-09',
        }
        merged = request(client, 'POST', '/api/inspection-documents/generate', token, json=payload)
        assert merged.status_code == 201, merged.text
        documents = merged.json()
        assert len(documents) == 2
        assert [item['source_id'] for item in documents] == source_ids
        assert [item['total_quantity'] for item in documents] == [3, 1]
        assert documents[0]['inspection_date'] == '2055-02-09'
        assert all(item['document_no'].startswith('JY20550209') for item in documents)
        assert documents[0]['source_dates'] == ['2055-02-01', '2055-02-03', '2055-02-05']
        assert [line['body_code'] for line in documents[0]['lines']] == ['MERGE-DAYS-2', 'MERGE-DAYS-4', 'MERGE-DAYS-1']
        assert [line['sequence'] for line in documents[0]['lines']] == [1, 2, 3]
        assert documents[1]['lines'][0]['sequence'] == 1
        assert request(client, 'GET', f"/api/inspection-documents/{documents[0]['id']}", token).json()['source_dates'] == documents[0]['source_dates']
        assert request(client, 'POST', '/api/inspection-documents/generate', token, json=payload).status_code == 409
        reopened = request(client, 'POST', '/api/inspection-documents/generate', token,
                           json={'inspection_date': '2055-02-01'}).json()
        assert [item['id'] for item in reopened] == [item['id'] for item in documents]
        days = request(client, 'GET', path, token).json()
        assert all(days[index]['generated'] and not days[index]['selectable'] for index in (0, 2, 4))
        new_id = code(0, 6, '2055-02-01', '2025-01-01')
        synced = request(client, 'POST', f"/api/inspection-documents/{documents[0]['id']}/sync", token)
        assert synced.status_code == 200, synced.text
        assert synced.json()['total_quantity'] == 4
        assert synced.json()['lines'][0]['bottle_id'] == new_id
        assert {first_id, second_id}.issubset({line['bottle_id'] for line in synced.json()['lines']})
        assert request(client, 'POST', f"/api/inspection-documents/{documents[0]['id']}/sync", token).json()['total_quantity'] == 4
        # Dates become available again only after every customer's document has been voided.
        for document in documents:
            assert request(client, 'POST', f"/api/inspection-documents/{document['id']}/void", token).status_code == 200
        assert request(client, 'GET', path, token).json()[0]['selectable']
        regenerated = request(client, 'POST', '/api/inspection-documents/generate', token, json=payload)
        assert regenerated.status_code == 201, regenerated.text
        assert not {item['document_no'] for item in documents}.intersection(item['document_no'] for item in regenerated.json())


def test_inspection_split_merges_customers_preserves_stock_and_keeps_empty_document_history():
    from backend.app.database import SessionLocal
    from backend.app.models import (Bottle, BottleSource, BottleSpec, BottleInbound, BottleOutbound,
                                    InspectionDocument, InspectionDocumentSplitHistory,
                                    InspectionDocumentExportRecord, UnverifiedBottle, User)
    day = date(2056, 4, 10)
    with TestClient(app) as client:
        token = client.post('/api/auth/login', json={'username': 'admin', 'password': 'Admin@123456'}).json()['access_token']
        with SessionLocal() as db:
            user = db.query(User).filter_by(username='admin').one()
            customers = [BottleSource(name=f'拆单客户{letter}') for letter in ('甲', '乙', '丙')]
            db.add_all(customers)
            spec = BottleSpec(name='拆单测试规格', volume=35.5, inspection_price=10)
            db.add(spec); db.flush()
            customer_ids = [item.id for item in customers]
            inbound = BottleInbound(document_no='RK20560410-SPLIT', business_date=day,
                                    source_id=customers[0].id, pending_quantity=5, handler='入库员', created_by=user.id)
            outbound = BottleOutbound(document_no='CK20560410-SPLIT', business_date=day,
                                      source_id=customers[0].id, customer=customers[0].name, handler='出库员', created_by=user.id)
            db.add_all([inbound, outbound]); db.flush()
            inbound_id, outbound_id = inbound.id, outbound.id
            for index in range(7):
                db.add(Bottle(body_code=f'SPLIT-RANGE-{index}', source_id=customers[0 if index < 5 else 1].id,
                              inbound_id=inbound.id if index < 5 else None, spec_id=spec.id,
                              manufacture_date=date(2025 - index, 1, 1), zone='GOOD',
                              coding_date=day, good_entry_date=day, coded_by=user.id))
            # A customer's document on a different date must not be a merge target.
            db.add(InspectionDocument(document_no='JY20560409-SPLIT-C', inspection_date=date(2056, 4, 9),
                                      source_id=customers[2].id, customer_name=customers[2].name, created_by=user.id))
            db.commit()
        generated = request(client, 'POST', '/api/inspection-documents/generate', token,
                            json={'inspection_date': day.isoformat()})
        assert generated.status_code == 201, generated.text
        original, existing = generated.json()
        assert original['total_quantity'] == 5 and existing['total_quantity'] == 2
        source_id = original['id']
        moved_line_ids = [line['id'] for line in original['lines'][1:3]]
        moved_bottle_ids = [line['bottle_id'] for line in original['lines'][1:3]]
        with SessionLocal() as db:
            bottle = db.get(Bottle, moved_bottle_ids[0])
            bottle.zone = 'OUTBOUND'; bottle.outbound_id = outbound_id; bottle.outbound_origin_zone = 'GOOD'
            db.add(UnverifiedBottle(inspection_document_id=source_id, inspection_line_id=moved_line_ids[0],
                                    bottle_id=bottle.id, sequence=2, body_code=bottle.body_code,
                                    manufacture_date=bottle.manufacture_date, manufacturer='', error_message='驳回测试',
                                    raw_fragment='气瓶列表第2行【驳回测试】', status='PENDING', created_by=user.id))
            db.add(InspectionDocumentExportRecord(document_id=source_id, kind=f'bottle:{moved_line_ids[0]}', export_date=day))
            db.commit()
            stock_before = {item.id: (item.source_id, item.inbound_id, item.outbound_id, item.zone, item.outbound_origin_zone)
                            for item in db.query(Bottle).filter(Bottle.body_code.like('SPLIT-RANGE-%')).all()}

        def preview(start, end, customer_index, target_id=None):
            response = request(client, 'POST', f'/api/inspection-documents/{source_id}/split-preview', token,
                               json={'start_sequence': start, 'end_sequence': end, 'customer_id': customer_ids[customer_index],
                                     'target_document_id': target_id})
            assert response.status_code == 200, response.text
            return response.json()

        def payload(data, start, end, customer_index, confirm_create=False, target_id=None):
            return {'start_sequence': start, 'end_sequence': end, 'customer_id': customer_ids[customer_index],
                    'target_document_id': target_id, 'expected_line_ids': data['line_ids'],
                    'expected_source_updated_at': data['source_updated_at'], 'expected_target_id': data['target']['id'] if data['target'] else None,
                    'expected_target_updated_at': data['target']['updated_at'] if data['target'] else None,
                    'confirm_create': confirm_create}

        path = f'/api/inspection-documents/{source_id}/split'
        for data in ({'start_sequence': 2, 'end_sequence': 1, 'customer_id': customer_ids[1]},
                     {'start_sequence': 0, 'end_sequence': 1, 'customer_id': customer_ids[1]}):
            assert request(client, 'POST', path + '-preview', token, json=data).status_code == 422
        assert request(client, 'POST', path + '-preview', token,
                       json={'start_sequence': 1, 'end_sequence': 6, 'customer_id': customer_ids[1]}).status_code == 400
        assert request(client, 'POST', path + '-preview', token,
                       json={'start_sequence': 1, 'end_sequence': 2, 'customer_id': customer_ids[0]}).status_code == 400
        first = preview(2, 3, 1)
        assert first['quantity'] == 2 and not first['will_create'] and first['target']['id'] == existing['id']
        assert first['lines'][0]['is_unverified']
        # A concurrent target modification invalidates the preview before any transfer.
        with SessionLocal() as db:
            db.get(InspectionDocument, existing['id']).updated_at = datetime.utcnow()
            db.commit()
        assert request(client, 'POST', path, token, json=payload(first, 2, 3, 1)).status_code == 409
        first = preview(2, 3, 1)
        merged = request(client, 'POST', path, token, json=payload(first, 2, 3, 1))
        assert merged.status_code == 200, merged.text
        data = merged.json()
        assert data['source']['total_quantity'] == 3 and data['target']['total_quantity'] == 4
        assert data['target']['id'] == existing['id'] and not data['history']['target_created']
        assert data['history']['quantity'] == 2 and [line['sequence'] for line in data['history']['snapshot']] == [2, 3]
        assert [line['sequence'] for line in data['target']['lines']] == [1, 2, 3, 4]
        assert set(moved_line_ids).issubset(line['id'] for line in data['target']['lines'])
        assert request(client, 'POST', path, token, json=payload(first, 2, 3, 1)).status_code in (400, 409)
        with SessionLocal() as db:
            record = db.query(UnverifiedBottle).filter_by(inspection_line_id=moved_line_ids[0]).one()
            assert record.inspection_document_id == existing['id']
            assert record.sequence == next(line['sequence'] for line in data['target']['lines'] if line['id'] == moved_line_ids[0])
            assert db.query(InspectionDocumentExportRecord).filter_by(document_id=existing['id'], kind=f'bottle:{moved_line_ids[0]}').count() == 1
        assert request(client, 'POST', f'/api/inspection-documents/{source_id}/sync', token).json()['total_quantity'] == 3
        code_records = request(client, 'GET', f'/api/bottles?coding_customer_id={customer_ids[1]}&date_from={day}&date_to={day}', token).json()
        assert code_records['total'] == 4
        for item in code_records['items']:
            assert item['effective_coding_customer_id'] == customer_ids[1]
        assert request(client, 'GET', f'/api/coding/summary?coding_customer_id={customer_ids[1]}&date_from={day}&date_to={day}', token).json()['total'] == 4
        second = preview(1, 3, 2)
        assert second['will_create'] and second['target'] is None
        assert request(client, 'POST', path, token, json=payload(second, 1, 3, 2)).status_code == 409
        emptied = request(client, 'POST', path, token, json=payload(second, 1, 3, 2, True))
        assert emptied.status_code == 200, emptied.text
        result = emptied.json()
        assert result['source']['status'] == 'SPLIT_EMPTY' and result['source']['document_no'] == original['document_no']
        assert result['source']['total_quantity'] == 0 and result['target']['total_quantity'] == 3
        assert result['target']['inspection_date'] == day.isoformat()
        assert result['history']['source_emptied'] and result['history']['target_created']
        assert request(client, 'POST', f'/api/inspection-documents/{source_id}/sync', token).status_code == 400
        history = request(client, 'GET', f'/api/inspection-documents/{source_id}/split-history', token).json()
        assert history['total'] == 2
        assert request(client, 'GET', f"/api/inspection-documents/{existing['id']}/split-history", token).json()['total'] == 1
        assert request(client, 'GET', f"/api/inspection-documents/{result['target']['id']}/split-history", token).json()['total'] == 1
        # Snapshot remains immutable when the current transferred line is edited.
        edited = request(client, 'PUT', f"/api/inspection-documents/{existing['id']}/lines/{moved_line_ids[1]}", token,
                         json={'manufacturer': '拆单后的修改'})
        assert edited.status_code == 200, edited.text
        with SessionLocal() as db:
            stock_after = {item.id: (item.source_id, item.inbound_id, item.outbound_id, item.zone, item.outbound_origin_zone)
                           for item in db.query(Bottle).filter(Bottle.body_code.like('SPLIT-RANGE-%')).all()}
            assert stock_after == stock_before
            assert db.get(BottleInbound, inbound_id).source_id == customer_ids[0]
            assert db.get(BottleOutbound, outbound_id).source_id == customer_ids[0]
            assert db.get(BottleOutbound, outbound_id).customer == '拆单客户甲'
            first_history = db.query(InspectionDocumentSplitHistory).filter_by(source_document_id=source_id).order_by(InspectionDocumentSplitHistory.id).first()
            assert all(item['manufacturer'] != '拆单后的修改' for item in first_history.snapshot)
        # History protects related documents against permanent deletion.
        new_id = result['target']['id']
        assert request(client, 'POST', f'/api/inspection-documents/{new_id}/void', token).status_code == 200
        assert request(client, 'DELETE', f'/api/inspection-documents/{new_id}', token).status_code == 409


def test_inspection_split_requires_target_choice_and_checks_permissions_and_source_revision():
    from backend.app.database import SessionLocal
    from backend.app.models import Bottle, BottleSource, BottleSpec, InspectionDocument, User
    day = date(2057, 6, 12)
    with TestClient(app) as client:
        token = client.post('/api/auth/login', json={'username': 'admin', 'password': 'Admin@123456'}).json()['access_token']
        with SessionLocal() as db:
            user_id = db.query(User.id).filter_by(username='admin').scalar()
            customers = [BottleSource(name=f'拆单目标选择{index}') for index in range(3)]
            spec = BottleSpec(name='拆单目标选择规格', volume=35.5)
            db.add_all(customers + [spec]); db.flush()
            customer_ids = [item.id for item in customers]
            for index in range(2):
                db.add(Bottle(body_code=f'SPLIT-CHOICE-{index}', source_id=customers[0].id, spec_id=spec.id,
                              manufacture_date=date(2023, 1, 1), zone='GOOD', good_entry_date=day, coding_date=day, coded_by=user_id))
            targets = [InspectionDocument(document_no=f'JY20570612-CHOICE-{index}', inspection_date=day,
                                          source_id=customers[1].id, customer_name=customers[1].name, created_by=user_id) for index in (1, 2)]
            # Simulate an older customer-name edit whose source_id was not updated.
            targets[1].source_id = customers[2].id
            db.add_all(targets); db.commit(); target_ids = [item.id for item in targets]
        generated = request(client, 'POST', '/api/inspection-documents/generate', token, json={'inspection_date': str(day)}).json()
        source = next(item for item in generated if item['source_id'] == customer_ids[0])
        path = f"/api/inspection-documents/{source['id']}/split"
        fields = {'start_sequence': 1, 'end_sequence': 1, 'customer_id': customer_ids[1]}
        preview = request(client, 'POST', path + '-preview', token, json=fields).json()
        assert preview['requires_target_selection'] and not preview['will_create']
        assert len(preview['target_candidates']) == 2

        def commit_payload(fields, preview):
            return dict(fields, expected_line_ids=preview['line_ids'], expected_source_updated_at=preview['source_updated_at'],
                        expected_target_id=preview['target']['id'] if preview['target'] else None,
                        expected_target_updated_at=preview['target']['updated_at'] if preview['target'] else None,
                        confirm_create=preview['will_create'])
        assert request(client, 'POST', path, token, json=commit_payload(fields, preview)).status_code == 409
        fields['target_document_id'] = target_ids[1]
        preview = request(client, 'POST', path + '-preview', token, json=fields).json()
        assert preview['target']['id'] == target_ids[1] and not preview['requires_target_selection']
        with SessionLocal() as db:
            db.get(InspectionDocument, source['id']).updated_at = datetime.utcnow()
            db.commit()
        assert request(client, 'POST', path, token, json=commit_payload(fields, preview)).status_code == 409
        editor = request(client, 'POST', '/api/users', token, json={
            'username': 'split_editor', 'name': '拆单编辑员', 'password': 'Editor@1234',
            'permissions': ['inspection_documents:view', 'inspection_documents:update'],
        })
        assert editor.status_code == 201, editor.text
        editor_token = client.post('/api/auth/login', json={'username': 'split_editor', 'password': 'Editor@1234'}).json()['access_token']
        preview = request(client, 'POST', path + '-preview', editor_token, json=fields).json()
        response = request(client, 'POST', path, editor_token, json=commit_payload(fields, preview))
        assert response.status_code == 200, response.text
        assert response.json()['target']['id'] == target_ids[1]
        fields = {'start_sequence': 1, 'end_sequence': 1, 'customer_id': customer_ids[2]}
        preview = request(client, 'POST', path + '-preview', editor_token, json=fields).json()
        assert preview['will_create']
        assert request(client, 'POST', path, editor_token, json=commit_payload(fields, preview)).status_code == 403
        assert request(client, 'GET', f"/api/inspection-documents/{source['id']}", token).json()['total_quantity'] == 1
        # Previously split customer assignments survive voiding and subsequent generation.
        target_id = target_ids[1]
        moved_id = response.json()['target']['lines'][0]['bottle_id']
        assert request(client, 'POST', f'/api/inspection-documents/{target_id}/void', token).status_code == 200
        regenerated = request(client, 'POST', '/api/inspection-documents/generate', token, json={'inspection_date': str(day)})
        assert regenerated.status_code == 201, regenerated.text
        assert any(item['source_id'] == customer_ids[1] and any(line['bottle_id'] == moved_id for line in item['lines']) for item in regenerated.json())


def test_restart_preserves_existing_inspection_sequence_and_rejection_mapping():
    from backend.app.database import SessionLocal
    from backend.app.main import initialize_database
    from backend.app.models import Bottle, BottleSource, BottleSpec, InspectionDocument, InspectionDocumentLine, UnverifiedBottle, User
    with TestClient(app):
        with SessionLocal() as db:
            user_id = db.query(User.id).filter_by(username='admin').scalar()
            source = BottleSource(name='升级序号保留客户')
            spec = BottleSpec(name='升级序号保留规格')
            db.add_all([source, spec]); db.flush()
            bottles = [Bottle(body_code=f'UPGRADE-SEQUENCE-{index}', source_id=source.id, spec_id=spec.id,
                              manufacture_date=date(year, 1, 1), coded_by=user_id, zone='GOOD')
                       for index, year in enumerate((2020, 2024))]
            db.add_all(bottles); db.flush()
            document = InspectionDocument(document_no='JY20580101-PRESERVE', inspection_date=date(2058, 1, 1),
                                          source_id=source.id, customer_name=source.name, created_by=user_id)
            document.lines = [InspectionDocumentLine(bottle_id=bottle.id, spec_id=spec.id, sequence=index + 1,
                                                     body_code=bottle.body_code, manufacture_date=bottle.manufacture_date,
                                                     manufacturer='', volume=0, inspection_result='GOOD')
                              for index, bottle in enumerate(bottles)]
            db.add(document); db.flush()
            line_ids = [line.id for line in document.lines]
            record = UnverifiedBottle(inspection_document_id=document.id, inspection_line_id=line_ids[0],
                                      bottle_id=bottles[0].id, sequence=1, body_code=bottles[0].body_code,
                                      manufacturer='', error_message='历史第一行驳回', raw_fragment='气瓶列表第1行',
                                      status='PENDING', created_by=user_id)
            db.add(record); db.commit(); document_id, record_id = document.id, record.id
        initialize_database()
        initialize_database()
        with SessionLocal() as db:
            preserved = db.get(InspectionDocument, document_id)
            assert [(line.id, line.sequence) for line in preserved.lines] == list(zip(line_ids, (1, 2)))
            assert db.get(UnverifiedBottle, record_id).sequence == 1
            assert preserved.document_no == 'JY20580101-PRESERVE'


def test_retry_documents_group_by_current_document_customer_preserve_stock_and_original_history():
    from backend.app.database import SessionLocal
    from backend.app.models import Bottle, BottleSource, BottleSpec, InspectionDocument, InspectionDocumentLine, UnverifiedBottle, User
    prefix = 'RETRY-CUSTOMER-VERIFY'
    with TestClient(app) as client:
        token = client.post('/api/auth/login', json={'username': 'admin', 'password': 'Admin@123456'}).json()['access_token']
        with SessionLocal() as db:
            user_id = db.query(User.id).filter_by(username='admin').scalar()
            customers = [BottleSource(name=f'重检分组{letter}客户') for letter in ('甲', '乙')]
            spec = BottleSpec(name='重检按客户分组规格', volume=35.5, wall_thickness=2.5)
            db.add_all(customers + [spec]); db.flush()
            customer_ids = [item.id for item in customers]
            documents = [InspectionDocument(document_no=f'{prefix}-{index}', inspection_date=date(2060, 2, index + 1),
                                             # The first original document simulates a customer-name edit with a stale source_id.
                                             source_id=customers[1].id if index in (0, 2) else customers[0].id if index == 1 else None,
                                             customer_name=customers[0].name if index < 2 else customers[1].name if index == 2 else '历史文字客户',
                                             inspector_name=f'检验员{index}', is_archived=index >= 3, created_by=user_id)
                         for index in range(5)]
            db.add_all(documents); db.flush()
            source_ids = [item.id for item in documents]
            records = []
            source_line_ids = []
            for index, document_index in enumerate((0, 0, 1, 2, 3, 4)):
                document = documents[document_index]
                bottle = None
                if document_index < 3:
                    bottle = Bottle(body_code=f'{prefix}-B-{index}', source_id=customers[1].id, spec_id=spec.id,
                                    manufacture_date=date(2020 + index, 1, 1), zone='GOOD', coded_by=user_id,
                                    coding_customer_id=customers[0].id if document_index < 2 else customers[1].id)
                    db.add(bottle); db.flush()
                line = InspectionDocumentLine(document_id=document.id, bottle_id=bottle.id if bottle else None,
                                              spec_id=spec.id, sequence=2 if index == 1 else 1,
                                              body_code=f'{prefix}-B-{index}', manufacture_date=date(2020 + index, 1, 1),
                                              manufacturer=f'厂家{index}', volume=35.5, inspector_name=document.inspector_name,
                                              inspection_result='GOOD')
                db.add(line); db.flush()
                source_line_ids.append(line.id)
                record = UnverifiedBottle(inspection_document_id=document.id, inspection_line_id=line.id,
                                          bottle_id=bottle.id if bottle else None, sequence=line.sequence,
                                          body_code=line.body_code, manufacture_date=line.manufacture_date, manufacturer=line.manufacturer,
                                          error_message=f'平台未通过{index}', raw_fragment=f'气瓶列表第{line.sequence}行',
                                          status='PENDING', created_by=user_id)
                db.add(record); records.append(record)
            db.commit()
            record_ids = [item.id for item in records]
            original_sequence = {item.id: [(line.id, line.sequence, line.body_code) for line in item.lines] for item in documents}
            original_stock = {bottle.id: (bottle.source_id, bottle.coding_customer_id, bottle.zone, bottle.inbound_id, bottle.outbound_id)
                              for bottle in db.query(Bottle).filter(Bottle.body_code.like(f'{prefix}%')).all()}
        retry_date = date(2062, 4, 5)
        response = request(client, 'POST', '/api/unverified-bottles/create-retry-document', token,
                           json={'document_no': prefix, 'inspection_date': retry_date.isoformat()})
        assert response.status_code == 201, response.text
        generated = response.json()
        assert len(generated) == 3
        assert all(item['inspection_date'] == retry_date.isoformat() for item in generated)
        assert all(item['document_no'].startswith('JY20620405-R') for item in generated)
        by_customer = {item['customer_name']: item for item in generated}
        assert {name: item['total_quantity'] for name, item in by_customer.items()} == {
            '重检分组甲客户': 3, '重检分组乙客户': 1, '历史文字客户': 2,
        }
        assert by_customer['重检分组甲客户']['source_id'] == customer_ids[0]
        assert by_customer['重检分组乙客户']['source_id'] == customer_ids[1]
        assert by_customer['历史文字客户']['source_id'] is None
        assert len({item['document_no'] for item in generated}) == 3
        assert all(item['is_retry'] and item['is_archived'] for item in generated)
        expected_sources = {'重检分组甲客户': set(source_line_ids[:3]), '重检分组乙客户': {source_line_ids[3]}, '历史文字客户': set(source_line_ids[4:])}
        retry_ids = []
        for name, summary in by_customer.items():
            detail = request(client, 'GET', f"/api/inspection-documents/{summary['id']}", token).json()
            retry_ids.append(summary['id'])
            assert {line['source_line_id'] for line in detail['lines']} == expected_sources[name]
            assert [line['sequence'] for line in detail['lines']] == list(range(1, len(detail['lines']) + 1))
            assert all(line['bottle_id'] is None for line in detail['lines'])
            assert [line['manufacture_date'] for line in detail['lines']] == sorted((line['manufacture_date'] for line in detail['lines']), reverse=True)
        assert request(client, 'POST', '/api/unverified-bottles/create-retry-document', token, json={'document_no': prefix}).status_code == 400
        with SessionLocal() as db:
            for source_id in source_ids:
                document = db.get(InspectionDocument, source_id)
                assert [(line.id, line.sequence, line.body_code) for line in document.lines] == original_sequence[source_id]
            assert all(db.get(UnverifiedBottle, record_id).status == 'RETRYING' for record_id in record_ids)
            assert {bottle.id: (bottle.source_id, bottle.coding_customer_id, bottle.zone, bottle.inbound_id, bottle.outbound_id)
                    for bottle in db.query(Bottle).filter(Bottle.body_code.like(f'{prefix}%')).all()} == original_stock
            assert all(db.get(UnverifiedBottle, record_id).error_message.startswith('平台未通过') for record_id in record_ids)
        # Voiding one customer's retry document only returns that customer's records to the pool.
        target = by_customer['重检分组甲客户']
        assert request(client, 'POST', f"/api/inspection-documents/{target['id']}/void", token).status_code == 200
        pool = request(client, 'GET', '/api/unverified-bottles', token, params={'document_no': prefix}).json()
        assert pool['total'] == 3 and {item['customer_name'] for item in pool['items']} == {'重检分组甲客户'}
        filtered = request(client, 'POST', '/api/unverified-bottles/create-retry-document', token,
                           json={'body_code': f'{prefix}-B-0'})
        assert filtered.status_code == 201 and len(filtered.json()) == 1 and filtered.json()[0]['total_quantity'] == 1
        assert filtered.json()[0]['customer_name'] == '重检分组甲客户'


def test_retry_customer_validation_is_atomic_and_retry_numbers_do_not_collide_with_gaps():
    from backend.app.api import china_today
    from backend.app.database import SessionLocal
    from backend.app.models import InspectionDocument, InspectionDocumentLine, UnverifiedBottle, User
    with TestClient(app) as client:
        token = client.post('/api/auth/login', json={'username': 'admin', 'password': 'Admin@123456'}).json()['access_token']
        with SessionLocal() as db:
            user_id = db.query(User.id).filter_by(username='admin').scalar()
            documents = [InspectionDocument(document_no=f'RETRY-MISSING-CUSTOMER-{index}', inspection_date=date(2061, 3, 1),
                                             customer_name='已设置重检客户' if index == 0 else '', is_archived=True, created_by=user_id)
                         for index in (0, 1)]
            db.add_all(documents); db.flush()
            record_ids = []
            for index, document in enumerate(documents):
                line = InspectionDocumentLine(document_id=document.id, sequence=1, body_code=f'RETRY-MISSING-{index}',
                                              manufacturer='', volume=0, inspection_result='GOOD')
                db.add(line); db.flush()
                record = UnverifiedBottle(inspection_document_id=document.id, inspection_line_id=line.id, sequence=1,
                                          body_code=line.body_code, manufacturer='', error_message='未通过', raw_fragment='第1行',
                                          status='PENDING', created_by=user_id)
                db.add(record); db.flush(); record_ids.append(record.id)
            missing_document_id = documents[1].id
            # A retained high suffix used to conflict with count+1 numbering after deletion.
            reserved = f'JY{china_today():%Y%m%d}-R99'
            db.add(InspectionDocument(document_no=reserved, inspection_date=china_today(), customer_name='保留的历史重检客户',
                                      is_retry=True, is_archived=True, status='VOID', created_by=user_id))
            db.commit()
            count_before = db.query(InspectionDocument).count()
        path = '/api/unverified-bottles/create-retry-document'
        response = request(client, 'POST', path, token, json={'document_no': 'RETRY-MISSING-CUSTOMER'})
        assert response.status_code == 400 and '未设置客户' in response.json()['detail']
        with SessionLocal() as db:
            assert db.query(InspectionDocument).count() == count_before
            assert all(db.get(UnverifiedBottle, record_id).status == 'PENDING' for record_id in record_ids)
            db.get(InspectionDocument, missing_document_id).customer_name = '另一重检客户'
            db.commit()
        response = request(client, 'POST', path, token, json={'document_no': 'RETRY-MISSING-CUSTOMER'})
        assert response.status_code == 201, response.text
        assert len(response.json()) == 2
        assert len({item['document_no'] for item in response.json()}) == 2
        assert all(int(item['document_no'].rsplit('-R', 1)[1]) > 99 for item in response.json())
