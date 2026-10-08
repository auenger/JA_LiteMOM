from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook
import pytest
import xlrd

from backend.app.inspection_excel import parse_full_inspection_excel
from backend.app.inspection_xls import AGE_HEADERS, FULL_HEADERS, inspection_age_table_rows, render_inspection_xls, split_inspection_lines_by_age


def _line(sequence: int, manufactured: date | None, *, unverified: bool = False):
    return SimpleNamespace(
        sequence=sequence, device_category="焊接气瓶", body_code=f"CODE-{sequence}",
        manufacture_date=manufactured, manufacturer="测试厂家", volume=Decimal("35.50"),
        next_inspection_date=None, inspection_result="GOOD", inspector_name=None,
        inspection_time=None, is_unverified=unverified,
    )


def _document():
    return SimpleNamespace(
        inspection_date=date(2026, 9, 12), customer_name="测试客户", inspector_name="测试员",
        lines=[
            _line(1, date(2024, 1, 1)),
            _line(2, date(2020, 1, 1), unverified=True),
            _line(3, date(2022, 9, 12)),
            _line(4, None),
            _line(5, date(2024, 1, 1)),
        ],
    )


def _sheet(content: bytes):
    assert content.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
    workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
    sheet = workbook.sheet_by_index(0)
    for row in range(sheet.nrows):
        for column in range(sheet.ncols):
            border = workbook.xf_list[sheet.cell(row, column).xf_index].border
            assert (border.left_line_style, border.right_line_style, border.top_line_style, border.bottom_line_style) == (1, 1, 1, 1)
    return sheet


def test_full_and_verified_exports_are_biff8_with_black_borders():
    document = _document()
    full_content, full_count = render_inspection_xls(document, "full")
    full = _sheet(full_content)
    assert full.name == "检验气瓶列表"
    assert full_count == 5
    assert tuple(full.row_values(1)) == FULL_HEADERS
    assert full.cell_value(0, 2) == "测试客户"
    assert [full.cell_value(row, 2) for row in range(2, 7)] == [f"CODE-{index}" for index in range(1, 6)]
    assert full.cell_value(2, 5) == "35.50"
    assert full.cell_value(2, 6) == "2026.09.12"
    assert full.cell_value(2, 7) == "2030.09.11"

    chosen_content, _ = render_inspection_xls(document, "full", date(2031, 10, 5))
    chosen = _sheet(chosen_content)
    assert [chosen.cell_value(row, 6) for row in range(2, 7)] == ["2031.10.05"] * 5
    assert [chosen.cell_value(row, 7) for row in range(2, 7)] == ["2035.10.04"] * 5
    customer, inspector, parsed = parse_full_inspection_excel(chosen_content, date(2031, 10, 5), "full.xls")
    assert customer == "测试客户"
    assert inspector == "测试员"
    assert len(parsed) == 5
    assert document.inspection_date == date(2026, 9, 12)

    verified_content, verified_count = render_inspection_xls(document, "verified")
    verified = _sheet(verified_content)
    assert verified_count == 4
    assert [verified.cell_value(row, 2) for row in range(2, 6)] == ["CODE-1", "CODE-3", "CODE-4", "CODE-5"]
    assert [verified.cell_value(row, 0) for row in range(2, 6)] == ["1", "2", "3", "4"]


def test_xlsx_with_renamed_data_sheet_is_accepted_only_when_headers_identify_one_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["充装单位名称", None, "测试客户"])
    sheet.append(list(FULL_HEADERS))
    sheet.append([1, "焊接气瓶", "001234", "2022-09", "测试厂家", "35.50", "2026.09.17", "2030.09.16", "合格", "测试员"])
    workbook.create_sheet("Sheet2")
    workbook.create_sheet("Sheet3")
    output = BytesIO()
    workbook.save(output)
    customer, inspector, rows = parse_full_inspection_excel(output.getvalue(), date(2026, 9, 17), "renamed.xlsx")
    assert (customer, inspector, len(rows), rows[0]["body_code"]) == ("测试客户", "测试员", 1, "001234")

    workbook.copy_worksheet(sheet).title = "another-copy"
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match="唯一一张"):
        parse_full_inspection_excel(output.getvalue(), date(2026, 9, 17), "ambiguous.xlsx")


def test_word_appendix_rows_keep_production_order_and_restart_sequence():
    document = _document()
    older, within = split_inspection_lines_by_age(document)
    regular = inspection_age_table_rows(document, older)
    assessment = inspection_age_table_rows(document, within)
    assert AGE_HEADERS == ("序号", "气瓶编号", "制造日期", "制造单位名称或代号", "公称容积（L）", "下次检验日期")
    assert [row[1] for row in regular] == ["CODE-4"]
    assert [row[1] for row in assessment] == ["CODE-1", "CODE-5", "CODE-3"]
    assert [row[0] for row in assessment] == ["1", "2", "3"]
    assert [row[2] for row in assessment] == ["2024.01", "2024.01", "2022.09"]
    assert [row[4:] for row in assessment] == [("35.50", "2030.09.11")] * 3
    assert assessment[0][0] == "1"


def test_word_age_groups_include_boundary_month_and_exclude_rejected_bottles():
    document = _document()
    document.lines.extend([_line(6, date(2022, 8, 31)), _line(7, date(2022, 9, 1))])
    older, within = split_inspection_lines_by_age(document)
    assert [row[1] for row in inspection_age_table_rows(document, older)] == ["CODE-6", "CODE-4"]
    assert [row[1] for row in inspection_age_table_rows(document, within)] == ["CODE-1", "CODE-5", "CODE-3", "CODE-7"]


def test_selected_date_controls_age_boundary_and_overrides_stored_next_date():
    document = _document()
    document.lines[0].next_inspection_date = date(2030, 9, 1)
    chosen_date = date(2027, 9, 14)
    older, within = split_inspection_lines_by_age(document, chosen_date)
    assert [line.body_code for line in older] == ["CODE-3", "CODE-4"]
    assert [line.body_code for line in within] == ["CODE-1", "CODE-5"]
    assert [row[-1] for row in inspection_age_table_rows(document, within, chosen_date)] == ["2031.09.13"] * 2
    workbook, _ = render_inspection_xls(document, "full", chosen_date)
    sheet = _sheet(workbook)
    assert [sheet.cell_value(row, 7) for row in range(2, sheet.nrows)] == ["2031.09.13"] * 5
